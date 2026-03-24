from __future__ import annotations

import json
import csv
import io
import logging
import os
import re
import shutil
import sqlite3
import hashlib
import time
import threading
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional
from urllib.parse import urlparse
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

BASE_DIR = Path(__file__).resolve().parents[1]
logger = logging.getLogger(__name__)
PROJECTS_ROOT = BASE_DIR / "data" / "workspace" / "projects"

RUN_TYPE_LABELS = {
    "all": "All",
    "exploratory": "Exploratory",
    "scenario": "Scenario",
    "spec_validation": "Spec Validation",
}
_EXCEL_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
EXPORT_RUN_LIMIT = 400


@dataclass
class ProjectInfo:
    slug: str
    name: str
    domain: str
    created_at: str


def _read_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def list_projects() -> List[ProjectInfo]:
    projects: List[ProjectInfo] = []
    if not PROJECTS_ROOT.exists():
        return projects
    directories: List[Path] = []
    try:
        with os.scandir(PROJECTS_ROOT) as entries:
            for entry in entries:
                try:
                    if entry.is_dir(follow_symlinks=False):
                        directories.append(Path(entry.path))
                except OSError:
                    continue
    except OSError:
        return projects
    for directory in sorted(directories):
        project_json = _read_json(directory / "project.json", {})
        slug = directory.name
        projects.append(
            ProjectInfo(
                slug=slug,
                name=str(project_json.get("name", slug)).strip() or slug,
                domain=str(project_json.get("domain", "")).strip(),
                created_at=str(project_json.get("created_at", "")).strip(),
            )
        )
    return projects


def _slugify_project_name(name: str) -> str:
    text = str(name or "").strip().lower()
    if not text:
        return "project"
    text = re.sub(r"[^a-z0-9]+", "-", text)
    text = text.strip("-")
    return text or "project"


def create_project(name: str, domain: str = "") -> ProjectInfo:
    raw_name = str(name or "").strip()
    if not raw_name:
        raise ValueError("프로젝트 이름은 필수입니다.")
    slug_base = _slugify_project_name(raw_name)
    slug = slug_base
    idx = 1
    while (PROJECTS_ROOT / slug).exists():
        idx += 1
        slug = f"{slug_base}-{idx}"
    root = PROJECTS_ROOT / slug
    root.mkdir(parents=True, exist_ok=True)
    now = datetime.now().isoformat()

    _write_json(
        root / "project.json",
        {
            "name": raw_name,
            "slug": slug,
            "domain": str(domain or "").strip(),
            "created_at": now,
        },
    )

    artifacts = root / "artifacts"
    artifacts.mkdir(parents=True, exist_ok=True)
    _write_json(
        artifacts / "sidebar_tree_meta.json",
        {
            "versions": [{"id": "ver_1_0", "label": "ver 1.0", "created_at": now}],
            "session_to_version": {},
            "session_labels": {},
        },
    )
    _write_json(artifacts / "product_ui_versions.json", {})
    _write_json(artifacts / "product_ui_settings.json", {})
    _write_json(artifacts / "product_ui_run_meta.json", {})
    _write_json(artifacts / "definition_specs.json", {"items": []})

    return ProjectInfo(slug=slug, name=raw_name, domain=str(domain or "").strip(), created_at=now)


def get_project_paths(project_slug: str) -> Dict[str, Path]:
    root = PROJECTS_ROOT / project_slug
    return {
        "root": root,
        "artifacts": root / "artifacts",
        "qa_sessions": root / "qa_sessions",
        "exports": root / "exports" / "event_definition_history",
        "definition_specs_dir": root / "definition_specs",
        "definition_specs_meta": root / "artifacts" / "definition_specs.json",
        "meta": root / "artifacts" / "sidebar_tree_meta.json",
        "version_notes": root / "artifacts" / "product_ui_versions.json",
        "settings": root / "artifacts" / "product_ui_settings.json",
        "run_meta": root / "artifacts" / "product_ui_run_meta.json",
    }


def _parse_iso(text: str) -> Optional[datetime]:
    raw = str(text or "").strip()
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except Exception:
        return None


def _slugify_filename(text: str) -> str:
    raw = str(text or "").strip()
    if not raw:
        return "definition.csv"
    raw = re.sub(r"[^\w\-.]+", "_", raw, flags=re.UNICODE)
    return raw[:120] or "definition.csv"


def _iter_export_runs(exports_root: Path) -> Iterator[Path]:
    try:
        project_dirs: List[str] = []
        with os.scandir(exports_root) as project_entries:
            for project_entry in project_entries:
                try:
                    if not project_entry.is_dir(follow_symlinks=False):
                        continue
                except OSError:
                    continue
                project_dirs.append(project_entry.path)
        for project_path in sorted(project_dirs, reverse=True):
                try:
                    run_dirs: List[str] = []
                    with os.scandir(project_path) as run_entries:
                        for run_entry in run_entries:
                            try:
                                if run_entry.is_dir(follow_symlinks=False):
                                    run_dirs.append(run_entry.path)
                            except OSError:
                                continue
                    for run_path in sorted(run_dirs, reverse=True):
                        yield Path(run_path)
                except OSError as exc:
                    logger.warning("collect_export_runs failed for %s: %s", project_path, exc)
    except OSError as exc:
        logger.warning("collect_export_runs failed for %s: %s", exports_root, exc)


def _collect_export_runs(exports_root: Path, max_runs: int = EXPORT_RUN_LIMIT) -> List[Path]:
    runs = [run_dir for run_dir in _iter_export_runs(exports_root)]
    def _safe_mtime(path: Path) -> float:
        try:
            return path.stat().st_mtime
        except OSError:
            return 0.0
    runs.sort(key=_safe_mtime, reverse=True)
    return runs[: int(max(1, max_runs))]


def _normalize_col_name(value: str) -> str:
    return str(value or "").strip().lower().replace(" ", "").replace("_", "")


def _extract_definition_summary(file_path: Path) -> Dict[str, Any]:
    if not file_path.exists():
        return {"event_names": [], "columns": [], "sample_rows": []}
    suffix = file_path.suffix.lower()
    key_candidates = {"eventname", "event", "이벤트명", "이벤트"}
    columns: List[str] = []
    sample_rows: List[Dict[str, str]] = []
    event_names: List[str] = []

    def append_event_names(rows: List[Dict[str, Any]], cols: List[str]) -> None:
        event_col = ""
        for c in cols:
            if _normalize_col_name(c) in key_candidates:
                event_col = c
                break
        if not event_col:
            return
        for row in rows:
            value = str((row or {}).get(event_col, "")).strip()
            if value and value not in event_names:
                event_names.append(value)

    if suffix in {".xlsx", ".xls"}:
        try:
            df = pd.read_excel(file_path, dtype=str)
            df = df.fillna("")
            columns = [str(c).strip() for c in list(df.columns) if str(c).strip()]
            rows_all = [
                {str(k).strip(): str(v).strip() for k, v in row.items() if str(k).strip()}
                for row in df.to_dict(orient="records")
            ]
            append_event_names(rows_all, columns)
            sample_rows = rows_all[:8]
        except Exception:
            return {"event_names": [], "columns": [], "sample_rows": []}
    else:
        content = b""
        try:
            content = file_path.read_bytes()
        except Exception:
            return {"event_names": [], "columns": [], "sample_rows": []}
        rows_text = None
        for enc in ("utf-8-sig", "utf-8", "cp949"):
            try:
                rows_text = content.decode(enc)
                break
            except Exception:
                continue
        if rows_text is None:
            return {"event_names": [], "columns": [], "sample_rows": []}
        try:
            reader = csv.DictReader(rows_text.splitlines())
            if not reader.fieldnames:
                return {"event_names": [], "columns": [], "sample_rows": []}
            columns = [str(c).strip() for c in reader.fieldnames if str(c).strip()]
            rows_all: List[Dict[str, str]] = []
            for row in reader:
                if not isinstance(row, dict):
                    continue
                rows_all.append(
                    {
                        str(k).strip(): str(v).strip()
                        for k, v in row.items()
                        if str(k).strip()
                    }
                )
            append_event_names(rows_all, columns)
            sample_rows = rows_all[:8]
        except Exception:
            return {"event_names": [], "columns": [], "sample_rows": []}

    # trim payload size
    trimmed_rows: List[Dict[str, str]] = []
    for row in sample_rows[:8]:
        item: Dict[str, str] = {}
        for idx, key in enumerate(columns[:10]):
            val = str((row or {}).get(key, "")).strip()
            item[key] = val[:160]
            if idx >= 9:
                break
        trimmed_rows.append(item)

    return {
        "event_names": event_names[:2000],
        "columns": columns[:50],
        "sample_rows": trimmed_rows,
    }


def load_definition_specs(project_slug: str) -> List[Dict[str, Any]]:
    paths = get_project_paths(project_slug)
    payload = _read_json(paths["definition_specs_meta"], {})
    rows = payload.get("items", []) if isinstance(payload, dict) else []
    out: List[Dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        if bool(row.get("deleted", False)):
            continue
        spec_id = str(row.get("id", "")).strip()
        if not spec_id:
            continue
        out.append(
            {
                "id": spec_id,
                "name": str(row.get("name", "")).strip() or spec_id,
                "filename": str(row.get("filename", "")).strip(),
                "uploaded_at": str(row.get("uploaded_at", "")).strip(),
                "event_count": int(row.get("event_count", 0) or 0),
                "event_names": row.get("event_names", []) if isinstance(row.get("event_names", []), list) else [],
                "columns": row.get("columns", []) if isinstance(row.get("columns", []), list) else [],
                "sample_rows": row.get("sample_rows", []) if isinstance(row.get("sample_rows", []), list) else [],
                "path": str(row.get("path", "")).strip(),
            }
        )
    out.sort(key=lambda x: _parse_iso(x.get("uploaded_at", "") or "") or datetime.min, reverse=True)
    return out


def add_definition_spec(project_slug: str, original_name: str, content: bytes) -> Dict[str, Any]:
    paths = get_project_paths(project_slug)
    spec_dir = paths["definition_specs_dir"]
    spec_dir.mkdir(parents=True, exist_ok=True)

    now = datetime.now()
    safe_name = _slugify_filename(original_name)
    if "." not in safe_name:
        safe_name = f"{safe_name}.csv"
    ext = Path(safe_name).suffix.lower()
    if ext not in {".csv", ".xlsx", ".xls"}:
        raise ValueError("정의서는 csv/xlsx/xls 형식만 업로드할 수 있습니다.")
    if ext in {".xlsx", ".xls"}:
        try:
            import openpyxl  # type: ignore  # noqa: F401
        except Exception as exc:
            raise ValueError("xlsx 업로드를 사용하려면 openpyxl 설치가 필요합니다.") from exc
    spec_id = f"spec_{now.strftime('%Y%m%d_%H%M%S')}_{uuid4().hex[:6]}"
    stored_name = f"{spec_id}_{safe_name}"
    full_path = spec_dir / stored_name
    full_path.write_bytes(content)

    summary = _extract_definition_summary(full_path)
    event_names = summary.get("event_names", []) if isinstance(summary.get("event_names", []), list) else []
    columns = summary.get("columns", []) if isinstance(summary.get("columns", []), list) else []
    sample_rows = summary.get("sample_rows", []) if isinstance(summary.get("sample_rows", []), list) else []
    item = {
        "id": spec_id,
        "name": Path(safe_name).stem,
        "filename": safe_name,
        "uploaded_at": now.isoformat(),
        "event_count": len(event_names),
        "event_names": event_names[:500],
        "columns": columns[:50],
        "sample_rows": sample_rows[:8],
        "path": str(full_path.relative_to(BASE_DIR)),
        "deleted": False,
    }

    meta = _read_json(paths["definition_specs_meta"], {})
    if not isinstance(meta, dict):
        meta = {}
    items = meta.get("items", []) if isinstance(meta.get("items", []), list) else []
    items.append(item)
    meta["items"] = items
    _write_json(paths["definition_specs_meta"], meta)
    return item


def _build_definition_template_event_entries() -> List[Dict[str, Any]]:
    return [
        {
            "state": "상용",
            "POC": "APP(Android, iOS) / MW",
            "no": "1",
            "action": "click",
            "object": "content",
            "event_name": "click_content",
            "event_type": "click",
            "description": "스토어홈 > 추천판 > 빅배너 클릭",
            "match_mode": "normal",
            "event 적합 검사": "Y",
            "ecommerce": "N",
            "params": [
                ("section_name", "main_display_big_banner"),
                ("section_index", "{index}"),
                ("section_title", "{title}"),
                ("content_id", "1234"),
                ("content_name", "{image}"),
                ("content_type", "banner"),
                ("banner_id", "{banner_id}"),
                ("ab_probs", "{ab_probs}"),
                ("landing_url", "{{landing_url}}"),
                ("page_id", "/main/musinsa/recommend"),
                ("page_link", "{{page_link}}"),
            ],
        },
        {
            "state": "상용",
            "POC": "APP(Android, iOS) / MW",
            "no": "2",
            "action": "impression",
            "object": "content",
            "event_name": "impression_content",
            "event_type": "impression",
            "description": "스토어홈 > 추천판 > 빅배너 노출",
            "match_mode": "normal",
            "event 적합 검사": "Y",
            "ecommerce": "N",
            "params": [
                ("section_name", "main_display_big_banner"),
                ("section_index", "{index}"),
                ("section_title", "{title}"),
                ("content_name", "{image}"),
                ("content_type", "banner"),
                ("banner_id", "{banner_id}"),
                ("page_id", "/main/musinsa/recommend"),
                ("page_link", "{{page_link}}"),
            ],
        },
        {
            "state": "상용",
            "POC": "APP(Android, iOS) / MW",
            "no": "3",
            "action": "click",
            "object": "button",
            "event_name": "click_button",
            "event_type": "click",
            "description": "스토어홈 > 추천판 > 전체보기 버튼 클릭",
            "match_mode": "normal",
            "event 적합 검사": "Y",
            "ecommerce": "N",
            "params": [
                ("section_name", "main_display_big_banner"),
                ("button_id", "all"),
                ("button_name", "전체보기 버튼"),
                ("landing_url", "{{landing_url}}"),
                ("page_id", "/main/musinsa/recommend"),
                ("page_link", "{{page_link}}"),
            ],
        },
    ]


def build_definition_template_rows() -> List[List[str]]:
    headers = [
        "state",
        "POC",
        "no",
        "action",
        "object",
        "event_name",
        "event_type",
        "description",
        "match_mode",
        "event 적합 검사",
        "ecommerce",
        "param_name",
        "param_value",
    ]
    rows: List[List[str]] = [headers]
    for item in _build_definition_template_event_entries():
        params = item.get("params", []) if isinstance(item.get("params", []), list) else []
        if not params:
            params = [("", "")]
        for param_name, param_value in params:
            rows.append(
                [
                    str(item.get("state", "")),
                    str(item.get("POC", "")),
                    str(item.get("no", "")),
                    str(item.get("action", "")),
                    str(item.get("object", "")),
                    str(item.get("event_name", "")),
                    str(item.get("event_type", "")),
                    str(item.get("description", "")),
                    str(item.get("match_mode", "")),
                    str(item.get("event 적합 검사", "")),
                    str(item.get("ecommerce", "")),
                    str(param_name or ""),
                    str(param_value or ""),
                ]
            )
    return rows


def build_definition_template_csv_bytes() -> bytes:
    rows = build_definition_template_rows()
    return pd.DataFrame(rows).to_csv(index=False, header=False).encode("utf-8-sig")


def build_definition_template_xlsx_bytes() -> bytes:
    rows = build_definition_template_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "definition_template"
    for row in rows:
        ws.append(row)

    if ws.max_column >= 3 and ws.max_row > 1:
        current_no = ""
        start_row = 2
        merge_columns = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11]
        for r in range(2, ws.max_row + 2):
            value = ""
            if r <= ws.max_row:
                value = str(ws.cell(row=r, column=3).value or "")
            if r == 2:
                current_no = value
                start_row = r
                continue
            if value != current_no or r > ws.max_row:
                end_row = r - 1
                if current_no and end_row > start_row:
                    for col in merge_columns:
                        ws.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)
                        ws.cell(row=start_row, column=col).alignment = Alignment(vertical="center", horizontal="center")
                start_row = r
                current_no = value

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def delete_definition_spec(project_slug: str, spec_id: str) -> bool:
    sid = str(spec_id or "").strip()
    if not sid:
        return False
    paths = get_project_paths(project_slug)
    meta = _read_json(paths["definition_specs_meta"], {})
    if not isinstance(meta, dict):
        return False
    items = meta.get("items", [])
    if not isinstance(items, list):
        return False
    changed = False
    next_items: List[Dict[str, Any]] = []
    for row in items:
        if not isinstance(row, dict):
            continue
        if str(row.get("id", "")).strip() == sid:
            changed = True
            rel = str(row.get("path", "")).strip()
            if rel:
                p = (BASE_DIR / rel).resolve()
                if p.exists() and p.is_file():
                    try:
                        p.unlink()
                    except Exception:
                        pass
            continue
        next_items.append(row)
    if changed:
        meta["items"] = next_items
        _write_json(paths["definition_specs_meta"], meta)
    return changed


def _suggest_display_name(version_id: str, created_at: str) -> str:
    dt = _parse_iso(created_at)
    if dt is None:
        raw = str(version_id or "").strip()
        digits = "".join(ch for ch in raw if ch.isdigit())
        if len(digits) >= 8:
            try:
                dt = datetime(
                    year=int(digits[0:4]),
                    month=int(digits[4:6]),
                    day=max(1, int(digits[6:8])),
                )
            except Exception:
                dt = None
    if dt is None:
        return str(version_id or "버전").strip() or "버전"
    return f"{dt.year} {dt.month}월 개편 QA"


def _read_session_rows(db_path: Path) -> List[Dict[str, Any]]:
    if not db_path.exists():
        return []
    query = """
        SELECT session_id, target_url, status, started_at, ended_at, captured_events, last_error, tester_name
        FROM qa_sessions
        ORDER BY started_at DESC
    """
    rows: List[Dict[str, Any]] = []
    try:
        with sqlite3.connect(str(db_path)) as conn:
            conn.row_factory = sqlite3.Row
            for row in conn.execute(query):
                rows.append(dict(row))
    except Exception:
        return []
    return rows


def _fetch_event_sample(db_path: Path, session_id: str, limit: int = 400) -> List[Dict[str, Any]]:
    if not db_path.exists():
        return []
    query = """
        SELECT captured_at, source, event_name, params_json
        FROM qa_events
        WHERE session_id = ?
        ORDER BY captured_at DESC
        LIMIT ?
    """
    rows: List[Dict[str, Any]] = []
    try:
        with sqlite3.connect(str(db_path)) as conn:
            conn.row_factory = sqlite3.Row
            for row in conn.execute(query, (session_id, int(limit))):
                item = dict(row)
                try:
                    item["params"] = json.loads(str(item.get("params_json", "{}")) or "{}")
                except Exception:
                    item["params"] = {}
                rows.append(item)
    except Exception:
        return []
    return rows


def _fetch_session_source_counts(db_path: Path, session_id: str) -> Dict[str, int]:
    if not db_path.exists():
        return {}
    query = """
        SELECT source, COUNT(*) AS cnt
        FROM qa_events
        WHERE session_id = ?
        GROUP BY source
    """
    out: Dict[str, int] = {}
    try:
        with sqlite3.connect(str(db_path)) as conn:
            conn.row_factory = sqlite3.Row
            for row in conn.execute(query, (session_id,)):
                source = str(row["source"] if isinstance(row, sqlite3.Row) else row[0]).strip()
                cnt = int((row["cnt"] if isinstance(row, sqlite3.Row) else row[1]) or 0)
                if source:
                    out[source] = cnt
    except Exception:
        return {}
    return out


def _infer_run_type(events: List[Dict[str, Any]]) -> str:
    if not events:
        return "exploratory"
    seen = " ".join(
        [
            f"{str(e.get('source', ''))} {str(e.get('event_name', ''))} {json.dumps(e.get('params', {}), ensure_ascii=False)}"
            for e in events[:200]
        ]
    ).lower()
    if any(token in seen for token in ["definition", "spec", "row_match", "validation_target_mode"]):
        return "spec_validation"
    if any(token in seen for token in ["scenario", "funnel", "transaction_id", "auto_crawl_scenario_group"]):
        return "scenario"
    return "exploratory"


def _to_qa_status(events: List[Dict[str, Any]], session_status: str) -> str:
    if str(session_status).strip().lower() in {"running", "stopping"}:
        return "Unchecked"
    if not events:
        return "Unchecked"

    text_blob = " ".join(
        [
            f"{str(e.get('source', ''))} {str(e.get('event_name', ''))} {json.dumps(e.get('params', {}), ensure_ascii=False)}"
            for e in events[:300]
        ]
    ).lower()
    ga_hits = sum(1 for e in events if str(e.get("source", "")).strip() == "ga_hit")

    if "access_challenge_detected" in text_blob or "captcha" in text_blob or "blocked" in text_blob:
        return "Blocked"
    if any(token in text_blob for token in ["fail_event_missing", "missing_parameter", "required_missing", "미도달"]):
        return "Missing"
    if any(token in text_blob for token in ["mismatch", "fail_param", "fail_section", "wrong_row_match", "불일치"]):
        return "Mismatch"
    if any(token in text_blob for token in ["retest", "재검증"]):
        return "Retest Needed"
    if ga_hits > 0:
        return "Matched"
    return "Unchecked"


def _analytics_source_summary(
    events: List[Dict[str, Any]],
    source_counts: Optional[Dict[str, int]] = None,
) -> Dict[str, Any]:
    counts = source_counts if isinstance(source_counts, dict) else {}
    ga4_hits = int(counts.get("ga_hit", 0) or 0)
    amplitude_hits = int(counts.get("amplitude_hit", 0) or 0)
    has_ga4 = ga4_hits > 0
    has_amplitude = amplitude_hits > 0
    if not counts:
        for e in events[:1000]:
            source = str(e.get("source", "")).strip().lower()
            if source == "ga_hit":
                has_ga4 = True
                ga4_hits += 1
            elif source == "amplitude_hit":
                has_amplitude = True
                amplitude_hits += 1
    label = "Unknown"
    if has_ga4 and has_amplitude:
        label = "GA4+Amplitude"
    elif has_ga4:
        label = "GA4"
    elif has_amplitude:
        label = "Amplitude"
    return {
        "has_ga4": has_ga4,
        "has_amplitude": has_amplitude,
        "ga4_hits": ga4_hits,
        "amplitude_hits": amplitude_hits,
        "analytics_source_label": label,
    }


def load_versions(project_slug: str, include_deleted: bool = False) -> List[Dict[str, Any]]:
    paths = get_project_paths(project_slug)
    meta = _read_json(paths["meta"], {})
    versions = meta.get("versions", []) if isinstance(meta.get("versions", []), list) else []
    version_notes = _read_json(paths["version_notes"], {})
    sessions = load_sessions(project_slug)
    session_count_by_version: Dict[str, int] = {}
    for sess in sessions:
        vid = str(sess.get("version_id", "")).strip()
        if not vid:
            continue
        session_count_by_version[vid] = session_count_by_version.get(vid, 0) + 1

    active_id = ""
    for key, row in version_notes.items() if isinstance(version_notes, dict) else []:
        if isinstance(row, dict) and bool(row.get("is_default", False)):
            active_id = str(key).strip()
            break

    out: List[Dict[str, Any]] = []
    for idx, row in enumerate(versions):
        if not isinstance(row, dict):
            continue
        vid = str(row.get("id", "")).strip()
        if not vid:
            continue
        note = version_notes.get(vid, {}) if isinstance(version_notes.get(vid, {}), dict) else {}
        deleted = bool(note.get("deleted", False))
        if deleted and not include_deleted:
            continue
        display_name = (
            str(note.get("display_name", "")).strip()
            or str(row.get("label", "")).strip()
            or _suggest_display_name(vid, str(note.get("created_at", "")).strip())
        )
        status = str(note.get("status", "")).strip() or ("Active" if vid == active_id else "Draft")
        if status not in {"Active", "Draft", "Archived"}:
            status = "Draft"
        if deleted:
            status = "Deleted"
        created_at = str(note.get("created_at", "")).strip() or str(row.get("created_at", "")).strip()
        if not created_at:
            created_at = datetime.now().isoformat()
        out.append(
            {
                "id": vid,
                "display_name": display_name,
                "created_at": created_at,
                "status": status,
                "is_default": bool(note.get("is_default", False)) or status == "Active",
                "change_reason": str(note.get("change_reason", "")).strip(),
                "definition_link": str(note.get("definition_link", "")).strip(),
                "site_change_memo": str(note.get("site_change_memo", "")).strip(),
                "session_count": int(session_count_by_version.get(vid, 0)),
                "sort_index": idx,
                "deleted": deleted,
            }
        )
    # Active > Draft > Archived, then original order
    order = {"Active": 0, "Draft": 1, "Archived": 2, "Deleted": 3}
    out.sort(key=lambda x: (order.get(str(x.get("status", "Draft")), 3), int(x.get("sort_index", 9999))))
    return out


def save_version_note(
    project_slug: str,
    version_id: str,
    display_name: str,
    status: str,
    change_reason: str,
    definition_link: str,
    site_change_memo: str,
    action: str = "",
) -> None:
    paths = get_project_paths(project_slug)
    payload = _read_json(paths["version_notes"], {})
    if not isinstance(payload, dict):
        payload = {}
    existing = payload.get(version_id, {}) if isinstance(payload.get(version_id, {}), dict) else {}
    normalized_status = str(status).strip() or str(existing.get("status", "")).strip() or "Draft"
    if normalized_status not in {"Active", "Draft", "Archived"}:
        normalized_status = "Draft"
    normalized_action = str(action or "").strip().lower()
    if normalized_action == "archive":
        normalized_status = "Archived"
    if normalized_action == "delete":
        # soft delete (trash)
        payload[version_id] = {
            **existing,
            "deleted": True,
            "status": "Archived",
            "is_default": False,
            "deleted_at": datetime.now().isoformat(),
            "updated_at": datetime.now().isoformat(),
        }
        _write_json(paths["version_notes"], payload)
        return
    if normalized_action == "restore":
        payload[version_id] = {
            **existing,
            "deleted": False,
            "status": str(existing.get("status", "")).strip() or "Draft",
            "updated_at": datetime.now().isoformat(),
        }
        _write_json(paths["version_notes"], payload)
        return
    if normalized_action == "hard_delete":
        payload.pop(version_id, None)
        meta_payload = _read_json(paths["meta"], {})
        if isinstance(meta_payload, dict):
            versions = meta_payload.get("versions", [])
            if isinstance(versions, list):
                meta_payload["versions"] = [
                    v for v in versions
                    if not (isinstance(v, dict) and str(v.get("id", "")).strip() == version_id)
                ]
            session_to_version = meta_payload.get("session_to_version", {})
            if isinstance(session_to_version, dict):
                meta_payload["session_to_version"] = {
                    sid: vid
                    for sid, vid in session_to_version.items()
                    if str(vid).strip() != version_id
                }
            _write_json(paths["meta"], meta_payload)
        _write_json(paths["version_notes"], payload)
        return
    if normalized_action == "set_default":
        for key, row in list(payload.items()):
            if not isinstance(row, dict):
                continue
            row["is_default"] = False
            if str(row.get("status", "")).strip() == "Active":
                row["status"] = "Draft"
            payload[key] = row
        normalized_status = "Active"

    payload[version_id] = {
        "created_at": str(existing.get("created_at", "")).strip() or datetime.now().isoformat(),
        "display_name": (
            str(display_name).strip()
            or str(existing.get("display_name", "")).strip()
            or _suggest_display_name(version_id, str(existing.get("created_at", "")).strip())
        ),
        "status": normalized_status,
        "is_default": normalized_action == "set_default" or bool(existing.get("is_default", False)),
        "change_reason": str(change_reason).strip(),
        "definition_link": str(definition_link).strip(),
        "site_change_memo": str(site_change_memo).strip(),
        "deleted": False,
        "updated_at": datetime.now().isoformat(),
    }
    if normalized_status == "Active":
        for key, row in list(payload.items()):
            if key == version_id or not isinstance(row, dict):
                continue
            row["is_default"] = False
            if str(row.get("status", "")).strip() == "Active":
                row["status"] = "Draft"
            payload[key] = row
    _write_json(paths["version_notes"], payload)


def load_settings(project_slug: str) -> Dict[str, Any]:
    paths = get_project_paths(project_slug)
    payload = _read_json(paths["settings"], {})
    if not isinstance(payload, dict):
        payload = {}
    saved_pages = payload.get("saved_start_pages", [])
    if not isinstance(saved_pages, list):
        saved_pages = []
    normalized_saved_pages: List[Dict[str, str]] = []
    for row in saved_pages:
        if not isinstance(row, dict):
            continue
        name = str(row.get("name", "")).strip()
        url = _normalize_url(str(row.get("url", "")).strip())
        if not name or not url:
            continue
        normalized_saved_pages.append({"name": name, "url": url})

    scenario_groups = payload.get("scenario_page_groups", [])
    if not isinstance(scenario_groups, list):
        scenario_groups = []
    normalized_groups: List[Dict[str, Any]] = []
    for group in scenario_groups:
        if not isinstance(group, dict):
            continue
        name = str(group.get("name", "")).strip()
        raw_urls = group.get("urls", [])
        urls: List[str] = []
        if isinstance(raw_urls, list):
            for item in raw_urls:
                parsed = _normalize_url(str(item).strip())
                if parsed and parsed not in urls:
                    urls.append(parsed)
        if name and urls:
            normalized_groups.append({"name": name, "urls": urls})

    return {
        "base_domain": _normalize_url(str(payload.get("base_domain", "")).strip()),
        "default_start_url": _normalize_url(str(payload.get("default_start_url", "")).strip()),
        "browser": str(payload.get("browser", "chromium")).strip() or "chromium",
        "viewport": str(payload.get("viewport", "Desktop 1440x900")).strip() or "Desktop 1440x900",
        "collection_option": str(payload.get("collection_option", "Auto Crawl")).strip() or "Auto Crawl",
        "judgement_rule": str(payload.get("judgement_rule", "Network First")).strip() or "Network First",
        "saved_start_pages": normalized_saved_pages,
        "scenario_page_groups": normalized_groups,
    }


def _normalize_url(url: str) -> str:
    text = str(url or "").strip()
    if not text:
        return ""
    parsed = urlparse(text)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return ""
    return text


def _parse_saved_pages(payload: Dict[str, Any]) -> List[Dict[str, str]]:
    direct = payload.get("saved_start_pages", [])
    if isinstance(direct, list):
        out: List[Dict[str, str]] = []
        for row in direct:
            if not isinstance(row, dict):
                continue
            name = str(row.get("name", "")).strip()
            url = _normalize_url(str(row.get("url", "")).strip())
            if name and url:
                out.append({"name": name, "url": url})
        return out[:100]

    raw_text = str(payload.get("saved_start_pages_text", "")).strip()
    out2: List[Dict[str, str]] = []
    if not raw_text:
        return out2
    for line in raw_text.splitlines():
        line = str(line).strip()
        if not line:
            continue
        if "|" in line:
            name, url = line.split("|", 1)
        else:
            name, url = line, line
        name = str(name).strip()
        url = _normalize_url(str(url).strip())
        if name and url:
            out2.append({"name": name, "url": url})
    return out2[:100]


def _parse_scenario_groups(payload: Dict[str, Any]) -> List[Dict[str, Any]]:
    direct = payload.get("scenario_page_groups", [])
    if isinstance(direct, list):
        out: List[Dict[str, Any]] = []
        for row in direct:
            if not isinstance(row, dict):
                continue
            name = str(row.get("name", "")).strip()
            raw_urls = row.get("urls", [])
            urls: List[str] = []
            if isinstance(raw_urls, list):
                for item in raw_urls:
                    parsed = _normalize_url(str(item).strip())
                    if parsed and parsed not in urls:
                        urls.append(parsed)
            if name and urls:
                out.append({"name": name, "urls": urls})
        return out[:80]

    raw_text = str(payload.get("scenario_page_groups_text", "")).strip()
    out2: List[Dict[str, Any]] = []
    if not raw_text:
        return out2
    for line in raw_text.splitlines():
        line = str(line).strip()
        if not line:
            continue
        if "|" not in line:
            continue
        name, urls_blob = line.split("|", 1)
        name = str(name).strip()
        urls: List[str] = []
        for token in str(urls_blob).split(","):
            parsed = _normalize_url(str(token).strip())
            if parsed and parsed not in urls:
                urls.append(parsed)
        if name and urls:
            out2.append({"name": name, "urls": urls})
    return out2[:80]


def save_settings(project_slug: str, payload: Dict[str, Any]) -> None:
    paths = get_project_paths(project_slug)
    sanitized = {
        "base_domain": _normalize_url(str(payload.get("base_domain", "")).strip()),
        "default_start_url": _normalize_url(str(payload.get("default_start_url", "")).strip()),
        "browser": str(payload.get("browser", "chromium")).strip() or "chromium",
        "viewport": str(payload.get("viewport", "Desktop 1440x900")).strip() or "Desktop 1440x900",
        "collection_option": str(payload.get("collection_option", "Auto Crawl")).strip() or "Auto Crawl",
        "judgement_rule": str(payload.get("judgement_rule", "Network First")).strip() or "Network First",
        "saved_start_pages": _parse_saved_pages(payload),
        "scenario_page_groups": _parse_scenario_groups(payload),
        "updated_at": datetime.now().isoformat(),
    }
    _write_json(paths["settings"], sanitized)


def load_sessions(project_slug: str) -> List[Dict[str, Any]]:
    paths = get_project_paths(project_slug)
    db_path = paths["qa_sessions"] / "qa_runs.db"
    rows = _read_session_rows(db_path)
    meta = _read_json(paths["meta"], {})
    run_meta = _read_json(paths["run_meta"], {})
    if not isinstance(run_meta, dict):
        run_meta = {}
    session_to_version = meta.get("session_to_version", {}) if isinstance(meta.get("session_to_version", {}), dict) else {}
    session_labels = meta.get("session_labels", {}) if isinstance(meta.get("session_labels", {}), dict) else {}

    out: List[Dict[str, Any]] = []
    for row in rows:
        session_id = str(row.get("session_id", "")).strip()
        if not session_id:
            continue
        meta_item = run_meta.get(session_id, {}) if isinstance(run_meta.get(session_id, {}), dict) else {}
        events = _fetch_event_sample(db_path, session_id, limit=300)
        source_counts = _fetch_session_source_counts(db_path, session_id)
        run_type = str(meta_item.get("run_type", "")).strip() or _infer_run_type(events)
        qa_status = _to_qa_status(events, str(row.get("status", "")))
        source_summary = _analytics_source_summary(events, source_counts=source_counts)
        out.append(
            {
                "session_id": session_id,
                "name": str(session_labels.get(session_id, "")).strip() or session_id,
                "version_id": str(meta_item.get("version_id", "")).strip()
                or str(session_to_version.get(session_id, "ver_1_0")).strip()
                or "ver_1_0",
                "version_name": str(meta_item.get("version_name", "")).strip() or "",
                "run_type": run_type,
                "run_type_label": RUN_TYPE_LABELS.get(run_type, "Exploratory"),
                "qa_mode": str(meta_item.get("qa_mode", "")).strip() or "-",
                "browser": str(meta_item.get("browser", "")).strip(),
                "viewport": str(meta_item.get("viewport", "")).strip(),
                "collection_option": str(meta_item.get("collection_option", "")).strip(),
                "analytics_source_mode": str(meta_item.get("analytics_source_mode", "both")).strip() or "both",
                "judgement_rule": str(meta_item.get("judgement_rule", "")).strip(),
                "tester": str(row.get("tester_name", "")).strip() or "-",
                "started_at": str(row.get("started_at", "")).strip(),
                "ended_at": str(row.get("ended_at", "")).strip(),
                "runtime_status": str(row.get("status", "")).strip() or "-",
                "captured_events": int(row.get("captured_events", 0) or 0),
                "target_url": str(row.get("target_url", "")).strip(),
                "qa_status": qa_status,
                "last_error": str(row.get("last_error", "")).strip(),
                "test_scope": str(meta_item.get("test_scope", "start_page")).strip() or "start_page",
                "start_page_mode": str(meta_item.get("start_page_mode", "use_default")).strip() or "use_default",
                "saved_page_name": str(meta_item.get("saved_page_name", "")).strip(),
                "scenario_group_name": str(meta_item.get("scenario_group_name", "")).strip(),
                "definition_spec_id": str(meta_item.get("definition_spec_id", "")).strip(),
                "definition_spec_name": str(meta_item.get("definition_spec_name", "")).strip(),
                "definition_event_target_count": int(meta_item.get("definition_event_target_count", 0) or 0),
                "definition_runtime_row_count": int(meta_item.get("definition_runtime_row_count", 0) or 0),
                "analytics_source_label": str(source_summary.get("analytics_source_label", "Unknown")),
                "ga4_hits": int(source_summary.get("ga4_hits", 0) or 0),
                "amplitude_hits": int(source_summary.get("amplitude_hits", 0) or 0),
            }
        )
    out.sort(key=lambda x: _parse_iso(x.get("started_at", "") or "") or datetime.min, reverse=True)
    return out


def _viewport_to_size(viewport: str) -> tuple[int, int]:
    text = str(viewport or "").strip().lower()
    if "tablet" in text or "768" in text:
        return (768, 1024)
    if "1920" in text:
        return (1920, 1080)
    if "390" in text or "mobile" in text:
        return (390, 844)
    return (1440, 900)


def start_session(
    project_slug: str,
    version_id: str,
    run_type: str,
    qa_mode: str,
    tester_name: str,
    start_mode: str,
    test_scope: str = "start_page",
    start_page_mode: str = "use_default",
    start_url: str = "",
    saved_page_name: str = "",
    scenario_group_name: str = "",
    definition_spec_id: str = "",
    viewport: str = "",
    analytics_source_mode: str = "both",
) -> Dict[str, Any]:
    from src.debug_runtime import start_debug_session

    paths = get_project_paths(project_slug)
    settings = load_settings(project_slug)
    base_domain = _normalize_url(str(settings.get("base_domain", "")).strip())
    default_start_url = _normalize_url(str(settings.get("default_start_url", "")).strip())
    saved_pages = settings.get("saved_start_pages", [])
    scenario_groups = settings.get("scenario_page_groups", [])
    definition_specs = load_definition_specs(project_slug)

    qa_mode_text = str(qa_mode or "").strip() or "전체 탐색 테스트"
    requested_run_type = str(run_type or "").strip().lower()
    requested_scope = str(test_scope or "start_page").strip().lower()
    requested_start_mode = str(start_mode or "manual").strip().lower()

    # QA 모드 단일 입력 정책:
    # - run_type / test_scope / start_mode 는 qa_mode 에서 파생한다.
    if qa_mode_text == "전체 탐색 테스트":
        requested_run_type = "exploratory"
        requested_scope = "site_wide"
        requested_start_mode = "auto"
    elif qa_mode_text == "시나리오 테스트":
        requested_run_type = "scenario"
        requested_scope = "scenario_group"
        requested_start_mode = "manual"
    elif qa_mode_text == "정의서 검증":
        requested_run_type = "spec_validation"
        requested_scope = "single_url"
        requested_start_mode = "manual"

    if requested_scope not in {"site_wide", "start_page", "scenario_group", "single_url"}:
        requested_scope = "start_page"

    requested_page_mode = str(start_page_mode or "use_default").strip().lower()
    if requested_page_mode not in {"use_default", "direct_input", "saved_page"}:
        requested_page_mode = "use_default"

    resolved_target_url = ""
    resolved_saved_page_name = str(saved_page_name or "").strip()
    resolved_scenario_group_name = str(scenario_group_name or "").strip()
    resolved_scenario_seed_urls: List[str] = []

    if requested_scope == "site_wide":
        resolved_target_url = base_domain or default_start_url
        if not resolved_target_url:
            raise ValueError("전체 사이트 탐색은 Settings의 기본 도메인이 필요합니다.")
    elif requested_scope == "single_url":
        resolved_target_url = _normalize_url(str(start_url or "").strip())
        if not resolved_target_url:
            raise ValueError("단일 URL 검증은 시작 URL을 정확히 입력해야 합니다.")
    elif requested_scope == "scenario_group":
        if not resolved_scenario_group_name and isinstance(scenario_groups, list) and scenario_groups:
            first_group = next((g for g in scenario_groups if isinstance(g, dict)), {})
            resolved_scenario_group_name = str(first_group.get("name", "")).strip()
        group = next(
            (
                g
                for g in scenario_groups
                if isinstance(g, dict) and str(g.get("name", "")).strip() == resolved_scenario_group_name
            ),
            None,
        )
        if not isinstance(group, dict):
            raise ValueError("선택한 시나리오 페이지군을 찾을 수 없습니다. Settings에서 먼저 저장해 주세요.")
        urls = group.get("urls", [])
        if not isinstance(urls, list) or not urls:
            raise ValueError("선택한 시나리오 페이지군에 유효한 URL이 없습니다.")
        resolved_scenario_seed_urls = [str(x).strip() for x in urls if _normalize_url(str(x).strip())]
        if not resolved_scenario_seed_urls:
            raise ValueError("선택한 시나리오 페이지군에 유효한 URL이 없습니다.")
        resolved_target_url = resolved_scenario_seed_urls[0]
    else:
        # start_page scope
        if requested_page_mode == "direct_input":
            resolved_target_url = _normalize_url(str(start_url or "").strip())
        elif requested_page_mode == "saved_page":
            item = next(
                (
                    p
                    for p in saved_pages
                    if isinstance(p, dict) and str(p.get("name", "")).strip() == resolved_saved_page_name
                ),
                None,
            )
            if isinstance(item, dict):
                resolved_target_url = _normalize_url(str(item.get("url", "")).strip())
        else:
            resolved_target_url = default_start_url or base_domain
        if not resolved_target_url:
            raise ValueError("시작 페이지를 결정할 수 없습니다. 세션 시작 옵션 또는 Settings를 확인해 주세요.")

    session_id = f"dbg_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{uuid4().hex[:8]}"
    session_dir = paths["qa_sessions"] / session_id
    output_file = session_dir / "debug_stream.jsonl"
    output_file.parent.mkdir(parents=True, exist_ok=True)
    selected_viewport = str(viewport or "").strip() or str(settings.get("viewport", "")).strip() or "Desktop 1440x900"
    source_mode = str(analytics_source_mode or "both").strip().lower()
    if source_mode not in {"both", "ga4", "amplitude"}:
        source_mode = "both"
    width, height = _viewport_to_size(selected_viewport)
    viewport_text = selected_viewport.lower()
    is_mobile_viewport = ("mobile" in viewport_text) or ("390" in viewport_text)
    browser_name = str(settings.get("browser", "chromium")).strip().lower()
    if browser_name not in {"chromium", "chrome"}:
        browser_name = "chromium"

    run_settings = {
        "environment": "prod",
        "browser_name": browser_name,
        "viewport_width": int(width),
        "viewport_height": int(height),
        "mobile_mode": bool(is_mobile_viewport),
        "mobile_device": "iPhone 13" if is_mobile_viewport else "",
        "auto_crawl_enabled": True,
        "auto_stop_after_crawl": False,
        "qa_mode": qa_mode_text,
        "auto_crawl_start_mode": "auto" if requested_start_mode == "auto" else "manual",
        "restrict_to_start_url": requested_scope in {"start_page", "single_url", "scenario_group"},
        "single_page_only": requested_scope == "single_url",
        "block_link_navigation": requested_scope == "single_url",
        "scenario_template": resolved_scenario_group_name if requested_scope == "scenario_group" else "",
        "analytics_source_mode": source_mode,
    }
    selected_spec_id = str(definition_spec_id or "").strip()
    selected_spec: Dict[str, Any] = {}
    if qa_mode_text == "정의서 검증":
        if not selected_spec_id and definition_specs:
            selected_spec_id = str(definition_specs[0].get("id", "")).strip()
        selected_spec = next((x for x in definition_specs if str(x.get("id", "")).strip() == selected_spec_id), {})
        if not selected_spec:
            raise ValueError("정의서 검증 모드는 정의서 업로드/선택이 필요합니다.")
        runtime_payload = _build_definition_runtime_payload(project_slug, selected_spec_id)
        runtime_event_names = runtime_payload.get("event_names", []) if isinstance(runtime_payload.get("event_names", []), list) else []
        run_settings["definition_file_name"] = str(selected_spec.get("filename", "")).strip()
        run_settings["definition_event_names"] = runtime_event_names or (
            selected_spec.get("event_names", []) if isinstance(selected_spec.get("event_names", []), list) else []
        )
        run_settings["definition_runtime_hints"] = (
            runtime_payload.get("runtime_hints", {})
            if isinstance(runtime_payload.get("runtime_hints", {}), dict)
            else {}
        )
        run_settings["definition_runtime_rows"] = (
            runtime_payload.get("runtime_rows", [])
            if isinstance(runtime_payload.get("runtime_rows", []), list)
            else []
        )
        run_settings["max_auto_clicks"] = 220
        run_settings["max_run_minutes"] = 40

    if requested_scope == "site_wide":
        run_settings["restrict_to_start_url"] = False
        run_settings["single_page_only"] = False
        run_settings["block_link_navigation"] = False
    if resolved_scenario_seed_urls:
        run_settings["scenario_seed_urls"] = resolved_scenario_seed_urls
    snapshot = start_debug_session(
        session_id=session_id,
        target_url=resolved_target_url,
        output_file=output_file,
        tester_name=str(tester_name).strip(),
        tester_note="product_ui",
        db_path=paths["qa_sessions"] / "qa_runs.db",
        launch_browser=True,
        run_settings=run_settings,
    )

    meta_payload = _read_json(paths["meta"], {})
    if not isinstance(meta_payload, dict):
        meta_payload = {}
    session_to_version = (
        meta_payload.get("session_to_version", {})
        if isinstance(meta_payload.get("session_to_version", {}), dict)
        else {}
    )
    session_to_version[session_id] = str(version_id).strip() or "ver_1_0"
    meta_payload["session_to_version"] = session_to_version
    save_sidebar_meta_path = paths["meta"]
    _write_json(save_sidebar_meta_path, meta_payload)

    run_meta = _read_json(paths["run_meta"], {})
    if not isinstance(run_meta, dict):
        run_meta = {}
    selected_version_name = str(version_id).strip() or "ver_1_0"
    for v in load_versions(project_slug):
        if str(v.get("id", "")).strip() == selected_version_name:
            selected_version_name = str(v.get("display_name", "")).strip() or selected_version_name
            break
    run_meta[session_id] = {
        "version_id": str(version_id).strip() or "ver_1_0",
        "version_name": selected_version_name,
        "run_type": requested_run_type if requested_run_type in {"exploratory", "scenario", "spec_validation"} else "exploratory",
        "qa_mode": qa_mode_text,
        "browser": str(settings.get("browser", "chromium")).strip() or "chromium",
        "viewport": selected_viewport,
        "mobile_mode": bool(is_mobile_viewport),
        "collection_option": str(settings.get("collection_option", "Auto Crawl")).strip() or "Auto Crawl",
        "analytics_source_mode": source_mode,
        "judgement_rule": str(settings.get("judgement_rule", "Network First")).strip() or "Network First",
        "tester_name": str(tester_name).strip(),
        "test_scope": requested_scope,
        "start_page_mode": requested_page_mode,
        "start_url": resolved_target_url,
        "saved_page_name": resolved_saved_page_name,
        "scenario_group_name": resolved_scenario_group_name,
        "definition_spec_id": selected_spec_id,
        "definition_spec_name": str(selected_spec.get("name", "")).strip(),
        "definition_event_target_count": int(
            len(run_settings.get("definition_event_names", []))
            if isinstance(run_settings.get("definition_event_names", []), list)
            else 0
        ),
        "definition_runtime_row_count": int(
            len(run_settings.get("definition_runtime_rows", []))
            if isinstance(run_settings.get("definition_runtime_rows", []), list)
            else 0
        ),
        "started_at": datetime.now().isoformat(),
    }
    _write_json(paths["run_meta"], run_meta)
    return snapshot


def stop_session(project_slug: str, session_id: str) -> Dict[str, Any]:
    from src.debug_runtime import stop_debug_session

    if not str(session_id or "").strip():
        return {}
    sid = str(session_id).strip()
    paths = get_project_paths(project_slug)
    db_path = paths["qa_sessions"] / "qa_runs.db"

    def _runtime_status() -> str:
        rows = _read_session_rows(db_path)
        row = next((r for r in rows if str(r.get("session_id", "")).strip() == sid), {})
        return str(row.get("status", "")).strip().lower()

    def _has_block_signal() -> bool:
        if not db_path.exists():
            return False
        try:
            with sqlite3.connect(str(db_path)) as conn:
                conn.row_factory = sqlite3.Row
                row = conn.execute(
                    """
                    SELECT COUNT(*) AS cnt
                    FROM qa_events
                    WHERE session_id = ?
                      AND (
                        event_name IN ('auto_crawl_blocked', 'auto_crawl_nav_blocked')
                        OR params_json LIKE '%access_challenge_detected%'
                        OR params_json LIKE '%captcha%'
                      )
                    """,
                    (sid,),
                ).fetchone()
                return int((row[0] if row else 0) or 0) > 0
        except Exception:
            return False

    def _has_export_csv_for_session() -> bool:
        exports_root = paths["exports"]
        if not exports_root.exists():
            return False
        for project_dir in exports_root.iterdir():
            if not project_dir.is_dir():
                continue
            for run_dir in project_dir.iterdir():
                if not run_dir.is_dir():
                    continue
                meta_path = run_dir / "meta.json"
                if not meta_path.exists():
                    continue
                meta = _read_json(meta_path, {})
                if str(meta.get("session_id", "")).strip() == sid and (run_dir / "qa_result.csv").exists():
                    return True
        return False

    def _has_review_export_for_session() -> bool:
        exports_root = paths["exports"]
        if not exports_root.exists():
            return False
        for project_dir in exports_root.iterdir():
            if not project_dir.is_dir():
                continue
            for run_dir in project_dir.iterdir():
                if not run_dir.is_dir():
                    continue
                meta_path = run_dir / "meta.json"
                if not meta_path.exists():
                    continue
                meta = _read_json(meta_path, {})
                if str(meta.get("session_id", "")).strip() == sid and (run_dir / "qa_result.xlsx").exists():
                    return True
        return False

    snapshot = stop_debug_session(sid)
    if not isinstance(snapshot, dict):
        snapshot = {}
    snapshot.setdefault("session_id", sid)
    run_meta = _read_json(paths["run_meta"], {})
    meta_item = run_meta.get(sid, {}) if isinstance(run_meta, dict) and isinstance(run_meta.get(sid, {}), dict) else {}
    qa_mode = str(meta_item.get("qa_mode", "")).strip()
    blocked = _has_block_signal()
    wait_deadline = time.time() + (8.0 if blocked else 3.0)
    while time.time() < wait_deadline:
        st = _runtime_status()
        if st and st not in {"running", "stopping"}:
            break
        time.sleep(0.5)

    try:
        if qa_mode == "정의서 검증":
            _finalize_definition_validation_exports(
                project_slug=project_slug,
                session_id=sid,
                include_review_exports=False,
            )
        else:
            _finalize_basic_session_exports(
                project_slug=project_slug,
                session_id=sid,
                include_review_exports=False,
            )
    except Exception:
        pass

    export_ready = _has_export_csv_for_session()
    review_ready = _has_review_export_for_session()
    finalize_queued = False
    if (not export_ready) or (not review_ready):
        finalize_queued = True

        def _finalize_worker() -> None:
            tries = 4 if blocked else 2
            for idx in range(tries):
                try:
                    if qa_mode == "정의서 검증":
                        _finalize_definition_validation_exports(
                            project_slug=project_slug,
                            session_id=sid,
                            include_review_exports=True,
                        )
                    else:
                        _finalize_basic_session_exports(
                            project_slug=project_slug,
                            session_id=sid,
                            include_review_exports=True,
                        )
                except Exception:
                    pass
                if _has_export_csv_for_session() and _has_review_export_for_session():
                    break
                if idx < tries - 1:
                    time.sleep(1.2 if blocked else 0.7)

        t = threading.Thread(target=_finalize_worker, daemon=True, name=f"finalize-{sid[:12]}")
        t.start()

    # 최신 상태를 다시 반영해서 반환
    try:
        st = _runtime_status()
        if st:
            snapshot["status"] = st
    except Exception:
        pass
    try:
        snapshot["finalize_block_mode"] = bool(blocked)
        snapshot["finalize_export_ready"] = bool(_has_export_csv_for_session())
        snapshot["finalize_queued"] = bool(finalize_queued)
    except Exception:
        pass
    return snapshot


def _is_system_event_name(name: str) -> bool:
    text = str(name or "").strip().lower()
    if not text:
        return True
    if "max_run_minutes_reached" in text:
        return True
    if text.startswith("manual_flow"):
        return True
    if text.startswith("qa_"):
        return True
    if text.startswith("auto_crawl_"):
        return True
    return False


def _read_definition_spec_df(project_slug: str, spec_id: str) -> pd.DataFrame:
    sid = str(spec_id or "").strip()
    if not sid:
        return pd.DataFrame()
    paths = get_project_paths(project_slug)
    payload = _read_json(paths["definition_specs_meta"], {})
    rows = payload.get("items", []) if isinstance(payload, dict) else []
    rel = ""
    for row in rows:
        if not isinstance(row, dict):
            continue
        if bool(row.get("deleted", False)):
            continue
        if str(row.get("id", "")).strip() == sid:
            rel = str(row.get("path", "")).strip()
            break
    if not rel:
        return pd.DataFrame()
    file_path = BASE_DIR / rel
    if not file_path.exists():
        return pd.DataFrame()
    try:
        if file_path.suffix.lower() in {".xlsx", ".xls"}:
            df = pd.read_excel(file_path, dtype=str)
        else:
            df = _read_csv_for_export(file_path)
        if df is None or df.empty:
            return pd.DataFrame()
        df = df.fillna("")
        # 헤더가 2줄(예: Unnamed + event_name 행)인 템플릿 자동 보정
        cols = [str(c).strip() for c in df.columns]
        unnamed_ratio = (
            sum(1 for c in cols if c.lower().startswith("unnamed")) / max(1, len(cols))
        )
        if unnamed_ratio >= 0.4 and len(df) >= 2:
            promote_idx = None
            for i in range(min(5, len(df))):
                row_vals = [str(v).strip().lower() for v in list(df.iloc[i].values)]
                joined = " ".join(row_vals)
                if "event_name" in joined or "이벤트명" in joined:
                    promote_idx = i
                    break
            if promote_idx is not None:
                new_cols: List[str] = []
                for j, val in enumerate(list(df.iloc[promote_idx].values)):
                    text = str(val).strip()
                    if not text:
                        text = f"col_{j+1}"
                    new_cols.append(text)
                df2 = df.iloc[promote_idx + 1 :].copy().reset_index(drop=True)
                df2.columns = new_cols
                # 완전 빈 행 제거
                df2 = df2[
                    df2.apply(lambda r: any(str(v).strip() for v in r.values), axis=1)
                ]
                if not df2.empty:
                    df = df2
        return df.fillna("")
    except Exception:
        return pd.DataFrame()


def _pick_col(df: pd.DataFrame, keys: List[str]) -> str:
    if df.empty:
        return ""
    normalized = {str(c): _normalize_col_name(str(c)) for c in df.columns}
    for key in keys:
        for orig, n in normalized.items():
            if n == key:
                return orig
    return ""


def _build_definition_runtime_payload(project_slug: str, spec_id: str) -> Dict[str, Any]:
    spec_df = _read_definition_spec_df(project_slug, spec_id)
    if spec_df.empty:
        return {"event_names": [], "runtime_hints": {}, "runtime_rows": []}

    def _pick_col_by_cell_values(keys: List[str]) -> str:
        keyset = {str(k).strip().lower() for k in keys if str(k).strip()}
        if not keyset:
            return ""
        for col in spec_df.columns:
            try:
                series = spec_df[col].fillna("").astype(str)
            except Exception:
                continue
            for value in series.head(300):
                if _normalize_col_name(str(value).strip()) in keyset:
                    return str(col)
        return ""

    event_col = _pick_col(spec_df, ["eventname", "event", "이벤트명", "이벤트", "event_name"])
    section_col = _pick_col(spec_df, ["sectionname", "section", "섹션명", "섹션"])
    page_col = _pick_col(spec_df, ["pageid", "page_id", "페이지id", "페이지"])
    desc_col = _pick_col(spec_df, ["description", "설명", "이벤트설명"])
    action_col = _pick_col(spec_df, ["action", "행동"])
    object_col = _pick_col(spec_df, ["object", "대상"])
    section_index_col = _pick_col(spec_df, ["sectionindex", "section_index", "index"])
    target_id_col = _pick_col(spec_df, ["targetid", "target_id"])
    button_id_col = _pick_col(spec_df, ["buttonid", "button_id"])
    content_id_col = _pick_col(spec_df, ["contentid", "content_id"])
    banner_id_col = _pick_col(spec_df, ["bannerid", "banner_id"])
    brand_id_col = _pick_col(spec_df, ["brandid", "brand_id"])
    category_id_col = _pick_col(spec_df, ["categoryid", "category_id"])
    filter_value_col = _pick_col(spec_df, ["filtervalue", "filter_value"])

    if not event_col:
        event_col = _pick_col_by_cell_values(["event_name", "eventname", "이벤트명", "event"])
    if not section_col:
        section_col = _pick_col_by_cell_values(["section_name", "sectionname", "로그상태", "로그 상태"])
    if not page_col:
        page_col = _pick_col_by_cell_values(["page_id", "pageid", "페이지id", "페이지"])

    def _clean_expected(raw: str) -> str:
        text = str(raw or "").strip()
        if not text:
            return ""
        lower = text.lower()
        if lower in {"nan", "none", "(not set)", "-"}:
            return ""
        return text

    def _is_noise_token(raw: str) -> bool:
        lower = str(raw or "").strip().lower()
        return lower in {
            "",
            "state",
            "상용",
            "추가",
            "poc",
            "parameter - 매개변수 매트릭스",
            "parameter",
            "param",
            "section_name",
            "sectionname",
            "page_id",
            "pageid",
            "event_name",
            "eventname",
            "이벤트명",
        }

    def _is_template_value(raw: str) -> bool:
        text = str(raw or "").strip()
        if not text:
            return False
        lower = text.lower()
        return ("{{" in text and "}}" in text) or lower.startswith("eg.") or "{" in text and "}" in text

    def _text_hints(description: str) -> List[str]:
        out: List[str] = []
        for token in re.split(r"[^0-9a-zA-Z가-힣_]+", str(description or "").strip().lower()):
            t = str(token).strip()
            if len(t) < 2:
                continue
            if t in {"click", "impression", "event", "이벤트", "클릭", "노출", "적용", "버튼"}:
                continue
            if t not in out:
                out.append(t)
            if len(out) >= 8:
                break
        return out

    event_names: List[str] = []
    runtime_rows: List[Dict[str, Any]] = []
    hints_target_ids: List[str] = []
    hints_section_names: List[str] = []
    hints_page_ids: List[str] = []
    hints_text_hints: List[str] = []

    if not event_col:
        return {"event_names": [], "runtime_hints": {}, "runtime_rows": []}

    for no, (_, row) in enumerate(spec_df.iterrows(), start=1):
        event_name = _clean_expected(str(row.get(event_col, "")))
        if not event_name or _is_noise_token(event_name):
            continue
        if event_name not in event_names:
            event_names.append(event_name)

        section_name = _clean_expected(str(row.get(section_col, ""))) if section_col else ""
        page_id = _clean_expected(str(row.get(page_col, ""))) if page_col else ""
        if _is_noise_token(section_name) or section_name == event_name:
            section_name = ""
        if _is_noise_token(page_id) or ("/" not in page_id and not page_id.startswith("http")):
            page_id = ""
        description = _clean_expected(str(row.get(desc_col, ""))) if desc_col else ""
        action = _clean_expected(str(row.get(action_col, ""))) if action_col else ""
        obj = _clean_expected(str(row.get(object_col, ""))) if object_col else ""
        section_index = _clean_expected(str(row.get(section_index_col, ""))) if section_index_col else ""

        expected_params: Dict[str, str] = {}
        if section_name:
            expected_params["section_name"] = section_name
        if page_id:
            expected_params["page_id"] = page_id
        if section_index and not _is_template_value(section_index):
            expected_params["section_index"] = section_index

        target_ids: List[str] = []
        for col_name, key_name in [
            (target_id_col, "target_id"),
            (button_id_col, "button_id"),
            (content_id_col, "content_id"),
            (banner_id_col, "banner_id"),
            (brand_id_col, "brand_id"),
            (category_id_col, "category_id"),
            (filter_value_col, "filter_value"),
        ]:
            if not col_name:
                continue
            value = _clean_expected(str(row.get(col_name, "")))
            if not value or _is_template_value(value):
                continue
            expected_params[key_name] = value
            if value not in target_ids:
                target_ids.append(value)

        text_hints = _text_hints(description)

        runtime_rows.append(
            {
                "definition_row_id": f"row_{no}",
                "no": str(no),
                "event_name": event_name,
                "page_id": page_id,
                "section_name": section_name,
                "canonical_key": f"{event_name}|{section_name}",
                "section_index": section_index,
                "action": action,
                "object": obj,
                "description": description,
                "target_ids": target_ids,
                "text_hints": text_hints,
                "expected_params": expected_params,
            }
        )

        if section_name and section_name not in hints_section_names:
            hints_section_names.append(section_name)
        if page_id and page_id not in hints_page_ids:
            hints_page_ids.append(page_id)
        for item in target_ids:
            if item not in hints_target_ids:
                hints_target_ids.append(item)
        for item in text_hints:
            if item not in hints_text_hints:
                hints_text_hints.append(item)

    return {
        "event_names": event_names,
        "runtime_hints": {
            "target_ids": hints_target_ids[:200],
            "section_names": hints_section_names[:200],
            "page_ids": hints_page_ids[:200],
            "text_hints": hints_text_hints[:300],
        },
        "runtime_rows": runtime_rows[:2000],
    }


def _finalize_definition_validation_exports(
    project_slug: str,
    session_id: str,
    include_review_exports: bool = True,
) -> None:
    sid = str(session_id or "").strip()
    if not sid:
        return
    paths = get_project_paths(project_slug)
    db_path = paths["qa_sessions"] / "qa_runs.db"
    if not db_path.exists():
        return

    run_meta = _read_json(paths["run_meta"], {})
    meta_item = run_meta.get(sid, {}) if isinstance(run_meta, dict) and isinstance(run_meta.get(sid, {}), dict) else {}
    qa_mode = str(meta_item.get("qa_mode", "")).strip()
    if qa_mode != "정의서 검증":
        return

    spec_id = str(meta_item.get("definition_spec_id", "")).strip()
    spec_df = _read_definition_spec_df(project_slug, spec_id)
    if spec_df.empty:
        return

    session_rows = _read_session_rows(db_path)
    session_row = next((r for r in session_rows if str(r.get("session_id", "")).strip() == sid), {})
    target_url = str(session_row.get("target_url", "")).strip()
    parsed = urlparse(target_url)
    compare_key = f"{parsed.scheme}://{parsed.netloc}{parsed.path}" if parsed.scheme and parsed.netloc else target_url
    url_hash = hashlib.md5(compare_key.encode("utf-8")).hexdigest()[:16] if compare_key else "no_url"
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    run_dir = paths["exports"] / url_hash / run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    events = _fetch_event_sample(db_path, sid, limit=5000)
    runtime_matched_row_ids_strict: set[str] = set()
    runtime_matched_row_ids_relaxed: set[str] = set()
    try:
        with sqlite3.connect(str(db_path)) as conn:
            conn.row_factory = sqlite3.Row
            for r in conn.execute(
                """
                SELECT params_json
                FROM qa_events
                WHERE session_id = ?
                  AND source = 'definition_match_debug'
                """,
                (sid,),
            ):
                params = _json_object((dict(r) if isinstance(r, sqlite3.Row) else {"params_json": r[0]}).get("params_json", "{}"))
                stage = str(params.get("stage", "")).strip()
                rid = str(params.get("definition_row_id", "")).strip()
                if not rid:
                    continue
                if stage == "row_matched":
                    runtime_matched_row_ids_strict.add(rid)
                elif stage in {"row_matched_relaxed", "row_matched_third"}:
                    runtime_matched_row_ids_relaxed.add(rid)
    except Exception:
        runtime_matched_row_ids_strict = set()
        runtime_matched_row_ids_relaxed = set()
    captured_raw = []
    for e in events:
        event_name = str(e.get("event_name", "")).strip()
        if _is_system_event_name(event_name):
            continue
        params = e.get("params", {}) if isinstance(e.get("params", {}), dict) else {}
        captured_raw.append(
            {
                "event_name": event_name,
                "params": params,
                "captured_at": str(e.get("captured_at", "")).strip(),
                "source": str(e.get("source", "")).strip(),
            }
        )

    def _pick_col_by_cell_values(keys: List[str]) -> str:
        keyset = {str(k).strip().lower() for k in keys if str(k).strip()}
        if not keyset:
            return ""
        for col in spec_df.columns:
            try:
                series = spec_df[col].fillna("").astype(str)
            except Exception:
                continue
            for value in series.head(300):
                token = _normalize_col_name(str(value).strip())
                if token in keyset:
                    return str(col)
        return ""

    event_col = _pick_col(spec_df, ["eventname", "event", "이벤트명", "이벤트", "event_name"])
    section_col = _pick_col(spec_df, ["sectionname", "section", "섹션명", "섹션"])
    if not section_col:
        section_col = _pick_col_by_cell_values(["sectionname", "section_name", "section", "로그상태", "로그 상태"])
    page_col = _pick_col(spec_df, ["pageid", "page_id", "페이지id", "페이지"])
    desc_col = _pick_col(spec_df, ["description", "설명", "이벤트설명"])
    action_col = _pick_col(spec_df, ["action", "행동"])
    object_col = _pick_col(spec_df, ["object", "대상"])
    detail_expected_cols: Dict[str, str] = {
        "button_id": _pick_col(spec_df, ["buttonid", "button_id"]),
        "button_name": _pick_col(spec_df, ["buttonname", "button_name"]),
        "content_id": _pick_col(spec_df, ["contentid", "content_id"]),
        "brand_id": _pick_col(spec_df, ["brandid", "brand_id"]),
        "banner_id": _pick_col(spec_df, ["bannerid", "banner_id"]),
        "filter_type": _pick_col(spec_df, ["filtertype", "filter_type"]),
        "filter_value": _pick_col(spec_df, ["filtervalue", "filter_value"]),
        "landing_url": _pick_col(spec_df, ["landingurl", "landing_url"]),
    }

    section_source_col = _pick_col(spec_df, ["sectionname", "로그상태", "로그 상태"])
    if not section_source_col:
        section_source_col = _pick_col_by_cell_values(["sectionname", "section_name", "로그상태", "로그상태", "로그 상태"])
    if not section_source_col:
        section_source_col = _pick_col_by_cell_values(["sectionname", "section_name", "로그상태", "로그상태", "로그 상태"])
    known_event_names: set[str] = set()
    if event_col:
        try:
            for raw in spec_df[event_col].fillna("").astype(str).tolist():
                ev = str(raw).strip()
                if ev and ev.lower() not in {"event_name", "이벤트명"}:
                    known_event_names.add(ev)
        except Exception:
            known_event_names = set()

    valid_sections: List[str] = []
    valid_sections_source_col = section_col or section_source_col
    if valid_sections_source_col:
        skip_tokens = {
            "state", "상용", "추가", "poc", "section_name", "sectionname", "로그 상태", "로그상태",
            "parameter - 매개변수 매트릭스", "param", "parameter", "page_id",
        }
        for raw in spec_df[valid_sections_source_col].fillna("").astype(str).tolist():
            sv = str(raw).strip()
            if not sv:
                continue
            lower = sv.lower()
            if lower in skip_tokens:
                continue
            if "://" in sv or sv.startswith("/"):
                continue
            if "{{" in sv or "}}" in sv:
                continue
            if sv in known_event_names:
                continue
            if sv.startswith("{") or sv.startswith("["):
                continue
            if sv not in valid_sections:
                valid_sections.append(sv)
    section_expected_text = (
        "section_name 값 존재"
        if not valid_sections
        else "section_name ∈ {" + ", ".join(valid_sections[:10]) + (" ..." if len(valid_sections) > 10 else "") + "}"
    )

    def _extract_section_name(params: Dict[str, Any]) -> str:
        if not isinstance(params, dict):
            return ""
        direct = str(params.get("section_name", "")).strip()
        if direct:
            return direct
        camel = str(params.get("sectionName", "")).strip()
        if camel:
            return camel
        nested = params.get("params", {})
        if isinstance(nested, dict):
            nested_sec = str(nested.get("section_name", "")).strip() or str(nested.get("sectionName", "")).strip()
            if nested_sec:
                return nested_sec
        canonical_key = str(params.get("canonical_key", "")).strip()
        if "|" in canonical_key:
            return str(canonical_key.split("|", 1)[1]).strip()
        return ""

    def _params_preview(params: Dict[str, Any]) -> str:
        if not isinstance(params, dict):
            return ""
        preferred = [
            "section_name",
            "button_id",
            "button_name",
            "content_id",
            "content_name",
            "category_id",
            "category_name",
            "brand_id",
            "filter_type",
            "filter_value",
            "page_id",
            "page_link",
            "landing_url",
            "index",
            "section_index",
        ]
        out: List[str] = []
        for k in preferred:
            if k not in params:
                continue
            v = str(params.get(k, "")).strip()
            if not v:
                continue
            out.append(f"{k}={v}")
        if out:
            return ", ".join(out[:12])
        fallback = []
        for k in sorted(params.keys()):
            v = str(params.get(k, "")).strip()
            if not v:
                continue
            fallback.append(f"{k}={v}")
        return ", ".join(fallback[:12])

    def _params_json(params: Dict[str, Any]) -> str:
        if not isinstance(params, dict):
            return "{}"
        try:
            return json.dumps(params, ensure_ascii=False, sort_keys=True)
        except Exception:
            return "{}"

    def _infer_action_from_hit(event_name: str, params: Dict[str, Any]) -> str:
        if not isinstance(params, dict):
            params = {}
        ev_type = str(params.get("event_type", "")).strip().lower()
        if ev_type:
            return ev_type
        ev = str(event_name or "").strip().lower()
        if ev.startswith("click_"):
            return "click"
        if ev.startswith("impression_"):
            return "impression"
        if ev.startswith("view_"):
            return "view"
        if ev.startswith("add_"):
            return "like"
        if ev.startswith("remove_"):
            return "unlike"
        return ""

    def _infer_object_from_hit(event_name: str, params: Dict[str, Any]) -> str:
        if not isinstance(params, dict):
            params = {}
        if str(params.get("button_id", "")).strip() or str(params.get("button_name", "")).strip():
            return "button"
        if str(params.get("content_id", "")).strip() or str(params.get("content_name", "")).strip():
            return "content"
        brand_id = str(params.get("brand_id", "")).strip()
        if brand_id and brand_id != "(not set)":
            return "brand"
        ev = str(event_name or "").strip().lower()
        if "item" in ev or "wishlist" in ev or "product" in ev:
            return "product"
        if "brand" in ev:
            return "brand"
        return ""

    def _hit_identity(hit: Dict[str, Any]) -> str:
        params = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
        event_name = str(hit.get("event_name", "")).strip()
        section = _extract_section_name(params)
        canonical_no = str(params.get("canonical_key_no", "")).strip()
        button_id = str(params.get("button_id", "")).strip()
        button_name = str(params.get("button_name", "")).strip()
        content_id = str(params.get("content_id", "")).strip()
        content_name = str(params.get("content_name", "")).strip()
        category_id = str(params.get("category_id", "")).strip()
        brand_id = str(params.get("brand_id", "")).strip()
        index_val = str(params.get("index", "")).strip()
        page_id = str(params.get("page_id", "")).strip()
        filter_type = str(params.get("filter_type", "")).strip()
        filter_value = str(params.get("filter_value", "")).strip()
        seed = "|".join(
            [
                event_name,
                section,
                canonical_no,
                button_id,
                button_name,
                content_id,
                content_name,
                category_id,
                brand_id,
                index_val,
                page_id,
                filter_type,
                filter_value,
            ]
        )
        return hashlib.md5(seed.encode("utf-8")).hexdigest()

    def _tokenize_description(text: str) -> List[str]:
        raw = str(text or "").strip().lower()
        if not raw:
            return []
        tokens = [t for t in re.split(r"[^0-9a-zA-Z가-힣_]+", raw) if t]
        stop = {"클릭", "노출", "적용", "열기", "닫기", "시", "및", "또는", "이벤트", "랭킹판", "상품", "브랜드"}
        out: List[str] = []
        for t in tokens:
            if len(t) < 2:
                continue
            if t in stop:
                continue
            if t not in out:
                out.append(t)
        return out

    def _hit_text(params: Dict[str, Any]) -> str:
        if not isinstance(params, dict):
            return ""
        keys = [
            "button_id",
            "button_name",
            "content_id",
            "content_name",
            "category_id",
            "category_name",
            "brand_id",
            "filter_type",
            "filter_value",
            "extra_info",
            "section_name",
            "section_title",
            "landing_url",
            "page_id",
            "page_link",
        ]
        vals = [str(params.get(k, "")).strip().lower() for k in keys if str(params.get(k, "")).strip()]
        return " ".join(vals)

    def _select_hit_by_description(candidates: List[Dict[str, Any]], description_text: str) -> Tuple[Optional[Dict[str, Any]], str]:
        if not candidates:
            return None, "후보 이벤트 없음"
        if len(candidates) == 1:
            return candidates[0], ""
        tokens = _tokenize_description(description_text)
        if not tokens:
            return None, "description 키워드가 없어 후보 선택 보류"

        scored: List[Tuple[int, int, Dict[str, Any]]] = []
        desc_norm = str(description_text or "").strip().lower()
        for idx, hit in enumerate(candidates):
            params = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
            text = _hit_text(params)
            score = 0
            if desc_norm and len(desc_norm) >= 3 and desc_norm in text:
                score += 5
            for tk in tokens:
                if tk in text:
                    score += 2
            if "툴팁" in desc_norm and str(params.get("button_id", "")).strip().lower() == "tooltip":
                score += 3
            if "필터" in desc_norm and (str(params.get("filter_type", "")).strip() or str(params.get("filter_value", "")).strip()):
                score += 2
            if "카테고리" in desc_norm and (str(params.get("category_id", "")).strip() or str(params.get("category_name", "")).strip()):
                score += 2
            if "브랜드" in desc_norm and str(params.get("brand_id", "")).strip():
                score += 2
            scored.append((score, idx, hit))

        scored.sort(key=lambda x: (x[0], -x[1]), reverse=True)
        top_score, _, top_hit = scored[0]
        second_score = scored[1][0] if len(scored) > 1 else -1
        if top_score < 2:
            return None, "description 점수가 낮아 후보 선택 보류"
        if top_score == second_score:
            return None, "description 동점 후보로 선택 보류"
        if second_score >= 0 and (top_score - second_score) < 2:
            return None, "description 점수 차이가 작아 선택 보류"
        return top_hit, ""

    captured: List[Dict[str, Any]] = []
    seen_capture_ids: set[str] = set()
    for hit in captured_raw:
        hid = _hit_identity(hit)
        if hid in seen_capture_ids:
            continue
        seen_capture_ids.add(hid)
        captured.append(hit)

    # 정의서 템플릿의 parameter matrix(section_name) 행을 이벤트 행 순서와 매핑
    expected_sections_by_row: Dict[int, str] = {}
    if section_source_col and event_col:
        header_idx = -1
        for ridx in range(len(spec_df)):
            token = str(spec_df.iloc[ridx].get(section_source_col, "")).strip().lower()
            if token in {"section_name", "sectionname"}:
                header_idx = ridx
                break
        event_row_indices: List[int] = []
        for ridx, row in spec_df.iterrows():
            ev = str(row.get(event_col, "")).strip()
            if not ev:
                continue
            if ev.lower() in {"event_name", "이벤트명"}:
                continue
            event_row_indices.append(int(ridx))
        matrix_sections: List[str] = []
        if header_idx >= 0:
            skip_tokens = {"state", "상용", "추가", "poc", "section_name", "sectionname"}
            for ridx in range(header_idx + 1, len(spec_df)):
                sec = str(spec_df.iloc[ridx].get(section_source_col, "")).strip()
                if not sec:
                    continue
                sec_lower = sec.lower()
                if sec_lower in skip_tokens:
                    continue
                if "parameter" in sec_lower:
                    continue
                if sec.startswith("{") or sec.startswith("["):
                    continue
                matrix_sections.append(sec)
        for i, ev_ridx in enumerate(event_row_indices):
            if i < len(matrix_sections):
                expected_sections_by_row[ev_ridx] = str(matrix_sections[i]).strip()

    expected_params_by_section: Dict[str, Dict[str, str]] = {}
    section_source_col_for_map = _pick_col(spec_df, ["sectionname", "로그상태", "로그 상태"])
    if not section_source_col_for_map:
        section_source_col_for_map = _pick_col_by_cell_values(["sectionname", "section_name", "로그상태", "로그상태", "로그 상태"])
    if section_source_col_for_map:
        header_idx_for_map = -1
        for ridx in range(len(spec_df)):
            sec_val = str(spec_df.iloc[ridx].get(section_source_col_for_map, "")).strip().lower()
            if sec_val in {"section_name", "sectionname"}:
                header_idx_for_map = ridx
                break
        if header_idx_for_map >= 0:
            header_row = spec_df.iloc[header_idx_for_map]
            ignore_norm = {
                "sectionname", "eventname", "description", "action", "object", "eventtype",
                "로그상태", "이벤트명", "설명", "action", "object",
            }
            param_columns_for_map: List[tuple[str, str]] = []
            for col in spec_df.columns:
                pv = str(header_row.get(col, "")).strip()
                if not pv:
                    continue
                pv_norm = _normalize_col_name(pv)
                if pv_norm in ignore_norm:
                    continue
                param_columns_for_map.append((col, pv))
            for ridx in range(header_idx_for_map + 1, len(spec_df)):
                row = spec_df.iloc[ridx]
                section_name = str(row.get(section_source_col_for_map, "")).strip()
                if not section_name:
                    continue
                sec_lower = section_name.lower()
                if sec_lower in {"state", "상용", "추가", "poc", "section_name"}:
                    continue
                if "parameter" in sec_lower:
                    continue
                if section_name.startswith("{{") or section_name.startswith("{"):
                    continue
                params_map: Dict[str, str] = {}
                for col, param_name in param_columns_for_map:
                    expected = str(row.get(col, "")).strip()
                    params_map[str(param_name).strip()] = expected
                if params_map:
                    expected_params_by_section[section_name] = params_map

    event_level_rows: List[Dict[str, Any]] = []
    summary_rows: List[Dict[str, Any]] = []
    matches_rows: List[Dict[str, Any]] = []
    unmatched_rows: List[Dict[str, Any]] = []
    used_hit_ids: set[str] = set()

    for idx, (row_idx, row) in enumerate(spec_df.iterrows(), start=1):
        expected_event = str(row.get(event_col, "") if event_col else "").strip()
        if not expected_event:
            continue
        if expected_event.lower() in {"event_name", "이벤트명"}:
            continue
        section_name = str(row.get(section_col, "") if section_col else "").strip()
        if (
            section_name.lower() in {"section_name", "sectionname", "page_id", "pageid", "event_name", "이벤트명", "state", "상용", "추가", "poc"}
            or section_name == expected_event
            or section_name.startswith("/")
            or "://" in section_name
            or "{{" in section_name
            or "}}" in section_name
        ):
            section_name = ""
        if not section_name:
            section_name = str(expected_sections_by_row.get(int(row_idx), "")).strip()
        page_id = str(row.get(page_col, "") if page_col else "").strip()
        description = str(row.get(desc_col, "") if desc_col else "").strip()
        expected_action = str(row.get(action_col, "") if action_col else "").strip().lower()
        expected_object = str(row.get(object_col, "") if object_col else "").strip().lower()
        expected_detail: Dict[str, str] = {}
        for key, col_name in detail_expected_cols.items():
            if not col_name:
                continue
            ev = str(row.get(col_name, "")).strip()
            if ev:
                expected_detail[key] = ev
        definition_row_id = f"row_{idx}"

        matched_hits = [c for c in captured if str(c.get("event_name", "")).strip() == expected_event]
        if page_id:
            page_filtered = []
            for hit in matched_hits:
                hp = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
                hp_page = str(hp.get("page_id", "")).strip()
                if hp_page == page_id:
                    page_filtered.append(hit)
            if page_filtered:
                matched_hits = page_filtered

        section_hits = []
        for hit in matched_hits:
            params = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
            sec = _extract_section_name(params)
            if sec:
                section_hits.append(hit)

        valid_section_hits: List[Dict[str, Any]] = []
        if section_name:
            for hit in matched_hits:
                params = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
                sec = _extract_section_name(params)
                if sec == section_name:
                    valid_section_hits.append(hit)
        elif valid_sections:
            allowed = set(valid_sections)
            for hit in section_hits:
                params = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
                sec = _extract_section_name(params)
                if sec in allowed:
                    valid_section_hits.append(hit)
        else:
            valid_section_hits = section_hits

        result = "Matched"
        missing_parameters = f"section_name = {section_name}" if section_name else section_expected_text
        extra_parameters = ""
        failure_reason = ""
        matched_capture_count = len(valid_section_hits)

        representative_params: Dict[str, Any] = {}
        chosen_hit: Optional[Dict[str, Any]] = None
        choose_reason = ""
        fresh_hits = [h for h in valid_section_hits if _hit_identity(h) not in used_hit_ids]
        if fresh_hits:
            stage_candidates = list(fresh_hits)

            # 2차: section/page/action/object 큰 범주 필터
            if page_id:
                page_filtered = []
                for hit in stage_candidates:
                    hp = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
                    if str(hp.get("page_id", "")).strip() == page_id:
                        page_filtered.append(hit)
                if page_filtered:
                    stage_candidates = page_filtered

            if expected_action:
                action_filtered = []
                for hit in stage_candidates:
                    hp = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
                    if _infer_action_from_hit(str(hit.get("event_name", "")).strip(), hp) == expected_action:
                        action_filtered.append(hit)
                if action_filtered:
                    stage_candidates = action_filtered

            if expected_object:
                object_filtered = []
                for hit in stage_candidates:
                    hp = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
                    if _infer_object_from_hit(str(hit.get("event_name", "")).strip(), hp) == expected_object:
                        object_filtered.append(hit)
                if object_filtered:
                    stage_candidates = object_filtered

            # 3차: 세부 속성 기반 필터
            for k, expected_v in expected_detail.items():
                detail_filtered = []
                for hit in stage_candidates:
                    hp = hit.get("params", {}) if isinstance(hit.get("params", {}), dict) else {}
                    if str(hp.get(k, "")).strip() == expected_v:
                        detail_filtered.append(hit)
                if detail_filtered:
                    stage_candidates = detail_filtered

            # 4차: description + 맥락 점수 선택
            chosen_hit, choose_reason = _select_hit_by_description(stage_candidates, description)
            if chosen_hit is None and len(stage_candidates) == 1:
                chosen_hit = stage_candidates[0]
                choose_reason = ""
        if len(matched_hits) == 0:
            result = "Missing"
            failure_reason = "이벤트가 수집되지 않음"
            unmatched_rows.append(
                {
                    "event_name": expected_event,
                    "section_name": section_name,
                    "page_id": page_id,
                    "reason": "NO_HIT",
                    "matched_identification": "",
                }
            )
        elif len(section_hits) == 0:
            result = "Missing"
            failure_reason = "section_name 파라미터 누락"
            representative_params = matched_hits[0].get("params", {}) if isinstance(matched_hits[0].get("params", {}), dict) else {}
        elif len(valid_section_hits) == 0:
            result = "Mismatch"
            first_params = matched_hits[0].get("params", {}) if isinstance(matched_hits[0].get("params", {}), dict) else {}
            extra_parameters = _extract_section_name(first_params)
            failure_reason = "section_name 값이 정의서 기대값과 불일치"
            representative_params = first_params
        elif chosen_hit is None:
            result = "Unchecked"
            first_params = valid_section_hits[0].get("params", {}) if isinstance(valid_section_hits[0].get("params", {}), dict) else {}
            extra_parameters = _extract_section_name(first_params)
            failure_reason = choose_reason or "후보 다수로 자동 판정 보류"
            representative_params = first_params
        else:
            first_params = chosen_hit.get("params", {}) if isinstance(chosen_hit.get("params", {}), dict) else {}
            extra_parameters = _extract_section_name(first_params)
            representative_params = first_params
            used_hit_ids.add(_hit_identity(chosen_hit))

        # 실시간 매칭 로그 보정:
        # - strict(row_matched): section_name까지 일치하는 경우에만 보정 허용
        # - relaxed/third: section 기대값이 없는 row에서만 보정 허용
        if result != "Matched":
            section_exact = bool(len(valid_section_hits) > 0) if section_name else True
            allow_runtime_sync = False
            if definition_row_id in runtime_matched_row_ids_strict and section_exact:
                allow_runtime_sync = True
            elif (
                (not section_name)
                and definition_row_id in runtime_matched_row_ids_relaxed
            ):
                allow_runtime_sync = True
            if allow_runtime_sync:
                result = "Matched"
                if matched_capture_count <= 0:
                    matched_capture_count = 1
                if not representative_params and matched_hits:
                    representative_params = (
                        matched_hits[0].get("params", {})
                        if isinstance(matched_hits[0].get("params", {}), dict)
                        else {}
                    )
                missing_parameters = ""
                failure_reason = "runtime strict 매칭 기준 보정"

        expected_params_for_row: Dict[str, str] = {}
        if section_name:
            expected_params_for_row["section_name"] = section_name
        if page_id:
            expected_params_for_row["page_id"] = page_id
        for k, v in expected_detail.items():
            vv = str(v).strip()
            if vv:
                expected_params_for_row[str(k).strip()] = vv
        matrix_expected = expected_params_by_section.get(section_name, {})
        if isinstance(matrix_expected, dict):
            for k, v in matrix_expected.items():
                kk = str(k).strip()
                vv = str(v).strip()
                if kk and vv:
                    expected_params_for_row[kk] = vv

        event_level_rows.append(
            {
                "no": idx,
                "definition_row_id": definition_row_id,
                "expected_identification": f"{expected_event}|{section_name}",
                "event_name": expected_event,
                "description": description,
                "section_name": section_name,
                "page_id": page_id,
                "result": result,
                "matched_capture_count": matched_capture_count,
                "missing_parameters": missing_parameters,
                "extra_parameters": extra_parameters,
                "failure_reason": failure_reason,
                "observed_param_preview": _params_preview(representative_params),
                "observed_params_json": _params_json(representative_params),
                "expected_params_json": _params_json(expected_params_for_row),
            }
        )
        if result != "Missing":
            matches_rows.append(
                {
                    "event_name": expected_event,
                    "result": result,
                    "missing_parameters": missing_parameters,
                    "section_name": section_name,
                    "matched_identification": f"{expected_event}|{section_name}",
                    "failure_reason": failure_reason,
                }
            )

    # Parameter matrix parser: supports templates where parameter expectations are defined by section_name rows.
    matrix_rows: List[Dict[str, Any]] = []
    section_source_col = _pick_col(spec_df, ["sectionname", "로그상태", "로그 상태"])
    if section_source_col:
        header_idx = -1
        for ridx in range(len(spec_df)):
            sec_val = str(spec_df.iloc[ridx].get(section_source_col, "")).strip().lower()
            if sec_val in {"section_name", "sectionname"}:
                header_idx = ridx
                break
        if header_idx >= 0:
            header_row = spec_df.iloc[header_idx]
            ignore_norm = {
                "sectionname", "eventname", "description", "action", "object", "eventtype",
                "로그상태", "이벤트명", "설명", "action", "object",
            }
            param_columns: List[tuple[str, str]] = []
            for col in spec_df.columns:
                pv = str(header_row.get(col, "")).strip()
                if not pv:
                    continue
                pv_norm = _normalize_col_name(pv)
                if pv_norm in ignore_norm:
                    continue
                param_columns.append((col, pv))

            rule_idx = 0
            for ridx in range(header_idx + 1, len(spec_df)):
                row = spec_df.iloc[ridx]
                section_name = str(row.get(section_source_col, "")).strip()
                if not section_name:
                    continue
                sec_lower = section_name.lower()
                if sec_lower in {"state", "상용", "추가", "poc", "section_name"}:
                    continue
                if "parameter" in sec_lower:
                    continue
                if section_name.startswith("{{") or section_name.startswith("{"):
                    continue

                row_has_rules = False
                for col, param_name in param_columns:
                    expected = str(row.get(col, "")).strip()
                    if expected:
                        row_has_rules = True
                        break
                if not row_has_rules:
                    continue

                section_hits = []
                for c in captured:
                    params = c.get("params", {}) if isinstance(c.get("params", {}), dict) else {}
                    if str(params.get("section_name", "")).strip() == section_name:
                        section_hits.append(c)
                inferred_event = str(section_hits[0].get("event_name", "")).strip() if section_hits else ""

                for col, param_name in param_columns:
                    expected = str(row.get(col, "")).strip()
                    if not expected:
                        continue
                    actual = ""
                    if section_hits:
                        params = section_hits[0].get("params", {}) if isinstance(section_hits[0].get("params", {}), dict) else {}
                        actual = str(params.get(param_name, "")).strip()

                    result = "Matched"
                    reason = ""
                    missing_parameters = ""
                    extra_parameters = ""
                    if not section_hits:
                        result = "Missing"
                        reason = f"section_name={section_name} 이벤트 미수집"
                        missing_parameters = param_name
                    else:
                        exp_lower = expected.lower()
                        if exp_lower in {"y", "yes", "true", "필수"}:
                            if not actual:
                                result = "Missing"
                                missing_parameters = param_name
                                reason = f"필수 파라미터 {param_name} 누락"
                        elif ("{{" in expected) or ("}}" in expected) or ("eg." in exp_lower):
                            if not actual:
                                result = "Missing"
                                missing_parameters = param_name
                                reason = f"템플릿 파라미터 {param_name} 값 누락"
                        elif actual and expected and actual != expected:
                            result = "Mismatch"
                            extra_parameters = f"{param_name}={actual}"
                            reason = f"{param_name} 기대값 불일치"
                        elif not actual and expected:
                            result = "Missing"
                            missing_parameters = param_name
                            reason = f"파라미터 {param_name} 값 누락"

                    rule_idx += 1
                    matrix_rows.append(
                        {
                            "no": rule_idx,
                            "definition_row_id": f"param_{ridx+1}_{rule_idx}",
                            "expected_identification": f"{inferred_event or 'event'}|{section_name}|{param_name}",
                            "event_name": inferred_event or "",
                            "description": f"{section_name} / {param_name}",
                            "section_name": section_name,
                            "page_id": "",
                            "result": result,
                            "matched_capture_count": len(section_hits),
                            "missing_parameters": missing_parameters or param_name,
                            "extra_parameters": extra_parameters,
                            "failure_reason": reason,
                        }
                    )
                    if result == "Missing":
                        unmatched_rows.append(
                            {
                                "event_name": inferred_event or "",
                                "section_name": section_name,
                                "page_id": "",
                                "reason": "NO_HIT" if not section_hits else "MISSING_PARAM",
                                "matched_identification": f"{section_name}|{param_name}",
                            }
                        )
                    else:
                        matches_rows.append(
                            {
                                "event_name": inferred_event or "",
                                "result": result,
                                "missing_parameters": param_name,
                                "section_name": section_name,
                                "matched_identification": f"{section_name}|{param_name}",
                                "failure_reason": reason,
                            }
                        )

    summary_rows = event_level_rows

    summary_df = pd.DataFrame(summary_rows)
    if summary_df.empty:
        return
    detail_df = summary_df.copy()
    matches_df = pd.DataFrame(matches_rows)
    unmatched_df = pd.DataFrame(unmatched_rows)

    _sanitize_excel_df(summary_df).to_csv(run_dir / "qa_result.csv", index=False, encoding="utf-8-sig")
    _sanitize_excel_df(detail_df).to_csv(run_dir / "qa_result_detail.csv", index=False, encoding="utf-8-sig")
    _sanitize_excel_df(matches_df).to_csv(run_dir / "definition_validation_matches.csv", index=False, encoding="utf-8-sig")
    _sanitize_excel_df(unmatched_df).to_csv(run_dir / "definition_validation_unmatched.csv", index=False, encoding="utf-8-sig")
    _sanitize_excel_df(summary_df).to_csv(run_dir / "definition_validation_summary.csv", index=False, encoding="utf-8-sig")

    meta = {
        "project_slug": project_slug,
        "target_url": target_url,
        "url_compare_key": compare_key,
        "url_hash": url_hash,
        "saved_at": datetime.now().isoformat(),
        "run_id": run_id,
        "session_id": sid,
        "bundle_scope": "definition_validation",
        "definition_file_name": str(meta_item.get("definition_spec_name", "")).strip(),
        "viewport": str(meta_item.get("viewport", "")).strip(),
        "browser": str(meta_item.get("browser", "")).strip() or "chromium",
        "captured_event_count": int(session_row.get("captured_events", 0) or len(captured)),
        "definition_row_count": int(len(summary_df)),
        "matched_row_count": int((summary_df["result"] == "Matched").sum()),
        "missing_row_count": int((summary_df["result"] == "Missing").sum()),
        "files": {
            "definition_validation_summary": "definition_validation_summary.csv",
            "definition_validation_matches": "definition_validation_matches.csv",
            "definition_validation_unmatched": "definition_validation_unmatched.csv",
        },
    }
    _write_json(run_dir / "meta.json", meta)
    if include_review_exports:
        session_obj = next((s for s in load_sessions(project_slug) if str(s.get("session_id", "")).strip() == sid), {})
        _ensure_review_exports(
            run_dir,
            session_obj if isinstance(session_obj, dict) else {},
            meta,
            generate_advanced=False,
        )


def _finalize_basic_session_exports(
    project_slug: str,
    session_id: str,
    include_review_exports: bool = False,
) -> None:
    sid = str(session_id or "").strip()
    if not sid:
        return
    paths = get_project_paths(project_slug)
    db_path = paths["qa_sessions"] / "qa_runs.db"
    if not db_path.exists():
        return

    session_rows = _read_session_rows(db_path)
    session_row = next((r for r in session_rows if str(r.get("session_id", "")).strip() == sid), {})
    if not session_row:
        return
    target_url = str(session_row.get("target_url", "")).strip()
    parsed = urlparse(target_url)
    compare_key = f"{parsed.scheme}://{parsed.netloc}{parsed.path}" if parsed.scheme and parsed.netloc else target_url
    url_hash = hashlib.md5(compare_key.encode("utf-8")).hexdigest()[:16] if compare_key else "no_url"

    # 이미 결과 CSV가 있으면 중복 생성은 피하되,
    # 요청이 include_review_exports=True 인 경우에는 기존 run에서 엑셀 생성만 이어서 수행한다.
    exports_root = paths["exports"] / url_hash
    if exports_root.exists():
        try:
            for run_dir in exports_root.iterdir():
                if not run_dir.is_dir():
                    continue
                meta_path = run_dir / "meta.json"
                if not meta_path.exists():
                    continue
                meta = _read_json(meta_path, {})
                if str(meta.get("session_id", "")).strip() == sid and (run_dir / "qa_result.csv").exists():
                    if include_review_exports and not (run_dir / "qa_result.xlsx").exists():
                        session_obj = next(
                            (s for s in load_sessions(project_slug) if str(s.get("session_id", "")).strip() == sid),
                            {},
                        )
                        _ensure_review_exports(
                            run_dir,
                            session_obj if isinstance(session_obj, dict) else {},
                            meta if isinstance(meta, dict) else {},
                            generate_advanced=False,
                        )
                    return
        except Exception:
            pass

    run_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    run_dir = paths["exports"] / url_hash / run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    def _basic_params_preview(params: Dict[str, Any]) -> str:
        if not isinstance(params, dict):
            return ""
        preferred = [
            "section_name",
            "button_id",
            "button_name",
            "content_id",
            "category_id",
            "brand_id",
            "filter_type",
            "filter_value",
            "page_id",
            "page_link",
            "landing_url",
            "index",
        ]
        out: List[str] = []
        for k in preferred:
            if k not in params:
                continue
            v = str(params.get(k, "")).strip()
            if v:
                out.append(f"{k}={v}")
        if out:
            return ", ".join(out[:10])
        fallback: List[str] = []
        for k in sorted(params.keys()):
            v = str(params.get(k, "")).strip()
            if v:
                fallback.append(f"{k}={v}")
        return ", ".join(fallback[:10])

    def _basic_params_json(params: Dict[str, Any]) -> str:
        try:
            return json.dumps(params if isinstance(params, dict) else {}, ensure_ascii=False, sort_keys=True)
        except Exception:
            return "{}"

    events = _fetch_event_sample(db_path, sid, limit=5000)
    rows: List[Dict[str, Any]] = []
    idx = 0
    for e in reversed(events):
        source = str(e.get("source", "")).strip()
        event_name = str(e.get("event_name", "")).strip()
        if not event_name:
            continue
        params = e.get("params", {}) if isinstance(e.get("params", {}), dict) else {}
        idx += 1
        rows.append(
            {
                "no": idx,
                "captured_at": str(e.get("captured_at", "")).strip(),
                "source": source,
                "event_name": event_name,
                "section_name": str(params.get("section_name", "")).strip() or str(params.get("sectionName", "")).strip(),
                "page_id": str(params.get("page_id", "")).strip(),
                "result": "Collected",
                "missing_parameters": "",
                "extra_parameters": "",
                "failure_reason": "",
                "observed_param_preview": _basic_params_preview(params),
                "observed_params_json": _basic_params_json(params),
                "expected_params_json": "{}",
            }
        )

    if not rows:
        rows.append(
            {
                "no": 1,
                "captured_at": "",
                "source": "",
                "event_name": "",
                "section_name": "",
                "page_id": "",
                "result": "No Data",
                "missing_parameters": "",
                "extra_parameters": "",
                "failure_reason": "수집 이벤트가 없어 기본 결과를 생성했습니다.",
                "observed_param_preview": "",
                "observed_params_json": "{}",
                "expected_params_json": "{}",
            }
        )

    summary_df = pd.DataFrame(rows)
    _sanitize_excel_df(summary_df).to_csv(run_dir / "qa_result.csv", index=False, encoding="utf-8-sig")
    _sanitize_excel_df(summary_df).to_csv(run_dir / "qa_result_detail.csv", index=False, encoding="utf-8-sig")

    run_meta = _read_json(paths["run_meta"], {})
    meta_item = run_meta.get(sid, {}) if isinstance(run_meta, dict) and isinstance(run_meta.get(sid, {}), dict) else {}
    meta = {
        "project_slug": project_slug,
        "target_url": target_url,
        "url_compare_key": compare_key,
        "url_hash": url_hash,
        "saved_at": datetime.now().isoformat(),
        "run_id": run_id,
        "session_id": sid,
        "bundle_scope": "basic_session",
        "qa_mode": str(meta_item.get("qa_mode", "")).strip(),
        "definition_file_name": "",
        "viewport": str(meta_item.get("viewport", "")).strip(),
        "browser": str(meta_item.get("browser", "")).strip() or "chromium",
        "captured_event_count": int(session_row.get("captured_events", 0) or len(rows)),
        "definition_row_count": int(len(summary_df)),
        "matched_row_count": int(len(summary_df)),
        "missing_row_count": 0,
        "files": {
            "definition_validation_summary": "qa_result.csv",
        },
    }
    _write_json(run_dir / "meta.json", meta)
    if include_review_exports:
        session_obj = next(
            (s for s in load_sessions(project_slug) if str(s.get("session_id", "")).strip() == sid),
            {},
        )
        _ensure_review_exports(
            run_dir,
            session_obj if isinstance(session_obj, dict) else {},
            meta,
            generate_advanced=False,
        )


def delete_session(project_slug: str, session_id: str) -> Dict[str, Any]:
    sid = str(session_id or "").strip()
    if not sid:
        return {"ok": False, "error": "empty_session_id"}
    paths = get_project_paths(project_slug)
    db_path = paths["qa_sessions"] / "qa_runs.db"

    # running session guard
    for s in load_sessions(project_slug):
        if str(s.get("session_id", "")).strip() == sid:
            runtime = str(s.get("runtime_status", "")).strip().lower()
            if runtime in {"running", "stopping"}:
                return {"ok": False, "error": "running_session_cannot_be_deleted"}
            break

    # delete db rows
    if db_path.exists():
        with sqlite3.connect(str(db_path), timeout=30) as conn:
            conn.execute("DELETE FROM qa_events WHERE session_id = ?", (sid,))
            conn.execute("DELETE FROM qa_ui_actions WHERE session_id = ?", (sid,))
            conn.execute("DELETE FROM qa_sessions WHERE session_id = ?", (sid,))
            conn.commit()

    # delete session folder
    session_dir = paths["qa_sessions"] / sid
    if session_dir.exists() and session_dir.is_dir():
        import shutil

        shutil.rmtree(session_dir, ignore_errors=True)

    # clean sidebar/run meta
    meta = _read_json(paths["meta"], {})
    if not isinstance(meta, dict):
        meta = {}
    run_meta = _read_json(paths["run_meta"], {})
    if not isinstance(run_meta, dict):
        run_meta = {}

    session_labels = meta.get("session_labels", {}) if isinstance(meta.get("session_labels", {}), dict) else {}
    session_to_version = meta.get("session_to_version", {}) if isinstance(meta.get("session_to_version", {}), dict) else {}
    if sid in session_labels:
        session_labels.pop(sid, None)
    if sid in session_to_version:
        session_to_version.pop(sid, None)
    if sid in run_meta:
        run_meta.pop(sid, None)
    meta["session_labels"] = session_labels
    meta["session_to_version"] = session_to_version
    _write_json(paths["meta"], meta)
    _write_json(paths["run_meta"], run_meta)

    # remove export runs bound to this session
    exports_root = paths["exports"]
    if exports_root.exists():
        import shutil

        for run_dir in _iter_export_runs(exports_root):
            info = _read_json(run_dir / "meta.json", {})
            if str(info.get("session_id", "")).strip() == sid:
                shutil.rmtree(run_dir, ignore_errors=True)

    return {"ok": True, "session_id": sid}


def _extract_path_label(raw_url: str) -> str:
    text = str(raw_url or "").strip()
    if not text:
        return ""
    try:
        u = urlparse(text)
        path = str(u.path or "/").strip() or "/"
        q_count = 0
        if u.query:
            q_count = len([x for x in str(u.query).split("&") if str(x).strip()])
        if q_count > 0:
            return f"{path} · query {q_count}개"
        return path
    except Exception:
        return text


def _export_title_from_session(session: Dict[str, Any], fallback_url: str = "") -> str:
    scope = str(session.get("test_scope", "")).strip()
    qa_mode = str(session.get("qa_mode", "")).strip()
    if scope == "site_wide":
        head = "전체 사이트 탐색"
    elif scope == "single_url":
        head = "단일 URL 검증"
    elif scope == "scenario_group":
        head = "시나리오 테스트"
    elif qa_mode == "정의서 검증":
        head = "정의서 검증"
    else:
        head = qa_mode or "세션 결과"

    context = ""
    if scope == "scenario_group":
        context = str(session.get("scenario_group_name", "")).strip()
    elif scope == "start_page":
        context = str(session.get("saved_page_name", "")).strip() or _extract_path_label(str(session.get("target_url", "")).strip() or fallback_url)
    elif scope in {"single_url", "site_wide"}:
        context = _extract_path_label(str(session.get("target_url", "")).strip() or fallback_url)
    if not context and qa_mode == "정의서 검증":
        context = str(session.get("definition_spec_name", "")).strip()
    if not context and fallback_url:
        context = _extract_path_label(fallback_url)

    version_name = str(session.get("version_name", "")).strip() or str(session.get("version_id", "")).strip() or "ver 1.0"
    if context:
        return f"{head} · {context} · {version_name}"
    return f"{head} · {version_name}"


def _read_csv_for_export(path: Path) -> pd.DataFrame:
    for enc in ["utf-8-sig", "utf-8", "cp949"]:
        try:
            return pd.read_csv(path, encoding=enc)
        except Exception:
            continue
    return pd.DataFrame()


def _sanitize_excel_df(df: pd.DataFrame) -> pd.DataFrame:
    if not isinstance(df, pd.DataFrame) or df.empty:
        return df
    out = df.copy()
    for col in out.columns:
        if pd.api.types.is_object_dtype(out[col]) or pd.api.types.is_string_dtype(out[col]):
            out[col] = out[col].map(
                lambda v: _EXCEL_ILLEGAL_RE.sub(" ", str(v)) if v is not None else ""
            )
    return out


def _event_summary_columns() -> List[str]:
    return [
        "event_no",
        "definition_row_id",
        "canonical_key",
        "event_name",
        "page_id",
        "scenario",
        "result",
        "matched_hit_count",
        "matched_target_id",
        "missing_required_params",
        "mismatched_params",
        "failure_reason",
        "observed_params",
    ]


def _param_detail_columns() -> List[str]:
    return [
        "event_no",
        "definition_row_id",
        "canonical_key",
        "event_name",
        "param_name",
        "expected_value",
        "actual_value",
        "status",
        "reason",
    ]


def _observed_columns() -> List[str]:
    return [
        "event_no",
        "canonical_key",
        "event_name",
        "page_id",
        "target_id",
        "selector",
        "annotation_id",
        "event_origin",
        "observed_params",
        "matched_definition_row_id",
        "matched_definition_no",
        "match_status",
        "candidate_reason",
    ]


def _event_block_columns() -> List[str]:
    return [
        "event_no",
        "event_name",
        "event_description",
        "section_name",
        "page_id",
        "definition_row_id",
        "param_name",
        "expected_value",
        "actual_value",
        "status",
    ]


def _to_event_summary_df(summary_df: pd.DataFrame) -> pd.DataFrame:
    cols = _event_summary_columns()
    if summary_df.empty:
        return pd.DataFrame(columns=cols)
    out = pd.DataFrame(columns=cols)
    out["event_no"] = summary_df.get("no", "")
    out["definition_row_id"] = summary_df.get("definition_row_id", "")
    out["canonical_key"] = summary_df.get("expected_identification", "")
    out["event_name"] = summary_df.get("event_name", "")
    out["page_id"] = summary_df.get("page_id", "")
    out["scenario"] = summary_df.get("section_name", "")
    out["result"] = summary_df.get("result", "")
    out["matched_hit_count"] = summary_df.get("matched_capture_count", "")
    out["matched_target_id"] = ""
    out["missing_required_params"] = summary_df.get("missing_parameters", "")
    out["mismatched_params"] = summary_df.get("extra_parameters", "")
    out["failure_reason"] = summary_df.get("failure_reason", "")
    out["observed_params"] = ""
    return out


def _to_param_detail_df(detail_df: pd.DataFrame) -> pd.DataFrame:
    cols = _param_detail_columns()
    if detail_df.empty:
        return pd.DataFrame(columns=cols)
    out = pd.DataFrame(columns=cols)
    out["event_no"] = detail_df.get("no", "")
    out["definition_row_id"] = detail_df.get("definition_row_id", "")
    out["canonical_key"] = detail_df.get("matched_identification", "")
    out["event_name"] = detail_df.get("event_name", "")
    out["param_name"] = ""
    out["expected_value"] = detail_df.get("missing_parameters", "")
    out["actual_value"] = detail_df.get("extra_parameters", "")
    out["status"] = detail_df.get("result", "")
    out["reason"] = detail_df.get("failure_reason", "")
    return out


def _to_observed_events_df(observed_df: pd.DataFrame) -> pd.DataFrame:
    cols = _observed_columns()
    if observed_df.empty:
        return pd.DataFrame(columns=cols)
    out = pd.DataFrame(columns=cols)
    out["event_no"] = observed_df.get("no", "")
    out["canonical_key"] = observed_df.get("class-attribute", "")
    out["event_name"] = observed_df.get("event name", "")
    out["page_id"] = observed_df.get("page_id", "")
    out["target_id"] = observed_df.get("target_id", "")
    out["selector"] = observed_df.get("selector", "")
    out["annotation_id"] = observed_df.get("annotation_no", "")
    out["event_origin"] = observed_df.get("status", "")
    out["observed_params"] = observed_df.get("params.value", "")
    out["matched_definition_row_id"] = ""
    out["matched_definition_no"] = ""
    out["match_status"] = ""
    out["candidate_reason"] = ""
    return out


def _to_event_block_df(unmatched_df: pd.DataFrame) -> pd.DataFrame:
    cols = _event_block_columns()
    if unmatched_df.empty:
        return pd.DataFrame(columns=cols)
    out = pd.DataFrame(columns=cols)
    out["event_no"] = ""
    out["event_name"] = unmatched_df.get("event_name", "")
    out["event_description"] = ""
    out["section_name"] = unmatched_df.get("section_name", "")
    out["page_id"] = unmatched_df.get("page_id", "")
    out["definition_row_id"] = ""
    out["param_name"] = unmatched_df.get("reason", "")
    out["expected_value"] = ""
    out["actual_value"] = unmatched_df.get("matched_identification", "")
    out["status"] = "UNMATCHED"
    return out


def _to_parameter_validation_df(df: pd.DataFrame) -> pd.DataFrame:
    cols = ["event_name", "schema_status", "parameter", "group", "value", "result", "severity", "issue_type", "reason"]
    if df.empty:
        return pd.DataFrame(columns=cols)
    out = pd.DataFrame(columns=cols)
    out["event_name"] = df.get("event_name", "")
    out["schema_status"] = df.get("result", "")
    out["parameter"] = df.get("missing_parameters", "")
    out["group"] = df.get("section_name", "")
    out["value"] = df.get("matched_identification", "")
    out["result"] = df.get("result", "")
    out["severity"] = ""
    out["issue_type"] = df.get("failure_reason", "")
    out["reason"] = df.get("failure_reason", "")
    return out


def _beautify_qa_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="E8EEF8")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="1E3A8A")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(1, min(ws.max_row, 120) + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                if value is None:
                    continue
                max_len = max(max_len, len(str(value)))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)
        for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 1000), min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)
    wb.close()


def _resolve_qa_template_path() -> Optional[Path]:
    candidates = [
        Path("/Users/havalovely/Desktop/ranking/001_qa_result_musinsa_ranking.xlsx"),
        Path("/Users/havalovely/Desktop/ranking/002_qa_result_beauty_ranking.xlsx"),
    ]
    for path in candidates:
        if path.exists() and path.is_file():
            return path
    return None


def _apply_qa_template_style(path: Path) -> bool:
    template_path = _resolve_qa_template_path()
    if template_path is None:
        return False
    out_wb = None
    tpl_wb = None
    try:
        out_wb = load_workbook(path)
        tpl_wb = load_workbook(template_path)
        for ws in out_wb.worksheets:
            try:
                if ws.title not in tpl_wb.sheetnames:
                    continue
                tpl = tpl_wb[ws.title]
                ws.freeze_panes = tpl.freeze_panes
                ws.auto_filter.ref = ws.dimensions if ws.max_row >= 1 and ws.max_column >= 1 else None
                if tpl.row_dimensions[1].height:
                    ws.row_dimensions[1].height = tpl.row_dimensions[1].height
                for key, dim in tpl.column_dimensions.items():
                    if dim.width:
                        ws.column_dimensions[key].width = dim.width
                max_col = min(max(1, ws.max_column), max(1, tpl.max_column))
                for col in range(1, max_col + 1):
                    src = tpl.cell(1, col)
                    dst = ws.cell(1, col)
                    dst.font = copy(src.font)
                    dst.fill = copy(src.fill)
                    dst.border = copy(src.border)
                    dst.alignment = copy(src.alignment)
                    dst.number_format = src.number_format
                    dst.protection = copy(src.protection)
            except Exception:
                continue
        out_wb.save(path)
        return True
    except Exception:
        return False
    finally:
        try:
            if tpl_wb is not None:
                tpl_wb.close()
        except Exception:
            pass
        try:
            if out_wb is not None:
                out_wb.close()
        except Exception:
            pass


def _normalize_result_text(value: Any) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return "UNCHECKED"
    if text in {"pass", "matched", "ok", "success"}:
        return "PASS"
    if text in {"unchecked"}:
        return "UNCHECKED"
    if text in {"warning", "warn", "retest needed"}:
        return "WARNING"
    if text in {"missing", "mismatch", "blocked", "fail", "failed", "error", "unmatched"}:
        return "FAIL"
    return "WARNING"


def _infer_issue_type(result: str, missing: str, extra: str, reason: str) -> str:
    r = str(reason or "").strip().lower()
    if result == "PASS":
        return ""
    if "no_hit" in r or "not captured" in r or "이벤트가 수집되지 않" in r:
        return "NO_HIT"
    if "불일치" in r or "mismatch" in r:
        return "MISMATCH"
    if str(missing or "").strip():
        return "MISSING_PARAM"
    if str(extra or "").strip():
        return "MISMATCH"
    if result == "UNCHECKED":
        return "UNCHECKED"
    return "MISMATCH"


def _pick_series(df: pd.DataFrame, candidates: List[str], default: str = "") -> pd.Series:
    for name in candidates:
        if name in df.columns:
            return df[name].fillna("").astype(str)
    return pd.Series([default] * len(df), index=df.index, dtype="string")


def _build_review_rows(summary_df: pd.DataFrame, detail_df: pd.DataFrame, unmatched_df: pd.DataFrame) -> pd.DataFrame:
    base = detail_df if not detail_df.empty else summary_df
    if base.empty:
        return pd.DataFrame(
            columns=[
                "No", "이벤트 설명", "이벤트명", "파라미터명", "정의서 기대값", "실제 수집값", "결과", "오류구분",
                "오류 설명", "정의서 기준 params", "정의서 누락 params", "정의서 외 params", "실제 추가 파라미터", "실제 params(json)",
                "발생 페이지", "section / selector", "증거", "QA 메모",
            ]
        )

    no_col = _pick_series(base, ["no", "event_no"])
    event_name = _pick_series(base, ["event_name", "event name"])
    section_name = _pick_series(base, ["section_name", "scenario"])
    page_id = _pick_series(base, ["page_id"])
    param_name_raw = _pick_series(base, ["param_name", "parameter", "missing_parameters"])
    expected_val = _pick_series(base, ["expected_value", "missing_parameters"])
    actual_val = _pick_series(base, ["actual_value", "extra_parameters", "matched_identification"])
    raw_result = _pick_series(base, ["result", "match_status"])
    failure_reason = _pick_series(base, ["failure_reason", "reason"])
    selector = _pick_series(base, ["selector"])
    matched_count = _pick_series(base, ["matched_capture_count", "matched_hit_count"])
    observed_param_preview = _pick_series(base, ["observed_param_preview"])
    observed_params_json = _pick_series(base, ["observed_params_json"])
    expected_params_json = _pick_series(base, ["expected_params_json"])

    result = raw_result.map(_normalize_result_text)

    def _derive_param_name(i: int) -> str:
        direct = str(param_name_raw.iloc[i]).strip()
        if "=" in direct:
            left = direct.split("=", 1)[0].strip()
            if left:
                direct = left
        if direct and direct.lower() not in {"event", "이벤트"}:
            return direct
        reason_text = str(failure_reason.iloc[i]).strip()
        m = re.search(r"필수\s*파라미터\s*([a-zA-Z0-9_]+)", reason_text)
        if m:
            return str(m.group(1)).strip()
        ev = str(expected_val.iloc[i]).strip()
        if "section_name" in ev.lower():
            return "section_name"
        if "=" in ev:
            return ev.split("=", 1)[0].strip() or "-"
        av = str(actual_val.iloc[i]).strip()
        if "=" in av:
            return av.split("=", 1)[0].strip() or "-"
        return "정의서 파라미터 미정의"

    derived_param_name = pd.Series([_derive_param_name(i) for i in range(len(base))], index=base.index, dtype="string")

    def _parse_json_object(raw: str) -> Dict[str, Any]:
        text = str(raw or "").strip()
        if not text:
            return {}
        try:
            data = json.loads(text)
            if isinstance(data, dict):
                return data
        except Exception:
            return {}
        return {}

    def _expected_keys(param_name: str, expected_text: str, observed: Dict[str, Any], expected_map: Dict[str, Any]) -> set[str]:
        keys: set[str] = set()
        pn = str(param_name or "").strip()
        if pn and pn not in {"-", "정의서 파라미터 미정의"}:
            keys.add(pn)
        exp = str(expected_text or "").strip()
        for m in re.finditer(r"([a-zA-Z_][a-zA-Z0-9_]*)\s*=", exp):
            keys.add(str(m.group(1)).strip())
        # 기대값에 자연어가 섞여 있어도, observed 키와 이름이 겹치는 토큰은 기대 키로 채택
        stop = {"and", "or", "true", "false", "null", "none", "not", "set"}
        exp_tokens = {t.lower() for t in re.findall(r"[a-zA-Z_][a-zA-Z0-9_]*", exp)}
        for key in observed.keys():
            lk = str(key).strip().lower()
            if lk and lk in exp_tokens and lk not in stop:
                keys.add(str(key).strip())
        for key in expected_map.keys():
            key_text = str(key).strip()
            if key_text:
                keys.add(key_text)
        return {k for k in keys if k}

    def _format_pairs(items: List[tuple[str, Any]], limit: int = 14) -> str:
        if not items:
            return "-"
        body = [f"{k}={v}" for k, v in items[:limit]]
        if len(items) > limit:
            body.append(f"...(+{len(items) - limit})")
        return "\n".join(body)

    params_defined_col: List[str] = []
    params_missing_col: List[str] = []
    params_extra_col: List[str] = []
    params_extra_preview_col: List[str] = []
    for i in range(len(base)):
        observed = _parse_json_object(str(observed_params_json.iloc[i]))
        expected_map = _parse_json_object(str(expected_params_json.iloc[i]))
        expected_keys = _expected_keys(str(derived_param_name.iloc[i]), str(expected_val.iloc[i]), observed, expected_map)
        ordered_expected_keys: List[str] = [str(k).strip() for k in expected_map.keys() if str(k).strip()]
        for key in sorted(expected_keys):
            if key not in ordered_expected_keys:
                ordered_expected_keys.append(key)
        in_def: List[tuple[str, Any]] = []
        missing_def: List[tuple[str, Any]] = []
        out_def: List[tuple[str, Any]] = []
        for key_text in ordered_expected_keys:
            if key_text in observed:
                in_def.append((key_text, observed.get(key_text)))
            else:
                exp_v = str(expected_map.get(key_text, "")).strip()
                miss_text = f"(missing){' expected:' + exp_v if exp_v else ''}"
                in_def.append((key_text, miss_text))
                missing_def.append((key_text, miss_text))
        for k, v in observed.items():
            key_text = str(k).strip()
            if not key_text:
                continue
            if key_text not in expected_keys:
                out_def.append((key_text, v))
        params_defined_col.append(_format_pairs(in_def))
        params_missing_col.append(_format_pairs(missing_def))
        params_extra_col.append(_format_pairs(out_def))
        params_extra_preview_col.append(_format_pairs(out_def, limit=8))

    issue_type = [
        _infer_issue_type(str(result.iloc[i]), str(expected_val.iloc[i]), str(actual_val.iloc[i]), str(failure_reason.iloc[i]))
        for i in range(len(base))
    ]

    error_desc = []
    for i in range(len(base)):
        rr = str(failure_reason.iloc[i]).strip()
        if rr:
            error_desc.append(rr)
            continue
        if result.iloc[i] == "PASS":
            error_desc.append("정의서 기대값과 실제 수집값이 일치")
        elif issue_type[i] == "MISSING_PARAM":
            error_desc.append(f"필수 파라미터 누락: {str(expected_val.iloc[i]).strip()}")
        elif issue_type[i] == "NO_HIT":
            error_desc.append("이벤트가 수집되지 않음")
        elif issue_type[i] == "UNCHECKED":
            error_desc.append("판정 근거 부족")
        else:
            error_desc.append("정의서 기대값과 실제값 불일치")

    evidence = []
    for i in range(len(base)):
        marks: List[str] = []
        if str(selector.iloc[i]).strip():
            marks.append("selector")
        if str(actual_val.iloc[i]).strip():
            marks.append("payload")
        if str(matched_count.iloc[i]).strip() and str(matched_count.iloc[i]).strip() not in {"0", ""}:
            marks.append("hit")
        evidence.append(", ".join(marks) if marks else "-")

    review = pd.DataFrame(
        {
            "No": no_col,
            "이벤트 설명": section_name.where(section_name.str.strip() != "", event_name),
            "이벤트명": event_name,
            "파라미터명": derived_param_name,
            "정의서 기대값": expected_val,
            "실제 수집값": actual_val,
            "결과": result,
            "오류구분": issue_type,
            "오류 설명": error_desc,
            "정의서 기준 params": params_defined_col,
            "정의서 누락 params": params_missing_col,
            "정의서 외 params": params_extra_col,
            "실제 추가 파라미터": params_extra_preview_col if any(x != "-" for x in params_extra_preview_col) else observed_param_preview,
            "실제 params(json)": observed_params_json,
            "발생 페이지": page_id,
            "section / selector": section_name.where(section_name.str.strip() != "", "") + selector.map(lambda x: f" / {x}" if str(x).strip() else ""),
            "증거": evidence,
            "QA 메모": "",
        }
    )

    # `qa_result.csv` / `qa_result_detail.csv`에 Missing이 이미 반영되어 있으므로,
    # unmatched를 여기서 재추가하면 KPI가 중복 집계된다.

    return _sanitize_excel_df(review.fillna(""))


def _apply_review_sheet_style(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="EEF3FF")
    defined_params_fill = PatternFill("solid", fgColor="EAF8EF")
    extra_params_fill = PatternFill("solid", fgColor="F3F4F6")
    result_colors = {
        "PASS": PatternFill("solid", fgColor="E9FBEF"),
        "FAIL": PatternFill("solid", fgColor="FDECEC"),
        "WARNING": PatternFill("solid", fgColor="FFF4DE"),
        "UNCHECKED": PatternFill("solid", fgColor="EFEFEF"),
    }
    issue_colors = {
        "NO_HIT": PatternFill("solid", fgColor="FDECEC"),
        "MISSING_PARAM": PatternFill("solid", fgColor="FFF4DE"),
        "MISMATCH": PatternFill("solid", fgColor="F8ECFF"),
        "UNCHECKED": PatternFill("solid", fgColor="EFEFEF"),
    }
    for ws in wb.worksheets:
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = Font(bold=True, color="1E3A8A")
            c.fill = header_fill
            c.alignment = Alignment(horizontal="left", vertical="center")
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(1, min(ws.max_row, 220) + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                max_len = max(max_len, len(str(value or "")))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 9), 44)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    if "Summary" in wb.sheetnames:
        sws = wb["Summary"]
        sws["A1"] = "QA Review Summary"
        sws["A1"].font = Font(bold=True, color="1E3A8A", size=13)

    if "QA Review" in wb.sheetnames:
        ws = wb["QA Review"]
        ws.freeze_panes = "E2"
        headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
        result_idx = headers.index("결과") + 1 if "결과" in headers else 0
        issue_idx = headers.index("오류구분") + 1 if "오류구분" in headers else 0
        event_idx = headers.index("이벤트명") + 1 if "이벤트명" in headers else 0
        defined_params_idx = headers.index("정의서 기준 params") + 1 if "정의서 기준 params" in headers else 0
        extra_params_idx = headers.index("정의서 외 params") + 1 if "정의서 외 params" in headers else 0
        prev_event = ""
        alt = False
        alt_fill = PatternFill("solid", fgColor="FAFAFA")
        if event_idx > 0:
            for r in range(2, ws.max_row + 1):
                event_name = str(ws.cell(r, event_idx).value or "").strip()
                if event_name and event_name != prev_event:
                    alt = not alt
                    prev_event = event_name
                if alt:
                    for c in range(1, ws.max_column + 1):
                        ws.cell(r, c).fill = alt_fill
                if result_idx > 0:
                    rv = str(ws.cell(r, result_idx).value or "").strip().upper()
                    if rv in result_colors:
                        ws.cell(r, result_idx).fill = result_colors[rv]
                if issue_idx > 0:
                    iv = str(ws.cell(r, issue_idx).value or "").strip().upper()
                    if iv in issue_colors:
                        ws.cell(r, issue_idx).fill = issue_colors[iv]
                if defined_params_idx > 0:
                    dv = str(ws.cell(r, defined_params_idx).value or "").strip()
                    if dv and dv != "-":
                        ws.cell(r, defined_params_idx).fill = defined_params_fill
                if extra_params_idx > 0:
                    ev = str(ws.cell(r, extra_params_idx).value or "").strip()
                    if ev and ev != "-":
                        ws.cell(r, extra_params_idx).fill = extra_params_fill
    wb.save(path)
    wb.close()


def _build_summary_sheet(meta: Dict[str, Any], review_df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    result_counts = {"PASS": 0, "FAIL": 0, "WARNING": 0, "UNCHECKED": 0}
    issue_counts: Dict[str, int] = {}
    if not review_df.empty:
        for value in review_df["결과"].fillna("").astype(str):
            key = str(value).strip().upper()
            if key in result_counts:
                result_counts[key] += 1
        for value in review_df["오류구분"].fillna("").astype(str):
            key = str(value).strip().upper()
            if not key:
                continue
            issue_counts[key] = issue_counts.get(key, 0) + 1
    total = len(review_df)
    definition_items = int(meta.get("definition_row_count", 0) or 0)
    captured_events = int(meta.get("captured_event_count", 0) or 0)
    one_line = (
        f"리뷰 판정 {total}개 항목 중 {result_counts['PASS']}개 PASS, {result_counts['FAIL']}개 FAIL, "
        f"{result_counts['WARNING']}개 WARNING, {result_counts['UNCHECKED']}개 UNCHECKED입니다."
    )

    meta_rows = pd.DataFrame(
        [
            ["Project", str(meta.get("project_name", "")).strip()],
            ["Page", str(meta.get("target_page", "")).strip()],
            ["Session", str(meta.get("session_name", "")).strip()],
            ["Session ID", str(meta.get("session_id", "")).strip()],
            ["Viewport", str(meta.get("viewport", "")).strip()],
            ["Browser", str(meta.get("browser", "")).strip()],
            ["QA Mode", str(meta.get("qa_mode", "")).strip()],
            ["Definition", str(meta.get("definition_name", "")).strip()],
            ["Version", str(meta.get("version_name", "")).strip()],
            ["Test Scope", str(meta.get("test_scope", "")).strip()],
            ["Generated At", str(meta.get("generated_at", "")).strip()],
        ],
        columns=["항목", "값"],
    )
    kpi_rows = pd.DataFrame(
        [
            ["정의서 검증 항목 수", definition_items if definition_items > 0 else total],
            ["실제 수집 이벤트 수", captured_events],
            ["리뷰 판정 항목 수", total],
            ["PASS", result_counts["PASS"]],
            ["FAIL", result_counts["FAIL"]],
            ["WARNING", result_counts["WARNING"]],
            ["UNCHECKED", result_counts["UNCHECKED"]],
        ],
        columns=["KPI", "값"],
    )
    issue_rows = pd.DataFrame([[k, v] for k, v in sorted(issue_counts.items())], columns=["오류유형", "건수"])
    one_line_df = pd.DataFrame([["해석", one_line]], columns=["항목", "값"])
    return meta_rows, kpi_rows, issue_rows, one_line_df


def _ensure_review_exports(
    run_dir: Path,
    session: Dict[str, Any],
    meta: Dict[str, Any],
    generate_advanced: bool = False,
) -> Dict[str, Optional[Path]]:
    existing_review = run_dir / "qa_review.xlsx"
    existing_result = run_dir / "qa_result.xlsx"
    existing_issues = run_dir / "qa_issues_only.xlsx"
    existing_raw = run_dir / "qa_raw_debug.xlsx"
    qa_mode_text = (
        str(session.get("qa_mode", "")).strip()
        or str(meta.get("qa_mode", "")).strip()
        or str(session.get("judgement_rule", "")).strip()
        or "-"
    )
    is_definition_mode = qa_mode_text == "정의서 검증"
    if existing_review.exists() and existing_result.exists():
        if not is_definition_mode:
            # 과거 포맷(정의서 시트)로 생성된 파일은 탐색/시나리오 전용 포맷으로 재생성한다.
            try:
                wb = load_workbook(existing_result, read_only=True)
                sheetnames = set(wb.sheetnames)
                wb.close()
                if "Collected Events" in sheetnames:
                    return {
                        "qa_review": existing_review,
                        "qa_result": existing_result,
                        "qa_issues_only": existing_issues if existing_issues.exists() else None,
                        "qa_raw_debug": existing_raw if existing_raw.exists() else None,
                    }
            except Exception:
                pass
        else:
            return {
                "qa_review": existing_review,
                "qa_result": existing_result,
                "qa_issues_only": existing_issues if existing_issues.exists() else None,
                "qa_raw_debug": existing_raw if existing_raw.exists() else None,
            }

    summary_csv = run_dir / "qa_result.csv"
    detail_csv = run_dir / "qa_result_detail.csv"
    definition_csv = run_dir / "event_definition.csv"
    unmatched_csv = run_dir / "definition_validation_unmatched.csv"
    matched_csv = run_dir / "definition_validation_matches.csv"
    if (
        not summary_csv.exists()
        and not detail_csv.exists()
        and not definition_csv.exists()
        and not unmatched_csv.exists()
        and not matched_csv.exists()
    ):
        return {"qa_review": None, "qa_result": None, "qa_issues_only": None, "qa_raw_debug": None}

    summary_df = _read_csv_for_export(summary_csv) if summary_csv.exists() else pd.DataFrame()
    detail_df = _read_csv_for_export(detail_csv) if detail_csv.exists() else pd.DataFrame()
    observed_df = _read_csv_for_export(definition_csv) if definition_csv.exists() else pd.DataFrame()
    unmatched_df = _read_csv_for_export(unmatched_csv) if unmatched_csv.exists() else pd.DataFrame()
    matched_df = _read_csv_for_export(matched_csv) if matched_csv.exists() else pd.DataFrame()

    review_df = _build_review_rows(summary_df, detail_df, unmatched_df) if is_definition_mode else pd.DataFrame()
    if is_definition_mode:
        issues_df = review_df[review_df["결과"].isin(["FAIL", "WARNING", "UNCHECKED"])].copy()
        extras_df = pd.DataFrame(columns=["No", "이벤트명", "파라미터명", "정의서 외 params", "실제 params(json)", "발생 페이지", "결과", "오류 설명"])
        missing_df = pd.DataFrame(columns=["No", "이벤트명", "파라미터명", "정의서 누락 params", "정의서 기대값", "실제 수집값", "결과", "오류 설명", "발생 페이지"])
        if "정의서 외 params" in review_df.columns:
            extras_df = review_df[review_df["정의서 외 params"].fillna("").astype(str).str.strip().ne("-")][
                ["No", "이벤트명", "파라미터명", "정의서 외 params", "실제 params(json)", "발생 페이지", "결과", "오류 설명"]
            ].copy()
        if "정의서 누락 params" in review_df.columns:
            missing_df = review_df[review_df["정의서 누락 params"].fillna("").astype(str).str.strip().ne("-")][
                ["No", "이벤트명", "파라미터명", "정의서 누락 params", "정의서 기대값", "실제 수집값", "결과", "오류 설명", "발생 페이지"]
            ].copy()
        review_export_df = review_df.drop(columns=["정의서 외 params"], errors="ignore")
        issues_export_df = issues_df.drop(columns=["정의서 외 params"], errors="ignore")
        if not issues_export_df.empty:
            issues_export_df.insert(0, "안내", "이 시트는 FAIL / WARNING / UNCHECKED 항목만 모아둔 검토용 시트입니다.")
    else:
        issues_df = pd.DataFrame()
        extras_df = pd.DataFrame()
        missing_df = pd.DataFrame()
        review_export_df = pd.DataFrame()
        issues_export_df = pd.DataFrame()

    summary_meta = {
        "project_name": str(meta.get("project_slug", "")).strip() or "default-project",
        "target_page": _extract_path_label(str(meta.get("target_url", "")).strip()),
        "session_name": str(session.get("name", "")).strip(),
        "session_id": str(meta.get("session_id", "")).strip(),
        "viewport": (
            str(session.get("viewport", "")).strip()
            or str(meta.get("viewport", "")).strip()
            or "Unknown"
        ),
        "browser": (
            str(session.get("browser", "")).strip()
            or str(meta.get("browser", "")).strip()
            or "chromium"
        ),
        "qa_mode": qa_mode_text,
        "definition_name": str(session.get("definition_spec_name", "")).strip() or str(meta.get("definition_file_name", "")).strip(),
        "version_name": str(session.get("version_name", "")).strip() or str(session.get("version_id", "")).strip() or "ver 1.0",
        "test_scope": str(session.get("test_scope", "")).strip() or str(meta.get("bundle_scope", "")).strip() or "-",
        "definition_row_count": int(meta.get("definition_row_count", 0) or 0),
        "captured_event_count": int(meta.get("captured_event_count", 0) or session.get("captured_events", 0) or 0),
        "generated_at": datetime.now().isoformat(),
    }
    if is_definition_mode:
        meta_rows, kpi_rows, issue_rows, one_line_df = _build_summary_sheet(summary_meta, review_df)
    else:
        simple_df = pd.DataFrame(
            {
                "No": summary_df.get("no", pd.Series(range(1, len(summary_df) + 1))),
                "수집시각": summary_df.get("captured_at", ""),
                "소스": summary_df.get("source", ""),
                "이벤트명": summary_df.get("event_name", ""),
                "section_name": summary_df.get("section_name", ""),
                "page_id": summary_df.get("page_id", ""),
                "결과": summary_df.get("result", ""),
                "파라미터 요약": summary_df.get("observed_param_preview", ""),
                "params(json)": summary_df.get("observed_params_json", ""),
            }
        )
        simple_df = _sanitize_excel_df(simple_df.fillna(""))
        event_counts = (
            simple_df.groupby("이벤트명", dropna=False)
            .size()
            .reset_index(name="hits")
            .sort_values("hits", ascending=False)
            if not simple_df.empty
            else pd.DataFrame(columns=["이벤트명", "hits"])
        )
        source_counts = (
            simple_df.groupby("소스", dropna=False)
            .size()
            .reset_index(name="hits")
            .sort_values("hits", ascending=False)
            if not simple_df.empty
            else pd.DataFrame(columns=["소스", "hits"])
        )
        meta_rows = pd.DataFrame(
            [
                ["Project", summary_meta["project_name"]],
                ["Page", summary_meta["target_page"]],
                ["Session", summary_meta["session_name"]],
                ["Session ID", summary_meta["session_id"]],
                ["Viewport", summary_meta["viewport"]],
                ["Browser", summary_meta["browser"]],
                ["QA Mode", summary_meta["qa_mode"]],
                ["Version", summary_meta["version_name"]],
                ["Test Scope", summary_meta["test_scope"]],
                ["Generated At", summary_meta["generated_at"]],
            ],
            columns=["항목", "값"],
        )
        kpi_rows = pd.DataFrame(
            [
                ["수집 hits", int(summary_meta["captured_event_count"])],
                ["고유 이벤트 수", int(event_counts["이벤트명"].nunique() if not event_counts.empty else 0)],
                ["source 종류 수", int(source_counts["소스"].nunique() if not source_counts.empty else 0)],
            ],
            columns=["KPI", "값"],
        )
        issue_rows = pd.DataFrame(columns=["오류유형", "건수"])
        one_line_df = pd.DataFrame(
            [["해석", f"탐색/시나리오 모드 결과: 총 {int(summary_meta['captured_event_count'])} hits 수집"]],
            columns=["항목", "값"],
        )

    qa_review = run_dir / "qa_review.xlsx"
    qa_result = run_dir / "qa_result.xlsx"
    qa_issues = run_dir / "qa_issues_only.xlsx"
    qa_raw = run_dir / "qa_raw_debug.xlsx"

    try:
        with pd.ExcelWriter(qa_review, engine="openpyxl") as writer:
            meta_rows.to_excel(writer, sheet_name="Summary", index=False, startrow=1)
            kpi_rows.to_excel(writer, sheet_name="Summary", index=False, startrow=len(meta_rows) + 4)
            issue_rows.to_excel(writer, sheet_name="Summary", index=False, startrow=len(meta_rows) + len(kpi_rows) + 7)
            one_line_df.to_excel(writer, sheet_name="Summary", index=False, startrow=len(meta_rows) + len(kpi_rows) + len(issue_rows) + 10)
            if is_definition_mode:
                review_export_df.to_excel(writer, sheet_name="QA Review", index=False)
                issues_export_df.to_excel(writer, sheet_name="Issues Only", index=False)
                extras_df.to_excel(writer, sheet_name="정의서 외 params", index=False)
                missing_df.to_excel(writer, sheet_name="정의서 누락 params", index=False)
            else:
                _sanitize_excel_df(summary_df).to_excel(writer, sheet_name="Collected Events", index=False)
                if "event_name" in summary_df.columns:
                    (
                        summary_df.groupby("event_name", dropna=False)
                        .size()
                        .reset_index(name="hits")
                        .sort_values("hits", ascending=False)
                    ).to_excel(writer, sheet_name="Event Counts", index=False)
                if "source" in summary_df.columns:
                    (
                        summary_df.groupby("source", dropna=False)
                        .size()
                        .reset_index(name="hits")
                        .sort_values("hits", ascending=False)
                    ).to_excel(writer, sheet_name="Source Counts", index=False)
        _apply_review_sheet_style(qa_review)
        shutil.copy2(qa_review, qa_result)
    except Exception:
        return {"qa_review": None, "qa_result": None, "qa_issues_only": None, "qa_raw_debug": None}

    try:
        with pd.ExcelWriter(qa_issues, engine="openpyxl") as writer:
            issues_export_df.to_excel(writer, sheet_name="Issues Only", index=False)
        _apply_review_sheet_style(qa_issues)
    except Exception:
        qa_issues = None  # type: ignore[assignment]

    if generate_advanced:
        try:
            with pd.ExcelWriter(qa_raw, engine="openpyxl") as writer:
                _sanitize_excel_df(summary_df).to_excel(writer, sheet_name="summary_raw", index=False)
                _sanitize_excel_df(detail_df).to_excel(writer, sheet_name="detail_raw", index=False)
                _sanitize_excel_df(matched_df).to_excel(writer, sheet_name="matched_raw", index=False)
                _sanitize_excel_df(unmatched_df).to_excel(writer, sheet_name="unmatched_raw", index=False)
                _sanitize_excel_df(observed_df).to_excel(writer, sheet_name="observed_raw", index=False)
            _beautify_qa_workbook(qa_raw)
        except Exception:
            qa_raw = None  # type: ignore[assignment]
    else:
        qa_raw = qa_raw if qa_raw.exists() else None  # type: ignore[assignment]

    return {
        "qa_review": qa_review if qa_review.exists() else None,
        "qa_result": qa_result if qa_result.exists() else None,
        "qa_issues_only": qa_issues if isinstance(qa_issues, Path) and qa_issues.exists() else None,
        "qa_raw_debug": qa_raw if isinstance(qa_raw, Path) and qa_raw.exists() else None,
    }


def load_results(project_slug: str) -> Dict[str, Any]:
    sessions = load_sessions(project_slug)
    recent_sessions = sessions[:20]
    recent_session_ids = {
        str(s.get("session_id", "")).strip()
        for s in recent_sessions
        if str(s.get("session_id", "")).strip()
    }
    session_by_id = {
        str(s.get("session_id", "")).strip(): s
        for s in sessions
        if str(s.get("session_id", "")).strip()
    }
    exports_root = get_project_paths(project_slug)["exports"]
    export_files: List[Dict[str, str]] = []
    seen_paths: set[str] = set()
    canonical_fallbacks = {
        "qa_result.csv": ["definition_validation_summary.csv"],
        "qa_result_detail.csv": ["definition_validation_matches.csv", "definition_validation_unmatched.csv"],
    }
    export_aliases = [("qa_result.xlsx", ["qa_result.xlsx", "qa_report.xlsx"])]
    if exports_root.exists():
        run_dirs = _collect_export_runs(exports_root)
        def _safe_mtime(path: Path) -> float:
            try:
                return path.stat().st_mtime
            except OSError:
                return 0.0
        # 경로 문자열이 아니라 실제 수정시각 기준 최신순 정렬
        run_dirs.sort(key=_safe_mtime, reverse=True)
        # 최신 세션 결과 누락 방지를 위해 지나친 선절단을 피한다.
        run_dirs = run_dirs[:200]
        for run_dir in run_dirs:
            if not run_dir.is_dir():
                continue
            meta = _read_json(run_dir / "meta.json", {})
            session_id = str(meta.get("session_id", "")).strip()
            run_id = str(meta.get("run_id", run_dir.name)).strip() or run_dir.name
            target_url = str(meta.get("target_url", "")).strip()
            sess = session_by_id.get(session_id, {})
            run_title = _export_title_from_session(sess, target_url)
            for canonical, candidates in canonical_fallbacks.items():
                canonical_path = run_dir / canonical
                if canonical_path.exists():
                    continue
                source_path: Optional[Path] = None
                for filename in candidates:
                    p = run_dir / filename
                    if p.exists():
                        source_path = p
                        break
                if source_path is not None:
                    try:
                        canonical_path.write_bytes(source_path.read_bytes())
                    except Exception:
                        pass
            generated = {
                "qa_review": run_dir / "qa_review.xlsx",
                "qa_result": run_dir / "qa_result.xlsx",
                "qa_issues_only": run_dir / "qa_issues_only.xlsx",
                "qa_raw_debug": run_dir / "qa_raw_debug.xlsx",
            }
            preferred = [
                ("qa_result.csv", run_dir / "qa_result.csv", "basic"),
                ("qa_review.xlsx", generated.get("qa_review"), "basic"),
                ("qa_result.xlsx", generated.get("qa_result"), "basic"),
                ("qa_issues_only.xlsx", generated.get("qa_issues_only"), "advanced"),
                ("qa_raw_debug.xlsx", generated.get("qa_raw_debug"), "advanced"),
            ]
            for display_name, path_obj, tier in preferred:
                if not isinstance(path_obj, Path) or not path_obj.exists():
                    continue
                rel_path = str(path_obj.relative_to(BASE_DIR))
                if rel_path in seen_paths:
                    continue
                seen_paths.add(rel_path)
                export_files.append(
                    {
                        "name": f"{run_title} · {run_id} · {display_name}",
                        "path": rel_path,
                        "tier": tier,
                        "session_id": session_id,
                        "run_id": run_id,
                    }
                )
            for display_name, candidates in export_aliases:
                picked: Optional[Path] = None
                for filename in candidates:
                    file_path = run_dir / filename
                    if file_path.exists():
                        picked = file_path
                        break
                if not picked:
                    continue
                rel_path = str(picked.relative_to(BASE_DIR))
                if rel_path in seen_paths:
                    continue
                seen_paths.add(rel_path)
                export_files.append(
                    {
                        "name": f"{run_title} · {run_id} · {display_name}",
                        "path": rel_path,
                        "tier": "advanced",
                        "session_id": session_id,
                        "run_id": run_id,
                    }
                )
    status_counts = {
        "Matched": 0,
        "Mismatch": 0,
        "Missing": 0,
        "Blocked": 0,
        "Unchecked": 0,
        "Retest Needed": 0,
    }
    for s in sessions:
        key = str(s.get("qa_status", "Unchecked"))
        if key not in status_counts:
            key = "Unchecked"
        status_counts[key] += 1

    destination_rate = {
        "GA hit 확인 세션 비율": 0,
        "전송 가능 세션 비율": 0,
    }
    total = len(sessions)
    if total > 0:
        matched_like = status_counts["Matched"]
        destination_rate["GA hit 확인 세션 비율"] = round((matched_like / total) * 100, 1)
        destination_rate["전송 가능 세션 비율"] = round(((total - status_counts["Blocked"]) / total) * 100, 1)

    # Results 메뉴는 최근 테스트 세션 중심으로 노출한다.
    export_files_recent = [
        item for item in export_files
        if str(item.get("session_id", "")).strip() in recent_session_ids
    ]
    export_session_ids = {
        str(item.get("session_id", "")).strip()
        for item in export_files_recent
        if str(item.get("session_id", "")).strip()
    }
    review_export_session_ids = {
        str(item.get("session_id", "")).strip()
        for item in export_files_recent
        if str(item.get("session_id", "")).strip()
        and str(item.get("path", "")).strip().lower().endswith((".xlsx", ".xls"))
    }
    collection_notices: List[str] = []
    # 최근 세션에서 결과가 없으면(종료 상태 기준) 빠른 CSV를 즉시 1회 생성 시도한다.
    # 목록 진입 시 "계속 대기"되는 체감을 줄이기 위한 보정 로직.
    for s in recent_sessions[:8]:
        sid = str(s.get("session_id", "")).strip()
        if not sid or sid in export_session_ids:
            continue
        runtime = str(s.get("runtime_status", "")).strip().lower()
        if runtime in {"running", "stopping", "auto_crawl_paused"}:
            continue
        try:
            paths = get_project_paths(project_slug)
            run_meta = _read_json(paths["run_meta"], {})
            meta_item = run_meta.get(sid, {}) if isinstance(run_meta, dict) and isinstance(run_meta.get(sid, {}), dict) else {}
            qa_mode = str(meta_item.get("qa_mode", "")).strip()
            if qa_mode == "정의서 검증":
                _finalize_definition_validation_exports(
                    project_slug=project_slug,
                    session_id=sid,
                    include_review_exports=False,
                )
            else:
                _finalize_basic_session_exports(project_slug=project_slug, session_id=sid)
        except Exception:
            pass

    # 자동 생성 시도 이후 export 목록을 다시 계산한다.
    if exports_root.exists():
        export_files = []
        seen_paths = set()
        run_dirs = _collect_export_runs(exports_root)
        def _safe_mtime(path: Path) -> float:
            try:
                return path.stat().st_mtime
            except OSError:
                return 0.0
        run_dirs.sort(key=_safe_mtime, reverse=True)
        run_dirs = run_dirs[:200]
        for run_dir in run_dirs:
            if not run_dir.is_dir():
                continue
            meta = _read_json(run_dir / "meta.json", {})
            session_id = str(meta.get("session_id", "")).strip()
            run_id = str(meta.get("run_id", run_dir.name)).strip() or run_dir.name
            target_url = str(meta.get("target_url", "")).strip()
            sess = session_by_id.get(session_id, {})
            run_title = _export_title_from_session(sess, target_url)
            preferred = [
                ("qa_result.csv", run_dir / "qa_result.csv", "basic"),
                ("qa_review.xlsx", run_dir / "qa_review.xlsx", "basic"),
                ("qa_result.xlsx", run_dir / "qa_result.xlsx", "basic"),
                ("qa_issues_only.xlsx", run_dir / "qa_issues_only.xlsx", "advanced"),
                ("qa_raw_debug.xlsx", run_dir / "qa_raw_debug.xlsx", "advanced"),
            ]
            for display_name, path_obj, tier in preferred:
                if not isinstance(path_obj, Path) or not path_obj.exists():
                    continue
                rel_path = str(path_obj.relative_to(BASE_DIR))
                if rel_path in seen_paths:
                    continue
                seen_paths.add(rel_path)
                export_files.append(
                    {
                        "name": f"{run_title} · {run_id} · {display_name}",
                        "path": rel_path,
                        "tier": tier,
                        "session_id": session_id,
                        "run_id": run_id,
                    }
                )
        export_files_recent = [
            item for item in export_files
            if str(item.get("session_id", "")).strip() in recent_session_ids
        ]
        export_session_ids = {
            str(item.get("session_id", "")).strip()
            for item in export_files_recent
            if str(item.get("session_id", "")).strip()
        }
        review_export_session_ids = {
            str(item.get("session_id", "")).strip()
            for item in export_files_recent
            if str(item.get("session_id", "")).strip()
            and str(item.get("path", "")).strip().lower().endswith((".xlsx", ".xls"))
        }

    for s in recent_sessions[:8]:
        sid = str(s.get("session_id", "")).strip()
        if not sid:
            continue
        runtime = str(s.get("runtime_status", "")).strip().lower()
        hits = int(s.get("captured_events", 0) or 0)
        if runtime in {"running", "stopping", "auto_crawl_paused"}:
            collection_notices.append(
                f"{sid}: 수집 진행중/종료 처리중입니다. (현재 {hits} hits)"
            )
            continue
        if sid not in export_session_ids:
            collection_notices.append(
                f"{sid}: 결과 파일 생성 대기 상태입니다. (수집 {hits} hits)"
            )
            continue
        if sid not in review_export_session_ids:
            collection_notices.append(
                f"{sid}: 빠른 결과(CSV)는 준비됨, 리뷰 엑셀 생성 진행중입니다."
            )

    return {
        "status_counts": status_counts,
        "destination_rate": destination_rate,
        "exports": export_files_recent[:60],
        "latest_sessions": recent_sessions,
        "collection_notices": collection_notices[:8],
    }


def load_project_overview(project_slug: str) -> Dict[str, Any]:
    sessions = load_sessions(project_slug)
    versions = load_versions(project_slug)

    latest_version = versions[0]["display_name"] if versions else "-"
    last_session = sessions[0] if sessions else {}

    missing_candidates: List[str] = []
    major_issues: List[str] = []
    for sess in sessions[:30]:
        status = str(sess.get("qa_status", "")).strip()
        if status == "Missing":
            missing_candidates.append(str(sess.get("name", "")).strip())
        if status in {"Missing", "Mismatch", "Blocked", "Retest Needed"}:
            major_issues.append(f"{sess.get('name', '-')}: {status}")

    destination_success = 0.0
    if sessions:
        ok = sum(1 for s in sessions if s.get("qa_status") in {"Matched", "Mismatch", "Missing", "Retest Needed"})
        destination_success = round((ok / len(sessions)) * 100, 1)

    return {
        "latest_version": latest_version,
        "recent_run_result": str(last_session.get("qa_status", "Unchecked")) if last_session else "Unchecked",
        "destination_success_rate": destination_success,
        "recent_missing_events": missing_candidates[:5],
        "major_issues": major_issues[:6],
        "last_tested_at": str(last_session.get("started_at", "-")) if last_session else "-",
        "session_count": len(sessions),
    }


def build_view_model(project_slug: str, session_filter: str = "all") -> Dict[str, Any]:
    projects = list_projects()
    if not projects:
        return {
            "projects": [],
            "selected_project": None,
            "overview": {},
            "versions": [],
            "sessions": [],
        "results": {},
        "settings": {},
        "definition_specs": [],
        "session_filter": session_filter,
        "run_type_labels": RUN_TYPE_LABELS,
    }

    selected = project_slug
    if not selected or selected not in {p.slug for p in projects}:
        selected = projects[0].slug

    all_sessions = load_sessions(selected)
    all_versions = load_versions(selected, include_deleted=True)
    active_versions = [v for v in all_versions if not bool(v.get("deleted", False))]
    deleted_versions = [v for v in all_versions if bool(v.get("deleted", False))]
    filtered_sessions = all_sessions
    if session_filter in {"exploratory", "scenario", "spec_validation"}:
        filtered_sessions = [s for s in all_sessions if str(s.get("run_type", "")) == session_filter]

    return {
        "projects": projects,
        "selected_project": selected,
        "overview": load_project_overview(selected),
        "versions": active_versions,
        "deleted_versions": deleted_versions,
        "sessions": filtered_sessions,
        "session_total": len(all_sessions),
        "results": load_results(selected),
        "settings": load_settings(selected),
        "definition_specs": load_definition_specs(selected),
        "session_filter": session_filter,
        "run_type_labels": RUN_TYPE_LABELS,
    }


def build_session_list_payload(project_slug: str, session_filter: str = "all") -> Dict[str, Any]:
    sessions = load_sessions(project_slug)
    if session_filter in {"exploratory", "scenario", "spec_validation"}:
        sessions = [s for s in sessions if str(s.get("run_type", "")) == session_filter]
    return {
        "project": project_slug,
        "session_filter": session_filter,
        "count": len(sessions),
        "sessions": sessions,
    }


def _json_object(value: object) -> Dict[str, Any]:
    if isinstance(value, dict):
        return value
    raw = str(value or "").strip()
    if not raw:
        return {}
    try:
        payload = json.loads(raw)
        if isinstance(payload, dict):
            return payload
    except Exception:
        return {}
    return {}


def load_session_diagnosis(
    project_slug: str, session_id: str, limit: int = 5000
) -> Dict[str, Any]:
    sid = str(session_id or "").strip()
    if not sid:
        return {"ok": False, "error": "empty_session_id", "session_id": ""}

    sessions = load_sessions(project_slug)
    session_row = next((s for s in sessions if str(s.get("session_id", "")).strip() == sid), None)
    if not session_row:
        return {"ok": False, "error": "session_not_found", "session_id": sid}

    db_path = get_project_paths(project_slug)["qa_sessions"] / "qa_runs.db"
    if not db_path.exists():
        return {"ok": False, "error": "db_not_found", "session_id": sid}

    query = """
        SELECT captured_at, source, event_name, params_json
        FROM qa_events
        WHERE session_id = ?
        ORDER BY captured_at DESC
        LIMIT ?
    """

    total_events = 0
    source_counts: Dict[str, int] = {}
    event_counts: Dict[str, int] = {}
    stage_counts: Dict[str, int] = {}
    auto_click_checks = 0
    auto_click_network_hit = 0
    auto_click_clicks = 0
    popup_close_failed = 0
    popup_close_confirmed = 0
    blocked_signals = 0
    footer_like_clicks = 0
    hypothesis_checks = 0
    hypothesis_met = 0

    try:
        with sqlite3.connect(str(db_path)) as conn:
            conn.row_factory = sqlite3.Row
            for row in conn.execute(query, (sid, int(max(1, limit)))):
                total_events += 1
                item = dict(row)
                source = str(item.get("source", "")).strip()
                event_name = str(item.get("event_name", "")).strip()
                params = _json_object(item.get("params_json", "{}"))

                if source:
                    source_counts[source] = int(source_counts.get(source, 0)) + 1
                if event_name:
                    event_counts[event_name] = int(event_counts.get(event_name, 0)) + 1

                if source == "definition_match_debug" and event_name == "definition_match_debug":
                    stage = str(params.get("stage", "")).strip()
                    if stage:
                        stage_counts[stage] = int(stage_counts.get(stage, 0)) + 1

                if source == "auto_crawl" and event_name == "auto_click":
                    auto_click_clicks += 1
                    pattern_key = str(params.get("pattern_key", "")).strip().lower()
                    if any(token in pattern_key for token in ["footer", "gnb", "copyright"]):
                        footer_like_clicks += 1

                if source == "auto_crawl" and event_name == "auto_click_hit_check":
                    auto_click_checks += 1
                    hit_ok = bool(params.get("network_hit_confirmed", False))
                    try:
                        hit_delta = int(params.get("network_hit_delta", 0) or 0)
                    except Exception:
                        hit_delta = 0
                    if hit_ok or hit_delta > 0:
                        auto_click_network_hit += 1

                if source == "auto_crawl" and event_name == "popup_close_failed":
                    popup_close_failed += 1
                if source == "auto_crawl" and event_name == "popup_close_confirmed":
                    popup_close_confirmed += 1
                if source == "auto_crawl" and event_name in {
                    "auto_crawl_blocked",
                    "auto_crawl_nav_blocked",
                }:
                    blocked_signals += 1
                if source == "auto_crawl" and event_name == "auto_click_hypothesis_check":
                    hypothesis_checks += 1
                    if bool(params.get("hypothesis_met", False)):
                        hypothesis_met += 1
    except Exception as exc:
        return {"ok": False, "error": str(exc), "session_id": sid}

    ga_hits = int(source_counts.get("ga_hit", 0))
    row_candidate_check = int(stage_counts.get("row_candidate_check", 0))
    row_not_matched = int(stage_counts.get("row_not_matched_stop_here", 0))
    row_matched_total = int(stage_counts.get("row_matched", 0)) + int(
        stage_counts.get("row_matched_relaxed", 0)
    ) + int(stage_counts.get("row_matched_third", 0))
    skipped_no_event_candidates = int(
        stage_counts.get("row_relaxed_skipped_no_event_candidates", 0)
    )

    findings: List[Dict[str, Any]] = []
    recommendations: List[str] = []

    auto_click_hit_rate = (
        round((auto_click_network_hit / auto_click_checks) * 100.0, 1)
        if auto_click_checks > 0
        else None
    )
    row_match_rate = (
        round((row_matched_total / row_candidate_check) * 100.0, 1)
        if row_candidate_check > 0
        else None
    )
    hypothesis_rate = (
        round((hypothesis_met / hypothesis_checks) * 100.0, 1)
        if hypothesis_checks > 0
        else None
    )

    if blocked_signals > 0:
        findings.append(
            {
                "severity": "high",
                "code": "BLOCK_SIGNAL",
                "message": "차단/챌린지 의심 신호가 감지되었습니다.",
                "evidence": {"blocked_signals": blocked_signals},
            }
        )
        recommendations.append("차단 의심 시 즉시 실패 처리하지 말고 재확인 대기 후 재시도 분기를 추가하세요.")

    if auto_click_checks >= 5 and auto_click_hit_rate is not None and auto_click_hit_rate < 20.0:
        findings.append(
            {
                "severity": "high",
                "code": "LOW_AUTO_CLICK_HIT_RATE",
                "message": "자동 클릭 대비 hit 확인 비율이 낮습니다.",
                "evidence": {
                    "auto_click_checks": auto_click_checks,
                    "auto_click_network_hit": auto_click_network_hit,
                    "auto_click_hit_rate": auto_click_hit_rate,
                },
            }
        )
        recommendations.append("클릭 전 상태 분류(정상/팝업/필터/로딩/차단)를 선행하고 상태별 클릭 후보를 분리하세요.")

    if row_candidate_check >= 10 and row_match_rate is not None and row_match_rate < 40.0:
        findings.append(
            {
                "severity": "medium",
                "code": "LOW_ROW_MATCH_RATE",
                "message": "정의서 row 후보 대비 최종 매칭 비율이 낮습니다.",
                "evidence": {
                    "row_candidate_check": row_candidate_check,
                    "row_matched_total": row_matched_total,
                    "row_not_matched_stop_here": row_not_matched,
                    "row_match_rate": row_match_rate,
                },
            }
        )
        recommendations.append("event_name 1차 일치 후 section/page/button/target 신호를 가중치 기반으로 결합해 보수적으로 완화 매칭하세요.")

    if row_candidate_check >= 10 and skipped_no_event_candidates > (row_candidate_check * 0.4):
        findings.append(
            {
                "severity": "medium",
                "code": "NO_EVENT_CANDIDATE_SPIKE",
                "message": "strict 실패 후 relaxed 후보(동일/동족 이벤트) 부재가 많습니다.",
                "evidence": {
                    "row_candidate_check": row_candidate_check,
                    "row_relaxed_skipped_no_event_candidates": skipped_no_event_candidates,
                },
            }
        )
        recommendations.append("hit 이벤트 키를 event_name/event/en 통합 추출로 맞추고, 수집 타이밍 지연(로딩 직후)을 보완하세요.")

    if popup_close_failed > popup_close_confirmed and popup_close_failed >= 3:
        findings.append(
            {
                "severity": "medium",
                "code": "POPUP_CLOSE_UNSTABLE",
                "message": "팝업 닫기 시도가 반복 실패하고 있습니다.",
                "evidence": {
                    "popup_close_failed": popup_close_failed,
                    "popup_close_confirmed": popup_close_confirmed,
                },
            }
        )
        recommendations.append("팝업 상태에서는 본문 클릭을 금지하고 팝업 처리 전용 루틴만 수행하세요.")

    if footer_like_clicks >= 3:
        findings.append(
            {
                "severity": "low",
                "code": "NON_MAIN_CLICK_DRIFT",
                "message": "footer/gnb 계열 클릭 비중이 관찰됩니다.",
                "evidence": {
                    "footer_like_clicks": footer_like_clicks,
                    "auto_click_clicks": auto_click_clicks,
                },
            }
        )
        recommendations.append("main 영역 가중치를 더 높이고 footer/gnb 섹션은 패널티를 강화하세요.")

    if hypothesis_checks > 0 and hypothesis_rate is not None and hypothesis_rate < 50.0:
        findings.append(
            {
                "severity": "medium",
                "code": "LOW_HYPOTHESIS_RATE",
                "message": "클릭 가설 충족률이 낮습니다.",
                "evidence": {
                    "hypothesis_checks": hypothesis_checks,
                    "hypothesis_met": hypothesis_met,
                    "hypothesis_rate": hypothesis_rate,
                },
            }
        )
        recommendations.append("가설 실패 연속 시 동일 패턴 재클릭을 감점하고 상태 재분류를 강제하세요.")

    if not findings:
        recommendations.append("명시적 이상 신호는 낮습니다. 다음 테스트에서 동일 시나리오를 1회 더 반복해 재현성만 확인하세요.")

    return {
        "ok": True,
        "session_id": sid,
        "generated_at": datetime.now().isoformat(),
        "metrics": {
            "total_events": total_events,
            "ga_hits": ga_hits,
            "source_counts": source_counts,
            "event_counts_top": sorted(
                [{"event_name": k, "count": v} for k, v in event_counts.items()],
                key=lambda x: int(x.get("count", 0)),
                reverse=True,
            )[:20],
            "definition_stage_counts": stage_counts,
            "auto_click_clicks": auto_click_clicks,
            "auto_click_checks": auto_click_checks,
            "auto_click_network_hit": auto_click_network_hit,
            "auto_click_hit_rate": auto_click_hit_rate,
            "row_match_rate": row_match_rate,
            "hypothesis_checks": hypothesis_checks,
            "hypothesis_met": hypothesis_met,
            "hypothesis_rate": hypothesis_rate,
        },
        "findings": findings,
        "recommendations": recommendations,
    }


def load_session_detail(project_slug: str, session_id: str, limit: int = 200) -> Dict[str, Any]:
    sid = str(session_id or "").strip()
    if not sid:
        return {"session_id": "", "summary": {}, "events": [], "suspicion": []}
    sessions = load_sessions(project_slug)
    session_row = next((s for s in sessions if str(s.get("session_id", "")).strip() == sid), None)
    if not session_row:
        return {"session_id": sid, "summary": {}, "events": [], "suspicion": []}

    db_path = get_project_paths(project_slug)["qa_sessions"] / "qa_runs.db"
    query = """
        SELECT captured_at, source, event_name, params_json, page_url, request_method
        FROM qa_events
        WHERE session_id = ?
        ORDER BY captured_at DESC
        LIMIT ?
    """
    events: List[Dict[str, Any]] = []
    unique_event_names: set[str] = set()
    definition_match_counts = {"strict_matched": 0, "relaxed_matched": 0, "total_matched": 0}
    if db_path.exists():
        try:
            with sqlite3.connect(str(db_path)) as conn:
                conn.row_factory = sqlite3.Row
                for row in conn.execute(query, (sid, int(max(1, limit)))):
                    item = dict(row)
                    try:
                        params = json.loads(str(item.get("params_json", "{}")) or "{}")
                    except Exception:
                        params = {}
                    preview_items: List[str] = []
                    if isinstance(params, dict):
                        for idx, (k, v) in enumerate(params.items()):
                            if idx >= 4:
                                break
                            preview_items.append(f"{k}={str(v)}")
                    events.append(
                        {
                            "captured_at": str(item.get("captured_at", "")).strip(),
                            "source": str(item.get("source", "")).strip(),
                            "event_name": str(item.get("event_name", "")).strip(),
                            "page_url": str(item.get("page_url", "")).strip(),
                            "request_method": str(item.get("request_method", "")).strip(),
                            "params_preview": ", ".join(preview_items) if preview_items else "-",
                        }
                    )
                try:
                    for row in conn.execute(
                        """
                        SELECT DISTINCT event_name
                        FROM qa_events
                        WHERE session_id = ?
                          AND event_name IS NOT NULL
                          AND TRIM(event_name) != ''
                        """,
                        (sid,),
                    ):
                        event_name = str(row[0] if isinstance(row, sqlite3.Row) else row[0]).strip()
                        if not event_name or _is_system_event_name(event_name):
                            continue
                        unique_event_names.add(event_name)
                except Exception:
                    unique_event_names = set()
                try:
                    row = conn.execute(
                        """
                        SELECT
                            SUM(
                                CASE
                                    WHEN source = 'definition_match_debug'
                                     AND (
                                        params_json LIKE '%"stage":"row_matched"%'
                                        OR params_json LIKE '%"stage": "row_matched"%'
                                     )
                                    THEN 1
                                    ELSE 0
                                END
                            ) AS strict_matched,
                            SUM(
                                CASE
                                    WHEN source = 'definition_match_debug'
                                     AND (
                                        params_json LIKE '%"stage":"row_matched_relaxed"%'
                                        OR params_json LIKE '%"stage": "row_matched_relaxed"%'
                                     )
                                    THEN 1
                                    ELSE 0
                                END
                            ) AS relaxed_matched,
                            SUM(
                                CASE
                                    WHEN source = 'definition_match_debug'
                                     AND (
                                        params_json LIKE '%"stage":"row_matched_third"%'
                                        OR params_json LIKE '%"stage": "row_matched_third"%'
                                     )
                                    THEN 1
                                    ELSE 0
                                END
                            ) AS third_matched
                        FROM qa_events
                        WHERE session_id = ?
                        """,
                        (sid,),
                    ).fetchone()
                    strict_matched = int((row[0] if row else 0) or 0)
                    relaxed_matched = int((row[1] if row else 0) or 0)
                    third_matched = int((row[2] if row else 0) or 0)
                    definition_match_counts = {
                        "strict_matched": strict_matched,
                        "relaxed_matched": relaxed_matched,
                        "third_matched": third_matched,
                        "total_matched": strict_matched + relaxed_matched + third_matched,
                    }
                except Exception:
                    definition_match_counts = {
                        "strict_matched": 0,
                        "relaxed_matched": 0,
                        "third_matched": 0,
                        "total_matched": 0,
                    }
        except Exception:
            events = []
            unique_event_names = set()
            definition_match_counts = {
                "strict_matched": 0,
                "relaxed_matched": 0,
                "third_matched": 0,
                "total_matched": 0,
            }

    suspicion: List[str] = []
    text_blob = " ".join(
        [
            f"{str(e.get('source', ''))} {str(e.get('event_name', ''))} {str(e.get('params_preview', ''))}"
            for e in events[:300]
        ]
    ).lower()
    if "captcha" in text_blob or "access_challenge_detected" in text_blob:
        suspicion.append("환경 차단 가능성: captcha/access challenge 감지")
    if "missing" in text_blob or "required_missing" in text_blob:
        suspicion.append("필수 이벤트/파라미터 누락 의심")
    if "mismatch" in text_blob or "fail_param" in text_blob or "fail_section" in text_blob:
        suspicion.append("값/매핑 불일치 의심")
    if not suspicion:
        suspicion.append("명시적 오류 신호가 크지 않음. 표본 검토 권장")

    def _empty_counts() -> Dict[str, int]:
        return {
            "Matched": 0,
            "Mismatch": 0,
            "Missing": 0,
            "Blocked": 0,
            "Unchecked": 0,
            "Retest Needed": 0,
        }

    def _result_to_bucket(result_value: str, reason_value: str, missing_value: str, extra_value: str) -> str:
        r = str(result_value or "").strip().lower()
        reason = str(reason_value or "").strip().lower()
        missing = str(missing_value or "").strip().lower()
        extra = str(extra_value or "").strip().lower()
        if "retest" in r or "retest" in reason:
            return "Retest Needed"
        if "blocked" in r or "captcha" in reason or "challenge" in reason:
            return "Blocked"
        if not r or r == "unchecked":
            return "Unchecked"
        if r == "missing":
            return "Missing"
        if r == "mismatch":
            return "Mismatch"
        if r in {"pass", "matched", "ok", "success"}:
            return "Matched"
        if "missing" in r:
            return "Missing"
        if "mismatch" in r:
            return "Mismatch"
        if "fail" in r or "warning" in r or "unmatched" in r:
            if any(token in reason for token in ["no_hit", "not captured", "수집되지 않", "미발생"]):
                return "Missing"
            if missing:
                return "Missing"
            if any(token in reason for token in ["mismatch", "불일치"]) or extra:
                return "Mismatch"
            if any(token in reason for token in ["blocked", "captcha", "challenge"]):
                return "Blocked"
            if any(token in reason for token in ["retest", "재검증"]):
                return "Retest Needed"
            return "Mismatch"
        return "Unchecked"

    session_counts = _empty_counts()
    latest_run: Optional[Path] = None
    latest_run_meta: Dict[str, Any] = {}
    latest_run_id = ""
    runtime_status = str(session_row.get("runtime_status", "-")).strip().lower()
    qa_mode_value = str(session_row.get("qa_mode", "")).strip()
    try:
        exports_root = get_project_paths(project_slug)["exports"]
        def _scan_latest_run() -> None:
            nonlocal latest_run, latest_run_meta, latest_run_id
            latest_key = ""
            latest_run = None
            latest_run_meta = {}
            latest_run_id = ""
            if not exports_root.exists():
                return
            for run_dir in _iter_export_runs(exports_root):
                meta = _read_json(run_dir / "meta.json", {})
                if str(meta.get("session_id", "")).strip() != sid:
                    continue
                run_key = str(meta.get("run_id", run_dir.name)).strip() or run_dir.name
                if run_key > latest_key:
                    latest_key = run_key
                    latest_run = run_dir
                    latest_run_meta = meta if isinstance(meta, dict) else {}
                    latest_run_id = run_key

        _scan_latest_run()
        # stop API를 거치지 않은 종료 세션 복구: 조회 시점에 자동 생성
        if (
            latest_run is None
            and qa_mode_value == "정의서 검증"
            and runtime_status not in {"running", "stopping"}
        ):
            try:
                _finalize_definition_validation_exports(
                    project_slug=project_slug,
                    session_id=sid,
                    include_review_exports=False,
                )
            except Exception:
                pass
            _scan_latest_run()
        if latest_run is not None:
            summary_csv = latest_run / "qa_result.csv"
            if summary_csv.exists():
                df = _read_csv_for_export(summary_csv)
                if not df.empty:
                    result_col = _pick_series(df, ["result"], "")
                    reason_col = _pick_series(df, ["failure_reason", "reason"], "")
                    missing_col = _pick_series(df, ["missing_parameters", "missing_parameter"], "")
                    extra_col = _pick_series(df, ["extra_parameters", "extra_parameter"], "")
                    for i in range(len(df)):
                        bucket = _result_to_bucket(
                            str(result_col.iloc[i]),
                            str(reason_col.iloc[i]),
                            str(missing_col.iloc[i]),
                            str(extra_col.iloc[i]),
                        )
                        session_counts[bucket] = int(session_counts.get(bucket, 0)) + 1
            else:
                session_counts["Unchecked"] = int(session_row.get("captured_events", 0) or 0)
        else:
            session_counts["Unchecked"] = int(session_row.get("captured_events", 0) or 0)
    except Exception:
        session_counts = _empty_counts()

    validation_state = "pending"
    if runtime_status in {"running", "stopping"}:
        validation_state = "in_progress"
    elif latest_run is not None and (latest_run / "qa_result.csv").exists():
        validation_state = "completed"
    elif str(session_row.get("qa_mode", "")).strip() != "정의서 검증":
        validation_state = "unavailable"

    found_unique_events = len(unique_event_names)
    found_event_hits = int(session_row.get("captured_events", 0) or 0)
    definition_target_events = int(session_row.get("definition_event_target_count", 0) or 0)
    definition_found_events = 0
    if qa_mode_value == "정의서 검증":
        spec_id = str(session_row.get("definition_spec_id", "")).strip()
        target_event_names: set[str] = set()
        if spec_id:
            runtime_payload = _build_definition_runtime_payload(project_slug, spec_id)
            target_event_names = {
                str(v).strip()
                for v in runtime_payload.get("event_names", [])
                if str(v).strip()
            } if isinstance(runtime_payload.get("event_names", []), list) else set()
        if target_event_names:
            definition_target_events = len(target_event_names)
            definition_found_events = len(unique_event_names.intersection(target_event_names))
        elif definition_target_events > 0:
            # target list를 복구하지 못한 경우에는 최소한 0/target 형태로라도 상태를 보여준다.
            definition_found_events = 0

    if qa_mode_value == "정의서 검증":
        target_total = max(0, int(definition_target_events))
        status_text = (
            f"자동수집 상태 · 정의서 이벤트 {definition_found_events}/{target_total}개 탐지"
            f" · 전체 이벤트 {found_unique_events}개 탐지 ({found_event_hits} hits)"
        )
    else:
        status_text = (
            f"자동수집 상태 · 전체 이벤트 {found_unique_events}개 탐지 ({found_event_hits} hits)"
        )

    summary = {
        "session_id": sid,
        "run_type": str(session_row.get("run_type_label", "-")).strip(),
        "qa_mode": str(session_row.get("qa_mode", "-")).strip(),
        "qa_status": str(session_row.get("qa_status", "Unchecked")).strip(),
        "runtime_status": str(session_row.get("runtime_status", "-")).strip(),
        "captured_events": int(session_row.get("captured_events", 0) or 0),
        "started_at": str(session_row.get("started_at", "-")).strip(),
        "target_url": str(session_row.get("target_url", "-")).strip(),
        "tester": str(session_row.get("tester", "-")).strip(),
        "last_error": str(session_row.get("last_error", "")).strip(),
        "validation_counts": session_counts,
        "validation_state": validation_state,
        "validation_updated_at": str(latest_run_meta.get("saved_at", "")).strip(),
        "validation_run_id": latest_run_id,
        "definition_match_counts": definition_match_counts,
        "found_unique_events": found_unique_events,
        "found_event_hits": found_event_hits,
        "definition_target_events": int(definition_target_events),
        "definition_found_events": int(definition_found_events),
        "auto_collect_status_text": status_text,
        "analytics_source_label": str(session_row.get("analytics_source_label", "Unknown")).strip() or "Unknown",
        "ga4_hits": int(session_row.get("ga4_hits", 0) or 0),
        "amplitude_hits": int(session_row.get("amplitude_hits", 0) or 0),
    }
    try:
        self_debug_payload = load_session_diagnosis(project_slug=project_slug, session_id=sid)
    except Exception as exc:
        self_debug_payload = {
            "ok": False,
            "session_id": sid,
            "error": f"self_debug_failed: {str(exc)}",
        }

    return {
        "session_id": sid,
        "summary": summary,
        "events": events,
        "suspicion": suspicion,
        "self_debug": self_debug_payload,
    }
