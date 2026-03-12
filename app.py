from __future__ import annotations

from datetime import date, datetime, timedelta
import hashlib
import html
import json
import os
import io
import re
import shutil
import sqlite3
import zipfile
from pathlib import Path
from typing import Dict, List, Tuple
from urllib import error as urllib_error
from urllib import parse as urllib_parse
from urllib import request as urllib_request
from urllib.parse import parse_qs, urlparse
from uuid import uuid4
from zoneinfo import ZoneInfo

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font

from src.debug_runtime import (
    get_debug_session_snapshot,
    load_debug_events,
    start_debug_session,
    stop_debug_session,
)
from src.flatten import flatten_events, normalize_param_name
from src.funnel import analyze_funnel_sequence
from src.issue_store import (
    DEFAULT_ISSUE_PATH,
    issues_to_dataframe,
    resolve_issue,
    upsert_auto_issues,
)
from src.qa_ingest_server import ensure_ingest_server
from src.reporting import (
    build_report_filename,
    events_to_csv_bytes,
    results_to_csv_bytes,
    results_to_pdf_bytes,
)
from src.scenario_validation import validate_single_scenario
from src.rules import run_qa_rules
from src.scoring import compute_integrity_score
from src.test_log_db import append_ui_action, get_session, init_test_log_db, list_recent_sessions


st.set_page_config(page_title="GA4 QA Reporter", layout="wide")
st.title("GA4 QA 리포터 (룰 기반)")
st.caption("실시간 DebugView 대체가 아닌, 로그 수집 + 자동 정리 + 룰 기반 판정을 위한 내부 QA 도구")
BASE_DIR = Path(__file__).resolve().parent
DEFAULT_OAUTH_REDIRECT_URI = "https://asknuggetdata.com/oauth2callback"
TEST_LOG_DB_PATH = Path("data/test_logs/qa_runs.db")
SCHEMA_STORE_PATH = Path("data/schemas/event_schemas.json")
WORKSPACE_ROOT = Path("data/workspace")
DEFAULT_PROJECT_SLUG = "default-project"
OAUTH_CONTEXT_DIR = Path("data/oauth_context")
CIRCLED_NUMBERS = {
    1: "①", 2: "②", 3: "③", 4: "④", 5: "⑤", 6: "⑥", 7: "⑦", 8: "⑧", 9: "⑨", 10: "⑩",
    11: "⑪", 12: "⑫", 13: "⑬", 14: "⑭", 15: "⑮", 16: "⑯", 17: "⑰", 18: "⑱", 19: "⑲", 20: "⑳",
}
ILLEGAL_EXCEL_CHAR_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")


def _slugify_project_name(text: str) -> str:
    raw = str(text or "").strip().lower()
    if not raw:
        return DEFAULT_PROJECT_SLUG
    # Keep unicode alnum chars (e.g. Korean), normalize separators to '-'
    normalized = "".join(ch if (ch.isalnum() or ch in {"-", "_"}) else "-" for ch in raw)
    normalized = re.sub(r"[-_]{2,}", "-", normalized)
    slug = normalized.strip("-_")
    return slug or DEFAULT_PROJECT_SLUG


def to_excel_safe_text(value: object) -> str:
    text = str(value if value is not None else "")
    return ILLEGAL_EXCEL_CHAR_RE.sub("", text)


def to_circled_number(n: int) -> str:
    return CIRCLED_NUMBERS.get(int(n), str(n))


def get_current_app_base_url() -> str:
    override = get_config_value("GA4_APP_BASE_URL", "").strip()
    if override:
        return override.rstrip("/")
    headers = _get_request_headers()
    host = str(headers.get("x-forwarded-host", "")).strip() or str(headers.get("host", "")).strip()
    proto = str(headers.get("x-forwarded-proto", "")).split(",")[0].strip().lower()
    if not host:
        return ""
    if proto not in {"http", "https"}:
        proto = "https"
    return f"{proto}://{host}".rstrip("/")


def save_oauth_context(state: str, context: Dict[str, object]) -> None:
    sid = str(state or "").strip()
    if not sid:
        return
    OAUTH_CONTEXT_DIR.mkdir(parents=True, exist_ok=True)
    payload = dict(context)
    payload["saved_at"] = datetime.now(tz=ZoneInfo("UTC")).isoformat()
    (OAUTH_CONTEXT_DIR / f"{sid}.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def pop_oauth_context(state: str) -> Dict[str, object]:
    sid = str(state or "").strip()
    if not sid:
        return {}
    path = OAUTH_CONTEXT_DIR / f"{sid}.json"
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            data = {}
    except Exception:
        data = {}
    try:
        path.unlink(missing_ok=True)
    except Exception:
        pass
    return data


def _workspace_projects_root() -> Path:
    return WORKSPACE_ROOT / "projects"


def list_workspace_projects() -> List[str]:
    root = _workspace_projects_root()
    if not root.exists():
        return [DEFAULT_PROJECT_SLUG]
    names = sorted([p.name for p in root.iterdir() if p.is_dir()])
    return names or [DEFAULT_PROJECT_SLUG]


def ensure_project_structure(project_slug: str) -> Dict[str, Path]:
    slug = _slugify_project_name(project_slug)
    project_root = _workspace_projects_root() / slug
    paths = {
        "project_root": project_root,
        "schemas_dir": project_root / "schemas",
        "qa_sessions_dir": project_root / "qa_sessions",
        "event_logs_dir": project_root / "event_logs",
        "tracking_dir": project_root / "tracking_plan",
        "exports_dir": project_root / "exports",
        "artifacts_dir": project_root / "artifacts",
        "schema_file": project_root / "schemas" / "event_schemas.json",
        "db_file": project_root / "qa_sessions" / "qa_runs.db",
        "action_map_file": project_root / "tracking_plan" / "action_object_map.csv",
    }
    for key in ["project_root", "schemas_dir", "qa_sessions_dir", "event_logs_dir", "tracking_dir", "exports_dir", "artifacts_dir"]:
        paths[key].mkdir(parents=True, exist_ok=True)
    return paths


def get_active_project_slug() -> str:
    try:
        slug = _slugify_project_name(st.session_state.get("qa_project_slug", DEFAULT_PROJECT_SLUG))
    except Exception:
        slug = DEFAULT_PROJECT_SLUG
    return slug


def get_active_project_paths() -> Dict[str, Path]:
    return ensure_project_structure(get_active_project_slug())


def get_active_schema_path() -> Path:
    return get_active_project_paths()["schema_file"]


def get_active_test_log_db_path() -> Path:
    return get_active_project_paths()["db_file"]


def get_active_action_map_path() -> Path:
    return get_active_project_paths()["action_map_file"]


def get_definition_validation_cache_path(project_slug: str | None = None) -> Path:
    slug = _slugify_project_name(project_slug or get_active_project_slug())
    return ensure_project_structure(slug)["artifacts_dir"] / "definition_validation_cache.json"


def save_definition_validation_cache(payload: Dict[str, object], project_slug: str | None = None) -> None:
    path = get_definition_validation_cache_path(project_slug)
    data = dict(payload)
    data["saved_at"] = datetime.now(tz=ZoneInfo("UTC")).isoformat()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_definition_validation_cache(project_slug: str | None = None) -> Dict[str, object]:
    path = get_definition_validation_cache_path(project_slug)
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def restore_definition_validation_cache_if_needed() -> None:
    if st.session_state.get("qa_uploaded_definition_schemas"):
        return
    cached = load_definition_validation_cache()
    schemas = cached.get("schemas", {})
    if not isinstance(schemas, dict) or not schemas:
        return
    st.session_state["qa_mode"] = str(cached.get("qa_mode", "정의서 검증")).strip() or "정의서 검증"
    st.session_state["qa_definition_validation_file_name"] = str(cached.get("file_name", "")).strip()
    st.session_state["qa_definition_validation_error"] = ""
    st.session_state["qa_uploaded_definition_schemas"] = schemas
    summary = cached.get("summary", [])
    st.session_state["qa_uploaded_definition_summary"] = summary if isinstance(summary, list) else []
    debug = cached.get("debug", {})
    st.session_state["qa_definition_validation_debug"] = debug if isinstance(debug, dict) else {}
    st.session_state["qa_definition_validation_loaded_key"] = str(cached.get("loaded_key", "")).strip()
    st.session_state["qa_definition_validation_target_mode"] = str(cached.get("target_mode", "선택하세요")).strip() or "선택하세요"
    st.session_state["qa_definition_validation_target_mode_prev"] = st.session_state["qa_definition_validation_target_mode"]
    st.session_state["qa_definition_validation_execution_mode"] = str(cached.get("execution_mode", "")).strip()
    st.session_state["qa_definition_validation_execution_session_id"] = str(cached.get("execution_session_id", "")).strip()


def _normalize_action_map_df(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(
            columns=[
                "event_name",
                "action",
                "object",
                "page_id",
                "section_name",
                "section_index",
                "description",
                "object_index",
            ]
        )
    rename = {str(c): normalize_param_name(str(c)) for c in df.columns}
    work = df.rename(columns=rename).copy()
    aliases = {
        "section": "section_name",
        "sectionid": "section_index",
        "index": "object_index",
        "event": "event_name",
    }
    for src, dst in aliases.items():
        if src in work.columns and dst not in work.columns:
            work[dst] = work[src]

    required_cols = [
        "event_name",
        "action",
        "object",
        "page_id",
        "section_name",
        "section_index",
        "description",
        "object_index",
    ]
    for col in required_cols:
        if col not in work.columns:
            work[col] = ""
    work = work[required_cols].copy()
    for col in required_cols:
        work[col] = work[col].astype(str).str.strip()
    work = work[work["event_name"] != ""].drop_duplicates(subset=["event_name"], keep="last").reset_index(drop=True)
    return work


def load_action_object_map(path: Path | None = None) -> pd.DataFrame:
    if path is None:
        path = get_active_action_map_path()
    if not path.exists():
        return _normalize_action_map_df(pd.DataFrame())
    try:
        raw = pd.read_csv(path)
    except Exception:
        return _normalize_action_map_df(pd.DataFrame())
    return _normalize_action_map_df(raw)


def save_action_object_map(df: pd.DataFrame, path: Path | None = None) -> None:
    if path is None:
        path = get_active_action_map_path()
    normalized = _normalize_action_map_df(df)
    path.parent.mkdir(parents=True, exist_ok=True)
    normalized.to_csv(path, index=False, encoding="utf-8-sig")


def action_map_template_csv_bytes() -> bytes:
    template = pd.DataFrame(
        [
            {
                "event_name": "click_button",
                "action": "click",
                "object": "button",
                "page_id": "/home",
                "section_name": "hero_banner",
                "section_index": "1",
                "description": "메인 배너 버튼 클릭",
                "object_index": "1",
            },
            {
                "event_name": "view_item_list",
                "action": "view",
                "object": "product",
                "page_id": "/home",
                "section_name": "product_list",
                "section_index": "2",
                "description": "상품 리스트 노출",
                "object_index": "{index}",
            },
        ]
    )
    return template.to_csv(index=False).encode("utf-8-sig")


def import_legacy_data_to_project(project_slug: str) -> Dict[str, int]:
    paths = ensure_project_structure(project_slug)
    target_db = paths["db_file"]
    init_test_log_db(target_db)

    counts = {
        "sessions": 0,
        "events": 0,
        "actions": 0,
        "streams": 0,
        "schemas": 0,
    }

    legacy_db = Path("data/test_logs/qa_runs.db")
    if legacy_db.exists():
        with sqlite3.connect(str(legacy_db)) as src, sqlite3.connect(str(target_db)) as dst:
            src.row_factory = sqlite3.Row
            dst.execute("PRAGMA journal_mode=WAL;")
            dst.execute("PRAGMA synchronous=NORMAL;")

            session_cols = [
                "session_id",
                "target_url",
                "status",
                "started_at",
                "ended_at",
                "captured_events",
                "last_error",
                "tester_name",
                "tester_note",
                "created_at",
            ]
            src_rows = src.execute(f"SELECT {', '.join(session_cols)} FROM qa_sessions").fetchall()
            for row in src_rows:
                vals = [row[c] for c in session_cols]
                dst.execute(
                    f"INSERT OR IGNORE INTO qa_sessions ({', '.join(session_cols)}) VALUES ({', '.join(['?'] * len(session_cols))})",
                    vals,
                )
            counts["sessions"] = int(len(src_rows))

            event_cols = [
                "session_id",
                "captured_at",
                "source",
                "event_name",
                "params_json",
                "page_url",
                "measurement_id",
                "client_id",
                "request_method",
            ]
            existing_event_count = int(dst.execute("SELECT COUNT(1) FROM qa_events").fetchone()[0] or 0)
            if existing_event_count == 0:
                event_rows = src.execute(f"SELECT {', '.join(event_cols)} FROM qa_events").fetchall()
                for row in event_rows:
                    vals = [row[c] for c in event_cols]
                    dst.execute(
                        f"INSERT INTO qa_events ({', '.join(event_cols)}) VALUES ({', '.join(['?'] * len(event_cols))})",
                        vals,
                    )
                counts["events"] = int(len(event_rows))

            action_cols = [
                "action_at",
                "session_id",
                "action_type",
                "actor_role",
                "actor_name",
                "detail_json",
                "remote_addr_hash",
                "user_agent",
            ]
            existing_action_count = int(dst.execute("SELECT COUNT(1) FROM qa_ui_actions").fetchone()[0] or 0)
            if existing_action_count == 0:
                action_rows = src.execute(f"SELECT {', '.join(action_cols)} FROM qa_ui_actions").fetchall()
                for row in action_rows:
                    vals = [row[c] for c in action_cols]
                    dst.execute(
                        f"INSERT INTO qa_ui_actions ({', '.join(action_cols)}) VALUES ({', '.join(['?'] * len(action_cols))})",
                        vals,
                    )
                counts["actions"] = int(len(action_rows))

            dst.commit()

    legacy_stream_dir = Path("data/debug_stream")
    if legacy_stream_dir.exists():
        for src_file in legacy_stream_dir.glob("dbg_*.jsonl"):
            sid = src_file.stem
            dst_file = paths["qa_sessions_dir"] / sid / "debug_stream.jsonl"
            dst_file.parent.mkdir(parents=True, exist_ok=True)
            if not dst_file.exists():
                shutil.copy2(src_file, dst_file)
                counts["streams"] += 1

    legacy_schema_file = Path("data/schemas/event_schemas.json")
    target_schema_file = paths["schema_file"]
    try:
        src_schema = {}
        dst_schema = {}
        if legacy_schema_file.exists():
            src_raw = json.loads(legacy_schema_file.read_text(encoding="utf-8"))
            src_schema = src_raw.get("events", src_raw) if isinstance(src_raw, dict) else {}
        if target_schema_file.exists():
            dst_raw = json.loads(target_schema_file.read_text(encoding="utf-8"))
            dst_schema = dst_raw.get("events", dst_raw) if isinstance(dst_raw, dict) else {}
        if isinstance(src_schema, dict) and src_schema:
            merged = dict(dst_schema) if isinstance(dst_schema, dict) else {}
            for k, v in src_schema.items():
                if k not in merged:
                    merged[k] = v
                    counts["schemas"] += 1
            target_schema_file.parent.mkdir(parents=True, exist_ok=True)
            target_schema_file.write_text(
                json.dumps({"events": merged, "updated_at": datetime.now(tz=ZoneInfo("UTC")).isoformat()}, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
    except Exception:
        pass

    return counts


def _bootstrap_dotenv_value(key: str) -> str:
    dotenv_path = BASE_DIR / ".env"
    if not dotenv_path.exists():
        return ""
    try:
        for raw in dotenv_path.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            if k.strip() == key:
                return v.strip().strip("'").strip('"')
    except Exception:
        return ""
    return ""


DISPLAY_TZ_NAME = str(os.getenv("QA_DISPLAY_TZ", "")).strip() or _bootstrap_dotenv_value("QA_DISPLAY_TZ")
if DISPLAY_TZ_NAME:
    try:
        LOCAL_TZ = ZoneInfo(DISPLAY_TZ_NAME)
    except Exception:
        LOCAL_TZ = datetime.now().astimezone().tzinfo
        DISPLAY_TZ_NAME = str(LOCAL_TZ)
else:
    LOCAL_TZ = datetime.now().astimezone().tzinfo
    DISPLAY_TZ_NAME = str(LOCAL_TZ)


def parse_csv_list(text: str) -> List[str]:
    return [item.strip() for item in text.split(",") if item.strip()]


def parse_required_params(text: str) -> Dict[str, List[str]]:
    mapping: Dict[str, List[str]] = {}
    for line in text.strip().splitlines():
        if not line.strip() or ":" not in line:
            continue
        event, params = line.split(":", 1)
        mapping[event.strip()] = parse_csv_list(params)
    return mapping


def load_event_schemas(path: Path | None = None) -> Dict[str, Dict[str, object]]:
    if path is None:
        path = get_active_schema_path()
    if not path.exists():
        return {}
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}

    src = raw.get("events", raw) if isinstance(raw, dict) else {}
    if not isinstance(src, dict):
        return {}

    out: Dict[str, Dict[str, object]] = {}
    for event_name, schema in src.items():
        if not isinstance(schema, dict):
            continue
        name = str(event_name).strip()
        if not name:
            continue
        required = [str(v).strip() for v in schema.get("required", []) if str(v).strip()]
        optional = [str(v).strip() for v in schema.get("optional", []) if str(v).strip()]
        out[name] = {
            "required": list(dict.fromkeys(required)),
            "optional": list(dict.fromkeys(optional)),
            "updated_at": str(schema.get("updated_at", "")).strip(),
        }
    return out


def save_event_schemas(schemas: Dict[str, Dict[str, object]], path: Path | None = None) -> None:
    if path is None:
        path = get_active_schema_path()
    payload = {
        "events": schemas,
        "updated_at": datetime.now(tz=ZoneInfo("UTC")).isoformat(),
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _is_required_flag(value: object) -> bool:
    return str(value or "").strip().lower() in {"y", "yes", "true", "1", "req", "required", "필수"}


def _clean_definition_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value).replace("\ufeff", "").replace("\xa0", " ").strip()


def _preserve_definition_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value).replace("\ufeff", "")


def _normalize_definition_header(value: object) -> str:
    text = _clean_definition_cell(value)
    if not text:
        return ""
    text = normalize_param_name(text).strip().lower()
    alias_map = {
        "event_name": "event_name",
        "이벤트명": "event_name",
        "영역순서": "no",
        "no": "no",
        "action": "action",
        "object": "object",
        "description": "description",
        "설명작성란": "description",
        "state": "state",
        "로그상태": "state",
        "poc": "poc",
        "구분": "poc",
        "event_적합_검사": "event_fit",
        "event적합검사": "event_fit",
        "적합검사": "event_fit",
        "ecommerce": "ecommerce",
        "영역에대한actionobject선택란": "",
        "event_설정정보": "",
    }
    return alias_map.get(text, text)


class DefinitionUploadError(ValueError):
    def __init__(self, message: str, debug_info: Dict[str, object] | None = None):
        super().__init__(message)
        self.debug_info = debug_info or {}


def _analyze_definition_matrix(raw_df: pd.DataFrame) -> Dict[str, object]:
    analysis: Dict[str, object] = {
        "raw_rows": int(len(raw_df)) if isinstance(raw_df, pd.DataFrame) else 0,
        "raw_cols": int(len(raw_df.columns)) if isinstance(raw_df, pd.DataFrame) else 0,
        "event_header_row_index": None,
        "param_header_row_index": None,
        "event_row_count": 0,
        "sample_event_names": [],
        "validation_reason": "",
        "schemas": {},
    }
    if raw_df.empty:
        analysis["validation_reason"] = "헤더 탐지 실패: raw DataFrame이 비어 있습니다."
        return analysis

    normalized_rows: List[List[str]] = []
    for _, row in raw_df.iterrows():
        normalized_rows.append([_normalize_definition_header(value) for value in row.tolist()])

    event_header_idx = -1
    param_header_idx = -1
    best_event_score = -1
    for idx, row in enumerate(normalized_rows):
        row_set = {cell for cell in row if cell}
        event_score = 0
        if "event_name" in row_set:
            event_score += 4
        for hint in ["action", "object", "no", "description", "event_fit", "ecommerce", "state", "poc"]:
            if hint in row_set:
                event_score += 1
        if event_score > best_event_score:
            best_event_score = event_score
            event_header_idx = idx if "event_name" in row_set else event_header_idx
        if idx <= event_header_idx:
            continue
        param_key_hits = len(
            [
                cell
                for cell in row_set
                if cell
                in {
                    "section_name",
                    "section_index",
                    "section_title",
                    "store_type",
                    "index",
                    "ab_bucket",
                    "ab_props",
                    "ds_uri",
                    "brand_id",
                    "content_id",
                    "content_name",
                    "content_type",
                    "content_no",
                    "banner_id",
                    "popup_title",
                    "button_id",
                    "button_name",
                    "category_id",
                    "category_name",
                    "filter_type",
                    "filter_value",
                    "liked_item_id",
                    "item_list_name",
                    "item_ab_bucket",
                    "item_ab_props",
                    "item_ds_uri",
                    "item_extra_info",
                    "item_used_grade",
                    "spc_code",
                    "ab_probs",
                    "landing_url",
                    "page_id",
                    "page_link",
                    "extra_info",
                }
            ]
        )
        if event_header_idx >= 0 and param_key_hits >= 2 and param_header_idx < 0:
            param_header_idx = idx
            break

    analysis["event_header_row_index"] = event_header_idx if event_header_idx >= 0 else None
    analysis["param_header_row_index"] = param_header_idx if param_header_idx >= 0 else None

    if event_header_idx < 0:
        analysis["validation_reason"] = "헤더 탐지 실패: event_name 컬럼이 있는 이벤트 헤더 행을 찾지 못했습니다."
        return analysis
    if param_header_idx < 0 or param_header_idx <= event_header_idx:
        analysis["validation_reason"] = "헤더 탐지 실패: 이벤트 헤더 아래에서 매개변수 매트릭스 헤더를 찾지 못했습니다."
        return analysis

    event_header_map: Dict[str, int] = {}
    for col_idx, cell in enumerate(normalized_rows[event_header_idx]):
        if cell and cell not in event_header_map:
            event_header_map[cell] = col_idx
    if "event_name" not in event_header_map:
        analysis["validation_reason"] = "헤더 탐지 실패: 선택된 이벤트 헤더 행에 event_name 컬럼이 없습니다."
        return analysis

    event_rows: List[Dict[str, str]] = []
    for row_idx in range(event_header_idx + 1, param_header_idx):
        row_values = raw_df.iloc[row_idx].tolist()
        event_name = _clean_definition_cell(row_values[event_header_map["event_name"]]) if event_header_map["event_name"] < len(row_values) else ""
        if not event_name:
            continue
        event_rows.append(
            {
                "event_name": event_name,
                "raw_row_index": str(row_idx),
                "no": _clean_definition_cell(row_values[event_header_map["no"]]) if "no" in event_header_map and event_header_map["no"] < len(row_values) else "",
                "action": _clean_definition_cell(row_values[event_header_map["action"]]) if "action" in event_header_map and event_header_map["action"] < len(row_values) else "",
                "object": _clean_definition_cell(row_values[event_header_map["object"]]) if "object" in event_header_map and event_header_map["object"] < len(row_values) else "",
                "description": _clean_definition_cell(row_values[event_header_map["description"]]) if "description" in event_header_map and event_header_map["description"] < len(row_values) else "",
                "state": _clean_definition_cell(row_values[event_header_map["state"]]) if "state" in event_header_map and event_header_map["state"] < len(row_values) else "",
                "poc": _clean_definition_cell(row_values[event_header_map["poc"]]) if "poc" in event_header_map and event_header_map["poc"] < len(row_values) else "",
                "event_fit": _clean_definition_cell(row_values[event_header_map["event_fit"]]) if "event_fit" in event_header_map and event_header_map["event_fit"] < len(row_values) else "",
                "ecommerce": _clean_definition_cell(row_values[event_header_map["ecommerce"]]) if "ecommerce" in event_header_map and event_header_map["ecommerce"] < len(row_values) else "",
            }
        )
    analysis["event_row_count"] = len(event_rows)
    analysis["sample_event_names"] = [row["event_name"] for row in event_rows[:5]]
    if not event_rows:
        analysis["validation_reason"] = "이벤트 행 추출 실패: event_name 아래에서 실제 이벤트 값을 1개도 찾지 못했습니다."
        return analysis

    param_headers = normalized_rows[param_header_idx]
    param_rows: List[List[object]] = []
    for row_idx in range(param_header_idx + 1, len(raw_df)):
        row_values = raw_df.iloc[row_idx].tolist()
        if not any(_clean_definition_cell(value) for value in row_values):
            continue
        param_rows.append(row_values)

    out: Dict[str, Dict[str, object]] = {}
    now_iso = datetime.now(tz=ZoneInfo("UTC")).isoformat()
    for idx, event_row in enumerate(event_rows):
        event_name = event_row["event_name"]
        param_values = param_rows[idx] if idx < len(param_rows) else []
        required_params: List[str] = []
        param_value_map: Dict[str, str] = {}
        for col_idx, header_name in enumerate(param_headers):
            if not header_name:
                continue
            if col_idx >= len(param_values):
                continue
            raw_cell_text = _preserve_definition_cell_text(param_values[col_idx])
            cell_text = _clean_definition_cell(param_values[col_idx])
            if not cell_text:
                continue
            param_name = normalize_param_name(header_name)
            if not param_name or param_name in required_params:
                continue
            required_params.append(param_name)
            param_value_map[param_name] = raw_cell_text
        event_meta = dict(event_row)
        event_meta["param_values"] = param_value_map
        _merge_definition_schema_entry(
            out,
            event_name,
            required_params,
            [],
            list(required_params),
            event_meta,
            now_iso,
        )
    analysis["schemas"] = out
    if not out:
        analysis["validation_reason"] = "validation 실패: 이벤트 스키마를 1개도 만들지 못했습니다."
        return analysis
    analysis["validation_reason"] = "ok"
    return analysis


def _extract_definition_matrix_schemas(raw_df: pd.DataFrame) -> Dict[str, Dict[str, object]]:
    return _analyze_definition_matrix(raw_df).get("schemas", {}) if isinstance(raw_df, pd.DataFrame) else {}


def _merge_definition_schema_entry(
    out: Dict[str, Dict[str, object]],
    event_name: str,
    required: List[str],
    optional: List[str],
    parameter_columns: List[str],
    meta: Dict[str, object],
    updated_at: str,
) -> None:
    key = str(event_name).strip()
    if not key:
        return
    entry = out.setdefault(
        key,
        {
            "required": [],
            "optional": [],
            "parameter_columns": [],
            "updated_at": updated_at,
            "meta": dict(meta),
            "cases": [],
        },
    )
    for bucket_name, values in {
        "required": required,
        "optional": optional,
        "parameter_columns": parameter_columns,
    }.items():
        bucket = entry.setdefault(bucket_name, [])
        for value in values:
            value_text = str(value).strip()
            if value_text and value_text not in bucket:
                bucket.append(value_text)
    entry["updated_at"] = updated_at
    entry.setdefault("meta", dict(meta))
    cases = entry.setdefault("cases", [])
    case_no = str(meta.get("no", "")).strip()
    case_desc = str(meta.get("description", "")).strip()
    raw_row_index = str(meta.get("raw_row_index", "")).strip()
    case_id = f"{key}::{case_no or raw_row_index or len(cases) + 1}::{case_desc or '-'}"
    case_entry = next((case for case in cases if str(case.get("case_id", "")).strip() == case_id), None)
    if not isinstance(case_entry, dict):
        case_entry = {
            "case_id": case_id,
            "required": [],
            "optional": [],
            "parameter_columns": [],
            "meta": dict(meta),
        }
        cases.append(case_entry)
    for bucket_name, values in {
        "required": required,
        "optional": optional,
        "parameter_columns": parameter_columns,
    }.items():
        bucket = case_entry.setdefault(bucket_name, [])
        for value in values:
            value_text = str(value).strip()
            if value_text and value_text not in bucket:
                bucket.append(value_text)
    if not isinstance(case_entry.get("meta"), dict) or not case_entry.get("meta"):
        case_entry["meta"] = dict(meta)


def _parse_definition_schemas_with_fallback(
    structured_df: pd.DataFrame | None = None,
    raw_df: pd.DataFrame | None = None,
) -> Tuple[Dict[str, Dict[str, object]], Dict[str, object]]:
    debug_info: Dict[str, object] = {
        "file_read_stage": "ok",
        "structured_rows": int(len(structured_df)) if isinstance(structured_df, pd.DataFrame) else 0,
        "structured_cols": int(len(structured_df.columns)) if isinstance(structured_df, pd.DataFrame) else 0,
        "raw_rows": int(len(raw_df)) if isinstance(raw_df, pd.DataFrame) else 0,
        "raw_cols": int(len(raw_df.columns)) if isinstance(raw_df, pd.DataFrame) else 0,
        "fallback_used": False,
        "event_header_row_index": None,
        "param_header_row_index": None,
        "parsed_event_count": 0,
        "sample_event_names": [],
        "failure_stage": "",
        "final_validation_reason": "",
    }
    schemas: Dict[str, Dict[str, object]] = {}
    if isinstance(structured_df, pd.DataFrame) and not structured_df.empty:
        try:
            schemas = parse_uploaded_definition_schemas(structured_df)
        except Exception:
            schemas = {}
    if schemas:
        debug_info["parsed_event_count"] = len(schemas)
        debug_info["sample_event_names"] = list(schemas.keys())[:5]
        debug_info["final_validation_reason"] = "ok"
        return schemas, debug_info
    debug_info["fallback_used"] = True
    if isinstance(raw_df, pd.DataFrame) and not raw_df.empty:
        analysis = _analyze_definition_matrix(raw_df)
        debug_info["event_header_row_index"] = analysis.get("event_header_row_index")
        debug_info["param_header_row_index"] = analysis.get("param_header_row_index")
        debug_info["parsed_event_count"] = len(analysis.get("schemas", {}))
        debug_info["sample_event_names"] = analysis.get("sample_event_names", [])
        debug_info["final_validation_reason"] = str(analysis.get("validation_reason", "")).strip()
        if analysis.get("schemas"):
            return analysis.get("schemas", {}), debug_info
        debug_info["failure_stage"] = (
            "헤더 탐지 실패"
            if "헤더 탐지 실패" in debug_info["final_validation_reason"]
            else "이벤트 행 추출 실패"
            if "이벤트 행 추출 실패" in debug_info["final_validation_reason"]
            else "validation 실패"
        )
        return {}, debug_info
    debug_info["failure_stage"] = "파일 읽기 실패"
    debug_info["final_validation_reason"] = "raw DataFrame을 만들지 못했습니다."
    return {}, debug_info


def build_schema_catalog_df(schemas: Dict[str, Dict[str, object]]) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    for event_name, schema in sorted((schemas or {}).items()):
        if not isinstance(schema, dict):
            continue
        cases = schema.get("cases", []) if isinstance(schema.get("cases", []), list) else []
        if not cases:
            cases = [
                {
                    "case_id": f"{str(event_name).strip()}::default",
                    "required": [str(v).strip() for v in schema.get("required", []) if str(v).strip()],
                    "optional": [str(v).strip() for v in schema.get("optional", []) if str(v).strip()],
                    "parameter_columns": [str(v).strip() for v in schema.get("parameter_columns", []) if str(v).strip()],
                    "meta": schema.get("meta", {}) if isinstance(schema.get("meta", {}), dict) else {},
                }
            ]
        for case in cases:
            meta = case.get("meta", {}) if isinstance(case.get("meta", {}), dict) else {}
            required = [str(v).strip() for v in case.get("required", []) if str(v).strip()]
            optional = [str(v).strip() for v in case.get("optional", []) if str(v).strip()]
            parameter_columns = [str(v).strip() for v in case.get("parameter_columns", []) if str(v).strip()]
            rows.append(
                {
                    "case_id": str(case.get("case_id", "")).strip(),
                    "event_name": str(event_name).strip(),
                    "no": str(meta.get("no", "")).strip(),
                    "action": str(meta.get("action", "")).strip(),
                    "object": str(meta.get("object", "")).strip(),
                    "description": str(meta.get("description", "")).strip(),
                    "required_count": len(required),
                    "optional_count": len(optional),
                    "parameter_count": len(parameter_columns),
                    "required": ", ".join(required),
                    "optional": ", ".join(optional),
                    "parameter_columns": ", ".join(parameter_columns),
                    "updated_at": str(schema.get("updated_at", "")).strip(),
                }
            )
    return pd.DataFrame(
        rows,
        columns=[
            "case_id",
            "event_name",
            "no",
            "action",
            "object",
            "description",
            "required_count",
            "optional_count",
            "parameter_count",
            "required",
            "optional",
            "parameter_columns",
            "updated_at",
        ],
    )


def build_definition_template_rows() -> List[List[str]]:
    event_headers = ["state", "POC", "no", "action", "object", "event_name", "description", "event 적합 검사", "ecommerce"]
    param_headers = [
        "section_name",
        "section_index",
        "section_title",
        "store_type",
        "index",
        "ab_bucket",
        "ab_props",
        "ds_uri",
        "brand_id",
        "content_id",
        "content_name",
        "content_type",
        "content_no",
        "banner_id",
        "popup_title",
        "button_id",
        "button_name",
        "category_id",
        "category_name",
        "filter_type",
        "filter_value",
        "liked_item_id",
        "item_list_name",
        "item_ab_bucket",
        "item_ab_props",
        "item_ds_uri",
        "item_extra_info",
        "item_used_grade",
        "spc_code",
        "ab_probs",
        "landing_url",
        "page_id",
        "page_link",
        "extra_info",
    ]
    width = max(len(event_headers), len(param_headers))

    def pad(values: List[str]) -> List[str]:
        return values + [""] * (width - len(values))

    rows: List[List[str]] = [
        pad(["EVENT - 설정정보"]),
        pad(["로그 상태", "POC 구분", "영역 순서", "action", "object", "이벤트명", "설명", "적합검사", "ecommerce 유무"]),
        pad(event_headers),
        pad(["상용", "APP(Android, iOS) / MW", "1", "click", "content", "click_content", "스토어홈 > 추천판 > 빅배너 클릭", "Y", "N"]),
        pad(["상용", "APP(Android, iOS) / MW", "2", "impression", "content", "impression_content", "스토어홈 > 추천판 > 빅배너 노출", "Y", "N"]),
        pad(["상용", "APP(Android, iOS) / MW", "3", "click", "button", "click_button", "스토어홈 > 추천판 > 전체보기 버튼 클릭", "Y", "N"]),
        pad([]),
        pad(["PARAMETER - 매개변수 매트릭스"]),
        pad(param_headers),
        pad([
            "main_display_big_banner", "{index}", "{title}", "musinsa", "{index}", "{ab_bucket}", "{ab_props}", "{ds_uri}", "",
            "1234", "{image}", "banner", "", "{banner_id}", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "{ab_probs}",
            "{{landing_url}}", "/main/musinsa/recommend", "{{page_link}}", ""
        ]),
        pad([
            "main_display_big_banner", "{index}", "{title}", "musinsa", "{index}", "{ab_bucket}", "{ab_props}", "{ds_uri}", "",
            "1234", "{image}", "banner", "", "{banner_id}", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "{ab_probs}",
            "{{landing_url}}", "/main/musinsa/recommend", "{{page_link}}", ""
        ]),
        pad([
            "main_display_big_banner", "{index}", "", "musinsa", "", "", "", "", "", "", "", "", "", "", "", "all", "전체보기 버튼",
            "", "", "", "", "", "", "", "", "", "", "", "", "", "{{landing_url}}", "/main/musinsa/recommend", "{{page_link}}", ""
        ]),
    ]
    return rows


def build_definition_template_csv_bytes() -> bytes:
    rows = build_definition_template_rows()
    df = pd.DataFrame(rows)
    return df.to_csv(index=False, header=False).encode("utf-8-sig")


def build_definition_template_xlsx_bytes() -> bytes:
    rows = build_definition_template_rows()
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, index=False, header=False, sheet_name="definition_template")
    return out.getvalue()


DEFINITION_MATCH_HINT_PRIORITY = [
    "button_id",
    "button_name",
    "section_name",
    "section_title",
    "content_id",
    "content_name",
    "content_type",
    "banner_id",
    "category_id",
    "category_name",
    "filter_type",
    "filter_value",
    "page_id",
]


def _is_definition_placeholder_value(value: object) -> bool:
    text = str(value or "").strip()
    if not text:
        return True
    if "{" in text or "}" in text:
        return True
    lowered = text.lower()
    return lowered in {"-", "na", "n/a", "null", "none"}


def _build_definition_identification_from_values(values: Dict[str, object]) -> str:
    if not isinstance(values, dict):
        return "-"
    bits: List[str] = []
    for key in DEFINITION_MATCH_HINT_PRIORITY:
        value_text = str(values.get(key, "")).strip()
        if not value_text or _is_definition_placeholder_value(value_text):
            continue
        if key == "filter_value" and any(bit.startswith("filter_type=") for bit in bits):
            bits.append(f"{key}={value_text}")
            continue
        bits.append(f"{key}={value_text}")
        if len(bits) >= 3:
            break
    return ", ".join(bits) if bits else "-"


def _iter_definition_cases(schemas: Dict[str, Dict[str, object]]) -> List[Dict[str, object]]:
    rows: List[Dict[str, object]] = []
    for event_name, schema in (schemas or {}).items():
        if not isinstance(schema, dict):
            continue
        cases = schema.get("cases", []) if isinstance(schema.get("cases", []), list) else []
        if not cases:
            cases = [
                {
                    "case_id": f"{str(event_name).strip()}::default",
                    "required": [str(v).strip() for v in schema.get("required", []) if str(v).strip()],
                    "optional": [str(v).strip() for v in schema.get("optional", []) if str(v).strip()],
                    "parameter_columns": [str(v).strip() for v in schema.get("parameter_columns", []) if str(v).strip()],
                    "meta": schema.get("meta", {}) if isinstance(schema.get("meta", {}), dict) else {},
                }
            ]
        for idx, case in enumerate(cases, start=1):
            meta = case.get("meta", {}) if isinstance(case.get("meta", {}), dict) else {}
            param_values = meta.get("param_values", {}) if isinstance(meta.get("param_values", {}), dict) else {}
            rows.append(
                {
                    "definition_row_id": str(case.get("case_id", "")).strip() or f"{str(event_name).strip()}::{idx}",
                    "event_name": str(event_name).strip(),
                    "no": str(meta.get("no", "")).strip(),
                    "action": str(meta.get("action", "")).strip(),
                    "object": str(meta.get("object", "")).strip(),
                    "description": str(meta.get("description", "")).strip(),
                    "required": [str(v).strip() for v in case.get("required", []) if str(v).strip()],
                    "optional": [str(v).strip() for v in case.get("optional", []) if str(v).strip()],
                    "parameter_columns": [str(v).strip() for v in case.get("parameter_columns", []) if str(v).strip()],
                    "param_values": {
                        str(k).strip(): _preserve_definition_cell_text(v)
                        for k, v in param_values.items()
                        if str(k).strip() and _preserve_definition_cell_text(v) != ""
                    },
                }
            )
    return rows


def _tokenize_match_text(text: object) -> set[str]:
    raw = str(text or "").strip().lower()
    if not raw:
        return set()
    tokens = re.split(r"[^0-9a-zA-Z가-힣]+", raw)
    return {token for token in tokens if token and len(token) > 1}


def _build_rt_event_identification(row: Dict[str, object]) -> str:
    values = {}
    params = row.get("전체 파라미터", {})
    if isinstance(params, dict):
        values.update({str(k).strip(): _to_text_value(v) for k, v in params.items()})
    for key in ["target_id", "section_name", "page_id"]:
        value_text = str(row.get(key, "")).strip()
        if value_text and key not in values:
            values[key] = value_text
    return _build_definition_identification_from_values(values)


def _normalize_match_value(value: object) -> str:
    text = _to_text_value(value)
    return "" if _is_missing_like_value(text) else text


def _is_auxiliary_common_key_value(key: str, value: object) -> bool:
    key_text = str(key or "").strip().lower()
    value_text = str(value or "").strip().lower()
    return (
        key_text in {"button_id", "button_name"}
        and value_text in {"more", "더보기"}
    )


def _build_definition_case_match_hints(case: Dict[str, object]) -> Dict[str, str]:
    param_values = case.get("param_values", {}) if isinstance(case.get("param_values", {}), dict) else {}
    hints: Dict[str, str] = {}
    for key in DEFINITION_MATCH_HINT_PRIORITY:
        value_text = str(param_values.get(key, "")).strip()
        if not value_text or _is_definition_placeholder_value(value_text):
            continue
        hints[key] = value_text
    return hints


def _build_rt_row_match_values(row: Dict[str, object]) -> Dict[str, str]:
    params = row.get("전체 파라미터", {})
    params = params if isinstance(params, dict) else {}
    row_values: Dict[str, str] = {str(k).strip(): _to_text_value(v) for k, v in params.items()}
    event_name = str(row.get("이벤트", "")).strip()
    action, obj = _infer_action_object(event_name)
    row_values.setdefault("action", action)
    row_values.setdefault("object", obj)
    for key in ["target_id", "section_name", "page_id"]:
        value_text = str(row.get(key, "")).strip()
        if value_text:
            row_values.setdefault(key, value_text)
    row_values.setdefault("description_context", " ".join(
        [
            str(row.get("target_id", "")).strip(),
            str(row.get("section_name", "")).strip(),
            str(row.get("page_id", "")).strip(),
            str(row.get("custom_parameters", "")).strip(),
            str(row.get("selector", "")).strip(),
        ]
    ).strip())
    return row_values


def _score_definition_case_match(case: Dict[str, object], row: Dict[str, object], duplicate_case_count: int) -> Dict[str, object]:
    row_values = _build_rt_row_match_values(row)
    hints = _build_definition_case_match_hints(case)
    expected_page_id = _normalize_match_value((case.get("param_values", {}) or {}).get("page_id", "") or case.get("page_id", ""))
    expected_section_name = _normalize_match_value((case.get("param_values", {}) or {}).get("section_name", "") or case.get("section_name", ""))
    actual_page_id = _normalize_match_value(row_values.get("page_id", ""))
    actual_section_name = _normalize_match_value(row_values.get("section_name", ""))
    if not expected_page_id or not expected_section_name:
        return {
            "accepted": False,
            "score": -999.0,
            "exact_matches": 0,
            "description_overlap": 0,
            "reason": "wrong_row_match",
        }
    if actual_page_id != expected_page_id or actual_section_name != expected_section_name:
        return {
            "accepted": False,
            "score": -999.0,
            "exact_matches": 0,
            "description_overlap": 0,
            "reason": "wrong_row_match",
        }
    exact_matches = 2
    mismatch_count = 0
    score = 10.0
    for key, weight in {
        "section_index": 4.0,
        "action": 2.0,
        "object": 2.0,
    }.items():
        expected = _normalize_match_value(str(case.get(key, "")).strip() or str((case.get("param_values", {}) or {}).get(key, "")).strip())
        if not expected or _is_definition_placeholder_value(expected):
            continue
        actual = _normalize_match_value(row_values.get(key, ""))
        if actual and actual == expected:
            exact_matches += 1
            score += weight
        elif actual:
            mismatch_count += 1
            score -= weight
    for key, expected in hints.items():
        expected = _normalize_match_value(expected)
        actual = _normalize_match_value(row_values.get(key, ""))
        if _is_auxiliary_common_key_value(key, expected):
            if actual and actual == expected:
                score += 0.5
            continue
        if actual and actual == expected:
            exact_matches += 1
            score += 4.0
        elif actual:
            mismatch_count += 1
            score -= 4.0
    desc_tokens = _tokenize_match_text(case.get("description", ""))
    row_tokens = _tokenize_match_text(row_values.get("description_context", ""))
    overlap = len(desc_tokens & row_tokens)
    if overlap:
        score += min(3.0, float(overlap))
    accepted = True
    reason = "ok"
    if duplicate_case_count > 1 and exact_matches <= 2 and overlap == 0:
        accepted = False
        reason = "wrong_row_match"
    elif mismatch_count > 0 and score <= 0:
        accepted = False
        reason = "wrong_row_match"
    return {
        "accepted": accepted,
        "score": score,
        "exact_matches": exact_matches,
        "description_overlap": overlap,
        "reason": reason,
    }


def validate_definition_case_params(case: Dict[str, object], params: Dict[str, object]) -> Dict[str, object]:
    param_map = params if isinstance(params, dict) else {}
    required = [str(v).strip() for v in case.get("required", []) if str(v).strip()]
    optional = [str(v).strip() for v in case.get("optional", []) if str(v).strip()]
    parameter_columns = [str(v).strip() for v in case.get("parameter_columns", []) if str(v).strip()]
    expected_keys = list(dict.fromkeys(required + optional + parameter_columns))
    missing = []
    for key in (required or parameter_columns):
        value_text = _normalize_match_value(param_map.get(key, ""))
        if key not in param_map or not value_text:
            missing.append(key)
    value_mismatch = []
    param_values = case.get("param_values", {}) if isinstance(case.get("param_values", {}), dict) else {}
    for key, expected_raw in param_values.items():
        expected = _normalize_match_value(expected_raw)
        if not expected or _is_definition_placeholder_value(expected) or _is_auxiliary_common_key_value(key, expected):
            continue
        actual = _normalize_match_value(param_map.get(key, ""))
        if actual and actual != expected:
            value_mismatch.append(f"{key}={expected}")
    extra = []
    for key, value in param_map.items():
        key_text = str(key).strip()
        if not key_text or _is_hidden_default_param_key(key_text):
            continue
        if expected_keys and key_text not in expected_keys:
            extra.append(key_text)
    if missing:
        return {
            "status": "FAIL",
            "failure_type": "missing_parameter",
            "message": f"missing_parameter: 필수 누락 {len(missing)}개",
            "missing_parameters": ", ".join(missing),
            "extra_parameters": ", ".join(extra),
        }
    if value_mismatch:
        return {
            "status": "FAIL",
            "failure_type": "value_mismatch",
            "message": f"value_mismatch: 기대값 불일치 {len(value_mismatch)}개",
            "missing_parameters": "-",
            "extra_parameters": ", ".join(extra),
        }
    if extra:
        return {
            "status": "WARN",
            "failure_type": "extra_parameter",
            "message": f"정의서 외 파라미터 {len(extra)}개",
            "missing_parameters": "-",
            "extra_parameters": ", ".join(extra),
        }
    return {
        "status": "PASS",
        "failure_type": "ok",
        "message": "정의서 row 충족",
        "missing_parameters": "-",
        "extra_parameters": "-",
    }


def build_definition_validation_frames(
    rt_events: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]],
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    summary_cols = [
        "definition_row_id",
        "no",
        "event_name",
        "description",
        "page_id",
        "section_name",
        "expected_identification",
        "matched_capture_count",
        "result",
        "failure_type",
        "failure_reason",
        "missing_parameters",
        "extra_parameters",
    ]
    matches_cols = [
        "definition_row_id",
        "no",
        "event_name",
        "description",
        "captured_event_no",
        "matched_at",
        "target_id",
        "section_name",
        "page_id",
        "matched_identification",
        "screenshot_file",
        "match_score",
        "result",
        "failure_type",
        "failure_reason",
        "missing_parameters",
        "extra_parameters",
    ]
    unmatched_cols = [
        "event_name",
        "captured_at",
        "target_id",
        "section_name",
        "page_id",
        "matched_identification",
        "failure_type",
        "reason",
    ]
    cases = _iter_definition_cases(schemas)
    if not cases:
        return (
            pd.DataFrame(columns=summary_cols),
            pd.DataFrame(columns=matches_cols),
            pd.DataFrame(columns=unmatched_cols),
        )
    cases_by_event: Dict[str, List[Dict[str, object]]] = {}
    for case in cases:
        cases_by_event.setdefault(str(case.get("event_name", "")).strip(), []).append(case)
    event_rows = rt_events.to_dict("records") if isinstance(rt_events, pd.DataFrame) and not rt_events.empty else []
    match_rows: List[Dict[str, object]] = []
    unmatched_rows: List[Dict[str, object]] = []
    for row in event_rows:
        event_name = str(row.get("이벤트", "")).strip()
        candidates = cases_by_event.get(event_name, [])
        if not candidates:
            unmatched_rows.append(
                {
                    "event_name": event_name,
                    "captured_at": str(row.get("시간", "-")).strip() or "-",
                    "target_id": str(row.get("target_id", "-")).strip() or "-",
                    "section_name": str(row.get("section_name", "-")).strip() or "-",
                    "page_id": str(row.get("page_id", "-")).strip() or "-",
                    "matched_identification": _build_rt_event_identification(row),
                    "failure_type": "wrong_row_match",
                    "reason": "정의서에 동일 event_name 없음",
                }
            )
            continue
        scored = [
            (
                case,
                _score_definition_case_match(case, row, len(candidates)),
            )
            for case in candidates
        ]
        scored.sort(key=lambda item: (float(item[1].get("score", 0.0)), int(item[1].get("exact_matches", 0))), reverse=True)
        best_case, best_score = scored[0]
        if not bool(best_score.get("accepted", False)):
            unmatched_rows.append(
                {
                    "event_name": event_name,
                    "captured_at": str(row.get("시간", "-")).strip() or "-",
                    "target_id": str(row.get("target_id", "-")).strip() or "-",
                    "section_name": str(row.get("section_name", "-")).strip() or "-",
                    "page_id": str(row.get("page_id", "-")).strip() or "-",
                    "matched_identification": _build_rt_event_identification(row),
                    "failure_type": str(best_score.get("reason", "wrong_row_match")).strip() or "wrong_row_match",
                    "reason": "정의서 내 동일 이벤트명은 있으나 위치/맥락 매칭 실패",
                }
            )
            continue
        params = row.get("전체 파라미터", {})
        params = params if isinstance(params, dict) else {}
        validation = validate_definition_case_params(best_case, params)
        match_rows.append(
            {
                "definition_row_id": str(best_case.get("definition_row_id", "")).strip(),
                "no": str(best_case.get("no", "")).strip(),
                "event_name": event_name,
                "description": str(best_case.get("description", "")).strip(),
                "captured_event_no": str(row.get("annotation_no", row.get("event no", ""))).strip() or "-",
                "matched_at": str(row.get("시간", "-")).strip() or "-",
                "target_id": str(row.get("target_id", "-")).strip() or "-",
                "section_name": str(row.get("section_name", "-")).strip() or "-",
                "page_id": str(row.get("page_id", "-")).strip() or "-",
                "matched_identification": _build_rt_event_identification(row),
                "screenshot_file": str(row.get("screenshot_path", row.get("selector_screenshot", "-"))).strip() or "-",
                "match_score": round(float(best_score.get("score", 0.0)), 2),
                "result": str(validation.get("status", "WARN")).upper(),
                "failure_type": str(validation.get("failure_type", "ok")).strip() or "ok",
                "failure_reason": str(validation.get("message", "-")).strip() or "-",
                "missing_parameters": str(validation.get("missing_parameters", "-")).strip() or "-",
                "extra_parameters": str(validation.get("extra_parameters", "-")).strip() or "-",
            }
        )
    matches_df = pd.DataFrame(match_rows, columns=matches_cols)
    summary_rows: List[Dict[str, object]] = []
    for case in cases:
        definition_row_id = str(case.get("definition_row_id", "")).strip()
        case_matches = matches_df[matches_df["definition_row_id"].astype(str) == definition_row_id].copy() if not matches_df.empty else pd.DataFrame()
        expected_identification = _build_definition_identification_from_values(case.get("param_values", {}))
        page_id = str((case.get("param_values", {}) or {}).get("page_id", "")).strip() or "-"
        section_name = str((case.get("param_values", {}) or {}).get("section_name", "")).strip() or "-"
        if case_matches.empty:
            summary_rows.append(
                {
                    "definition_row_id": definition_row_id,
                    "no": str(case.get("no", "")).strip(),
                    "event_name": str(case.get("event_name", "")).strip(),
                    "description": str(case.get("description", "")).strip(),
                    "page_id": page_id,
                    "section_name": section_name,
                    "expected_identification": expected_identification,
                    "matched_capture_count": 0,
                    "result": "FAIL",
                    "failure_type": "wrong_row_match",
                    "failure_reason": "wrong_row_match: 정의서 row 미충족",
                    "missing_parameters": ", ".join(case.get("required", []) or case.get("parameter_columns", [])) or "-",
                    "extra_parameters": "-",
                }
            )
            continue
        ranked = {"PASS": 3, "WARN": 2, "FAIL": 1}
        case_matches["rank"] = case_matches["result"].astype(str).map(ranked).fillna(0)
        case_matches = case_matches.sort_values(by=["rank", "match_score"], ascending=[False, False])
        best_match = case_matches.iloc[0].to_dict()
        summary_rows.append(
            {
                "definition_row_id": definition_row_id,
                "no": str(case.get("no", "")).strip(),
                "event_name": str(case.get("event_name", "")).strip(),
                "description": str(case.get("description", "")).strip(),
                "page_id": page_id,
                "section_name": section_name,
                "expected_identification": expected_identification,
                "matched_capture_count": int(len(case_matches)),
                "result": str(best_match.get("result", "WARN")).upper(),
                "failure_type": str(best_match.get("failure_type", "ok")).strip() or "ok",
                "failure_reason": str(best_match.get("failure_reason", "-")).strip() or "-",
                "missing_parameters": str(best_match.get("missing_parameters", "-")).strip() or "-",
                "extra_parameters": str(best_match.get("extra_parameters", "-")).strip() or "-",
            }
        )
    return (
        pd.DataFrame(summary_rows, columns=summary_cols),
        matches_df,
        pd.DataFrame(unmatched_rows, columns=unmatched_cols),
    )


def build_definition_case_validation_df(rt_events: pd.DataFrame, schemas: Dict[str, Dict[str, object]]) -> pd.DataFrame:
    summary_df, _, _ = build_definition_validation_frames(rt_events, schemas)
    return summary_df


def build_definition_runtime_hints(schemas: Dict[str, Dict[str, object]]) -> Dict[str, List[str]]:
    hints = {
        "target_ids": [],
        "section_names": [],
        "page_ids": [],
        "text_hints": [],
    }
    seen = {key: set() for key in hints.keys()}
    for case in _iter_definition_cases(schemas):
        param_values = case.get("param_values", {}) if isinstance(case.get("param_values", {}), dict) else {}
        for key in ["button_id", "banner_id", "content_id", "category_id", "filter_value"]:
            value_text = str(param_values.get(key, "")).strip()
            if value_text and not _is_definition_placeholder_value(value_text) and value_text not in seen["target_ids"]:
                seen["target_ids"].add(value_text)
                hints["target_ids"].append(value_text)
        for key in ["section_name"]:
            value_text = str(param_values.get(key, "")).strip()
            if value_text and not _is_definition_placeholder_value(value_text) and value_text not in seen["section_names"]:
                seen["section_names"].add(value_text)
                hints["section_names"].append(value_text)
        for key in ["page_id"]:
            value_text = str(param_values.get(key, "")).strip()
            if value_text and not _is_definition_placeholder_value(value_text) and value_text not in seen["page_ids"]:
                seen["page_ids"].add(value_text)
                hints["page_ids"].append(value_text)
        for key in ["button_name", "content_name", "category_name", "section_title"]:
            value_text = str(param_values.get(key, "")).strip()
            if value_text and not _is_definition_placeholder_value(value_text) and value_text not in seen["text_hints"]:
                seen["text_hints"].add(value_text)
                hints["text_hints"].append(value_text)
    return hints


def parse_uploaded_definition_schemas(df: pd.DataFrame) -> Dict[str, Dict[str, object]]:
    if df.empty:
        return {}

    out: Dict[str, Dict[str, object]] = {}
    now_iso = datetime.now(tz=ZoneInfo("UTC")).isoformat()
    columns = {str(col).strip().lower(): col for col in df.columns}

    if "event name" in columns and "params.key" in columns:
        event_col = columns["event name"]
        param_col = columns["params.key"]
        req_col = columns.get("req")
        current_event = ""
        current_meta: Dict[str, object] = {}
        for row_index, raw_row in enumerate(df.to_dict("records")):
            event_name = str(raw_row.get(event_col, "")).strip()
            if event_name:
                current_event = event_name
                current_meta = {"raw_row_index": str(row_index)}
                _merge_definition_schema_entry(out, current_event, [], [], [], current_meta, now_iso)
            if not current_event:
                continue
            param_name = normalize_param_name(str(raw_row.get(param_col, "")).strip())
            if not param_name:
                continue
            bucket = "required" if _is_required_flag(raw_row.get(req_col, "")) else "optional"
            _merge_definition_schema_entry(
                out,
                current_event,
                [param_name] if bucket == "required" else [],
                [param_name] if bucket == "optional" else [],
                [param_name],
                current_meta,
                now_iso,
            )
        return out

    if "event_name" in columns and ("required" in columns or "optional" in columns):
        event_col = columns["event_name"]
        required_col = columns.get("required")
        optional_col = columns.get("optional")
        for row_index, raw_row in enumerate(df.to_dict("records")):
            event_name = str(raw_row.get(event_col, "")).strip()
            if not event_name:
                continue
            required = [
                normalize_param_name(v)
                for v in parse_csv_list(str(raw_row.get(required_col, "")))
            ] if required_col else []
            optional = [
                normalize_param_name(v)
                for v in parse_csv_list(str(raw_row.get(optional_col, "")))
            ] if optional_col else []
            meta = {"raw_row_index": str(row_index)}
            _merge_definition_schema_entry(
                out,
                event_name,
                [v for v in required if v],
                [v for v in optional if v and v not in required],
                [v for v in list(dict.fromkeys([*required, *optional])) if v],
                meta,
                now_iso,
            )
        return out

    if "event_name" in columns and "parameter" in columns:
        event_col = columns["event_name"]
        param_col = columns["parameter"]
        req_col = columns.get("req") or columns.get("required")
        for row_index, raw_row in enumerate(df.to_dict("records")):
            event_name = str(raw_row.get(event_col, "")).strip()
            param_name = normalize_param_name(str(raw_row.get(param_col, "")).strip())
            if not event_name or not param_name:
                continue
            bucket = "required" if _is_required_flag(raw_row.get(req_col, "")) else "optional"
            _merge_definition_schema_entry(
                out,
                event_name,
                [param_name] if bucket == "required" else [],
                [param_name] if bucket == "optional" else [],
                [param_name],
                {"raw_row_index": str(row_index)},
                now_iso,
            )
        return out

    raise ValueError("지원 형식이 아닙니다. event_definition.csv 또는 event_name/required/optional 컬럼 형식을 사용하세요.")


def _read_uploaded_csv_candidates(raw_bytes: bytes) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[str, object]]:
    encodings = ["utf-8-sig", "utf-8", "cp949", "euc-kr"]
    debug: Dict[str, object] = {"read_attempts": [], "selected_encoding": "", "failure_stage": "", "final_validation_reason": ""}
    structured_df = pd.DataFrame()
    raw_df = pd.DataFrame()
    last_error = ""
    for encoding in encodings:
        try:
            structured_df = pd.read_csv(io.BytesIO(raw_bytes), encoding=encoding)
            debug["read_attempts"].append(f"structured:{encoding}:ok")
            debug["selected_encoding"] = encoding
            break
        except Exception as exc:
            last_error = str(exc)
            debug["read_attempts"].append(f"structured:{encoding}:fail")
    for encoding in encodings:
        try:
            raw_df = pd.read_csv(io.BytesIO(raw_bytes), encoding=encoding, header=None)
            debug["read_attempts"].append(f"raw:{encoding}:ok")
            if not debug["selected_encoding"]:
                debug["selected_encoding"] = encoding
            break
        except Exception as exc:
            last_error = str(exc)
            debug["read_attempts"].append(f"raw:{encoding}:fail")
    if structured_df.empty and raw_df.empty and last_error:
        debug["failure_stage"] = "파일 읽기 실패"
        debug["final_validation_reason"] = f"CSV 읽기 실패: {last_error}"
    return structured_df, raw_df, debug


def load_uploaded_definition_schemas(uploaded_file) -> Tuple[Dict[str, Dict[str, object]], pd.DataFrame, Dict[str, object]]:
    if uploaded_file is None:
        return {}, pd.DataFrame(columns=["event_name", "no", "action", "object", "required_count", "optional_count", "required", "optional", "updated_at"]), {}
    raw_bytes = uploaded_file.getvalue()
    if not raw_bytes:
        return {}, pd.DataFrame(columns=["event_name", "no", "action", "object", "required_count", "optional_count", "required", "optional", "updated_at"]), {
            "failure_stage": "파일 읽기 실패",
            "final_validation_reason": "업로드 파일이 비어 있습니다.",
        }
    suffix = Path(str(getattr(uploaded_file, "name", "") or "")).suffix.lower()
    schemas: Dict[str, Dict[str, object]] = {}
    debug_info: Dict[str, object] = {
        "file_name": str(getattr(uploaded_file, "name", "")).strip(),
        "file_type": suffix,
        "fallback_used": False,
        "failure_stage": "",
        "final_validation_reason": "",
    }
    if suffix in {".xlsx", ".xls"}:
        try:
            excel = pd.ExcelFile(io.BytesIO(raw_bytes))
        except Exception as exc:
            debug_info["failure_stage"] = "파일 읽기 실패"
            debug_info["final_validation_reason"] = f"엑셀 파일을 열지 못했습니다: {exc}"
            return {}, pd.DataFrame(columns=["event_name", "no", "action", "object", "required_count", "optional_count", "required", "optional", "updated_at"]), debug_info
        sheet_debugs: List[Dict[str, object]] = []
        for sheet_name in excel.sheet_names:
            source_df = pd.read_excel(excel, sheet_name=sheet_name)
            raw_sheet_df = pd.read_excel(excel, sheet_name=sheet_name, header=None)
            sheet_schemas, sheet_debug = _parse_definition_schemas_with_fallback(source_df, raw_sheet_df)
            sheet_debug["sheet_name"] = sheet_name
            sheet_debugs.append(sheet_debug)
            if sheet_schemas:
                schemas.update(sheet_schemas)
                debug_info.update(sheet_debug)
                break
        debug_info["sheet_debugs"] = sheet_debugs
    else:
        source_df, raw_df, read_debug = _read_uploaded_csv_candidates(raw_bytes)
        debug_info.update(read_debug)
        schemas, parse_debug = _parse_definition_schemas_with_fallback(source_df, raw_df)
        debug_info.update(parse_debug)
    if not schemas:
        failure_stage = str(debug_info.get("failure_stage", "")).strip() or "validation 실패"
        final_reason = str(debug_info.get("final_validation_reason", "")).strip() or "유효한 이벤트 테이블을 만들지 못했습니다."
        raise DefinitionUploadError(f"{failure_stage}: {final_reason}", debug_info=debug_info)
    debug_info["parsed_event_count"] = len(schemas)
    debug_info["sample_event_names"] = list(schemas.keys())[:5]
    debug_info["final_validation_reason"] = "ok"
    return schemas, build_schema_catalog_df(schemas), debug_info


def infer_event_schema_from_rows(
    rt_events: pd.DataFrame,
    event_name: str,
    include_system_params: bool = False,
) -> Dict[str, object]:
    if rt_events.empty:
        return {"required": [], "optional": [], "sample_size": 0, "excluded_system_params": 0}
    subset = rt_events[rt_events["이벤트"].astype(str) == str(event_name)].copy()
    if subset.empty:
        return {"required": [], "optional": [], "sample_size": 0, "excluded_system_params": 0}

    nonmissing_count: Dict[str, int] = {}
    seen_count: Dict[str, int] = {}
    total = int(len(subset))
    excluded_system_params = 0

    for _, row in subset.iterrows():
        params = row.get("전체 파라미터")
        if not isinstance(params, dict):
            continue
        row_seen: set[str] = set()
        row_nonmissing: set[str] = set()
        for key, value in params.items():
            key_text = str(key).strip()
            if not key_text:
                continue
            if _is_system_param_key(key_text) and (not include_system_params):
                excluded_system_params += 1
                continue
            row_seen.add(key_text)
            if not _is_missing_like_value(_to_text_value(value)):
                row_nonmissing.add(key_text)
        for key in row_seen:
            seen_count[key] = seen_count.get(key, 0) + 1
        for key in row_nonmissing:
            nonmissing_count[key] = nonmissing_count.get(key, 0) + 1

    all_keys = sorted(seen_count.keys())
    required = [k for k in all_keys if nonmissing_count.get(k, 0) >= total]
    optional = [k for k in all_keys if k not in required]
    return {
        "required": required,
        "optional": optional,
        "sample_size": total,
        "excluded_system_params": int(excluded_system_params),
    }


def validate_schema_for_event(
    event_name: str,
    params: Dict[str, object],
    schemas: Dict[str, Dict[str, object]],
) -> Dict[str, object]:
    schema = schemas.get(str(event_name).strip(), {})
    required = [str(v).strip() for v in schema.get("required", []) if str(v).strip()] if isinstance(schema, dict) else []
    optional = [str(v).strip() for v in schema.get("optional", []) if str(v).strip()] if isinstance(schema, dict) else []
    parameter_columns = _schema_parameter_columns(event_name, schemas)
    if not required and not optional and parameter_columns:
        required = list(parameter_columns)
    checks: List[Dict[str, object]] = []

    if not required and not optional:
        return {
            "status": "WARN",
            "message": "기준표 없음(스키마 없음)",
            "checks": checks,
        }

    missing_required: List[str] = []
    for key in required:
        value = _to_text_value(params.get(key, ""))
        ok = (key in params) and (not _is_missing_like_value(value))
        if not ok:
            missing_required.append(key)
        checks.append({"param": key, "ok": ok, "value": value or "-", "group": "required"})

    for key in optional:
        value = _to_text_value(params.get(key, ""))
        checks.append(
            {
                "param": key,
                "ok": (key in params) and (not _is_missing_like_value(value)),
                "value": value or "-",
                "group": "optional",
            }
        )

    if missing_required:
        return {
            "status": "FAIL",
            "message": f"필수 누락 {len(missing_required)}개",
            "checks": checks,
        }
    return {
        "status": "PASS",
        "message": "필수 파라미터 정상",
        "checks": checks,
    }


def summarize_schema_validation(rt_events: pd.DataFrame, schemas: Dict[str, Dict[str, object]]) -> Dict[str, int]:
    out = {"PASS": 0, "WARN": 0, "FAIL": 0}
    if rt_events.empty:
        return out
    for row in rt_events.to_dict("records"):
        params = row.get("전체 파라미터", {})
        if not isinstance(params, dict):
            params = {}
        status = str(
            validate_schema_for_event(
                event_name=str(row.get("이벤트", "")),
                params=params,
                schemas=schemas,
            ).get("status", "WARN")
        ).upper()
        if status not in out:
            status = "WARN"
        out[status] += 1
    return out


def build_schema_validation_by_event(rt_events: pd.DataFrame, schemas: Dict[str, Dict[str, object]]) -> pd.DataFrame:
    if rt_events.empty:
        return pd.DataFrame(columns=["event_name", "PASS", "WARN", "FAIL", "total"])
    rows: List[Dict[str, object]] = []
    for row in rt_events.to_dict("records"):
        event_name = str(row.get("이벤트", "")).strip()
        if not event_name:
            continue
        params = row.get("전체 파라미터", {})
        if not isinstance(params, dict):
            params = {}
        status = str(
            validate_schema_for_event(
                event_name=event_name,
                params=params,
                schemas=schemas,
            ).get("status", "WARN")
        ).upper()
        if status not in {"PASS", "WARN", "FAIL"}:
            status = "WARN"
        rows.append({"event_name": event_name, "status": status, "count": 1})
    if not rows:
        return pd.DataFrame(columns=["event_name", "PASS", "WARN", "FAIL", "total"])
    work = pd.DataFrame(rows)
    pivot = (
        work.pivot_table(index="event_name", columns="status", values="count", aggfunc="sum", fill_value=0)
        .reset_index()
    )
    for col in ["PASS", "WARN", "FAIL"]:
        if col not in pivot.columns:
            pivot[col] = 0
    pivot["total"] = pivot["PASS"] + pivot["WARN"] + pivot["FAIL"]
    return pivot[["event_name", "PASS", "WARN", "FAIL", "total"]].sort_values(
        by=["FAIL", "WARN", "total", "event_name"], ascending=[False, False, False, True]
    )


def build_issue_summary(rt_events: pd.DataFrame, schemas: Dict[str, Dict[str, object]]) -> pd.DataFrame:
    if rt_events.empty:
        return pd.DataFrame(columns=["event_name", "status", "reason", "count"])

    issue_rows: List[Dict[str, object]] = []
    for row in rt_events.to_dict("records"):
        event_name = str(row.get("이벤트", "")).strip()
        if not event_name:
            continue

        reasons: List[str] = []
        event_status = str(row.get("상태", "OK")).upper()
        if event_status == "ERROR":
            reasons.append("이벤트 판정 오류(ERROR)")
        elif event_status == "WARN":
            reasons.append("이벤트 판정 경고(WARN)")

        missing_group = row.get("값없음 그룹", {})
        if isinstance(missing_group, dict) and missing_group:
            reasons.append(f"값 비어있음 {len(missing_group)}개")
        suspicious_group = row.get("의심 그룹", {})
        if isinstance(suspicious_group, dict) and suspicious_group:
            reasons.append(f"항상 동일한 값 {len(suspicious_group)}개")

        params = row.get("전체 파라미터", {})
        if not isinstance(params, dict):
            params = {}
        schema_check = validate_schema_for_event(event_name=event_name, params=params, schemas=schemas)
        schema_status = str(schema_check.get("status", "WARN")).upper()
        if schema_status in {"WARN", "FAIL"}:
            reasons.append(f"스키마: {schema_check.get('message', '-')}")

        if not reasons:
            continue

        merged_status = "FAIL" if ("FAIL" in {schema_status} or event_status == "ERROR") else "WARN"
        issue_rows.append(
            {
                "event_name": event_name,
                "status": merged_status,
                "reason": " | ".join(dict.fromkeys(reasons)),
                "count": 1,
            }
        )

    if not issue_rows:
        return pd.DataFrame(columns=["event_name", "status", "reason", "count"])

    out = pd.DataFrame(issue_rows)
    out = (
        out.groupby(["event_name", "status", "reason"], as_index=False)["count"]
        .sum()
        .sort_values(by=["status", "count", "event_name"], ascending=[True, False, True])
    )
    return out


def build_tracking_plan_df(rt_events: pd.DataFrame, schemas: Dict[str, Dict[str, object]]) -> pd.DataFrame:
    event_counts: Dict[str, int] = {}
    if not rt_events.empty:
        for name in rt_events["이벤트"].astype(str).tolist():
            key = str(name).strip()
            if not key:
                continue
            event_counts[key] = event_counts.get(key, 0) + 1

    all_event_names = sorted(list(dict.fromkeys(list(schemas.keys()) + list(event_counts.keys()))))
    rows: List[Dict[str, object]] = []
    for event_name in all_event_names:
        schema = schemas.get(event_name, {})
        required = [str(v).strip() for v in schema.get("required", []) if str(v).strip()] if isinstance(schema, dict) else []
        optional = [str(v).strip() for v in schema.get("optional", []) if str(v).strip()] if isinstance(schema, dict) else []
        rows.append(
            {
                "event_name": event_name,
                "triggered_count": int(event_counts.get(event_name, 0)),
                "schema_status": "ready" if (required or optional) else "missing",
                "required": ", ".join(required),
                "optional": ", ".join(optional),
            }
        )
    return pd.DataFrame(rows)


def dataframe_to_csv_bytes(df: pd.DataFrame) -> bytes:
    if df.empty:
        return b""
    return df.to_csv(index=False).encode("utf-8-sig")


def dataframes_to_excel_bytes(sheets: Dict[str, pd.DataFrame]) -> bytes:
    if not sheets:
        return b""
    out = io.BytesIO()
    try:
        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            for sheet_name, df in sheets.items():
                safe_sheet = str(sheet_name or "Sheet1")[:31]
                df_to_write = df if isinstance(df, pd.DataFrame) else pd.DataFrame()
                if not df_to_write.empty:
                    df_to_write = df_to_write.copy()
                    df_to_write.columns = [to_excel_safe_text(col) for col in df_to_write.columns]
                    for col in df_to_write.columns:
                        series = df_to_write[col]
                        if pd.api.types.is_object_dtype(series) or pd.api.types.is_string_dtype(series):
                            df_to_write[col] = series.map(lambda value: to_excel_safe_text(value) if value is not None else "")
                df_to_write.to_excel(writer, index=False, sheet_name=safe_sheet)
        return out.getvalue()
    except Exception as exc:
        fallback = io.BytesIO()
        with pd.ExcelWriter(fallback, engine="openpyxl") as writer:
            pd.DataFrame(
                [
                    {"status": "export_failed", "reason": to_excel_safe_text(exc)},
                    {"status": "sheets", "reason": ", ".join([to_excel_safe_text(name) for name in sheets.keys()])},
                ]
            ).to_excel(writer, index=False, sheet_name="ExportError")
        return fallback.getvalue()


def _parse_pipe_kv_text(text: object) -> Dict[str, str]:
    out: Dict[str, str] = {}
    for chunk in str(text or "").split(" | "):
        if "=" not in chunk:
            continue
        key, value = chunk.split("=", 1)
        key_text = str(key).strip()
        value_text = str(value).strip()
        if key_text and key_text not in out:
            out[key_text] = value_text
    return out


def _build_event_definition_lookup(event_definition_df: pd.DataFrame) -> Dict[str, Dict[str, object]]:
    lookup: Dict[str, Dict[str, object]] = {}
    current_event_no = ""
    for row in event_definition_df.to_dict("records") if isinstance(event_definition_df, pd.DataFrame) and not event_definition_df.empty else []:
        row_no = str(row.get("no", "")).strip()
        if row_no:
            current_event_no = row_no
            lookup[current_event_no] = {
                "event_name": str(row.get("event name", "")).strip(),
                "action": str(row.get("action", "")).strip(),
                "object": str(row.get("object", "")).strip(),
                "rows": [],
            }
        if not current_event_no or current_event_no not in lookup:
            continue
        param_key = str(row.get("params.key", "")).strip()
        if not param_key:
            continue
        lookup[current_event_no]["rows"].append(
            {
                "params.key": param_key,
                "req": str(row.get("req", "")).strip(),
                "params.value": str(row.get("params.value", "")).strip(),
            }
        )
    return lookup


def build_qa_review_excel_bytes(
    event_definition_df: pd.DataFrame,
    event_detail_df: pd.DataFrame,
    definition_validation_summary_df: pd.DataFrame | None = None,
    bundle_scope: str = "custom_event_review",
) -> bytes:
    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)
    if str(bundle_scope or "").strip() == "definition_validation" and isinstance(definition_validation_summary_df, pd.DataFrame):
        ws = wb.create_sheet("Definition Summary")
        ws["A1"] = to_excel_safe_text("Definition Validation Summary")
        ws["A1"].font = Font(bold=True, size=14)
        if definition_validation_summary_df.empty:
            ws["A3"] = to_excel_safe_text("데이터가 없습니다.")
        else:
            cols = [col for col in ["definition_row_id", "no", "event_name", "description", "section_name", "matched_capture_count", "result", "failure_reason", "missing_parameters", "extra_parameters"] if col in definition_validation_summary_df.columns]
            for col_idx, col_name in enumerate(cols, start=1):
                cell = ws.cell(row=3, column=col_idx, value=to_excel_safe_text(col_name))
                cell.font = Font(bold=True)
            for row_idx, (_, row) in enumerate(definition_validation_summary_df[cols].fillna("-").iterrows(), start=4):
                for col_idx, col_name in enumerate(cols, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=to_excel_safe_text(row.get(col_name, "-")))
    if event_detail_df.empty:
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    detail_lookup = _build_event_definition_lookup(event_definition_df)
    work = event_detail_df.copy()
    work["sheet_group"] = work["screenshot_file"].astype(str).replace("-", "").str.strip()
    work.loc[work["sheet_group"] == "", "sheet_group"] = work["event no"].astype(str).radd("event_")
    grouped = list(work.groupby("sheet_group", sort=False))
    for sheet_index, (_, group_df) in enumerate(grouped, start=1):
        first = group_df.iloc[0].to_dict()
        event_label = str(first.get("custom_event", first.get("event_name", "review"))).strip() or "review"
        sheet_name = f"{sheet_index:02d}_{event_label}"[:31]
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = to_excel_safe_text("QA Review Sheet")
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = to_excel_safe_text(f"group={sheet_index} | event_count={len(group_df)}")
        cursor = 4
        screenshot_rel = str(first.get("screenshot_file", "")).strip()
        screenshot_abs = BASE_DIR / screenshot_rel if screenshot_rel and screenshot_rel != "-" else None
        if screenshot_abs and screenshot_abs.exists():
            try:
                img = OpenpyxlImage(str(screenshot_abs))
                if getattr(img, "width", 0) and img.width > 900:
                    ratio = 900.0 / float(img.width)
                    img.width = int(img.width * ratio)
                    img.height = int(img.height * ratio)
                ws.add_image(img, f"A{cursor}")
                cursor += 24
            except Exception:
                ws[f"A{cursor}"] = to_excel_safe_text(f"image load failed: {screenshot_rel}")
                cursor += 2
        else:
            ws[f"A{cursor}"] = to_excel_safe_text("screenshot not available")
            cursor += 2
        for _, event_row in group_df.iterrows():
            row_dict = event_row.to_dict()
            event_no = str(row_dict.get("event no", "")).strip()
            ws[f"A{cursor}"] = to_excel_safe_text(f"Event #{event_no or '-'}")
            ws[f"A{cursor}"].font = Font(bold=True)
            cursor += 1
            summary_pairs = [
                ("event_name", str(row_dict.get("custom_event", row_dict.get("event_name", "-"))).strip() or "-"),
                ("target_id", str(row_dict.get("target_id", "-")).strip() or "-"),
                ("section_name", str(row_dict.get("section_name", "-")).strip() or "-"),
                ("page_id", str(row_dict.get("page_id", "-")).strip() or "-"),
                ("status", str(row_dict.get("status", "-")).strip() or "-"),
            ]
            for key, value in summary_pairs:
                ws.cell(row=cursor, column=1, value=to_excel_safe_text(key))
                ws.cell(row=cursor, column=2, value=to_excel_safe_text(value))
                cursor += 1
            cursor += 1
            headers = ["event definition", "parameter definition", "actual value", "class-attribute", "status"]
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=cursor, column=col_idx, value=to_excel_safe_text(header))
                cell.font = Font(bold=True)
            cursor += 1
            actual_map = {}
            actual_map.update(_parse_pipe_kv_text(row_dict.get("custom_parameters", "")))
            actual_map.update(_parse_pipe_kv_text(row_dict.get("extra_custom_parameters", "")))
            param_rows = detail_lookup.get(event_no, {}).get("rows", [])
            if not param_rows:
                param_rows = [{"params.key": key, "req": "", "params.value": ""} for key in actual_map.keys()]
            for idx, param_row in enumerate(param_rows):
                param_key = str(param_row.get("params.key", "")).strip()
                actual_value = actual_map.get(param_key, "-")
                param_status = "missing" if _is_missing_like_value(str(actual_value)) else str(row_dict.get("status", "-")).strip() or "-"
                ws.cell(row=cursor, column=1, value=to_excel_safe_text(str(row_dict.get("custom_event", row_dict.get("event_name", "-"))) if idx == 0 else ""))
                definition_text = f"{param_key} ({str(param_row.get('req', '')).strip() or 'N'})"
                expected_text = str(param_row.get("params.value", "")).strip() or "-"
                ws.cell(row=cursor, column=2, value=to_excel_safe_text(f"{definition_text} / expected={expected_text}"))
                ws.cell(row=cursor, column=3, value=to_excel_safe_text(actual_value))
                ws.cell(row=cursor, column=4, value=to_excel_safe_text(str(row_dict.get("class_attribute", "-")).strip() or "-"))
                ws.cell(row=cursor, column=5, value=to_excel_safe_text(param_status))
                cursor += 1
            cursor += 2
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 46
        ws.column_dimensions["C"].width = 32
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 18
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_log_definition_excel_bytes(
    event_definition_df: pd.DataFrame,
    event_detail_df: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]] | None = None,
    definition_matches_df: pd.DataFrame | None = None,
) -> bytes:
    review_df, _, _records, _unmatched = build_log_definition_export_tables(
        event_definition_df=event_definition_df,
        event_detail_df=event_detail_df,
        schemas=schemas,
        definition_matches_df=definition_matches_df,
    )
    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)
    if review_df.empty:
        ws = wb.create_sheet("log_definition")
        ws["A1"] = to_excel_safe_text("생성 가능한 로그정의서 데이터가 없습니다.")
        ws["A2"] = to_excel_safe_text("이벤트 매칭 성공 건이 있으면 파라미터 FAIL/WARN이어도 로그정의서를 생성합니다.")
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    detail_map: Dict[str, Dict[str, object]] = {}
    if isinstance(event_detail_df, pd.DataFrame) and not event_detail_df.empty:
        for row in event_detail_df.to_dict("records"):
            event_no = str(row.get("event no", "")).strip()
            if event_no and event_no not in detail_map:
                detail_map[event_no] = row

    grouped_rows: List[Tuple[str, pd.DataFrame]] = []
    work = review_df.copy()
    work["sheet_group"] = work["screenshot_file"].astype(str).replace("-", "").str.strip()
    current_group = ""
    group_keys: List[str] = []
    for row in work.to_dict("records"):
        row_no = str(row.get("no", "")).strip()
        if row_no:
            current_group = str(row.get("screenshot_file", "")).strip() or f"event_{row_no}"
            if current_group not in group_keys:
                group_keys.append(current_group)
        row["sheet_group"] = current_group or str(row.get("screenshot_file", "")).strip() or "group"
    work = pd.DataFrame(work.to_dict("records"))
    for group_key in group_keys or sorted(work["sheet_group"].astype(str).unique().tolist()):
        grouped_rows.append((group_key, work[work["sheet_group"].astype(str) == str(group_key)].copy()))

    created_sheet = False
    for idx, (group_key, group_df) in enumerate(grouped_rows, start=1):
        if group_df.empty:
            continue
        event_start_rows = group_df[group_df["no"].astype(str).str.strip() != ""].copy()
        first_no = str(event_start_rows.iloc[0]["no"]).strip() if not event_start_rows.empty else ""
        detail = detail_map.get(first_no, {})
        first_group_row = group_df.iloc[0].to_dict() if not group_df.empty else {}
        section_name = str(detail.get("section_name", "")).strip() or str(first_group_row.get("section_name", "")).strip() or "review"
        sheet_name = to_excel_safe_text(f"{section_name}_{idx}")[:31] or f"review_{idx}"
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = to_excel_safe_text("로그정의서 생성")
        ws["A1"].font = Font(bold=True, size=14)
        screenshot_rel = str(detail.get("screenshot_file", "")).strip() or str(first_group_row.get("screenshot_file", "")).strip()
        cursor = 3
        screenshot_abs = BASE_DIR / screenshot_rel if screenshot_rel and screenshot_rel != "-" else None
        if screenshot_abs and screenshot_abs.exists():
            try:
                img = OpenpyxlImage(str(screenshot_abs))
                if getattr(img, "width", 0) and img.width > 900:
                    ratio = 900.0 / float(img.width)
                    img.width = int(img.width * ratio)
                    img.height = int(img.height * ratio)
                ws.add_image(img, f"A{cursor}")
                cursor += 24
            except Exception:
                ws[f"A{cursor}"] = to_excel_safe_text(f"image load failed: {screenshot_rel}")
                cursor += 2
        else:
            ws[f"A{cursor}"] = to_excel_safe_text("screenshot not available")
            cursor += 2
        headers = [
            "action",
            "object",
            "event_name",
            "params.key",
            "required",
            "params.value (as-is)",
            "params.value",
            "class_attribute",
            "qa_status",
        ]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=cursor, column=col_idx, value=to_excel_safe_text(header))
            cell.font = Font(bold=True)
        cursor += 1
        for _, row in group_df.iterrows():
            for col_idx, header in enumerate(headers, start=1):
                ws.cell(row=cursor, column=col_idx, value=to_excel_safe_text(row.get(header, "")))
            cursor += 1
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 8
        ws.column_dimensions["F"].width = 28
        ws.column_dimensions["G"].width = 28
        ws.column_dimensions["H"].width = 22
        ws.column_dimensions["I"].width = 14
        created_sheet = True
    if not created_sheet:
        ws = wb.create_sheet("log_definition")
        ws["A1"] = to_excel_safe_text("생성 가능한 로그정의서 시트가 없습니다.")
        ws["A2"] = to_excel_safe_text("정의서 매칭 성공 건이 없거나 그룹화 가능한 screenshot 정보가 없습니다.")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _sanitize_export_filename_part(text: object, default: str = "item") -> str:
    raw = str(text or "").strip().lower()
    safe = re.sub(r"[^0-9a-zA-Z가-힣]+", "_", raw).strip("_")
    return safe[:48] or default


def build_annotation_export_relpath(no: object, identification: object) -> str:
    no_text = str(no or "").strip() or "0"
    ident_text = _sanitize_export_filename_part(identification, default="event")
    return f"annotations/{no_text}_{ident_text}.png"


def normalize_expected_param_value(raw_value: object) -> str:
    text = str(raw_value or "").strip()
    if not text:
        return ""
    if ">>" in text:
        return str(text.split(">>")[-1]).strip()
    return text


def _build_definition_case_lookup(schemas: Dict[str, Dict[str, object]]) -> Dict[str, Dict[str, object]]:
    return {
        str(case.get("definition_row_id", "")).strip(): case
        for case in _iter_definition_cases(schemas)
        if str(case.get("definition_row_id", "")).strip()
    }


def _build_captured_event_case_map(
    definition_matches_df: pd.DataFrame | None,
    schemas: Dict[str, Dict[str, object]] | None,
) -> Dict[str, Dict[str, object]]:
    if not isinstance(definition_matches_df, pd.DataFrame) or definition_matches_df.empty:
        return {}
    case_lookup = _build_definition_case_lookup(schemas or {})
    if not case_lookup:
        return {}
    ranked = definition_matches_df.copy()
    rank_map = {"PASS": 3, "WARN": 2, "FAIL": 1}
    ranked["rank"] = ranked.get("result", pd.Series(dtype=str)).astype(str).map(rank_map).fillna(0)
    ranked["match_score"] = pd.to_numeric(ranked.get("match_score", 0), errors="coerce").fillna(0)
    ranked = ranked.sort_values(by=["rank", "match_score"], ascending=[False, False])
    out: Dict[str, Dict[str, object]] = {}
    for row in ranked.to_dict("records"):
        captured_event_no = str(row.get("captured_event_no", "")).strip()
        definition_row_id = str(row.get("definition_row_id", "")).strip()
        if not captured_event_no or captured_event_no in out:
            continue
        matched_case = case_lookup.get(definition_row_id)
        if matched_case:
            out[captured_event_no] = matched_case
    return out


def _build_detail_param_map(detail_row: Dict[str, object]) -> Dict[str, str]:
    out: Dict[str, str] = {}
    out.update(_parse_pipe_kv_text(detail_row.get("custom_parameters", "")))
    out.update(_parse_pipe_kv_text(detail_row.get("extra_custom_parameters", "")))
    raw_payload = str(detail_row.get("raw_payload_json", "")).strip()
    if raw_payload:
        try:
            payload_obj = json.loads(raw_payload)
            if isinstance(payload_obj, dict):
                for key, value in payload_obj.items():
                    key_text = str(key).strip()
                    if key_text and key_text not in out:
                        out[key_text] = _to_text_value(value)
        except Exception:
            pass
    return out


def build_log_definition_export_tables(
    event_definition_df: pd.DataFrame,
    event_detail_df: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]] | None = None,
    definition_matches_df: pd.DataFrame | None = None,
) -> Tuple[pd.DataFrame, pd.DataFrame, List[Dict[str, object]], List[Dict[str, object]]]:
    final_columns = [
        "no",
        "이벤트 설명",
        "action",
        "object",
        "event_name",
        "section_name",
        "section_index",
        "params.key",
        "required",
        "params.value (as-is)",
        "params.value",
        "actual_value",
        "qa_status",
        "issue_type",
        "trigger",
        "qa_note",
        "screenshot_file",
        "selector",
        "class_attribute",
    ]
    support_columns = [
        "no",
        "event_name",
        "identification",
        "target_id",
        "section_name",
        "page_id",
        "status",
        "custom_parameters",
        "extra_custom_parameters",
        "raw_payload_json",
        "raw_selector",
        "screenshot_path",
        "class_attribute",
        "bbox_x",
        "bbox_y",
        "bbox_width",
        "bbox_height",
    ]
    if event_definition_df.empty or event_detail_df.empty:
        return pd.DataFrame(columns=final_columns), pd.DataFrame(columns=support_columns), [], []

    detail_map: Dict[str, Dict[str, object]] = {}
    for row in event_detail_df.to_dict("records"):
        event_no = str(row.get("event no", "")).strip()
        if event_no and event_no not in detail_map:
            detail_map[event_no] = row
    case_lookup = _build_definition_case_lookup(schemas or {})
    if isinstance(definition_matches_df, pd.DataFrame) and case_lookup:
        ranked = definition_matches_df.copy()
        rank_map = {"PASS": 3, "WARN": 2, "FAIL": 1}
        ranked["rank"] = ranked.get("result", pd.Series(dtype=str)).astype(str).map(rank_map).fillna(0)
        ranked["match_score"] = pd.to_numeric(ranked.get("match_score", 0), errors="coerce").fillna(0)
        ranked = ranked.sort_values(by=["rank", "match_score"], ascending=[False, False])

        final_rows: List[Dict[str, object]] = []
        support_rows: List[Dict[str, object]] = []
        html_records: List[Dict[str, object]] = []
        unmatched_records: List[Dict[str, object]] = []
        matched_definition_ids: set[str] = set()
        for match_row in ranked.groupby("definition_row_id", as_index=False).first().to_dict("records") if not ranked.empty else []:
            definition_row_id = str(match_row.get("definition_row_id", "")).strip()
            captured_event_no = str(match_row.get("captured_event_no", "")).strip()
            if not definition_row_id or not captured_event_no:
                continue
            case = case_lookup.get(definition_row_id, {})
            detail = detail_map.get(captured_event_no, {})
            if not case or not detail:
                continue
            matched_definition_ids.add(definition_row_id)
            param_values = case.get("param_values", {}) if isinstance(case.get("param_values", {}), dict) else {}
            required_keys = {str(v).strip() for v in case.get("required", []) if str(v).strip()}
            parameter_keys = list(
                dict.fromkeys(
                    [str(v).strip() for v in case.get("required", []) if str(v).strip()]
                    + [str(v).strip() for v in case.get("optional", []) if str(v).strip()]
                    + [str(v).strip() for v in case.get("parameter_columns", []) if str(v).strip()]
                    + [str(v).strip() for v in param_values.keys() if str(v).strip()]
                )
            )
            actual_map = _build_detail_param_map(detail)
            display_no = captured_event_no
            parameter_rows: List[Dict[str, object]] = []
            for param_name in parameter_keys:
                expected_raw = str(param_values.get(param_name, "")).strip()
                expected_value = normalize_expected_param_value(expected_raw) or "-"
                actual_value = str(actual_map.get(param_name, "")).strip()
                issue_type = "ok"
                qa_status = str(match_row.get("result", detail.get("status", "-"))).strip() or "-"
                if param_name in required_keys and _is_missing_like_value(actual_value):
                    qa_status = "FAIL"
                    issue_type = "missing_parameter"
                elif expected_value != "-" and not _is_definition_placeholder_value(expected_value) and actual_value and actual_value != expected_value:
                    qa_status = "WARN"
                    issue_type = "value_mismatch"
                elif _is_missing_like_value(actual_value):
                    qa_status = "WARN"
                    issue_type = "missing_parameter"
                parameter_rows.append(
                    {
                        "params.key": param_name,
                        "req": "Y" if param_name in required_keys else "N",
                        "params.value (as-is)": expected_raw or "-",
                        "params.value": expected_value,
                        "actual_value": actual_value or "-",
                        "status": qa_status,
                        "issue_type": issue_type,
                    }
                )
                final_rows.append(
                    {
                        "no": display_no,
                        "이벤트 설명": str(case.get("description", detail.get("identification", "-"))).strip() or "-",
                        "action": str(case.get("action", detail.get("action", "-"))).strip() or "-",
                        "object": str(case.get("object", detail.get("object", "-"))).strip() or "-",
                        "event_name": str(case.get("event_name", detail.get("custom_event", "-"))).strip() or "-",
                        "section_name": str(detail.get("section_name", "-")).strip() or "-",
                        "section_index": actual_map.get("section_index", "-") or "-",
                        "params.key": param_name,
                        "required": "Y" if param_name in required_keys else "N",
                        "params.value (as-is)": expected_raw or "-",
                        "params.value": expected_value,
                        "actual_value": actual_value or "-",
                        "qa_status": qa_status,
                        "issue_type": issue_type,
                        "trigger": str(match_row.get("matched_identification", detail.get("identification", "-"))).strip() or "-",
                        "qa_note": str(match_row.get("failure_reason", detail.get("status", "-"))).strip() or "-",
                        "screenshot_file": str(detail.get("screenshot_file", "-")).strip() or "-",
                        "selector": str(detail.get("selector", "-")).strip() or "-",
                        "class_attribute": str(detail.get("class_attribute", "-")).strip() or "-",
                    }
                )
            support_rows.append(
                {
                    "no": display_no,
                    "event_name": str(case.get("event_name", detail.get("custom_event", "-"))).strip() or "-",
                    "identification": str(case.get("description", detail.get("identification", "-"))).strip() or "-",
                    "target_id": str(detail.get("target_id", "-")).strip() or "-",
                    "section_name": str(detail.get("section_name", "-")).strip() or "-",
                    "page_id": str(detail.get("page_id", "-")).strip() or "-",
                    "status": str(match_row.get("result", detail.get("status", "-"))).strip() or "-",
                    "custom_parameters": str(detail.get("custom_parameters", "-")).strip() or "-",
                    "extra_custom_parameters": str(detail.get("extra_custom_parameters", "-")).strip() or "-",
                    "raw_payload_json": str(detail.get("raw_payload_json", "{}")).strip() or "{}",
                    "raw_selector": str(detail.get("selector", "-")).strip() or "-",
                    "screenshot_path": str(detail.get("screenshot_file", "-")).strip() or "-",
                    "class_attribute": str(detail.get("class_attribute", "-")).strip() or "-",
                    "bbox_x": detail.get("bbox_x", 0),
                    "bbox_y": detail.get("bbox_y", 0),
                    "bbox_width": detail.get("bbox_width", 0),
                    "bbox_height": detail.get("bbox_height", 0),
                }
            )
            html_records.append(
                {
                    "no": display_no,
                    "event_name": str(case.get("event_name", detail.get("custom_event", "-"))).strip() or "-",
                    "identification": str(case.get("description", detail.get("identification", "-"))).strip() or "-",
                    "record_type": "definition_match",
                    "definition_row_id": definition_row_id,
                    "definition_no": str(case.get("no", "")).strip() or "-",
                    "match_status": str(match_row.get("result", detail.get("status", "-"))).strip() or "-",
                    "target_id": str(detail.get("target_id", "-")).strip() or "-",
                    "section_name": str(detail.get("section_name", "-")).strip() or "-",
                    "page_id": str(detail.get("page_id", "-")).strip() or "-",
                    "status": str(match_row.get("result", detail.get("status", "-"))).strip() or "-",
                    "custom_parameters": str(detail.get("custom_parameters", "-")).strip() or "-",
                    "extra_custom_parameters": str(detail.get("extra_custom_parameters", "-")).strip() or "-",
                    "raw_payload_json": str(detail.get("raw_payload_json", "{}")).strip() or "{}",
                    "raw_selector": str(detail.get("selector", "-")).strip() or "-",
                    "screenshot_file": str(detail.get("screenshot_file", "-")).strip() or "-",
                    "parameter_rows": parameter_rows,
                }
            )
        for definition_row_id, case in case_lookup.items():
            if definition_row_id in matched_definition_ids:
                continue
            unmatched_records.append(
                {
                    "definition_row_id": definition_row_id,
                    "definition_no": str(case.get("no", "")).strip() or "-",
                    "event_name": str(case.get("event_name", "")).strip() or "-",
                    "description": str(case.get("description", "")).strip() or "-",
                    "page_id": str((case.get("param_values", {}) or {}).get("page_id", "")).strip() or "-",
                    "section_name": str((case.get("param_values", {}) or {}).get("section_name", "")).strip() or "-",
                    "match_status": "UNMATCHED",
                }
            )
        return (
            pd.DataFrame(final_rows, columns=final_columns),
            pd.DataFrame(support_rows, columns=support_columns),
            html_records,
            unmatched_records,
        )
    matched_case_by_event_no = _build_captured_event_case_map(definition_matches_df, schemas)

    current_event: Dict[str, object] = {}
    final_rows: List[Dict[str, object]] = []
    support_rows: List[Dict[str, object]] = []
    html_record_map: Dict[str, Dict[str, object]] = {}
    added_support_nos: set[str] = set()
    for row in event_definition_df.to_dict("records"):
        row_no = str(row.get("no", "")).strip()
        if row_no:
            current_event = dict(row)
            current_event["no"] = row_no
        if not current_event:
            continue
        event_no = str(current_event.get("no", "")).strip()
        detail = detail_map.get(event_no, {})
        if not detail:
            continue
        actual_map = _build_detail_param_map(detail)
        matched_case = matched_case_by_event_no.get(event_no, {})
        matched_case_param_values = matched_case.get("param_values", {}) if isinstance(matched_case.get("param_values", {}), dict) else {}
        param_name = str(row.get("params.key", "")).strip()
        if not param_name:
            continue
        expected_raw = (
            str(matched_case_param_values.get(param_name, "")).strip()
            or str(row.get("params.value (as-is)", "")).strip()
            or str(row.get("params.value", "")).strip()
        )
        expected_value = normalize_expected_param_value(expected_raw) or str(row.get("params.value", "")).strip()
        actual_value = actual_map.get(param_name, "")
        issue_type = "ok"
        qa_status = str(detail.get("status", "-")).strip() or "-"
        if _is_missing_like_value(str(actual_value)):
            qa_status = "FAIL"
            issue_type = "missing_parameter"
        elif expected_value and not _is_definition_placeholder_value(expected_value) and actual_value and expected_value != actual_value:
            qa_status = "WARN"
            issue_type = "value_mismatch"
        matched_description = str(matched_case.get("description", "")).strip()
        event_description = matched_description or str(detail.get("identification", detail.get("custom_event", "-"))).strip() or "-"
        required_flag = str(row.get("req", "")).strip()
        if not required_flag and matched_case:
            required_flag = "Y" if param_name in {str(v).strip() for v in matched_case.get("required", []) if str(v).strip()} else "N"
        final_rows.append(
            {
                "no": event_no,
                "이벤트 설명": event_description,
                "action": str(current_event.get("action", detail.get("action", "-"))).strip() or "-",
                "object": str(current_event.get("object", detail.get("object", "-"))).strip() or "-",
                "event_name": str(current_event.get("event name", detail.get("custom_event", "-"))).strip() or "-",
                "section_name": str(detail.get("section_name", "-")).strip() or "-",
                "section_index": actual_map.get("section_index", "-") or "-",
                "params.key": param_name,
                "required": required_flag or "N",
                "params.value (as-is)": expected_raw or "-",
                "params.value": expected_value or "-",
                "actual_value": actual_value or "-",
                "qa_status": qa_status,
                "issue_type": issue_type,
                "trigger": str(detail.get("identification", "-")).strip() or "-",
                "qa_note": str(detail.get("status", "-")).strip() or "-",
                "screenshot_file": str(detail.get("screenshot_file", "-")).strip() or "-",
                "selector": str(detail.get("selector", "-")).strip() or "-",
                "class_attribute": str(detail.get("class_attribute", "-")).strip() or "-",
            }
        )
        if event_no not in added_support_nos:
            support_rows.append(
                {
                    "no": event_no,
                    "event_name": str(detail.get("custom_event", detail.get("event_name", "-"))).strip() or "-",
                    "identification": str(detail.get("identification", "-")).strip() or "-",
                    "target_id": str(detail.get("target_id", "-")).strip() or "-",
                    "section_name": str(detail.get("section_name", "-")).strip() or "-",
                    "page_id": str(detail.get("page_id", "-")).strip() or "-",
                    "status": str(detail.get("status", "-")).strip() or "-",
                    "custom_parameters": str(detail.get("custom_parameters", "-")).strip() or "-",
                    "extra_custom_parameters": str(detail.get("extra_custom_parameters", "-")).strip() or "-",
                    "raw_payload_json": str(detail.get("raw_payload_json", "{}")).strip() or "{}",
                    "raw_selector": str(detail.get("selector", "-")).strip() or "-",
                    "screenshot_path": str(detail.get("screenshot_file", "-")).strip() or "-",
                    "class_attribute": str(detail.get("class_attribute", "-")).strip() or "-",
                    "bbox_x": detail.get("bbox_x", 0),
                    "bbox_y": detail.get("bbox_y", 0),
                    "bbox_width": detail.get("bbox_width", 0),
                    "bbox_height": detail.get("bbox_height", 0),
                }
            )
            html_record_map[event_no] = (
                {
                    "no": event_no,
                    "event_name": str(detail.get("custom_event", detail.get("event_name", "-"))).strip() or "-",
                    "identification": event_description,
                    "record_type": "captured_event",
                    "definition_row_id": "-",
                    "definition_no": "-",
                    "match_status": str(detail.get("status", "-")).strip() or "-",
                    "target_id": str(detail.get("target_id", "-")).strip() or "-",
                    "section_name": str(detail.get("section_name", "-")).strip() or "-",
                    "page_id": str(detail.get("page_id", "-")).strip() or "-",
                    "status": str(detail.get("status", "-")).strip() or "-",
                    "custom_parameters": str(detail.get("custom_parameters", "-")).strip() or "-",
                    "extra_custom_parameters": str(detail.get("extra_custom_parameters", "-")).strip() or "-",
                    "raw_payload_json": str(detail.get("raw_payload_json", "{}")).strip() or "{}",
                    "raw_selector": str(detail.get("selector", "-")).strip() or "-",
                    "screenshot_file": str(detail.get("screenshot_file", "-")).strip() or "-",
                    "parameter_rows": [],
                }
            )
            added_support_nos.add(event_no)
        html_record_map.setdefault(
            event_no,
            {
                "no": event_no,
                "event_name": str(current_event.get("event name", detail.get("custom_event", "-"))).strip() or "-",
                "identification": event_description,
                "record_type": "captured_event",
                "definition_row_id": "-",
                "definition_no": "-",
                "match_status": str(detail.get("status", "-")).strip() or "-",
                "target_id": str(detail.get("target_id", "-")).strip() or "-",
                "section_name": str(detail.get("section_name", "-")).strip() or "-",
                "page_id": str(detail.get("page_id", "-")).strip() or "-",
                "status": str(detail.get("status", "-")).strip() or "-",
                "custom_parameters": str(detail.get("custom_parameters", "-")).strip() or "-",
                "extra_custom_parameters": str(detail.get("extra_custom_parameters", "-")).strip() or "-",
                "raw_payload_json": str(detail.get("raw_payload_json", "{}")).strip() or "{}",
                "raw_selector": str(detail.get("selector", "-")).strip() or "-",
                "screenshot_file": str(detail.get("screenshot_file", "-")).strip() or "-",
                "parameter_rows": [],
            },
        )
        html_record_map[event_no]["parameter_rows"].append(
            {
                "params.key": param_name,
                "req": required_flag or "N",
                "params.value (as-is)": expected_raw or "-",
                "params.value": expected_value or "-",
                "actual_value": actual_value or "-",
                "status": qa_status,
            }
        )
    return (
        pd.DataFrame(final_rows, columns=final_columns),
        pd.DataFrame(support_rows, columns=support_columns),
        list(html_record_map.values()),
        [],
    )


def build_log_definition_review_html(records: List[Dict[str, object]], unmatched_records: List[Dict[str, object]] | None = None) -> str:
    unmatched_records = unmatched_records or []
    if not records and not unmatched_records:
        return """<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>event_definition_review</title></head><body><p>데이터가 없습니다.</p></body></html>"""
    first_no = str(records[0].get("no", "")).strip() if records else "1"
    nav_html = "".join(
        [
            f'<button type="button" class="annotation-chip{" active" if idx == 0 else ""}" data-select-no="{html.escape(str(row.get("no", "")))}">{html.escape(str(row.get("no", "")))}</button>'
            for idx, row in enumerate(records)
        ]
    )
    card_html = "".join(
        [
            (
                lambda param_rows_html, issue_summary_html: f"""
            <article class="event-card{' active' if idx == 0 else ''}" id="card-{html.escape(str(row.get('no', '')))}" data-no="{html.escape(str(row.get('no', '')))}">
              <button type="button" class="card-button" data-select-no="{html.escape(str(row.get('no', '')))}">
                <div class="card-head">
                  <div class="event-number">#{html.escape(str(row.get('no', '')))}</div>
                  <div>
                    <h2>{html.escape(str(row.get('event_name', '-')))}</h2>
                    <p>{html.escape(str(row.get('identification', '-')))}</p>
                  </div>
                  <span class="status">{html.escape(str(row.get('status', '-')))}</span>
                </div>
                <section class="card-section">
                  <h3>매칭 구분</h3>
                  <div class="issue-summary">
                    <span class="issue-chip ok">{html.escape('정의서 매칭 성공 row' if str(row.get('record_type', '')).strip() == 'definition_match' else '캡처 이벤트')}</span>
                    <span class="issue-chip subtle">definition_row_id={html.escape(str(row.get('definition_row_id', '-')))}</span>
                    <span class="issue-chip subtle">definition_no={html.escape(str(row.get('definition_no', '-')))}</span>
                    <span class="issue-chip subtle">match_status={html.escape(str(row.get('match_status', row.get('status', '-'))))}</span>
                  </div>
                </section>
                <section class="card-section">
                  <h3>위치 정보</h3>
                  <dl class="meta-grid">
                    <dt>target_id</dt><dd>{html.escape(str(row.get('target_id', '-')))}</dd>
                    <dt>section_name</dt><dd>{html.escape(str(row.get('section_name', '-')))}</dd>
                    <dt>page_id</dt><dd>{html.escape(str(row.get('page_id', '-')))}</dd>
                  </dl>
                </section>
                <section class="card-section">
                  <h3>핵심 파라미터</h3>
                  <p class="key-param-block">{html.escape(str(row.get('custom_parameters', '-')))}</p>
                </section>
                <section class="card-section">
                  <h3>이슈 요약</h3>
                  <div class="issue-summary">{issue_summary_html}</div>
                </section>
              </button>
              <details><summary>검수 상세</summary>
                <div class="detail-stack">
                  <section class="card-section">
                    <h3>추가 파라미터</h3>
                    <p class="detail-text">{html.escape(str(row.get('extra_custom_parameters', '-')))}</p>
                  </section>
                  <section class="card-section">
                    <h3>raw selector</h3>
                    <p class="detail-text mono">{html.escape(str(row.get('raw_selector', '-')))}</p>
                  </section>
                  <section class="card-section">
                    <h3>screenshot path</h3>
                    <p class="detail-text mono">{html.escape(str(row.get('screenshot_file', '-')))}</p>
                  </section>
                  <section class="card-section">
                    <h3>parameter review</h3>
                    {param_rows_html}
                  </section>
                </div>
              </details>
              <details><summary>raw payload</summary><pre>{html.escape(str(row.get('raw_payload_json', '{}')))}</pre></details>
            </article>
            """
            )(
                (
                    "<table class=\"param-table\"><thead><tr><th>params.key</th><th>req</th><th>params.value (as-is)</th><th>params.value</th><th>actual</th><th>status</th></tr></thead><tbody>"
                    + "".join(
                        [
                            "<tr>"
                            f"<td>{html.escape(str(param_row.get('params.key', '-')))}</td>"
                            f"<td>{html.escape(str(param_row.get('req', '-')))}</td>"
                            f"<td>{html.escape(str(param_row.get('params.value (as-is)', '-')))}</td>"
                            f"<td>{html.escape(str(param_row.get('params.value', '-')))}</td>"
                            f"<td>{html.escape(str(param_row.get('actual_value', '-')))}</td>"
                            f"<td>{html.escape(str(param_row.get('status', '-')))}</td>"
                            "</tr>"
                            for param_row in (row.get("parameter_rows", []) if isinstance(row.get("parameter_rows", []), list) else [])
                        ]
                    )
                    + "</tbody></table>"
                ),
                (
                    (
                        lambda param_rows: (
                            f"<span class=\"issue-chip {'ok' if not [p for p in param_rows if str(p.get('status', '-')).strip().upper() not in {'PASS', '정상', 'OK', '-'}] else 'warn'}\">"
                            + (
                                "정상"
                                if not [p for p in param_rows if str(p.get('status', '-')).strip().upper() not in {'PASS', '정상', 'OK', '-'}]
                                else f"이슈 {len([p for p in param_rows if str(p.get('status', '-')).strip().upper() not in {'PASS', '정상', 'OK', '-'}])}건"
                            )
                            + "</span>"
                            + (
                                "".join(
                                    [
                                        f"<span class=\"issue-chip subtle\">{html.escape(str(p.get('params.key', '-')))}: {html.escape(str(p.get('status', '-')))}</span>"
                                        for p in param_rows
                                        if str(p.get('status', '-')).strip().upper() not in {'PASS', '정상', 'OK', '-'}
                                    ][:4]
                                )
                                or "<span class=\"issue-chip subtle\">누락/불일치 없음</span>"
                            )
                        )
                    )(
                        row.get("parameter_rows", []) if isinstance(row.get("parameter_rows", []), list) else []
                    )
                )
            )
            for idx, row in enumerate(records)
        ]
    )
    event_payload = json.dumps(records, ensure_ascii=False)
    unmatched_html = (
        "<table class=\"param-table\"><thead><tr><th>definition_row_id</th><th>definition no</th><th>event_name</th><th>description</th><th>page_id</th><th>section_name</th><th>status</th></tr></thead><tbody>"
        + "".join(
            [
                "<tr>"
                f"<td>{html.escape(str(row.get('definition_row_id', '-')))}</td>"
                f"<td>{html.escape(str(row.get('definition_no', '-')))}</td>"
                f"<td>{html.escape(str(row.get('event_name', '-')))}</td>"
                f"<td>{html.escape(str(row.get('description', '-')))}</td>"
                f"<td>{html.escape(str(row.get('page_id', '-')))}</td>"
                f"<td>{html.escape(str(row.get('section_name', '-')))}</td>"
                f"<td>{html.escape(str(row.get('match_status', 'UNMATCHED')))}</td>"
                "</tr>"
                for row in unmatched_records
            ]
        )
        + "</tbody></table>"
    ) if unmatched_records else "<p>없음</p>"
    return f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>event_definition_review</title>
  <style>
    body {{ margin: 0; background: #f3efe8; color: #18212f; font-family: Arial, sans-serif; }}
    main {{ max-width: 1500px; margin: 0 auto; padding: 24px; }}
    .layout {{ display: grid; grid-template-columns: minmax(560px,1fr) minmax(420px,0.9fr); gap: 20px; align-items: stretch; }}
    .panel {{ background: #fffdf8; border: 1px solid #d8d1c3; border-radius: 18px; padding: 18px; display: flex; flex-direction: column; min-height: min(82vh, 960px); }}
    .annotation-stage {{ flex: 1; min-height: 0; overflow: auto; padding-right: 4px; }}
    .annotation-stage img {{ width: 100%; max-height: 100%; object-fit: contain; border-radius: 12px; border: 1px solid #d8d1c3; }}
    .annotation-chip {{ margin-right: 8px; margin-bottom: 8px; border-radius: 999px; border: 1px solid #d8d1c3; background: #fff; padding: 8px 12px; cursor: pointer; }}
    .annotation-chip.active {{ border-color: #0f766e; color: #0f766e; }}
    .event-card {{ border: 1px solid #d8d1c3; border-radius: 14px; margin-bottom: 12px; background: #fff; }}
    .event-card.active {{ border-color: #0f766e; box-shadow: 0 10px 24px rgba(15,118,110,0.12); }}
    .card-button {{ width: 100%; text-align: left; border: 0; background: transparent; padding: 16px; color: inherit; cursor: pointer; }}
    .card-head {{ display: grid; grid-template-columns: auto 1fr auto; gap: 12px; align-items: start; }}
    .event-number {{ width: 40px; height: 40px; border-radius: 999px; background: #ecfeff; color: #0f766e; display:flex; align-items:center; justify-content:center; font-weight:700; }}
    .card-section {{ margin-top: 14px; }}
    .card-section h3 {{ margin: 0 0 8px; font-size: 13px; color:#6b7280; }}
    .meta-grid {{ display:grid; grid-template-columns: 110px 1fr; gap: 6px 10px; }}
    .key-param-block, .detail-text {{ margin:0; line-height:1.5; word-break: break-word; }}
    .mono {{ font-family: ui-monospace, SFMono-Regular, Menlo, monospace; font-size: 12px; }}
    .issue-summary {{ display:flex; flex-wrap: wrap; gap: 8px; }}
    .issue-chip {{ display:inline-flex; align-items:center; border-radius:999px; padding:6px 10px; font-size:12px; border:1px solid #d8d1c3; background:#fff; }}
    .issue-chip.ok {{ background:#ecfeff; color:#0f766e; border-color:#99f6e4; }}
    .issue-chip.warn {{ background:#fff7ed; color:#b45309; border-color:#fdba74; }}
    .issue-chip.subtle {{ background:#f8f4eb; color:#475569; }}
    .detail-stack {{ padding: 0 16px 16px; }}
    dl {{ display:grid; grid-template-columns: 130px 1fr; gap: 8px 12px; }}
    dt {{ color:#6b7280; }}
    dd {{ margin:0; word-break: break-word; }}
    .param-table {{ width:100%; border-collapse: collapse; margin-top: 12px; font-size: 13px; }}
    .param-table th, .param-table td {{ border: 1px solid #d8d1c3; padding: 8px; vertical-align: top; text-align: left; }}
    .param-table th {{ background: #f8f4eb; }}
    pre {{ overflow:auto; background:#0f172a; color:#e2e8f0; padding:12px; border-radius:12px; }}
    #cards {{ flex: 1; min-height: 0; overflow: auto; padding-right: 4px; }}
    @media (max-width: 1100px) {{ .layout {{ grid-template-columns: 1fr; }} .panel {{ min-height: auto; }} #cards, .annotation-stage {{ overflow: visible; }} }}
  </style>
</head>
<body>
  <main>
    <section class="layout">
      <section class="panel">
        <h1>Annotations</h1>
        <div id="annotation-nav">{nav_html}</div>
        <div class="annotation-stage">
          {'<img id="annotation-image" src="' + html.escape(str(records[0].get('screenshot_file', ''))) + '" alt="annotation">' if records else '<p>매칭 성공한 annotation이 없습니다.</p>'}
        </div>
      </section>
      <section class="panel">
        <h1>Event Definition Cards</h1>
        <div id="cards">{card_html or '<p>매칭 성공한 definition row가 없습니다.</p>'}</div>
      </section>
    </section>
    <section class="panel" style="margin-top:20px; min-height:auto;">
      <h1>Unmatched Definition Rows</h1>
      {unmatched_html}
    </section>
  </main>
  <script>
    const RECORDS = {event_payload};
    const image = document.getElementById("annotation-image");
    const cards = Array.from(document.querySelectorAll(".event-card"));
    const buttons = Array.from(document.querySelectorAll("[data-select-no]"));
    function setActive(no) {{
      const record = RECORDS.find((item) => String(item.no) === String(no));
      if (!record || !image) return;
      image.src = record.screenshot_file;
      cards.forEach((card) => card.classList.toggle("active", card.dataset.no === String(no)));
      buttons.forEach((button) => button.classList.toggle("active", button.dataset.selectNo === String(no)));
      const card = document.getElementById(`card-${{no}}`);
      if (card) card.scrollIntoView({{ behavior: "smooth", block: "nearest" }});
    }}
    buttons.forEach((button) => button.addEventListener("click", () => setActive(button.dataset.selectNo)));
    if (RECORDS.length > 0) setActive("{html.escape(first_no)}");
  </script>
</body>
</html>
"""


def build_log_definition_export_zip_bytes(
    event_definition_df: pd.DataFrame,
    event_detail_df: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]] | None = None,
    definition_matches_df: pd.DataFrame | None = None,
) -> bytes:
    review_df, support_df, html_records, unmatched_records = build_log_definition_export_tables(
        event_definition_df,
        event_detail_df,
        schemas=schemas,
        definition_matches_df=definition_matches_df,
    )
    review_records_for_html: List[Dict[str, object]] = []
    annotation_file_map: Dict[str, bytes] = {}
    for row in html_records:
        no = str(row.get("no", "")).strip() or "0"
        ident = str(row.get("identification", "")).strip() or str(row.get("target_id", "")).strip() or "event"
        src_rel = str(row.get("screenshot_file", "")).strip()
        export_name = f"{no}_{_sanitize_export_filename_part(ident, default='event')}.png"
        export_rel = f"annotations/{export_name}"
        src_abs = BASE_DIR / src_rel if src_rel and src_rel != "-" else None
        if src_abs and src_abs.exists() and src_abs.is_file():
            annotation_file_map[export_rel] = src_abs.read_bytes()
        row_copy = dict(row)
        row_copy["screenshot_file"] = export_rel
        review_records_for_html.append(row_copy)
        if not review_df.empty:
            review_df.loc[review_df["no"].astype(str) == no, "screenshot_file"] = export_rel
        if not support_df.empty:
            support_df.loc[support_df["no"].astype(str) == no, "screenshot_path"] = export_rel
    html_bytes = build_log_definition_review_html(review_records_for_html, unmatched_records=unmatched_records).encode("utf-8")
    file_map = {
        "event_definition_review.html": html_bytes,
        "event_definition.csv": dataframe_to_csv_bytes(review_df),
        "event_support.csv": dataframe_to_csv_bytes(support_df),
    }
    file_map.update(annotation_file_map)
    return build_zip_bytes(file_map)


def build_qa_result_export_df(
    base_df: pd.DataFrame,
    event_detail_df: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]] | None = None,
    definition_matches_df: pd.DataFrame | None = None,
    review_html_filename: str = "event_definition_review.html",
) -> pd.DataFrame:
    cols = [
        "no",
        "이벤트 설명",
        "이벤트 이름",
        "매개변수명",
        "매개변수 값(예시)",
        "PASS/FAIL 체크",
        "테스트 시각",
        "데이터 수집값",
        "오류 원인",
        "트리거 시점",
        "QA 내용",
        "matched_identification",
        "screenshot_file",
        "screenshot_link",
        "review_html_anchor",
    ]
    if base_df.empty:
        return pd.DataFrame(columns=cols)

    detail_records = event_detail_df.to_dict("records") if isinstance(event_detail_df, pd.DataFrame) and not event_detail_df.empty else []
    detail_by_event_no: Dict[str, Dict[str, object]] = {}
    detail_by_event: Dict[str, Dict[str, object]] = {}
    for row in detail_records:
        event_no = str(row.get("event no", "")).strip()
        event_name = str(row.get("custom_event", row.get("event_name", ""))).strip()
        if event_no and event_no not in detail_by_event_no:
            detail_by_event_no[event_no] = row
        if event_name and event_name not in detail_by_event:
            detail_by_event[event_name] = row

    schema_store = schemas if isinstance(schemas, dict) else {}
    case_lookup = _build_definition_case_lookup(schema_store)
    out = base_df.copy()
    best_match_by_definition_row: Dict[str, Dict[str, object]] = {}
    if isinstance(definition_matches_df, pd.DataFrame) and not definition_matches_df.empty and "definition_row_id" in out.columns:
        ranked = definition_matches_df.copy()
        rank_map = {"PASS": 3, "WARN": 2, "FAIL": 1}
        ranked["rank"] = ranked["result"].astype(str).map(rank_map).fillna(0)
        ranked["match_score"] = pd.to_numeric(ranked.get("match_score", 0), errors="coerce").fillna(0)
        ranked = ranked.sort_values(by=["rank", "match_score"], ascending=[False, False])
        best = ranked.groupby("definition_row_id", as_index=False).first()
        best["review_html_anchor"] = best["captured_event_no"].astype(str).apply(lambda v: f"{review_html_filename}#card-{v}" if str(v).strip() and str(v).strip() != "-" else review_html_filename)
        best["screenshot_link"] = best.apply(
            lambda row: build_annotation_export_relpath(row.get("captured_event_no", "0"), row.get("matched_identification", ""))
            if str(row.get("screenshot_file", "")).strip() and str(row.get("screenshot_file", "")).strip() != "-"
            else "-",
            axis=1,
        )
        out = out.merge(
            best[["definition_row_id", "matched_identification", "screenshot_file", "screenshot_link", "review_html_anchor"]],
            on="definition_row_id",
            how="left",
        )
        best_match_by_definition_row = {
            str(row.get("definition_row_id", "")).strip(): row
            for row in best.to_dict("records")
            if str(row.get("definition_row_id", "")).strip()
        }
    else:
        matched_identifications = []
        screenshot_files = []
        screenshot_links = []
        review_anchors = []
        for _, row in out.iterrows():
            detail = detail_by_event.get(str(row.get("event_name", "")).strip(), {})
            event_no = str(detail.get("event no", "")).strip() or "0"
            identification = str(detail.get("identification", "-")).strip() or "-"
            matched_identifications.append(identification)
            screenshot_file = str(detail.get("screenshot_file", "-")).strip() or "-"
            screenshot_files.append(screenshot_file)
            screenshot_links.append(build_annotation_export_relpath(event_no, identification) if screenshot_file != "-" else "-")
            review_anchors.append(f"{review_html_filename}#card-{event_no}" if event_no != "0" else review_html_filename)
        out["matched_identification"] = matched_identifications
        out["screenshot_file"] = screenshot_files
        out["screenshot_link"] = screenshot_links
        out["review_html_anchor"] = review_anchors

    for col in ["description", "failure_reason", "missing_parameters", "extra_parameters", "matched_identification", "screenshot_file", "screenshot_link", "review_html_anchor"]:
        if col not in out.columns:
            out[col] = "-"
        out[col] = out[col].fillna("-").astype(str)
    if "no" not in out.columns:
        out["no"] = "-"
    vertical_rows: List[Dict[str, object]] = []
    for row in out.to_dict("records"):
        definition_row_id = str(row.get("definition_row_id", "")).strip()
        event_name = str(row.get("event_name", "")).strip() or "-"
        description = str(row.get("description", "-")).strip() or "-"
        result = str(row.get("result", row.get("schema_status", "-"))).strip().upper() or "-"
        failure_reason = str(row.get("failure_reason", "-")).strip() or "-"
        missing_parameters = str(row.get("missing_parameters", "-")).strip() or "-"
        extra_parameters = str(row.get("extra_parameters", "-")).strip() or "-"
        matched_identification = str(row.get("matched_identification", "-")).strip() or "-"
        screenshot_file = str(row.get("screenshot_file", "-")).strip() or "-"
        screenshot_link = str(row.get("screenshot_link", "-")).strip() or "-"
        review_html_anchor = str(row.get("review_html_anchor", review_html_filename)).strip() or review_html_filename
        best_match = best_match_by_definition_row.get(definition_row_id, {})
        trigger_at = str(best_match.get("matched_at", row.get("matched_at", "-"))).strip() or "-"
        test_at = trigger_at
        matched_event_no = str(best_match.get("captured_event_no", "")).strip()
        detail = detail_by_event_no.get(matched_event_no, {}) if matched_event_no else detail_by_event.get(event_name, {})
        actual_param_map = _build_detail_param_map(detail) if detail else {}
        qa_note = "정상" if result == "PASS" else failure_reason

        vertical_rows.append(
            {
                "no": str(row.get("no", "-")).strip() or "-",
                "이벤트 설명": description,
                "이벤트 이름": event_name,
                "매개변수명": "",
                "매개변수 값(예시)": "",
                "PASS/FAIL 체크": result,
                "테스트 시각": test_at,
                "데이터 수집값": matched_identification,
                "오류 원인": failure_reason if failure_reason != "-" else "",
                "트리거 시점": trigger_at,
                "QA 내용": qa_note,
                "matched_identification": matched_identification,
                "screenshot_file": screenshot_file,
                "screenshot_link": screenshot_link,
                "review_html_anchor": review_html_anchor,
            }
        )

        parameter_rows: List[Dict[str, object]] = []
        if definition_row_id and definition_row_id in case_lookup:
            case = case_lookup[definition_row_id]
            expected_map = case.get("param_values", {}) if isinstance(case.get("param_values", {}), dict) else {}
            parameter_keys = list(
                dict.fromkeys(
                    [str(v).strip() for v in case.get("required", []) if str(v).strip()]
                    + [str(v).strip() for v in case.get("optional", []) if str(v).strip()]
                    + [str(v).strip() for v in case.get("parameter_columns", []) if str(v).strip()]
                    + [str(v).strip() for v in expected_map.keys() if str(v).strip()]
                )
            )
            required_keys = {str(v).strip() for v in case.get("required", []) if str(v).strip()}
            for key in parameter_keys:
                expected_raw = str(expected_map.get(key, "")).strip()
                expected_norm = normalize_expected_param_value(expected_raw)
                actual_value = str(actual_param_map.get(key, "")).strip()
                if key in required_keys and _is_missing_like_value(actual_value):
                    param_result = "FAIL"
                    param_reason = "missing_parameter"
                    param_note = "필수 파라미터 누락"
                elif expected_norm and not _is_definition_placeholder_value(expected_norm) and actual_value and actual_value != expected_norm:
                    param_result = "FAIL"
                    param_reason = "value_mismatch"
                    param_note = "기대값과 실제값 불일치"
                elif _is_missing_like_value(actual_value):
                    param_result = "WARN"
                    param_reason = ""
                    param_note = "수집값 없음"
                else:
                    param_result = "PASS"
                    param_reason = ""
                    param_note = "정상"
                parameter_rows.append(
                    {
                        "no": "",
                        "이벤트 설명": "",
                        "이벤트 이름": "",
                        "매개변수명": key,
                        "매개변수 값(예시)": expected_raw or expected_norm or "-",
                        "PASS/FAIL 체크": param_result,
                        "테스트 시각": test_at,
                        "데이터 수집값": actual_value or "-",
                        "오류 원인": param_reason,
                        "트리거 시점": trigger_at,
                        "QA 내용": param_note,
                        "matched_identification": "",
                        "screenshot_file": "",
                        "screenshot_link": "",
                        "review_html_anchor": "",
                    }
                )
            extra_keys = [key.strip() for key in extra_parameters.split(",") if key.strip() and key.strip() != "-"]
            for key in extra_keys:
                parameter_rows.append(
                    {
                        "no": "",
                        "이벤트 설명": "",
                        "이벤트 이름": "",
                        "매개변수명": key,
                        "매개변수 값(예시)": "-",
                        "PASS/FAIL 체크": "WARN",
                        "테스트 시각": test_at,
                        "데이터 수집값": str(actual_param_map.get(key, "-")).strip() or "-",
                        "오류 원인": "extra_parameter",
                        "트리거 시점": trigger_at,
                        "QA 내용": "정의서 외 추가 파라미터",
                        "matched_identification": "",
                        "screenshot_file": "",
                        "screenshot_link": "",
                        "review_html_anchor": "",
                    }
                )
        else:
            schema_obj = schema_store.get(event_name, {})
            schema_param_keys = list(
                dict.fromkeys(
                    [str(v).strip() for v in schema_obj.get("required", []) if str(v).strip()]
                    + [str(v).strip() for v in schema_obj.get("optional", []) if str(v).strip()]
                    + [str(v).strip() for v in actual_param_map.keys() if str(v).strip()]
                )
            )
            required_keys = {str(v).strip() for v in schema_obj.get("required", []) if str(v).strip()}
            for key in schema_param_keys:
                actual_value = str(actual_param_map.get(key, "")).strip()
                if key in required_keys and _is_missing_like_value(actual_value):
                    param_result = "FAIL"
                    param_reason = "missing_parameter"
                    param_note = "필수 파라미터 누락"
                elif _is_missing_like_value(actual_value):
                    param_result = "WARN"
                    param_reason = ""
                    param_note = "수집값 없음"
                else:
                    param_result = "PASS"
                    param_reason = ""
                    param_note = "정상"
                parameter_rows.append(
                    {
                        "no": "",
                        "이벤트 설명": "",
                        "이벤트 이름": "",
                        "매개변수명": key,
                        "매개변수 값(예시)": actual_value or "-",
                        "PASS/FAIL 체크": param_result,
                        "테스트 시각": test_at,
                        "데이터 수집값": actual_value or "-",
                        "오류 원인": param_reason,
                        "트리거 시점": trigger_at,
                        "QA 내용": param_note,
                        "matched_identification": "",
                        "screenshot_file": "",
                        "screenshot_link": "",
                        "review_html_anchor": "",
                    }
                )
        vertical_rows.extend(parameter_rows)
    return pd.DataFrame(vertical_rows, columns=cols)


def build_qa_summary_export_df(
    event_detail_df: pd.DataFrame,
    definition_validation_summary_df: pd.DataFrame | None = None,
    bundle_scope: str = "custom_event_review",
) -> pd.DataFrame:
    if str(bundle_scope or "").strip() == "definition_validation" and isinstance(definition_validation_summary_df, pd.DataFrame):
        return definition_validation_summary_df.copy()
    if event_detail_df.empty:
        return pd.DataFrame(
            columns=["event_no", "custom_event", "identification", "target_id", "section_name", "page_id", "status", "custom_parameters"]
        )
    rows: List[Dict[str, object]] = []
    for row in event_detail_df.to_dict("records"):
        rows.append(
            {
                "event_no": str(row.get("event no", "")).strip(),
                "custom_event": str(row.get("custom_event", row.get("event_name", ""))).strip(),
                "identification": str(row.get("identification", "-")).strip() or "-",
                "target_id": str(row.get("target_id", "-")).strip() or "-",
                "section_name": str(row.get("section_name", "-")).strip() or "-",
                "page_id": str(row.get("page_id", "-")).strip() or "-",
                "status": str(row.get("status", "-")).strip() or "-",
                "custom_parameters": str(row.get("custom_parameters", "-")).strip() or "-",
            }
        )
    return pd.DataFrame(rows)


def build_qa_debug_export_df(
    event_detail_df: pd.DataFrame,
    definition_validation_matches_df: pd.DataFrame | None = None,
    definition_validation_unmatched_df: pd.DataFrame | None = None,
    bundle_scope: str = "custom_event_review",
) -> pd.DataFrame:
    if str(bundle_scope or "").strip() == "definition_validation":
        parts: List[pd.DataFrame] = []
        if isinstance(definition_validation_matches_df, pd.DataFrame) and not definition_validation_matches_df.empty:
            matches = definition_validation_matches_df.copy()
            matches.insert(0, "record_type", "matched")
            parts.append(matches)
        if isinstance(definition_validation_unmatched_df, pd.DataFrame) and not definition_validation_unmatched_df.empty:
            unmatched = definition_validation_unmatched_df.copy()
            unmatched.insert(0, "record_type", "unmatched")
            parts.append(unmatched)
        if not parts:
            return pd.DataFrame(columns=["record_type"])
        return pd.concat(parts, ignore_index=True)
    return event_detail_df.drop(columns=["raw_payload_json"], errors="ignore").copy() if isinstance(event_detail_df, pd.DataFrame) else pd.DataFrame()


def build_zip_bytes(file_map: Dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for file_name, payload in file_map.items():
            if not payload:
                continue
            zf.writestr(str(file_name), payload)
    return out.getvalue()


def build_event_logs_export_df(rt_events: pd.DataFrame) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    if rt_events.empty:
        return pd.DataFrame(
            columns=[
                "timestamp",
                "custom_event",
                "custom_parameter",
                "target_id",
                "selector",
                "screen_name",
                "screenshot_path",
                "bounding_box",
                "annotation_no",
                "page_id",
                "status",
            ]
        )
    for row in rt_events.to_dict("records"):
        event_name = str(row.get("custom_event", row.get("이벤트", ""))).strip()
        timestamp = str(row.get("시간", "-"))
        params = row.get("전체 파라미터", {})
        params = params if isinstance(params, dict) else {}
        custom_parameter = str(row.get("custom_parameters", "")).strip() or "-"
        bounding_box = {
            "x": int(pd.to_numeric(row.get("bbox_x", 0), errors="coerce") or 0),
            "y": int(pd.to_numeric(row.get("bbox_y", 0), errors="coerce") or 0),
            "width": int(pd.to_numeric(row.get("bbox_width", 0), errors="coerce") or 0),
            "height": int(pd.to_numeric(row.get("bbox_height", 0), errors="coerce") or 0),
        }
        rows.append(
            {
                "timestamp": timestamp,
                "custom_event": event_name,
                "custom_parameter": custom_parameter,
                "target_id": str(row.get("target_id", "-")).strip() or "-",
                "selector": str(row.get("selector", "-")).strip() or "-",
                "screen_name": str(row.get("screen_name", row.get("page_id", "-"))).strip() or "-",
                "screenshot_path": str(row.get("screenshot_path", row.get("selector_screenshot", "-"))).strip() or "-",
                "bounding_box": json.dumps(bounding_box, ensure_ascii=False),
                "annotation_no": str(row.get("annotation_no", "")).strip() or "-",
                "page_id": str(row.get("page_id", _to_text_value(params.get("page_id", ""))) or "-").strip() or "-",
                "status": str(row.get("status", row.get("상태", "-"))).strip() or "-",
            }
        )
    return pd.DataFrame(rows)


def build_parameter_validation_export_df(rt_events: pd.DataFrame, schemas: Dict[str, Dict[str, object]]) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    for row in rt_events.to_dict("records"):
        event_name = str(row.get("이벤트", "")).strip()
        params = row.get("전체 파라미터", {})
        if not isinstance(params, dict):
            params = {}
        check = validate_schema_for_event(event_name=event_name, params=params, schemas=schemas)
        schema_required = set()
        schema_optional = set()
        schema_parameter_columns = set()
        schema_obj = schemas.get(event_name, {})
        if isinstance(schema_obj, dict):
            schema_required = {str(v).strip() for v in schema_obj.get("required", []) if str(v).strip()}
            schema_optional = {str(v).strip() for v in schema_obj.get("optional", []) if str(v).strip()}
            schema_parameter_columns = {str(v).strip() for v in schema_obj.get("parameter_columns", []) if str(v).strip()}
        known_schema_params = schema_required | schema_optional | schema_parameter_columns

        for item in check.get("checks", []) if isinstance(check.get("checks", []), list) else []:
            if not isinstance(item, dict):
                continue
            param_name = str(item.get("param", "-"))
            value_text = str(item.get("value", "-"))
            result = "PASS" if bool(item.get("ok", False)) else "FAIL"
            issue_type = "required_missing" if (str(item.get("group", "")) == "required" and result == "FAIL") else "schema_check"
            rows.append(
                {
                    "event_name": event_name,
                    "schema_status": str(check.get("status", "WARN")),
                    "parameter": param_name,
                    "group": str(item.get("group", "-")),
                    "value": value_text,
                    "result": result,
                    "severity": "FAIL" if result == "FAIL" else "PASS",
                    "issue_type": issue_type,
                    "reason": "필수 누락" if issue_type == "required_missing" else "-",
                }
            )

        for key, value in params.items():
            key_text = str(key).strip()
            if not key_text or _is_system_param_key(key_text):
                continue
            value_text = _to_text_value(value)
            if known_schema_params and key_text not in known_schema_params:
                rows.append(
                    {
                        "event_name": event_name,
                        "schema_status": str(check.get("status", "WARN")),
                        "parameter": key_text,
                        "group": "extra",
                        "value": value_text,
                        "result": "WARN",
                        "severity": "WARN",
                        "issue_type": "unexpected_parameter",
                        "reason": "스키마에 정의되지 않은 파라미터",
                    }
                )
            valid_ok, valid_msg = validate_param_format(key_text, value_text)
            if not valid_ok:
                rows.append(
                    {
                        "event_name": event_name,
                        "schema_status": str(check.get("status", "WARN")),
                        "parameter": key_text,
                        "group": "format",
                        "value": value_text,
                        "result": "FAIL",
                        "severity": "FAIL",
                        "issue_type": "type_validation_fail",
                        "reason": valid_msg,
                    }
                )
    if not rows:
        return pd.DataFrame(
            columns=["event_name", "schema_status", "parameter", "group", "value", "result", "severity", "issue_type", "reason"]
        )
    return pd.DataFrame(rows)


def build_event_order_validation_df(rt_events: pd.DataFrame, expected_steps: List[str]) -> pd.DataFrame:
    cols = ["timestamp", "event_name", "status", "reason", "expected_prev_step", "actual_step_pointer"]
    steps = [str(v).strip() for v in expected_steps if str(v).strip()]
    if rt_events.empty or not steps:
        return pd.DataFrame(columns=cols)

    work = rt_events.copy()
    work["captured_at"] = pd.to_datetime(work["captured_at"], errors="coerce")
    work = work.sort_values("captured_at")
    step_index = {name: idx for idx, name in enumerate(steps)}
    pointer = 0
    rows: List[Dict[str, object]] = []

    for _, row in work.iterrows():
        event_name = str(row.get("이벤트", "")).strip()
        if event_name not in step_index:
            continue
        idx = step_index[event_name]
        ts = format_local_time(row.get("captured_at"))
        if idx < pointer:
            rows.append(
                {
                    "timestamp": ts,
                    "event_name": event_name,
                    "status": "WARN",
                    "reason": "이미 지난 step의 이벤트가 재발생",
                    "expected_prev_step": steps[pointer - 1] if pointer > 0 else "-",
                    "actual_step_pointer": pointer,
                }
            )
        elif idx == pointer:
            rows.append(
                {
                    "timestamp": ts,
                    "event_name": event_name,
                    "status": "PASS",
                    "reason": "순서 정상",
                    "expected_prev_step": steps[pointer - 1] if pointer > 0 else "-",
                    "actual_step_pointer": pointer,
                }
            )
            pointer += 1
        else:
            rows.append(
                {
                    "timestamp": ts,
                    "event_name": event_name,
                    "status": "FAIL",
                    "reason": "선행 step 누락 상태에서 발생",
                    "expected_prev_step": steps[pointer] if pointer < len(steps) else "-",
                    "actual_step_pointer": pointer,
                }
            )

    if pointer < len(steps):
        rows.append(
            {
                "timestamp": "-",
                "event_name": steps[pointer],
                "status": "FAIL",
                "reason": "필수 순서 step 미도달",
                "expected_prev_step": steps[pointer - 1] if pointer > 0 else "-",
                "actual_step_pointer": pointer,
            }
        )
    return pd.DataFrame(rows, columns=cols)


def build_event_qa_report_export_df(
    rt_events: pd.DataFrame,
    issue_df: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]],
    parameter_validation_df: pd.DataFrame,
    order_validation_df: pd.DataFrame,
) -> pd.DataFrame:
    qa_by_event = build_schema_validation_by_event(rt_events, schemas)
    issue_reason_map: Dict[str, str] = {}
    if not issue_df.empty:
        for event_name, group in issue_df.groupby("event_name"):
            reasons = [str(v).strip() for v in group["reason"].tolist() if str(v).strip()]
            issue_reason_map[str(event_name).strip()] = " | ".join(dict.fromkeys(reasons))

    param_fail_map: Dict[str, int] = {}
    param_warn_map: Dict[str, int] = {}
    if not parameter_validation_df.empty:
        work = parameter_validation_df.copy()
        work["event_name"] = work["event_name"].astype(str).str.strip()
        for event_name, group in work.groupby("event_name"):
            param_fail_map[event_name] = int((group["severity"].astype(str) == "FAIL").sum())
            param_warn_map[event_name] = int((group["severity"].astype(str) == "WARN").sum())

    order_fail_map: Dict[str, int] = {}
    order_warn_map: Dict[str, int] = {}
    if not order_validation_df.empty:
        owork = order_validation_df.copy()
        owork["event_name"] = owork["event_name"].astype(str).str.strip()
        for event_name, group in owork.groupby("event_name"):
            order_fail_map[event_name] = int((group["status"].astype(str) == "FAIL").sum())
            order_warn_map[event_name] = int((group["status"].astype(str) == "WARN").sum())

    event_error_map: Dict[str, int] = {}
    if not rt_events.empty:
        ework = rt_events.copy()
        ework["이벤트"] = ework["이벤트"].astype(str).str.strip()
        for event_name, group in ework.groupby("이벤트"):
            event_error_map[event_name] = int((group["상태"].astype(str).str.upper() == "ERROR").sum())

    rows: List[Dict[str, object]] = []
    if qa_by_event.empty:
        return pd.DataFrame(columns=["event_name", "schema_status", "parameter_issue", "test_count", "result"])
    for _, row in qa_by_event.iterrows():
        event_name = str(row.get("event_name", "")).strip()
        pass_n = int(row.get("PASS", 0))
        warn_n = int(row.get("WARN", 0))
        fail_n = int(row.get("FAIL", 0))
        total = int(row.get("total", 0))
        param_fail = int(param_fail_map.get(event_name, 0))
        param_warn = int(param_warn_map.get(event_name, 0))
        order_fail = int(order_fail_map.get(event_name, 0))
        order_warn = int(order_warn_map.get(event_name, 0))
        event_error = int(event_error_map.get(event_name, 0))

        # QA 상태 기준
        # FAIL: 스키마 FAIL 또는 필수/타입 FAIL 또는 순서 FAIL 또는 이벤트 ERROR
        # WARN: FAIL이 아니고 스키마 WARN 또는 optional/extra 경고 또는 순서 WARN
        # PASS: 위 조건이 없을 때
        if fail_n > 0 or param_fail > 0 or order_fail > 0 or event_error > 0:
            schema_status = "FAIL"
        elif warn_n > 0 or param_warn > 0 or order_warn > 0:
            schema_status = "WARN"
        else:
            schema_status = "PASS"
        parameter_issue_text = issue_reason_map.get(event_name, "-")
        if parameter_issue_text == "-" and param_fail > 0:
            parameter_issue_text = f"파라미터 FAIL {param_fail}개"
        if parameter_issue_text == "-" and param_warn > 0:
            parameter_issue_text = f"파라미터 WARN {param_warn}개"

        rows.append(
            {
                "event_name": event_name,
                "schema_status": schema_status,
                "parameter_issue": parameter_issue_text,
                "test_count": total,
                "result": schema_status,
                "pass_count": pass_n,
                "warn_count": warn_n,
                "fail_count": fail_n,
                "parameter_fail_count": param_fail,
                "parameter_warn_count": param_warn,
                "order_fail_count": order_fail,
                "order_warn_count": order_warn,
                "event_error_count": event_error,
                "qa_status_criteria": "FAIL=스키마FAIL/필수·타입FAIL/순서FAIL/이벤트ERROR, WARN=스키마WARN/추가파라미터/순서WARN, PASS=그외",
            }
        )
    return pd.DataFrame(rows).sort_values(by=["result", "test_count", "event_name"], ascending=[True, False, True])


def build_definition_validation_issue_df(
    rt_events: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]],
    issue_df: pd.DataFrame,
) -> pd.DataFrame:
    case_df = build_definition_case_validation_df(rt_events, schemas)
    if case_df.empty:
        return pd.DataFrame(columns=["no", "event_name", "description", "status", "reason", "count"])
    issue_rows: List[Dict[str, object]] = []
    for row in case_df.to_dict("records"):
        result = str(row.get("result", "WARN")).upper()
        if result == "PASS":
            continue
        issue_rows.append(
            {
                "no": str(row.get("no", "")).strip(),
                "event_name": str(row.get("event_name", "")).strip(),
                "description": str(row.get("description", "")).strip(),
                "status": result,
                "reason": str(row.get("failure_reason", "-")).strip() or "-",
                "count": 1,
            }
        )
    if not issue_rows:
        return pd.DataFrame(columns=["no", "event_name", "description", "status", "reason", "count"])
    return pd.DataFrame(issue_rows).sort_values(by=["status", "event_name", "no"], ascending=[True, True, True]).reset_index(drop=True)


def build_definition_validation_report_df(
    rt_events: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]],
    issue_df: pd.DataFrame,
    parameter_validation_df: pd.DataFrame,
    order_validation_df: pd.DataFrame,
) -> pd.DataFrame:
    case_df = build_definition_case_validation_df(rt_events, schemas)
    if case_df.empty:
        return pd.DataFrame(
            columns=[
                "definition_row_id",
                "no",
                "event_name",
                "description",
                "expected_identification",
                "page_id",
                "section_name",
                "schema_status",
                "matched_capture_count",
                "result",
                "failure_reason",
                "missing_parameters",
                "extra_parameters",
            ]
        )
    report_df = case_df.rename(columns={"result": "schema_status"}).copy()
    report_df["result"] = report_df["schema_status"]
    return report_df.sort_values(by=["result", "event_name", "no"], ascending=[True, True, True]).reset_index(drop=True)


def summarize_event_qa_report(report_df: pd.DataFrame) -> Dict[str, int]:
    out = {"PASS": 0, "WARN": 0, "FAIL": 0}
    if report_df.empty or "result" not in report_df.columns:
        return out
    for value in report_df["result"].astype(str):
        key = value.strip().upper()
        if key in out:
            out[key] += 1
    return out


def _infer_action_object(event_name: str) -> Tuple[str, str]:
    name = str(event_name).strip().lower()
    action = name.split("_", 1)[0] if "_" in name else name
    if action in {"", "gtm.click", "gtm.linkclick"}:
        action = "click"
    if "impression" in name:
        action = "impression"
    obj = "content"
    if "button" in name:
        obj = "button"
    elif "product" in name:
        obj = "product"
    elif "popup" in name:
        obj = "popup"
    elif "link" in name:
        obj = "link"
    return action or "event", obj


def build_tracking_plan_export_df(
    rt_events: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]],
    action_map_df: pd.DataFrame | None = None,
) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    event_names = sorted(list(dict.fromkeys(list(schemas.keys()) + rt_events.get("이벤트", pd.Series(dtype=str)).astype(str).tolist())))
    action_map: Dict[str, Dict[str, str]] = {}
    if isinstance(action_map_df, pd.DataFrame) and not action_map_df.empty:
        norm_map_df = _normalize_action_map_df(action_map_df)
        for _, map_row in norm_map_df.iterrows():
            ev = str(map_row.get("event_name", "")).strip()
            if not ev:
                continue
            action_map[ev] = {
                "action": str(map_row.get("action", "")).strip(),
                "object": str(map_row.get("object", "")).strip(),
                "page_id": str(map_row.get("page_id", "")).strip(),
                "section_name": str(map_row.get("section_name", "")).strip(),
                "section_index": str(map_row.get("section_index", "")).strip(),
                "description": str(map_row.get("description", "")).strip(),
                "object_index": str(map_row.get("object_index", "")).strip(),
            }
    event_meta: Dict[str, Dict[str, str]] = {}
    if not rt_events.empty:
        for _, row in rt_events.iterrows():
            ev = str(row.get("이벤트", "")).strip()
            if not ev:
                continue
            params = row.get("전체 파라미터", {})
            if not isinstance(params, dict):
                params = {}
            meta = event_meta.setdefault(ev, {})
            for key in ["page_id", "section_name", "section_index", "action", "object", "description", "index"]:
                val = _to_text_value(params.get(key, ""))
                if val and (key not in meta or not meta[key]):
                    meta[key] = val

    for event_name in event_names:
        if not str(event_name).strip():
            continue
        schema = schemas.get(str(event_name).strip(), {})
        required = [str(v).strip() for v in schema.get("required", []) if str(v).strip()] if isinstance(schema, dict) else []
        optional = [str(v).strip() for v in schema.get("optional", []) if str(v).strip()] if isinstance(schema, dict) else []
        params = required + [v for v in optional if v not in required]
        action, obj = _infer_action_object(str(event_name))
        meta = event_meta.get(str(event_name).strip(), {})
        map_override = action_map.get(str(event_name).strip(), {})
        page_id = map_override.get("page_id") or meta.get("page_id", "unknown")
        section_name = map_override.get("section_name") or meta.get("section_name", "general")
        section_index = map_override.get("section_index") or meta.get("section_index", "-")
        action_value = map_override.get("action") or meta.get("action", action)
        object_value = map_override.get("object") or meta.get("object", obj)
        description = map_override.get("description") or meta.get("description", "-")
        object_index = map_override.get("object_index") or meta.get("index", "-")
        if not params:
            params = ["-"]
        for p in params:
            rows.append(
                {
                    "page_id": page_id,
                    "section_name": section_name,
                    "section_index": section_index,
                    "action": action_value,
                    "object": object_value,
                    "event_name": str(event_name).strip(),
                    "description": description,
                    "object_index": object_index,
                    "parameter": p,
                }
            )
    if not rows:
        return pd.DataFrame(
            columns=[
                "page_id",
                "section_name",
                "section_index",
                "action",
                "object",
                "event_name",
                "description",
                "object_index",
                "parameter",
            ]
        )
    return pd.DataFrame(rows)


def build_screen_definition_export_df(tracking_plan_df: pd.DataFrame) -> pd.DataFrame:
    if tracking_plan_df.empty:
        return pd.DataFrame(columns=["section_index", "section_name", "UI_type", "event"])
    unique_sections = (
        tracking_plan_df[["section_index", "section_name", "object", "event_name"]]
        .drop_duplicates()
        .fillna("-")
    )
    rows: List[Dict[str, object]] = []
    for idx, row in unique_sections.reset_index(drop=True).iterrows():
        event_name = str(row.get("event_name", "")).strip()
        action, obj = _infer_action_object(event_name)
        sec_idx = str(row.get("section_index", "")).strip() or str(idx + 1)
        sec_name = str(row.get("section_name", "")).strip() or f"section_{idx+1}"
        ui_type = str(row.get("object", "")).strip() or obj
        rows.append(
            {
                "section_index": sec_idx,
                "section_name": sec_name,
                "UI_type": ui_type,
                "event": event_name,
                "action": action,
            }
        )
    return pd.DataFrame(rows)


def build_screen_marked_png_bytes(screen_definition_df: pd.DataFrame) -> bytes:
    try:
        from PIL import Image, ImageDraw
    except Exception:
        return b""

    width, height = 1280, 720
    image = Image.new("RGB", (width, height), color=(248, 250, 252))
    draw = ImageDraw.Draw(image)
    draw.rectangle([20, 20, width - 20, 90], outline=(30, 64, 175), width=3)
    draw.text((40, 45), "Screen Marked Preview (QA)", fill=(30, 64, 175))

    rows = screen_definition_df.to_dict("records")[:8] if not screen_definition_df.empty else []
    top = 130
    box_h = 64
    circle_r = 18
    for idx, row in enumerate(rows, start=1):
        y1 = top + (idx - 1) * (box_h + 14)
        y2 = y1 + box_h
        draw.rectangle([80, y1, width - 40, y2], outline=(220, 38, 38), width=3)
        cx, cy = 54, y1 + int(box_h / 2)
        draw.ellipse([cx - circle_r, cy - circle_r, cx + circle_r, cy + circle_r], fill=(30, 64, 175), outline=(30, 64, 175), width=2)
        draw.text((cx - 6, cy - 10), str(idx), fill=(255, 255, 255))
        marker = to_circled_number(idx)
        label = f"{marker} {row.get('section_name', '-')} | event: {row.get('event', '-')}"
        draw.text((96, y1 + 20), label, fill=(17, 24, 39))

    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


EVENT_DEFINITION_PARAM_PRIORITY = [
    "button_id",
    "button_name",
    "section_name",
    "section_index",
    "section_title",
    "index",
    "content_id",
    "content_name",
    "content_type",
    "banner_id",
    "category_id",
    "category_name",
    "filter_type",
    "filter_value",
    "applied_tab",
    "page_id",
    "page_link",
    "extra_info",
]
EVENT_DEFINITION_HIDDEN_KEYS = {
    "gtm",
    "gcd",
    "uaa",
    "uab",
    "uafvl",
    "uap",
    "uapv",
    "uaw",
    "up.client_id",
    "page_hostname",
}


def _status_to_ko(status: str) -> str:
    normalized = str(status or "").strip().upper()
    return {"OK": "정상", "WARN": "경고", "ERROR": "오류", "INFO": "정보"}.get(normalized, "정상")


def _build_url_compare_key(url: str) -> str:
    parsed = urlparse(str(url or "").strip())
    scheme = parsed.scheme or "https"
    host = parsed.netloc or ""
    path = parsed.path or "/"
    return f"{scheme}://{host}{path}"


def _build_url_hash(url_compare_key: str) -> str:
    return hashlib.sha256(str(url_compare_key).encode("utf-8")).hexdigest()[:16]


def _extract_page_id(page_url: str, params: Dict[str, object]) -> str:
    page_id = _to_text_value(params.get("page_id", ""))
    if page_id:
        return page_id
    location_url = _to_text_value(params.get("page_location", page_url))
    try:
        parsed = urlparse(location_url)
        return parsed.path or "/"
    except Exception:
        return "/"


def _build_event_meta(row: pd.Series, params: Dict[str, object], auto_ctx: Dict[str, object] | None = None) -> Dict[str, object]:
    ctx = auto_ctx or {}
    page_url = _to_text_value(row.get("page_url", "")) or _to_text_value(params.get("page_location", ""))
    page_id = _extract_page_id(page_url, params)
    section_name = (
        _to_text_value(params.get("section_name", ""))
        or _to_text_value(params.get("qa_section_name", ""))
        or _to_text_value(ctx.get("section_name", ""))
        or "global"
    )
    target_id = (
        _to_text_value(params.get("qa_target_id", ""))
        or _to_text_value(params.get("target_id", ""))
        or _to_text_value(ctx.get("target_id", ""))
    )
    if not target_id:
        button_id = _to_text_value(params.get("button_id", ""))
        if button_id:
            target_id = f"button_id:{button_id}"
    if not target_id and section_name:
        target_id = f"data-section-name:{section_name}"
    selector = (
        _to_text_value(params.get("qa_selector", ""))
        or _to_text_value(params.get("selector", ""))
        or _to_text_value(ctx.get("selector", ""))
    )
    class_attribute = (
        _to_text_value(params.get("qa_class_attribute", ""))
        or _to_text_value(params.get("class_attribute", ""))
        or _to_text_value(ctx.get("class_attribute", ""))
    )
    screen_state = (
        _to_text_value(params.get("qa_screen_state", ""))
        or _to_text_value(params.get("screen_state", ""))
        or _to_text_value(ctx.get("screen_state", ""))
        or "default"
    )
    screenshot_file = (
        _to_text_value(params.get("qa_selector_screenshot", ""))
        or _to_text_value(params.get("qa_screenshot_path", ""))
        or _to_text_value(params.get("screenshot_path", ""))
        or _to_text_value(params.get("selector_screenshot", ""))
        or _to_text_value(ctx.get("screenshot_path", ""))
        or _to_text_value(ctx.get("selector_screenshot", ""))
    )
    raw_screenshot_file = (
        _to_text_value(params.get("qa_raw_screenshot_file", ""))
        or _to_text_value(params.get("raw_screenshot_file", ""))
        or _to_text_value(ctx.get("raw_screenshot_file", ""))
    )
    screen_name = (
        _to_text_value(params.get("qa_screen_name", ""))
        or _to_text_value(params.get("screen_name", ""))
        or _to_text_value(ctx.get("screen_name", ""))
        or page_id
        or "default"
    )
    annotation_no = (
        _to_text_value(params.get("qa_annotation_no", ""))
        or _to_text_value(params.get("annotation_no", ""))
        or _to_text_value(ctx.get("annotation_no", ""))
    )
    return {
        "page_url": page_url,
        "page_id": page_id,
        "screen_name": screen_name,
        "section_name": section_name,
        "target_id": target_id or "-",
        "selector": selector or "-",
        "class_attribute": class_attribute or "-",
        "screen_state": screen_state or "default",
        "annotation_no": annotation_no or "-",
        "screenshot_path": screenshot_file,
        "screenshot_file": screenshot_file,
        "raw_screenshot_file": raw_screenshot_file,
        "bbox_x": _to_text_value(params.get("qa_bbox_x", params.get("bbox_x", ctx.get("bbox_x", 0)))),
        "bbox_y": _to_text_value(params.get("qa_bbox_y", params.get("bbox_y", ctx.get("bbox_y", 0)))),
        "bbox_width": _to_text_value(params.get("qa_bbox_width", params.get("bbox_width", ctx.get("bbox_width", 0)))),
        "bbox_height": _to_text_value(params.get("qa_bbox_height", params.get("bbox_height", ctx.get("bbox_height", 0)))),
    }


def _choose_export_params(params: Dict[str, object], schema_obj: Dict[str, object]) -> List[Tuple[str, str]]:
    required = {str(v).strip() for v in schema_obj.get("required", []) if str(v).strip()} if isinstance(schema_obj, dict) else set()
    visible_keys = _sort_custom_param_keys(
        [key for key in list(required) + list(params.keys()) if not _is_hidden_default_param_key(str(key))]
    )
    rows: List[Tuple[str, str]] = []
    for key in visible_keys:
        if key in params:
            rows.append((key, "Y" if key in required else "N"))
    return rows


def _build_event_identity(event_name: str, params: Dict[str, object], meta: Dict[str, object]) -> Tuple[str, ...]:
    return (
        str(meta.get("page_id", "")).strip(),
        str(event_name).strip(),
        _to_text_value(params.get("section_name", meta.get("section_name", ""))),
        _to_text_value(params.get("button_id", "")),
        _to_text_value(params.get("filter_type", "")),
        _to_text_value(params.get("filter_value", "")),
        _to_text_value(meta.get("target_id", "")),
        _to_text_value(meta.get("selector", "")),
        _to_text_value(meta.get("screen_state", "")),
    )


def _find_nearest_auto_context(auto_rows: List[Dict[str, object]], captured_at: pd.Timestamp, page_id: str) -> Dict[str, object]:
    if pd.isna(captured_at):
        return {}
    for item in reversed(auto_rows):
        item_ts = item.get("captured_at")
        if pd.isna(item_ts):
            continue
        delta = abs((captured_at - item_ts).total_seconds())
        if delta > 8:
            continue
        if str(item.get("page_id", "")).strip() and str(item.get("page_id", "")).strip() != str(page_id).strip():
            continue
        return item
    return {}


def _build_short_selector(selector: str) -> str:
    raw = str(selector or "").strip()
    if not raw:
        return "-"
    parts = [part.strip() for part in raw.split(">") if part.strip()]
    short_parts = parts[-2:] if len(parts) >= 2 else parts
    short_selector = " > ".join(short_parts)
    return short_selector[:160] if len(short_selector) > 160 else short_selector


def _build_identification_summary(custom_params: Dict[str, str], meta: Dict[str, object]) -> str:
    def has_value(key: str) -> bool:
        return bool(str(custom_params.get(key, "")).strip())

    if has_value("button_id"):
        return f"button_id={custom_params['button_id']}"
    if has_value("button_name"):
        return f"button_name={custom_params['button_name']}"
    if has_value("filter_type") and has_value("filter_value"):
        return f"filter_type={custom_params['filter_type']}, filter_value={custom_params['filter_value']}"
    if has_value("content_id") and has_value("content_name"):
        return f"content_id={custom_params['content_id']}, content_name={custom_params['content_name']}"
    if has_value("banner_id"):
        return f"banner_id={custom_params['banner_id']}"
    short_selector = _build_short_selector(str(meta.get("selector", "")))
    if short_selector != "-":
        return f"selector={short_selector}"
    return "식별 조건 보완 필요"


def _evaluate_inspection_status(meta: Dict[str, object], custom_items: List[Tuple[str, str]]) -> str:
    target_id = str(meta.get("target_id", "")).strip().lower()
    if not target_id or target_id in {"-", "(not set)", "not set", "none", "null"}:
        return "식별자 보완 필요"
    if not custom_items:
        return "확인 필요"
    bbox_values = [
        str(meta.get("bbox_x", "0")).strip(),
        str(meta.get("bbox_y", "0")).strip(),
        str(meta.get("bbox_width", "0")).strip(),
        str(meta.get("bbox_height", "0")).strip(),
    ]
    has_bbox = False
    try:
        has_bbox = float(bbox_values[2]) > 0 and float(bbox_values[3]) > 0
    except Exception:
        has_bbox = False
    screenshot_file = str(meta.get("screenshot_file", "")).strip()
    selector = str(meta.get("selector", "")).strip()
    if not screenshot_file or not selector or not has_bbox:
        return "매핑 오류 의심"
    return "정상"


def build_event_definition_exports(
    raw_debug_df: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]],
) -> Tuple[pd.DataFrame, pd.DataFrame, List[str]]:
    event_def_columns = [
        "no",
        "action",
        "object",
        "event name",
        "params.key",
        "req",
        "params.value (as-is)",
        "params.value",
        "class-attribute",
        "status",
        "page_id",
        "section_name",
        "target_id",
        "selector",
        "screen_state",
        "annotation_no",
        "screenshot_file",
    ]
    support_columns = [
        "page_id",
        "section_name",
        "event no",
        "annotation_no",
        "target_id",
        "selector",
        "screenshot_file",
        "raw_screenshot_file",
        "bbox_x",
        "bbox_y",
        "bbox_width",
        "bbox_height",
        "screen_state",
        "note",
    ]
    if raw_debug_df.empty:
        return pd.DataFrame(columns=event_def_columns), pd.DataFrame(columns=support_columns), []

    work = raw_debug_df.copy()
    work["captured_at"] = pd.to_datetime(work["captured_at"], errors="coerce")
    work = work.sort_values("captured_at")

    auto_rows: List[Dict[str, object]] = []
    for _, row in work[work["source"].astype(str) == "auto_crawl"].iterrows():
        params = row.get("params", {})
        if not isinstance(params, dict):
            params = {}
        auto_rows.append(
            {
                "captured_at": row.get("captured_at"),
                "page_id": _extract_page_id(_to_text_value(row.get("page_url", "")), params),
                "section_name": _to_text_value(params.get("section_name", "")),
                "target_id": _to_text_value(params.get("target_id", "")),
                "selector": _to_text_value(params.get("selector", "")),
                "class_attribute": _to_text_value(params.get("class_attribute", "")),
                "screen_state": _to_text_value(params.get("screen_state", "")),
                "selector_screenshot": _to_text_value(params.get("selector_screenshot", "")),
                "raw_screenshot_file": _to_text_value(params.get("raw_screenshot_file", "")),
                "bbox_x": _to_text_value(params.get("bbox_x", 0)),
                "bbox_y": _to_text_value(params.get("bbox_y", 0)),
                "bbox_width": _to_text_value(params.get("bbox_width", 0)),
                "bbox_height": _to_text_value(params.get("bbox_height", 0)),
            }
        )

    event_rows: List[Dict[str, object]] = []
    support_rows: List[Dict[str, object]] = []
    asset_paths: List[str] = []
    seen_identities: set[Tuple[str, ...]] = set()
    event_no = 0

    ga_rows = work[work["source"].astype(str) == "ga_hit"]
    for _, row in ga_rows.iterrows():
        event_name = str(row.get("event_name", "")).strip()
        if not event_name:
            continue
        params = row.get("params", {})
        if not isinstance(params, dict):
            params = {}
        custom_event_name, custom_items = extract_custom_event_payload(event_name, params, schemas)
        if not custom_event_name:
            continue
        top_custom_items, _ = split_custom_param_items(custom_items)
        page_id = _extract_page_id(_to_text_value(row.get("page_url", "")), params)
        auto_ctx = _find_nearest_auto_context(auto_rows, row.get("captured_at"), page_id)
        meta = _build_event_meta(row, params, auto_ctx=auto_ctx)
        custom_param_map = {key: value for key, value in top_custom_items}
        identity = _build_event_identity(custom_event_name, custom_param_map, meta)
        if identity in seen_identities:
            continue
        seen_identities.add(identity)
        event_no += 1

        status = _evaluate_inspection_status(meta, top_custom_items)
        action, obj = _infer_action_object(custom_event_name)
        schema_obj = schemas.get(event_name, {})
        export_params = [(key, "Y" if key in {str(v).strip() for v in schema_obj.get("required", []) if str(v).strip()} else "N") for key, _ in top_custom_items]
        annotation_no = str(event_no) if str(meta.get("screenshot_file", "")).strip() else ""
        screenshot_file = str(meta.get("screenshot_file", "")).strip()
        raw_screenshot_file = str(meta.get("raw_screenshot_file", "")).strip()
        if screenshot_file:
            asset_paths.append(screenshot_file)
        if raw_screenshot_file:
            asset_paths.append(raw_screenshot_file)

        event_rows.append(
            {
                "no": event_no,
                "action": action,
                "object": obj,
                "event name": custom_event_name,
                "params.key": "",
                "req": "",
                "params.value (as-is)": "",
                "params.value": "",
                "class-attribute": meta["class_attribute"],
                "status": status,
                "page_id": meta["page_id"],
                "section_name": meta["section_name"],
                "target_id": meta["target_id"],
                "selector": meta["selector"],
                "screen_state": meta["screen_state"],
                "annotation_no": annotation_no,
                "screenshot_file": screenshot_file,
            }
        )
        for param_key, req_flag in export_params:
            value_text = custom_param_map.get(param_key, "")
            event_rows.append(
                {
                    "no": "",
                    "action": "",
                    "object": "",
                    "event name": "",
                    "params.key": param_key,
                    "req": req_flag,
                    "params.value (as-is)": value_text,
                    "params.value": value_text,
                    "class-attribute": meta["class_attribute"],
                    "status": status,
                    "page_id": meta["page_id"],
                    "section_name": meta["section_name"],
                    "target_id": meta["target_id"],
                    "selector": meta["selector"],
                    "screen_state": meta["screen_state"],
                    "annotation_no": annotation_no,
                    "screenshot_file": screenshot_file,
                }
            )

        support_rows.append(
            {
                "page_id": meta["page_id"],
                "section_name": meta["section_name"],
                "event no": event_no,
                "annotation_no": annotation_no,
                "target_id": meta["target_id"],
                "selector": meta["selector"],
                "screenshot_file": screenshot_file or "-",
                "raw_screenshot_file": raw_screenshot_file or "-",
                "bbox_x": meta["bbox_x"] or 0,
                "bbox_y": meta["bbox_y"] or 0,
                "bbox_width": meta["bbox_width"] or 0,
                "bbox_height": meta["bbox_height"] or 0,
                "screen_state": meta["screen_state"],
                "note": status if status != "정상" else "-",
            }
        )

    return (
        pd.DataFrame(event_rows, columns=event_def_columns),
        pd.DataFrame(support_rows, columns=support_columns),
        sorted(dict.fromkeys([p for p in asset_paths if str(p).strip()])),
    )


def build_event_review_detail_df(
    raw_debug_df: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]],
) -> pd.DataFrame:
    detail_columns = [
        "event no",
        "annotation_no",
        "custom_event",
        "custom_parameters",
        "extra_custom_parameters",
        "event_name",
        "action",
        "object",
        "status",
        "page_id",
        "section_name",
        "target_id",
        "selector",
        "screen_state",
        "class_attribute",
        "identification",
        "short_selector",
        "key_params",
        "screenshot_file",
        "raw_screenshot_file",
        "bbox_x",
        "bbox_y",
        "bbox_width",
        "bbox_height",
        "raw_payload_json",
    ]
    if raw_debug_df.empty:
        return pd.DataFrame(columns=detail_columns)

    work = raw_debug_df.copy()
    work["captured_at"] = pd.to_datetime(work["captured_at"], errors="coerce")
    work = work.sort_values("captured_at")

    auto_rows: List[Dict[str, object]] = []
    for _, row in work[work["source"].astype(str) == "auto_crawl"].iterrows():
        params = row.get("params", {})
        if not isinstance(params, dict):
            params = {}
        auto_rows.append(
            {
                "captured_at": row.get("captured_at"),
                "page_id": _extract_page_id(_to_text_value(row.get("page_url", "")), params),
                "section_name": _to_text_value(params.get("section_name", "")),
                "target_id": _to_text_value(params.get("target_id", "")),
                "selector": _to_text_value(params.get("selector", "")),
                "class_attribute": _to_text_value(params.get("class_attribute", "")),
                "screen_state": _to_text_value(params.get("screen_state", "")),
                "selector_screenshot": _to_text_value(params.get("selector_screenshot", "")),
                "raw_screenshot_file": _to_text_value(params.get("raw_screenshot_file", "")),
                "bbox_x": _to_text_value(params.get("bbox_x", 0)),
                "bbox_y": _to_text_value(params.get("bbox_y", 0)),
                "bbox_width": _to_text_value(params.get("bbox_width", 0)),
                "bbox_height": _to_text_value(params.get("bbox_height", 0)),
            }
        )

    rows: List[Dict[str, object]] = []
    seen_identities: set[Tuple[str, ...]] = set()
    event_no = 0
    ga_rows = work[work["source"].astype(str) == "ga_hit"]
    for _, row in ga_rows.iterrows():
        event_name = str(row.get("event_name", "")).strip()
        if not event_name:
            continue
        params = row.get("params", {})
        if not isinstance(params, dict):
            params = {}
        custom_event_name, custom_items = extract_custom_event_payload(event_name, params, schemas)
        if not custom_event_name:
            continue
        top_custom_items, extra_custom_items = split_custom_param_items(custom_items)
        page_id = _extract_page_id(_to_text_value(row.get("page_url", "")), params)
        auto_ctx = _find_nearest_auto_context(auto_rows, row.get("captured_at"), page_id)
        meta = _build_event_meta(row, params, auto_ctx=auto_ctx)
        custom_param_map = {key: value for key, value in top_custom_items}
        identity = _build_event_identity(custom_event_name, custom_param_map, meta)
        if identity in seen_identities:
            continue
        seen_identities.add(identity)
        event_no += 1

        status = _evaluate_inspection_status(meta, top_custom_items)
        action, obj = _infer_action_object(custom_event_name)
        key_params = [f"{param_key}={value_text}" for param_key, value_text in top_custom_items if value_text]
        extra_params = [f"{param_key}={value_text}" for param_key, value_text in extra_custom_items if value_text]
        identification_text = _build_identification_summary(custom_param_map, meta)
        rows.append(
            {
                "event no": event_no,
                "annotation_no": str(event_no) if str(meta.get("screenshot_file", "")).strip() else "",
                "custom_event": custom_event_name,
                "custom_parameters": " | ".join(key_params) if key_params else "-",
                "extra_custom_parameters": " | ".join(extra_params) if extra_params else "-",
                "event_name": custom_event_name,
                "action": action,
                "object": obj,
                "status": status,
                "page_id": meta["page_id"],
                "section_name": meta["section_name"],
                "target_id": meta["target_id"],
                "selector": meta["selector"],
                "screen_state": meta["screen_state"],
                "class_attribute": meta["class_attribute"],
                "identification": identification_text,
                "short_selector": _build_short_selector(meta["selector"]),
                "key_params": " | ".join(key_params) if key_params else "-",
                "custom_parameters": " | ".join(key_params) if key_params else "-",
                "screenshot_file": str(meta.get("screenshot_file", "")).strip() or "-",
                "raw_screenshot_file": str(meta.get("raw_screenshot_file", "")).strip() or "-",
                "bbox_x": meta["bbox_x"] or 0,
                "bbox_y": meta["bbox_y"] or 0,
                "bbox_width": meta["bbox_width"] or 0,
                "bbox_height": meta["bbox_height"] or 0,
                "raw_payload_json": json.dumps(params, ensure_ascii=False, sort_keys=True),
            }
        )

    return pd.DataFrame(rows, columns=detail_columns)


def build_event_definition_bundle_html(
    meta: Dict[str, object],
    event_definition_df: pd.DataFrame,
    event_support_df: pd.DataFrame,
    event_detail_df: pd.DataFrame,
) -> str:
    event_records = event_detail_df.to_dict("records") if not event_detail_df.empty else []
    cards_html: List[str] = []
    nav_html: List[str] = []
    for idx, row in enumerate(event_records):
        event_no = str(row.get("event no", "")).strip()
        annotation_no = str(row.get("annotation_no", event_no)).strip() or event_no
        custom_event = html.escape(str(row.get("custom_event", row.get("event_name", ""))))
        status = html.escape(str(row.get("status", "")))
        status_class = (
            "warn" if status in {"확인 필요", "식별자 보완 필요"} else "error" if status == "매핑 오류 의심" else "ok"
        )
        custom_parameters = html.escape(str(row.get("custom_parameters", "-")))
        cards_html.append(
            f"""
            <article
              class="event-card{' active' if idx == 0 else ''}"
              id="event-{html.escape(event_no)}"
              data-event-no="{html.escape(event_no)}"
              data-annotation-no="{html.escape(annotation_no)}"
            >
              <button type="button" class="card-button" data-select-event="{html.escape(event_no)}">
                <div class="event-card-head">
                  <div class="event-number">#{html.escape(annotation_no)}</div>
                  <div>
                    <h2>{custom_event}</h2>
                    <p class="event-ident">{html.escape(str(row.get("identification", "-")))}</p>
                  </div>
                  <span class="status {status_class}">{status}</span>
                </div>
                <dl class="event-meta">
                  <dt>target_id</dt><dd>{html.escape(str(row.get("target_id", "-")))}</dd>
                  <dt>section_name</dt><dd>{html.escape(str(row.get("section_name", "-")))}</dd>
                  <dt>page_id</dt><dd>{html.escape(str(row.get("page_id", "-")))}</dd>
                  <dt>custom_parameters</dt><dd>{custom_parameters}</dd>
                </dl>
              </button>
              <details class="event-debug">
                <summary>상세 보기</summary>
                <dl class="event-debug-meta">
                  <dt>extra custom</dt><dd>{html.escape(str(row.get("extra_custom_parameters", "-")))}</dd>
                  <dt>short selector</dt><dd>{html.escape(str(row.get("short_selector", "-")))}</dd>
                  <dt>raw selector</dt><dd>{html.escape(str(row.get("selector", "-")))}</dd>
                  <dt>bbox</dt><dd>{html.escape(f"x={row.get('bbox_x', 0)}, y={row.get('bbox_y', 0)}, w={row.get('bbox_width', 0)}, h={row.get('bbox_height', 0)}")}</dd>
                </dl>
                <pre>{html.escape(str(row.get("raw_payload_json", "{}")))}</pre>
              </details>
            </article>
            """
        )
        nav_html.append(
            f'<button type="button" class="annotation-chip{" active" if idx == 0 else ""}" data-select-event="{html.escape(event_no)}">{html.escape(annotation_no)}</button>'
        )

    event_payload = json.dumps(event_records, ensure_ascii=False)
    summary_count = int((event_definition_df["no"].astype(str).str.strip() != "").sum()) if not event_definition_df.empty else 0

    return f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Event Definition Review Report</title>
  <style>
    :root {{
      --bg: #f3efe8;
      --panel: #fffdf8;
      --line: #d8d1c3;
      --ink: #18212f;
      --muted: #6b7280;
      --accent: #0f766e;
      --accent-weak: rgba(15, 118, 110, 0.14);
      --warn: #9a3412;
      --warn-weak: rgba(154, 52, 18, 0.12);
      --error: #b91c1c;
      --error-weak: rgba(185, 28, 28, 0.12);
    }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; background: linear-gradient(180deg, #f8f4ec 0%, #ede7dd 100%); color: var(--ink); font-family: "Helvetica Neue", Arial, sans-serif; }}
    a {{ color: inherit; }}
    main {{ max-width: 1600px; margin: 0 auto; padding: 24px; }}
    .hero {{ background: var(--panel); border: 1px solid var(--line); border-radius: 18px; padding: 22px; margin-bottom: 18px; }}
    .hero h1 {{ margin: 0 0 10px; font-size: 28px; }}
    .hero p {{ margin: 6px 0; }}
    .file-links {{ display: flex; flex-wrap: wrap; gap: 12px; margin-top: 14px; }}
    .file-links a {{ text-decoration: none; border: 1px solid var(--line); border-radius: 999px; padding: 8px 12px; background: #fff; }}
    .layout {{ display: grid; grid-template-columns: minmax(560px, 1.15fr) minmax(380px, 0.85fr); gap: 18px; align-items: start; }}
    .panel {{ background: var(--panel); border: 1px solid var(--line); border-radius: 18px; }}
    .annotation-panel {{ position: sticky; top: 20px; padding: 18px; }}
    .annotation-header {{ display: flex; justify-content: space-between; gap: 12px; align-items: flex-start; margin-bottom: 12px; }}
    .annotation-header h2 {{ margin: 0; font-size: 20px; }}
    .annotation-sub {{ color: var(--muted); font-size: 14px; }}
    .annotation-stage {{ position: relative; border: 1px solid var(--line); border-radius: 16px; overflow: hidden; background: #f8fafc; min-height: 420px; }}
    .annotation-stage img {{ display: block; width: 100%; height: auto; }}
    .overlay-layer {{ position: absolute; inset: 0; }}
    .overlay-box {{
      position: absolute;
      border: 2px solid rgba(148, 163, 184, 0.6);
      background: rgba(148, 163, 184, 0.12);
      border-radius: 8px;
      cursor: pointer;
      transition: all 160ms ease;
    }}
    .overlay-box.active {{
      border-color: var(--accent);
      background: var(--accent-weak);
      box-shadow: 0 0 0 2px rgba(255,255,255,0.85) inset;
      z-index: 3;
    }}
    .overlay-label {{
      position: absolute;
      top: -12px;
      left: -2px;
      min-width: 28px;
      height: 28px;
      padding: 0 8px;
      border-radius: 999px;
      display: flex;
      align-items: center;
      justify-content: center;
      background: rgba(100, 116, 139, 0.82);
      color: #fff;
      font-size: 12px;
      font-weight: 700;
    }}
    .overlay-box.active .overlay-label {{ background: var(--accent); }}
    .annotation-chips {{ display: flex; flex-wrap: wrap; gap: 8px; margin: 14px 0; }}
    .annotation-chip {{
      border: 1px solid var(--line);
      background: #fff;
      color: var(--ink);
      border-radius: 999px;
      padding: 7px 11px;
      cursor: pointer;
      opacity: 0.55;
    }}
    .annotation-chip.active {{ opacity: 1; border-color: var(--accent); color: var(--accent); }}
    .annotation-footer {{ display: flex; flex-wrap: wrap; gap: 10px; align-items: center; font-size: 14px; color: var(--muted); }}
    .annotation-footer a {{ text-decoration: none; border-bottom: 1px solid currentColor; }}
    .cards-panel {{ padding: 18px; }}
    .cards-panel h2 {{ margin: 0 0 10px; font-size: 20px; }}
    .cards-list {{ display: grid; gap: 12px; }}
    .event-card {{ border: 1px solid var(--line); border-radius: 16px; overflow: hidden; background: #fff; transition: border-color 140ms ease, box-shadow 140ms ease, transform 140ms ease; }}
    .event-card.active {{ border-color: var(--accent); box-shadow: 0 14px 34px rgba(15, 118, 110, 0.12); transform: translateY(-1px); }}
    .card-button {{ width: 100%; background: transparent; border: 0; text-align: left; padding: 16px; cursor: pointer; color: inherit; }}
    .event-card-head {{ display: grid; grid-template-columns: auto 1fr auto; gap: 12px; align-items: start; }}
    .event-card-head h2 {{ margin: 0 0 6px; font-size: 18px; }}
    .event-ident {{ margin: 0; color: var(--muted); font-size: 14px; }}
    .event-number {{ width: 42px; height: 42px; border-radius: 999px; background: #ecfeff; color: var(--accent); display: flex; align-items: center; justify-content: center; font-weight: 700; }}
    .status {{ border-radius: 999px; padding: 5px 10px; font-size: 12px; font-weight: 700; white-space: nowrap; }}
    .status.ok {{ background: var(--accent-weak); color: var(--accent); }}
    .status.warn {{ background: var(--warn-weak); color: var(--warn); }}
    .status.error {{ background: var(--error-weak); color: var(--error); }}
    .event-meta, .event-debug-meta {{ display: grid; grid-template-columns: 124px 1fr; gap: 8px 12px; margin: 14px 0 0; }}
    .event-meta dt, .event-debug-meta dt {{ color: var(--muted); }}
    .event-meta dd, .event-debug-meta dd {{ margin: 0; word-break: break-word; }}
    .event-debug {{ border-top: 1px solid var(--line); padding: 0 16px 16px; }}
    .event-debug summary {{ cursor: pointer; padding-top: 12px; color: var(--muted); }}
    .event-debug pre {{ margin: 12px 0 0; padding: 12px; border-radius: 12px; background: #0f172a; color: #e2e8f0; overflow: auto; font-size: 12px; }}
    .empty-stage {{ min-height: 420px; display: flex; align-items: center; justify-content: center; color: var(--muted); }}
    @media (max-width: 1180px) {{
      .layout {{ grid-template-columns: 1fr; }}
      .annotation-panel {{ position: static; }}
    }}
  </style>
</head>
<body>
  <main>
    <section class="hero">
      <h1>Event Definition Review Report</h1>
      <p>HTML 하나만 열어도 어느 위치를 테스트했는지, 어떤 custom event가 수집됐는지, 핵심 custom parameter가 무엇인지 빠르게 검수할 수 있도록 구성했습니다.</p>
      <p>project={html.escape(str(meta.get("project_slug", "-")))} | session={html.escape(str(meta.get("session_id", "-")))} | run_id={html.escape(str(meta.get("run_id", "-")))} | events={summary_count}</p>
      <p>scope={html.escape(str(meta.get("bundle_scope", "custom_event_review")))} | definition={html.escape(str(meta.get("definition_file_name", "-") or "-"))} | definition_events={html.escape(str(meta.get("definition_event_count", 0)))}</p>
      <div class="file-links">
        <a href="event_definition.csv">1. 이벤트 정의 CSV</a>
        <a href="event_support.csv">2. annotation 지원 CSV</a>
        <a href="event_review_details.csv">3. 이벤트별 상세 검수 산출물</a>
        <a href="qa_review.xlsx">4. QA Review Excel</a>
      </div>
    </section>
    <section class="layout">
      <section class="panel annotation-panel">
        <div class="annotation-header">
          <div>
            <h2>Annotation View</h2>
            <div class="annotation-sub">왼쪽은 사용자가 실제로 본 viewport 캡처 기반 bbox 오버레이, 오른쪽 카드는 이벤트 상세입니다.</div>
          </div>
          <div class="annotation-sub" id="annotation-meta">선택된 이벤트 없음</div>
        </div>
        <div class="annotation-stage" id="annotation-stage">
          <img id="annotation-image" alt="annotation stage" />
          <div class="overlay-layer" id="overlay-layer"></div>
          <div class="empty-stage" id="annotation-empty">표시할 annotation 이미지가 없습니다.</div>
        </div>
        <div class="annotation-chips" id="annotation-nav">
          {''.join(nav_html)}
        </div>
        <div class="annotation-footer">
          <a id="annotation-link" href="#" target="_blank" rel="noopener">annotation 이미지 열기</a>
          <a id="raw-link" href="#" target="_blank" rel="noopener">원본 viewport 스크린샷 열기</a>
        </div>
      </section>
      <section class="panel cards-panel">
        <h2>Event Cards</h2>
        <div class="cards-list">
          {''.join(cards_html)}
        </div>
      </section>
    </section>
  </main>
  <script>
    const EVENT_DATA = {event_payload};
    const annotationImage = document.getElementById("annotation-image");
    const annotationEmpty = document.getElementById("annotation-empty");
    const overlayLayer = document.getElementById("overlay-layer");
    const annotationMeta = document.getElementById("annotation-meta");
    const annotationLink = document.getElementById("annotation-link");
    const rawLink = document.getElementById("raw-link");
    const navButtons = Array.from(document.querySelectorAll("[data-select-event]"));
    const cards = Array.from(document.querySelectorAll(".event-card"));
    let activeEventNo = EVENT_DATA.length ? String(EVENT_DATA[0]["event no"]) : "";

    function getEventRecord(eventNo) {{
      return EVENT_DATA.find((item) => String(item["event no"]) === String(eventNo)) || null;
    }}

    function getGroupItems(rawScreenshotFile) {{
      return EVENT_DATA.filter((item) => String(item.raw_screenshot_file || "") === String(rawScreenshotFile || ""));
    }}

    function renderOverlay(groupItems, activeItem) {{
      overlayLayer.innerHTML = "";
      if (!groupItems.length) return;
      const naturalWidth = annotationImage.naturalWidth || annotationImage.clientWidth || 1;
      const naturalHeight = annotationImage.naturalHeight || annotationImage.clientHeight || 1;
      const displayWidth = annotationImage.clientWidth || 1;
      const displayHeight = annotationImage.clientHeight || 1;
      groupItems.forEach((item) => {{
        const x = Number(item.bbox_x || 0);
        const y = Number(item.bbox_y || 0);
        const w = Number(item.bbox_width || 0);
        const h = Number(item.bbox_height || 0);
        if (!w || !h) return;
        const box = document.createElement("button");
        box.type = "button";
        box.className = "overlay-box" + (String(item["event no"]) === String(activeItem["event no"]) ? " active" : "");
        box.style.left = `${{(x / naturalWidth) * displayWidth}}px`;
        box.style.top = `${{(y / naturalHeight) * displayHeight}}px`;
        box.style.width = `${{(w / naturalWidth) * displayWidth}}px`;
        box.style.height = `${{(h / naturalHeight) * displayHeight}}px`;
        box.dataset.selectEvent = String(item["event no"]);
        const label = document.createElement("span");
        label.className = "overlay-label";
        label.textContent = String(item.annotation_no || item["event no"] || "");
        box.appendChild(label);
        box.addEventListener("click", () => setActiveEvent(String(item["event no"]), true));
        overlayLayer.appendChild(box);
      }});
    }}

    function syncSelection(eventNo) {{
      cards.forEach((card) => {{
        card.classList.toggle("active", card.dataset.eventNo === String(eventNo));
      }});
      navButtons.forEach((button) => {{
        button.classList.toggle("active", button.dataset.selectEvent === String(eventNo));
      }});
    }}

    function setActiveEvent(eventNo, scrollCard) {{
      const item = getEventRecord(eventNo);
      if (!item) return;
      activeEventNo = String(eventNo);
      syncSelection(activeEventNo);
      const rawScreenshotFile = String(item.raw_screenshot_file || "");
      const annotationFile = String(item.screenshot_file || "");
      const hasStageImage = rawScreenshotFile && rawScreenshotFile !== "-";
      annotationImage.style.display = hasStageImage ? "block" : "none";
      annotationEmpty.style.display = hasStageImage ? "none" : "flex";
      annotationMeta.textContent = `#${{item.annotation_no || item["event no"]}} · ${{item.custom_event || item.event_name || "-"}} · ${{item.status || "-"}}`;
      annotationLink.href = annotationFile && annotationFile !== "-" ? annotationFile : "#";
      rawLink.href = hasStageImage ? rawScreenshotFile : "#";
      if (!hasStageImage) {{
        overlayLayer.innerHTML = "";
        return;
      }}
      annotationImage.onload = () => {{
        renderOverlay(getGroupItems(rawScreenshotFile), item);
      }};
      annotationImage.src = rawScreenshotFile;
      if (annotationImage.complete) {{
        renderOverlay(getGroupItems(rawScreenshotFile), item);
      }}
      if (scrollCard) {{
        const card = document.getElementById(`event-${{item["event no"]}}`);
        if (card) card.scrollIntoView({{ behavior: "smooth", block: "center" }});
      }}
    }}

    navButtons.forEach((button) => {{
      button.addEventListener("click", () => setActiveEvent(button.dataset.selectEvent, true));
    }});
    cards.forEach((card) => {{
      const button = card.querySelector(".card-button");
      if (button) {{
        button.addEventListener("click", () => setActiveEvent(card.dataset.eventNo, false));
      }}
    }});
    window.addEventListener("resize", () => {{
      const item = getEventRecord(activeEventNo);
      if (!item || !annotationImage.src) return;
      renderOverlay(getGroupItems(String(item.raw_screenshot_file || "")), item);
    }});
    if (EVENT_DATA.length) {{
      setActiveEvent(activeEventNo, false);
    }}
  </script>
</body>
</html>
"""


def _build_simple_html_table(df: pd.DataFrame, columns: List[str]) -> str:
    if df.empty:
        return "<p class='empty'>데이터가 없습니다.</p>"
    head = "".join([f"<th>{html.escape(str(col))}</th>" for col in columns])
    body_rows: List[str] = []
    for _, row in df[columns].fillna("-").iterrows():
        cells = "".join([f"<td>{html.escape(str(row.get(col, '-')))}</td>" for col in columns])
        body_rows.append(f"<tr>{cells}</tr>")
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body_rows)}</tbody></table>"


def build_definition_validation_bundle_html(
    meta: Dict[str, object],
    summary_df: pd.DataFrame,
    matches_df: pd.DataFrame,
    unmatched_df: pd.DataFrame,
    event_detail_df: pd.DataFrame,
) -> str:
    pass_count = int((summary_df.get("result", pd.Series(dtype=str)).astype(str) == "PASS").sum()) if not summary_df.empty else 0
    warn_count = int((summary_df.get("result", pd.Series(dtype=str)).astype(str) == "WARN").sum()) if not summary_df.empty else 0
    fail_count = int((summary_df.get("result", pd.Series(dtype=str)).astype(str) == "FAIL").sum()) if not summary_df.empty else 0
    matched_rows = int((pd.to_numeric(summary_df.get("matched_capture_count", 0), errors="coerce").fillna(0) > 0).sum()) if not summary_df.empty else 0
    missing_rows = int((pd.to_numeric(summary_df.get("matched_capture_count", 0), errors="coerce").fillna(0) == 0).sum()) if not summary_df.empty else 0
    summary_html = _build_simple_html_table(
        summary_df,
        [col for col in ["definition_row_id", "no", "event_name", "description", "section_name", "matched_capture_count", "result", "failure_reason", "missing_parameters", "extra_parameters"] if col in summary_df.columns],
    )
    matches_html = _build_simple_html_table(
        matches_df,
        [col for col in ["definition_row_id", "no", "event_name", "description", "matched_at", "matched_identification", "match_score", "result", "failure_reason"] if col in matches_df.columns],
    )
    unmatched_html = _build_simple_html_table(
        unmatched_df,
        [col for col in ["event_name", "captured_at", "matched_identification", "reason"] if col in unmatched_df.columns],
    )
    card_html = _build_simple_html_table(
        event_detail_df.drop(columns=["raw_payload_json"], errors="ignore"),
        [col for col in ["event no", "custom_event", "identification", "target_id", "section_name", "page_id", "status", "custom_parameters"] if col in event_detail_df.columns],
    )
    return f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Definition Validation Report</title>
  <style>
    body {{ margin: 0; background: #f5f1ea; color: #1f2937; font-family: Arial, sans-serif; }}
    main {{ max-width: 1480px; margin: 0 auto; padding: 24px; }}
    .hero, .panel {{ background: #fffdf8; border: 1px solid #d7d1c7; border-radius: 18px; padding: 20px; margin-bottom: 18px; }}
    .metrics {{ display: grid; grid-template-columns: repeat(6, minmax(0, 1fr)); gap: 12px; margin-top: 14px; }}
    .metric {{ background: #fff; border: 1px solid #e5dfd4; border-radius: 14px; padding: 12px; }}
    .metric strong {{ display: block; font-size: 24px; margin-top: 6px; }}
    .file-links {{ display: flex; flex-wrap: wrap; gap: 10px; margin-top: 14px; }}
    .file-links a {{ text-decoration: none; color: #0f766e; border-bottom: 1px solid currentColor; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 14px; }}
    th, td {{ border-bottom: 1px solid #ece5d9; padding: 10px 8px; text-align: left; vertical-align: top; }}
    th {{ background: #f8f4ec; }}
    .empty {{ color: #6b7280; }}
    @media (max-width: 1024px) {{ .metrics {{ grid-template-columns: repeat(2, minmax(0, 1fr)); }} }}
  </style>
</head>
<body>
  <main>
    <section class="hero">
      <h1>Definition Validation Report</h1>
      <p>정의서의 각 row가 이번 세션에서 충족됐는지를 기준으로 판정한 리포트입니다.</p>
      <p>session id={html.escape(str(meta.get("session_id", "-")))} | definition file={html.escape(str(meta.get("definition_file_name", "-") or "-"))}</p>
      <div class="metrics">
        <div class="metric">definition rows<strong>{len(summary_df)}</strong></div>
        <div class="metric">matched rows<strong>{matched_rows}</strong></div>
        <div class="metric">missing rows<strong>{missing_rows}</strong></div>
        <div class="metric">PASS<strong>{pass_count}</strong></div>
        <div class="metric">WARN<strong>{warn_count}</strong></div>
        <div class="metric">FAIL<strong>{fail_count}</strong></div>
      </div>
      <div class="file-links">
        <a href="definition_validation_summary.csv">definition_validation_summary.csv</a>
        <a href="definition_validation_matches.csv">definition_validation_matches.csv</a>
        <a href="definition_validation_unmatched.csv">definition_validation_unmatched.csv</a>
        <a href="qa_review.xlsx">qa_review.xlsx</a>
      </div>
    </section>
    <section class="panel">
      <h2>Definition Row Summary</h2>
      {summary_html}
    </section>
    <section class="panel">
      <h2>Definition Row Matches</h2>
      {matches_html}
    </section>
    <section class="panel">
      <h2>Unmatched Captures</h2>
      {unmatched_html}
    </section>
    <section class="panel">
      <h2>Captured Event Cards (Secondary)</h2>
      {card_html}
    </section>
  </main>
</body>
</html>
"""


def save_event_definition_bundle(
    raw_debug_df: pd.DataFrame,
    schemas: Dict[str, Dict[str, object]],
    qa_report_xlsx: bytes,
    target_url: str,
    session_id: str,
    scoped_event_names: set[str] | None = None,
    definition_file_name: str = "",
    bundle_scope: str = "custom_event_review",
    persist_run: bool = True,
) -> Tuple[bytes, bytes, pd.DataFrame, pd.DataFrame, Dict[str, object]]:
    scoped_debug_df = filter_debug_df_for_event_scope(raw_debug_df, scoped_event_names)
    event_definition_df, event_support_df, asset_paths = build_event_definition_exports(scoped_debug_df, schemas)
    event_detail_df = build_event_review_detail_df(scoped_debug_df, schemas)
    definition_validation_summary_df, definition_validation_matches_df, definition_validation_unmatched_df = build_definition_validation_frames(
        build_realtime_event_rows(
            filter_debug_events_by_view(scoped_debug_df, "히트(collect)"),
            allowed_events=list(scoped_event_names) if scoped_event_names else list(schemas.keys()),
            unknown_event_policy="무시",
            schemas=schemas,
            scoped_event_names=scoped_event_names,
        ) if str(bundle_scope or "").strip() == "definition_validation" else pd.DataFrame(),
        schemas,
    ) if str(bundle_scope or "").strip() == "definition_validation" else (pd.DataFrame(), pd.DataFrame(), pd.DataFrame())
    event_detail_export_df = (
        event_detail_df.drop(columns=["raw_payload_json"], errors="ignore").copy()
        if not event_detail_df.empty
        else event_detail_df.copy()
    )
    qa_review_xlsx = build_qa_review_excel_bytes(
        event_definition_df=event_definition_df,
        event_detail_df=event_detail_df,
        definition_validation_summary_df=definition_validation_summary_df,
        bundle_scope=bundle_scope,
    )
    log_definition_xlsx = build_log_definition_excel_bytes(
        event_definition_df=event_definition_df,
        event_detail_df=event_detail_df,
        schemas=schemas,
        definition_matches_df=definition_validation_matches_df,
    )
    qa_summary_df = build_qa_summary_export_df(
        event_detail_df=event_detail_df,
        definition_validation_summary_df=definition_validation_summary_df,
        bundle_scope=bundle_scope,
    )
    qa_debug_df = build_qa_debug_export_df(
        event_detail_df=event_detail_df,
        definition_validation_matches_df=definition_validation_matches_df,
        definition_validation_unmatched_df=definition_validation_unmatched_df,
        bundle_scope=bundle_scope,
    )
    qa_summary_csv = dataframe_to_csv_bytes(qa_summary_df)
    qa_debug_csv = dataframe_to_csv_bytes(qa_debug_df)
    compare_key = _build_url_compare_key(target_url)
    url_hash = _build_url_hash(compare_key)
    run_id = datetime.now(tz=ZoneInfo("UTC")).strftime("%Y%m%d_%H%M%S_%f")
    exports_root = get_active_project_paths()["exports_dir"] / "event_definition_history" / url_hash
    run_dir = exports_root / run_id
    event_definition_csv = dataframe_to_csv_bytes(event_definition_df)
    event_support_csv = dataframe_to_csv_bytes(event_support_df)
    event_detail_csv = dataframe_to_csv_bytes(event_detail_export_df)
    definition_validation_summary_csv = dataframe_to_csv_bytes(definition_validation_summary_df)
    definition_validation_matches_csv = dataframe_to_csv_bytes(definition_validation_matches_df)
    definition_validation_unmatched_csv = dataframe_to_csv_bytes(definition_validation_unmatched_df)

    baseline_path = exports_root / "baseline_run_id.txt"
    latest_path = exports_root / "latest_run_id.txt"
    auto_baseline = bool(persist_run and (not baseline_path.exists()))

    meta = {
        "project_slug": get_active_project_slug(),
        "target_url": target_url,
        "url_compare_key": compare_key,
        "url_hash": url_hash,
        "saved_at": datetime.now(tz=ZoneInfo("UTC")).isoformat(),
        "run_id": run_id,
        "session_id": session_id,
        "event_rows": int(len(definition_validation_summary_df)) if str(bundle_scope or "").strip() == "definition_validation" else int(len(event_definition_df)),
        "event_count": int(len(definition_validation_summary_df)) if str(bundle_scope or "").strip() == "definition_validation" else int((event_definition_df["no"].astype(str).str.strip() != "").sum()) if not event_definition_df.empty else 0,
        "support_rows": int(len(definition_validation_matches_df)) if str(bundle_scope or "").strip() == "definition_validation" else int(len(event_support_df)),
        "detail_rows": int(len(definition_validation_summary_df)) if str(bundle_scope or "").strip() == "definition_validation" else int(len(event_detail_export_df)),
        "bundle_scope": str(bundle_scope or "custom_event_review").strip() or "custom_event_review",
        "definition_file_name": str(definition_file_name or "").strip(),
        "definition_event_count": int(len(scoped_event_names or [])) if scoped_event_names else 0,
        "definition_row_count": int(len(definition_validation_summary_df)),
        "matched_row_count": int((pd.to_numeric(definition_validation_summary_df.get("matched_capture_count", 0), errors="coerce").fillna(0) > 0).sum()) if not definition_validation_summary_df.empty else 0,
        "missing_row_count": int((pd.to_numeric(definition_validation_summary_df.get("matched_capture_count", 0), errors="coerce").fillna(0) == 0).sum()) if not definition_validation_summary_df.empty else 0,
        "files": {
            "final_report": "qa_result.xlsx",
            "log_definition_review": "log_definition_review.xlsx",
            "summary_csv": "qa_summary.csv",
            "debug_csv": "qa_debug_details.csv",
        },
        "auto_baseline": auto_baseline,
    }
    if str(bundle_scope or "").strip() == "definition_validation":
        meta["files"].update(
            {
                "definition_summary_source": "definition_validation_summary.csv",
            }
        )
    index_html = (
        build_definition_validation_bundle_html(
            meta,
            definition_validation_summary_df,
            definition_validation_matches_df,
            definition_validation_unmatched_df,
            event_detail_df,
        )
        if str(bundle_scope or "").strip() == "definition_validation"
        else build_event_definition_bundle_html(meta, event_definition_df, event_support_df, event_detail_df)
    )
    meta_json = json.dumps(meta, ensure_ascii=False, indent=2).encode("utf-8")
    basic_bundle_bytes = build_zip_bytes(
        {
            "qa_result.xlsx": qa_report_xlsx,
            "log_definition_review.xlsx": log_definition_xlsx,
            "qa_summary.csv": qa_summary_csv,
            "qa_debug_details.csv": qa_debug_csv,
        }
    )
    debug_file_map = {
        "event_definition.csv": event_definition_csv,
        "event_support.csv": event_support_csv,
        "event_review_details.csv": event_detail_csv,
        "meta.json": meta_json,
        "index.html": index_html.encode("utf-8"),
    }
    if str(bundle_scope or "").strip() == "definition_validation":
        debug_file_map.update(
            {
                "definition_validation_summary.csv": definition_validation_summary_csv,
                "definition_validation_matches.csv": definition_validation_matches_csv,
                "definition_validation_unmatched.csv": definition_validation_unmatched_csv,
            }
        )
    if qa_report_xlsx:
        debug_file_map["qa_report.xlsx"] = qa_report_xlsx
    if qa_review_xlsx:
        debug_file_map["qa_review.xlsx"] = qa_review_xlsx
    if log_definition_xlsx:
        debug_file_map["log_definition_review.xlsx"] = log_definition_xlsx
    debug_bundle_out = io.BytesIO()
    with zipfile.ZipFile(debug_bundle_out, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for file_name, payload in debug_file_map.items():
            if payload:
                zf.writestr(file_name, payload)
        for rel_path in asset_paths:
            abs_path = BASE_DIR / rel_path
            if abs_path.exists() and abs_path.is_file():
                zf.write(abs_path, arcname=rel_path)
    debug_bundle_bytes = debug_bundle_out.getvalue()

    if persist_run:
        run_dir.mkdir(parents=True, exist_ok=True)
        (run_dir / "event_definition.csv").write_bytes(event_definition_csv)
        (run_dir / "event_support.csv").write_bytes(event_support_csv)
        (run_dir / "event_review_details.csv").write_bytes(event_detail_csv)
        if qa_review_xlsx:
            (run_dir / "qa_review.xlsx").write_bytes(qa_review_xlsx)
        if log_definition_xlsx:
            (run_dir / "log_definition_review.xlsx").write_bytes(log_definition_xlsx)
        if qa_summary_csv:
            (run_dir / "qa_summary.csv").write_bytes(qa_summary_csv)
        if qa_debug_csv:
            (run_dir / "qa_debug_details.csv").write_bytes(qa_debug_csv)
        if str(bundle_scope or "").strip() == "definition_validation":
            (run_dir / "definition_validation_summary.csv").write_bytes(definition_validation_summary_csv)
            (run_dir / "definition_validation_matches.csv").write_bytes(definition_validation_matches_csv)
            (run_dir / "definition_validation_unmatched.csv").write_bytes(definition_validation_unmatched_csv)
        (run_dir / "meta.json").write_bytes(meta_json)
        (run_dir / "index.html").write_text(index_html, encoding="utf-8")
        if not baseline_path.exists():
            baseline_path.write_text(run_id, encoding="utf-8")
        latest_path.write_text(run_id, encoding="utf-8")

    return basic_bundle_bytes, debug_bundle_bytes, event_definition_df, event_support_df, meta


def get_event_definition_history_root(target_url: str) -> tuple[Path, str, str]:
    compare_key = _build_url_compare_key(target_url)
    url_hash = _build_url_hash(compare_key)
    exports_root = get_active_project_paths()["exports_dir"] / "event_definition_history" / url_hash
    return exports_root, compare_key, url_hash


def list_event_definition_runs(target_url: str) -> pd.DataFrame:
    columns = ["run_id", "saved_at", "event_count", "session_id", "state", "target_url", "run_dir"]
    if not str(target_url or "").strip():
        return pd.DataFrame(columns=columns)
    exports_root, _, _ = get_event_definition_history_root(target_url)
    if not exports_root.exists():
        return pd.DataFrame(columns=columns)
    baseline_run_id = ""
    latest_run_id = ""
    baseline_path = exports_root / "baseline_run_id.txt"
    latest_path = exports_root / "latest_run_id.txt"
    if baseline_path.exists():
        baseline_run_id = baseline_path.read_text(encoding="utf-8").strip()
    if latest_path.exists():
        latest_run_id = latest_path.read_text(encoding="utf-8").strip()
    rows: List[Dict[str, object]] = []
    for run_dir in sorted([p for p in exports_root.iterdir() if p.is_dir()], reverse=True):
        meta_path = run_dir / "meta.json"
        if not meta_path.exists():
            continue
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
        except Exception:
            meta = {}
        run_id = str(meta.get("run_id", run_dir.name)).strip() or run_dir.name
        state_parts: List[str] = []
        if run_id == baseline_run_id:
            state_parts.append("Baseline")
        if run_id == latest_run_id:
            state_parts.append("Latest")
        if not state_parts:
            state_parts.append("Candidate")
        rows.append(
            {
                "run_id": run_id,
                "saved_at": str(meta.get("saved_at", "")).strip(),
                "event_count": int(meta.get("event_count", 0) or 0),
                "session_id": str(meta.get("session_id", "")).strip(),
                "state": " / ".join(state_parts),
                "target_url": str(meta.get("target_url", target_url)).strip(),
                "run_dir": str(run_dir),
            }
        )
    if not rows:
        return pd.DataFrame(columns=columns)
    out = pd.DataFrame(rows)
    out["saved_at_dt"] = pd.to_datetime(out["saved_at"], errors="coerce", utc=True)
    out = out.sort_values(["saved_at_dt", "run_id"], ascending=[False, False]).drop(columns=["saved_at_dt"])
    return out.reset_index(drop=True)


def set_event_definition_baseline(target_url: str, run_id: str) -> None:
    exports_root, _, _ = get_event_definition_history_root(target_url)
    exports_root.mkdir(parents=True, exist_ok=True)
    (exports_root / "baseline_run_id.txt").write_text(str(run_id or "").strip(), encoding="utf-8")


def load_event_definition_run_detail_df(target_url: str, run_id: str) -> pd.DataFrame:
    columns = [
        "event no",
        "custom_event",
        "identification",
        "target_id",
        "section_name",
        "page_id",
        "status",
        "custom_parameters",
    ]
    if not str(target_url or "").strip() or not str(run_id or "").strip():
        return pd.DataFrame(columns=columns)
    exports_root, _, _ = get_event_definition_history_root(target_url)
    detail_path = exports_root / str(run_id).strip() / "event_review_details.csv"
    if not detail_path.exists():
        return pd.DataFrame(columns=columns)
    try:
        return pd.read_csv(detail_path).fillna("")
    except Exception:
        return pd.DataFrame(columns=columns)


def compare_event_definition_runs(baseline_df: pd.DataFrame, candidate_df: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "change_type",
        "custom_event",
        "identification",
        "target_id",
        "section_name",
        "page_id",
        "baseline_status",
        "candidate_status",
        "baseline_parameters",
        "candidate_parameters",
    ]
    if baseline_df.empty and candidate_df.empty:
        return pd.DataFrame(columns=columns)

    def build_map(df: pd.DataFrame) -> Dict[Tuple[str, ...], Dict[str, object]]:
        mapping: Dict[Tuple[str, ...], Dict[str, object]] = {}
        if df.empty:
            return mapping
        for row in df.to_dict("records"):
            key = (
                str(row.get("custom_event", row.get("event_name", ""))).strip(),
                str(row.get("identification", "")).strip(),
                str(row.get("target_id", "")).strip(),
                str(row.get("section_name", "")).strip(),
                str(row.get("page_id", "")).strip(),
            )
            if any(key):
                mapping[key] = row
        return mapping

    base_map = build_map(baseline_df)
    cand_map = build_map(candidate_df)
    keys = sorted(set(base_map.keys()) | set(cand_map.keys()))
    rows: List[Dict[str, object]] = []
    for key in keys:
        base_row = base_map.get(key, {})
        cand_row = cand_map.get(key, {})
        if not base_row:
            change_type = "added"
        elif not cand_row:
            change_type = "removed"
        else:
            baseline_status = str(base_row.get("status", "")).strip()
            candidate_status = str(cand_row.get("status", "")).strip()
            baseline_parameters = str(base_row.get("custom_parameters", "")).strip()
            candidate_parameters = str(cand_row.get("custom_parameters", "")).strip()
            if baseline_status != candidate_status or baseline_parameters != candidate_parameters:
                change_type = "changed"
            else:
                continue
        rows.append(
            {
                "change_type": change_type,
                "custom_event": key[0],
                "identification": key[1],
                "target_id": key[2],
                "section_name": key[3],
                "page_id": key[4],
                "baseline_status": str(base_row.get("status", "")).strip() or "-",
                "candidate_status": str(cand_row.get("status", "")).strip() or "-",
                "baseline_parameters": str(base_row.get("custom_parameters", "")).strip() or "-",
                "candidate_parameters": str(cand_row.get("custom_parameters", "")).strip() or "-",
            }
        )
    return pd.DataFrame(rows, columns=columns)


def build_tracking_package_zip_bytes(
    tracking_plan_xlsx: bytes,
    screen_definition_xlsx: bytes,
    screen_marked_png: bytes,
    qa_report_xlsx: bytes,
) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        if tracking_plan_xlsx:
            zf.writestr("tracking_plan.xlsx", tracking_plan_xlsx)
        if screen_definition_xlsx:
            zf.writestr("screen_definition.xlsx", screen_definition_xlsx)
        if screen_marked_png:
            zf.writestr("screen_marked.png", screen_marked_png)
        if qa_report_xlsx:
            zf.writestr("qa_report.xlsx", qa_report_xlsx)
    return out.getvalue()


def load_dotenv_values(dotenv_path: Path | None = None) -> Dict[str, str]:
    values: Dict[str, str] = {}
    if dotenv_path is None:
        dotenv_path = BASE_DIR / ".env"
    if not dotenv_path.exists():
        return values

    for raw in dotenv_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        if not key:
            continue
        values[key] = value.strip().strip("'").strip('"')
    return values


def get_config_value(key: str, default: str = "") -> str:
    env_val = os.getenv(key, "").strip()
    if env_val:
        return env_val
    try:
        secret_val = str(st.secrets.get(key, "")).strip()
        if secret_val:
            return secret_val
    except Exception:
        pass
    dotenv_val = load_dotenv_values().get(key, "").strip()
    if dotenv_val:
        return dotenv_val
    return default


def is_truthy(value: object) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on"}


def get_novnc_popup_url() -> str:
    configured = get_config_value("QA_NOVNC_PUBLIC_URL", "").strip()
    if configured:
        return configured
    # Nginx /vnc/ reverse-proxy uses the /vnc/websockify backend path.
    return "/vnc/vnc.html?autoconnect=1&resize=remote&path=vnc/websockify"


def auto_open_popup_window(url: str, popup_name: str = "qa_debug_popup") -> None:
    target_url = str(url or "").strip()
    if not target_url:
        return
    safe_url = json.dumps(target_url)
    safe_name = json.dumps(popup_name)
    components.html(
        f"""
        <script>
          (() => {{
            const w = window.parent || window;
            const features = "width=1320,height=900,resizable=yes,scrollbars=yes";
            w.open({safe_url}, {safe_name}, features);
          }})();
        </script>
        """,
        height=0,
        width=0,
    )


def normalize_debug_target_url(raw_url: str) -> str:
    text = str(raw_url or "").strip()
    if not text:
        return ""
    parsed = urlparse(text)
    if parsed.scheme and parsed.netloc:
        return text
    if (not parsed.scheme) and parsed.path and ("." in parsed.path):
        guess = f"https://{parsed.path}"
        g = urlparse(guess)
        if g.scheme and g.netloc:
            return guess
    return ""


QA_MODE_OPTIONS = ["전체 이벤트 테스트", "시나리오 테스트", "Tracking plan 검증", "정의서 검증"]
QA_ENVIRONMENT_OPTIONS = ["prod", "staging", "custom"]
QA_BROWSER_OPTIONS = ["Chrome", "Chromium"]
QA_VIEWPORT_PRESETS = {
    "Desktop 1440 x 900": {"width": 1440, "height": 900, "mobile_mode": False, "mobile_device": ""},
    "Laptop 1280 x 720": {"width": 1280, "height": 720, "mobile_mode": False, "mobile_device": ""},
    "iPhone 13": {"width": 390, "height": 844, "mobile_mode": True, "mobile_device": "iPhone 13"},
}
QA_SCENARIO_TEMPLATES = {
    "상품 구매 흐름": {
        "steps": "view_item,add_to_cart,begin_checkout,purchase",
        "required_events": "view_item,add_to_cart,begin_checkout,purchase",
        "funnel_steps": "view_item,add_to_cart,begin_checkout,purchase",
        "key_mode_label": "transaction_id 기준",
    },
    "회원가입 흐름": {
        "steps": "page_view,sign_up",
        "required_events": "page_view,sign_up",
        "funnel_steps": "page_view,sign_up",
        "key_mode_label": "qa_debug_session_id 기준",
    },
    "랜딩 → 클릭 흐름": {
        "steps": "page_view,click_button,click_content",
        "required_events": "page_view,click_button,click_content",
        "funnel_steps": "page_view,click_button,click_content",
        "key_mode_label": "qa_debug_session_id 기준",
    },
}


def build_debug_run_settings() -> Dict[str, object]:
    viewport_key = str(st.session_state.get("qa_viewport_preset", "Desktop 1440 x 900")).strip()
    viewport = QA_VIEWPORT_PRESETS.get(viewport_key, QA_VIEWPORT_PRESETS["Desktop 1440 x 900"])
    browser_label = str(st.session_state.get("qa_browser_name", "Chrome")).strip()
    browser_name = "chromium" if browser_label == "Chromium" else "chrome"
    uploaded_schemas = st.session_state.get("qa_uploaded_definition_schemas", {})
    return {
        "environment": st.session_state.get("qa_environment", "prod"),
        "browser_name": browser_name,
        "viewport_width": int(viewport["width"]),
        "viewport_height": int(viewport["height"]),
        "mobile_mode": bool(viewport["mobile_mode"]),
        "mobile_device": str(viewport.get("mobile_device", "")).strip(),
        "auto_crawl_enabled": True,
        "auto_stop_after_crawl": True,
        "max_auto_clicks": int(st.session_state.get("qa_auto_crawl_max_clicks", 80)),
        "click_interval_ms": int(st.session_state.get("qa_auto_crawl_click_interval_ms", 1200)),
        "wait_after_click_ms": int(st.session_state.get("qa_auto_crawl_wait_after_click_ms", 1200)),
        "block_link_navigation": bool(st.session_state.get("qa_auto_block_link_nav", True)),
        "single_page_only": bool(st.session_state.get("qa_single_page_only", True)),
        "qa_mode": st.session_state.get("qa_mode", "전체 이벤트 테스트"),
        "scenario_template": st.session_state.get("qa_scenario_template", ""),
        "definition_file_name": st.session_state.get("qa_definition_validation_file_name", ""),
        "definition_event_names": sorted(list(uploaded_schemas.keys())),
        "definition_runtime_hints": build_definition_runtime_hints(uploaded_schemas if isinstance(uploaded_schemas, dict) else {}),
    }


def save_test_setup_snapshot() -> None:
    st.session_state["qa_test_setup_snapshot"] = {
        "qa_project_slug": st.session_state.get("qa_project_slug", DEFAULT_PROJECT_SLUG),
        "qa_project_slug_selectbox": st.session_state.get(
            "qa_project_slug_selectbox",
            st.session_state.get("qa_project_slug", DEFAULT_PROJECT_SLUG),
        ),
        "qa_project_domain": st.session_state.get("qa_project_domain", ""),
        "qa_environment": st.session_state.get("qa_environment", "prod"),
        "qa_debug_target_url": st.session_state.get("qa_debug_target_url", ""),
        "qa_browser_name": st.session_state.get("qa_browser_name", "Chrome"),
        "qa_viewport_preset": st.session_state.get("qa_viewport_preset", "Desktop 1440 x 900"),
        "qa_tester_name": st.session_state.get("qa_tester_name", ""),
        "qa_tester_note": st.session_state.get("qa_tester_note", ""),
    }


def restore_test_setup_snapshot_if_needed(force: bool = False) -> None:
    snapshot = st.session_state.get("qa_test_setup_snapshot", {})
    if not isinstance(snapshot, dict) or not snapshot:
        return
    text_keys = [
        "qa_project_domain",
        "qa_debug_target_url",
        "qa_tester_name",
        "qa_tester_note",
    ]
    select_keys = {
        "qa_project_slug": DEFAULT_PROJECT_SLUG,
        "qa_project_slug_selectbox": DEFAULT_PROJECT_SLUG,
        "qa_environment": "prod",
        "qa_browser_name": "Chrome",
        "qa_viewport_preset": "Desktop 1440 x 900",
    }
    for key in text_keys:
        current = str(st.session_state.get(key, "")).strip()
        saved = str(snapshot.get(key, "")).strip()
        if saved and (force or (not current)):
            st.session_state[key] = saved
    for key, default in select_keys.items():
        current = str(st.session_state.get(key, default)).strip()
        saved = str(snapshot.get(key, "")).strip()
        if saved and (force or current == str(default).strip()):
            st.session_state[key] = saved


def _get_request_headers() -> Dict[str, str]:
    try:
        raw_headers = st.context.headers
    except Exception:
        return {}
    out: Dict[str, str] = {}
    try:
        for k, v in raw_headers.items():
            out[str(k).strip().lower()] = str(v)
    except Exception:
        return {}
    return out


def _extract_remote_addr() -> str:
    headers = _get_request_headers()
    xff = str(headers.get("x-forwarded-for", "")).strip()
    if xff:
        return xff.split(",")[0].strip()
    xr = str(headers.get("x-real-ip", "")).strip()
    if xr:
        return xr
    return ""


def _hash_remote_addr(remote_addr: str) -> str:
    raw = str(remote_addr or "").strip()
    if not raw:
        return ""
    salt = get_config_value("QA_ACTION_LOG_SALT", "qa_action_salt")
    return hashlib.sha256(f"{salt}:{raw}".encode("utf-8")).hexdigest()[:20]


def log_ui_action(action_type: str, detail: Dict[str, object] | None = None) -> None:
    action = str(action_type or "").strip()
    if not action:
        return
    try:
        active_db_path = get_active_test_log_db_path()
        init_test_log_db(active_db_path)
        headers = _get_request_headers()
        actor_name = str(
            st.session_state.get("qa_access_user", "")
            or st.session_state.get("qa_tester_name", "")
            or "-"
        ).strip()
        append_ui_action(
            active_db_path,
            {
                "action_at": datetime.now(tz=ZoneInfo("UTC")).isoformat(),
                "session_id": str(st.session_state.get("qa_debug_session_id", "")).strip(),
                "action_type": action,
                "actor_role": str(st.session_state.get("qa_access_role", "guest")).strip(),
                "actor_name": actor_name,
                "detail": detail or {},
                "remote_addr_hash": _hash_remote_addr(_extract_remote_addr()),
                "user_agent": str(headers.get("user-agent", "")).strip(),
            },
        )
    except Exception:
        return


def enforce_access_gate() -> None:
    bypass = is_truthy(get_config_value("QA_APP_ACCESS_BYPASS", "1"))
    if bypass:
        if not st.session_state.get("qa_access_authenticated", False):
            st.session_state["qa_access_authenticated"] = True
            st.session_state["qa_access_role"] = "admin"
            st.session_state["qa_access_user"] = "bypass_access"
        return

    enabled = is_truthy(get_config_value("QA_APP_ACCESS_ENABLED", "0"))
    if not enabled:
        if not st.session_state.get("qa_access_authenticated", False):
            st.session_state["qa_access_authenticated"] = True
            st.session_state["qa_access_role"] = "admin"
            st.session_state["qa_access_user"] = "open_access"
        return

    if st.session_state.get("qa_access_authenticated", False):
        return

    tester_code = get_config_value("QA_TESTER_ACCESS_CODE", "").strip()
    admin_code = get_config_value("QA_ADMIN_ACCESS_CODE", "").strip()

    st.subheader("접근 제한")
    st.caption("테스트 전용 접속 코드가 필요합니다.")
    with st.form("qa_access_form", clear_on_submit=False):
        input_user = st.text_input("테스터 식별자", placeholder="예: tester_01")
        input_code = st.text_input("접속 코드", type="password")
        submitted = st.form_submit_button("입장")

    if submitted:
        code = str(input_code or "").strip()
        role = ""
        if admin_code and code == admin_code:
            role = "admin"
        elif tester_code and code == tester_code:
            role = "tester"
        elif (not tester_code) and (not admin_code):
            log_ui_action("access_login_error", {"reason": "codes_not_configured"})
            st.error("서버 접근 코드가 설정되지 않았습니다. 관리자에게 문의하세요.")
            st.stop()

        if role:
            st.session_state["qa_access_authenticated"] = True
            st.session_state["qa_access_role"] = role
            st.session_state["qa_access_user"] = str(input_user or "").strip() or f"{role}_user"
            log_ui_action("access_login_success", {"role": role})
            st.rerun()
        else:
            log_ui_action("access_login_failed", {"user": str(input_user or "").strip()})
            st.error("접속 코드가 올바르지 않습니다.")
    st.stop()


def probe_ingest_health(url: str, timeout_sec: float = 1.0) -> bool:
    target = str(url or "").strip()
    if not target:
        return False
    try:
        req = urllib_request.Request(target, method="GET")
        with urllib_request.urlopen(req, timeout=float(timeout_sec)) as resp:
            body = resp.read().decode("utf-8", errors="ignore").strip()
            if resp.status != 200:
                return False
            try:
                parsed = json.loads(body) if body else {}
                if isinstance(parsed, dict):
                    return bool(parsed.get("ok", False))
            except Exception:
                pass
            return '"ok":true' in body.lower().replace(" ", "")
    except Exception:
        return False


def get_google_access_token(token_path: Path) -> str:
    if not token_path.is_absolute():
        token_path = (BASE_DIR / token_path).resolve()
    if not token_path.exists():
        raise RuntimeError(f"OAuth 토큰 파일이 없습니다: {token_path}")

    token_payload = json.loads(token_path.read_text(encoding="utf-8"))
    access_token = str(token_payload.get("token", "")).strip()
    expiry = pd.to_datetime(token_payload.get("expiry"), errors="coerce", utc=True)
    now_utc = pd.Timestamp.now(tz="UTC")
    if access_token and pd.notna(expiry) and expiry > now_utc + pd.Timedelta(minutes=1):
        return access_token

    refresh_token = str(token_payload.get("refresh_token", "")).strip()
    client_id = str(token_payload.get("client_id", "")).strip()
    client_secret = str(token_payload.get("client_secret", "")).strip()
    token_uri = str(token_payload.get("token_uri", "https://oauth2.googleapis.com/token")).strip()
    if not refresh_token or not client_id or not client_secret:
        raise RuntimeError("OAuth refresh 토큰 정보가 부족합니다. token.json을 다시 발급하세요.")

    payload = urllib_parse.urlencode(
        {
            "grant_type": "refresh_token",
            "refresh_token": refresh_token,
            "client_id": client_id,
            "client_secret": client_secret,
        }
    ).encode("utf-8")
    req = urllib_request.Request(
        token_uri,
        data=payload,
        method="POST",
        headers={"Content-Type": "application/x-www-form-urlencoded"},
    )
    try:
        with urllib_request.urlopen(req, timeout=20) as resp:
            refreshed = json.loads(resp.read().decode("utf-8"))
    except urllib_error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"OAuth 토큰 갱신 실패: {exc.code} {detail}") from exc
    except Exception as exc:
        raise RuntimeError(f"OAuth 토큰 갱신 실패: {exc}") from exc

    new_token = str(refreshed.get("access_token", "")).strip()
    if not new_token:
        raise RuntimeError("OAuth 토큰 갱신 결과에 access_token이 없습니다.")

    token_payload["token"] = new_token
    expires_in = refreshed.get("expires_in")
    if expires_in is not None:
        try:
            token_payload["expiry"] = (
                pd.Timestamp.now(tz="UTC") + pd.Timedelta(seconds=int(expires_in))
            ).isoformat()
        except Exception:
            pass
    try:
        token_path.write_text(json.dumps(token_payload, ensure_ascii=False), encoding="utf-8")
    except Exception:
        # 토큰 파일 저장 실패는 치명적이지 않으므로 무시한다.
        pass

    return new_token


def _ga4_api_request(
    *,
    method: str,
    url: str,
    access_token: str,
    body: Dict[str, object] | None = None,
) -> Dict[str, object]:
    payload = None
    headers = {"Authorization": f"Bearer {access_token}"}
    if body is not None:
        payload = json.dumps(body).encode("utf-8")
        headers["Content-Type"] = "application/json"

    req = urllib_request.Request(url, method=method.upper(), headers=headers, data=payload)
    try:
        with urllib_request.urlopen(req, timeout=30) as resp:
            raw = resp.read().decode("utf-8")
            return json.loads(raw) if raw else {}
    except urllib_error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"GA API 호출 실패: {exc.code} {detail}") from exc
    except Exception as exc:
        raise RuntimeError(f"GA API 호출 실패: {exc}") from exc


def _resolve_project_path(path_text: str) -> Path:
    path_obj = Path(path_text).expanduser()
    if not path_obj.is_absolute():
        path_obj = (BASE_DIR / path_obj).resolve()
    return path_obj


def _load_google_client_config(client_secrets_file: str) -> Dict[str, str]:
    # 1) Streamlit secrets.toml ([auth], [auth.google]) 우선 지원
    #    예:
    #    [auth]
    #    redirect_uri = "https://asknuggetdata.com/oauth2callback"
    #    [auth.google]
    #    client_id = "..."
    #    client_secret = "..."
    try:
        auth_conf = st.secrets.get("auth", {})
        google_conf = auth_conf.get("google", {}) if hasattr(auth_conf, "get") else {}
        sec_client_id = str(google_conf.get("client_id", "")).strip() if hasattr(google_conf, "get") else ""
        sec_client_secret = str(google_conf.get("client_secret", "")).strip() if hasattr(google_conf, "get") else ""
        sec_metadata_url = str(google_conf.get("server_metadata_url", "")).strip() if hasattr(google_conf, "get") else ""

        if sec_client_id and sec_client_secret:
            redirect_uri_override = get_config_value("GA4_OAUTH_REDIRECT_URI", "").strip()
            sec_redirect_uri = str(auth_conf.get("redirect_uri", "")).strip() if hasattr(auth_conf, "get") else ""
            dynamic_redirect_uri = ""
            if is_truthy(get_config_value("GA4_OAUTH_DYNAMIC_REDIRECT", "0")):
                current_base = get_current_app_base_url()
                if current_base:
                    dynamic_redirect_uri = f"{current_base}/oauth2callback"
            redirect_uri = (
                redirect_uri_override
                or sec_redirect_uri
                or dynamic_redirect_uri
                or DEFAULT_OAUTH_REDIRECT_URI
            )
            auth_uri = "https://accounts.google.com/o/oauth2/v2/auth"
            token_uri = "https://oauth2.googleapis.com/token"

            if sec_metadata_url:
                try:
                    req = urllib_request.Request(sec_metadata_url, method="GET")
                    with urllib_request.urlopen(req, timeout=10) as resp:
                        meta = json.loads(resp.read().decode("utf-8"))
                    auth_uri = str(meta.get("authorization_endpoint", auth_uri)).strip() or auth_uri
                    token_uri = str(meta.get("token_endpoint", token_uri)).strip() or token_uri
                except Exception:
                    # metadata 조회 실패 시 기본 엔드포인트 사용
                    pass

            return {
                "client_id": sec_client_id,
                "client_secret": sec_client_secret,
                "auth_uri": auth_uri,
                "token_uri": token_uri,
                "redirect_uri": redirect_uri,
            }
    except Exception:
        pass

    # 2) 기존 client_secret.json 방식 fallback
    cfg_path = _resolve_project_path(client_secrets_file)
    if not cfg_path.exists():
        raise RuntimeError(
            f"OAuth client secret 파일이 없습니다: {cfg_path} "
            "(또는 .streamlit/secrets.toml의 [auth.google] client_id/client_secret를 설정하세요.)"
        )
    raw = json.loads(cfg_path.read_text(encoding="utf-8"))
    base = raw.get("web") or raw.get("installed")
    if not isinstance(base, dict):
        raise RuntimeError("client_secret.json 형식이 올바르지 않습니다. web/installed 설정이 필요합니다.")

    redirect_uris = [str(u).strip() for u in (base.get("redirect_uris") or []) if str(u).strip()]
    dynamic_redirect_uri = ""
    if is_truthy(get_config_value("GA4_OAUTH_DYNAMIC_REDIRECT", "0")):
        current_base = get_current_app_base_url()
        if current_base:
            dynamic_redirect_uri = f"{current_base}/oauth2callback"
    redirect_uri_override = get_config_value("GA4_OAUTH_REDIRECT_URI", "").strip()

    def _is_local_redirect(uri_text: str) -> bool:
        try:
            host = (urlparse(uri_text).hostname or "").strip().lower()
        except Exception:
            host = ""
        return host in {"localhost", "127.0.0.1", "::1"} or host.endswith(".local")

    if redirect_uri_override:
        redirect_uri = redirect_uri_override
    elif dynamic_redirect_uri:
        redirect_uri = dynamic_redirect_uri
    else:
        https_non_local = [
            u for u in redirect_uris if u.lower().startswith("https://") and not _is_local_redirect(u)
        ]
        redirect_uri = https_non_local[0] if https_non_local else (redirect_uris[0] if redirect_uris else DEFAULT_OAUTH_REDIRECT_URI)
    if not redirect_uri:
        raise RuntimeError("client_secret.json에 redirect_uris가 없습니다.")

    return {
        "client_id": str(base.get("client_id", "")).strip(),
        "client_secret": str(base.get("client_secret", "")).strip(),
        "auth_uri": str(base.get("auth_uri", "https://accounts.google.com/o/oauth2/v2/auth")).strip(),
        "token_uri": str(base.get("token_uri", "https://oauth2.googleapis.com/token")).strip(),
        "redirect_uri": redirect_uri,
    }


def build_google_oauth_url(client_secrets_file: str, state: str) -> str:
    conf = _load_google_client_config(client_secrets_file)
    if not conf["client_id"] or not conf["client_secret"]:
        raise RuntimeError("client_id/client_secret가 비어 있습니다.")
    scopes = ["https://www.googleapis.com/auth/analytics.readonly"]
    query = urllib_parse.urlencode(
        {
            "client_id": conf["client_id"],
            "redirect_uri": conf["redirect_uri"],
            "response_type": "code",
            "scope": " ".join(scopes),
            "access_type": "offline",
            "include_granted_scopes": "true",
            "prompt": "consent",
            "state": state,
        }
    )
    return f"{conf['auth_uri']}?{query}"


def get_google_oauth_redirect_uri(client_secrets_file: str) -> str:
    conf = _load_google_client_config(client_secrets_file)
    return str(conf.get("redirect_uri", "")).strip()


def exchange_google_oauth_code(
    *,
    client_secrets_file: str,
    code: str,
    token_file: str,
) -> None:
    conf = _load_google_client_config(client_secrets_file)
    payload = urllib_parse.urlencode(
        {
            "code": code,
            "client_id": conf["client_id"],
            "client_secret": conf["client_secret"],
            "redirect_uri": conf["redirect_uri"],
            "grant_type": "authorization_code",
        }
    ).encode("utf-8")
    req = urllib_request.Request(
        conf["token_uri"],
        data=payload,
        method="POST",
        headers={"Content-Type": "application/x-www-form-urlencoded"},
    )
    try:
        with urllib_request.urlopen(req, timeout=30) as resp:
            token_resp = json.loads(resp.read().decode("utf-8"))
    except urllib_error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        if "redirect_uri_mismatch" in detail:
            raise RuntimeError(
                "Google 로그인 토큰 교환 실패: redirect_uri_mismatch. "
                "Google Cloud OAuth 설정에 현재 redirect URI를 추가하세요 "
                "(예: https://asknuggetdata.com/oauth2callback)."
            ) from exc
        raise RuntimeError(f"Google 로그인 토큰 교환 실패: {exc.code} {detail}") from exc
    except Exception as exc:
        raise RuntimeError(f"Google 로그인 토큰 교환 실패: {exc}") from exc

    access_token = str(token_resp.get("access_token", "")).strip()
    if not access_token:
        raise RuntimeError("로그인 토큰 교환 응답에 access_token이 없습니다.")

    token_path = _resolve_project_path(token_file)
    existing: Dict[str, object] = {}
    if token_path.exists():
        try:
            existing = json.loads(token_path.read_text(encoding="utf-8"))
        except Exception:
            existing = {}

    merged = dict(existing)
    merged["token"] = access_token
    if str(token_resp.get("refresh_token", "")).strip():
        merged["refresh_token"] = str(token_resp.get("refresh_token", "")).strip()
    merged["token_uri"] = conf["token_uri"]
    merged["client_id"] = conf["client_id"]
    merged["client_secret"] = conf["client_secret"]
    merged["scopes"] = token_resp.get("scope", "").split() or merged.get("scopes", [])
    expires_in = token_resp.get("expires_in")
    if expires_in is not None:
        try:
            merged["expiry"] = (
                pd.Timestamp.now(tz="UTC") + pd.Timedelta(seconds=int(expires_in))
            ).isoformat()
        except Exception:
            pass
    merged["universe_domain"] = "googleapis.com"

    token_path.parent.mkdir(parents=True, exist_ok=True)
    token_path.write_text(json.dumps(merged, ensure_ascii=False), encoding="utf-8")


def _query_param_text(name: str) -> str:
    value = st.query_params.get(name)
    if isinstance(value, list):
        return str(value[0]).strip() if value else ""
    return str(value).strip() if value is not None else ""


def process_google_oauth_callback_if_present() -> None:
    error_text = _query_param_text("error")
    code = _query_param_text("code")
    if not error_text and not code:
        return

    state = _query_param_text("state")
    expected_state = str(st.session_state.get("qa_oauth_state", "")).strip()
    token_file = str(
        st.session_state.get("qa_report_token_file", get_config_value("GA4_TOKEN_FILE", "token.json"))
    ).strip() or "token.json"
    client_file = str(
        st.session_state.get("qa_report_client_secret_file", get_config_value("GA4_CLIENT_SECRETS_FILE", "client_secret.json"))
    ).strip() or "client_secret.json"

    try:
        if error_text:
            raise RuntimeError(f"Google 로그인 실패: {error_text}")
        if expected_state and state and state != expected_state:
            raise RuntimeError("Google 로그인 state 검증에 실패했습니다. 다시 시도하세요.")
        if not code:
            raise RuntimeError("Google 로그인 인가 코드가 없습니다.")
        exchange_google_oauth_code(
            client_secrets_file=client_file,
            code=code,
            token_file=token_file,
        )
        resume_state = state or expected_state
        oauth_ctx = pop_oauth_context(resume_state)
        restore_keys = [
            "qa_project_slug",
            "qa_project_domain",
            "qa_debug_session_id",
            "qa_debug_output_file",
            "qa_debug_started_at",
            "qa_debug_target_url",
            "qa_tester_name",
            "qa_tester_note",
            "required_event_text_input",
            "unknown_event_policy",
            "qa_report_property_id",
        ]
        for key in restore_keys:
            if key in oauth_ctx:
                st.session_state[key] = oauth_ctx.get(key)
        if str(st.session_state.get("qa_debug_session_id", "")).strip():
            st.session_state["qa_oauth_resume_notice"] = (
                "Google 로그인 후 기존 실시간 QA 세션을 복원했습니다. "
                "실시간 데이터 새로고침으로 이어서 확인하세요."
            )
        log_ui_action("oauth_callback_success")
        st.session_state["qa_oauth_notice"] = f"Google 로그인 완료. 토큰 저장: {token_file}"
        st.session_state["qa_oauth_error"] = ""
        st.session_state["qa_oauth_auth_url"] = ""
        st.session_state["qa_oauth_state"] = ""
        st.session_state["qa_ui_focus_after_oauth"] = "report"
    except Exception as exc:
        log_ui_action("oauth_callback_error", {"error": str(exc)})
        st.session_state["qa_oauth_error"] = to_user_error_message(exc)
        st.session_state["qa_oauth_notice"] = ""
    finally:
        st.query_params.clear()
        st.rerun()


@st.cache_data(ttl=600, show_spinner=False)
def fetch_ga4_property_list(token_file: str) -> pd.DataFrame:
    access_token = get_google_access_token(_resolve_project_path(token_file))
    rows: List[Dict[str, str]] = []
    page_token = ""

    while True:
        url = "https://analyticsadmin.googleapis.com/v1alpha/accountSummaries?pageSize=200"
        if page_token:
            url = f"{url}&pageToken={urllib_parse.quote(page_token)}"
        resp = _ga4_api_request(method="GET", url=url, access_token=access_token)
        for acc in resp.get("accountSummaries", []) or []:
            account_name = str(acc.get("displayName", "")).strip() or "-"
            account_ref = str(acc.get("account", "")).strip()
            for prop in acc.get("propertySummaries", []) or []:
                prop_ref = str(prop.get("property", "")).strip()
                prop_id = prop_ref.split("/")[-1] if "/" in prop_ref else prop_ref
                prop_name = str(prop.get("displayName", "")).strip() or "-"
                if not prop_id:
                    continue
                rows.append(
                    {
                        "property_id": prop_id,
                        "property_name": prop_name,
                        "account_name": account_name,
                        "account_ref": account_ref,
                        "label": f"{prop_name} ({prop_id}) · {account_name}",
                    }
                )
        page_token = str(resp.get("nextPageToken", "")).strip()
        if not page_token:
            break

    out = pd.DataFrame(rows)
    if out.empty:
        return pd.DataFrame(columns=["property_id", "property_name", "account_name", "account_ref", "label"])
    return out.drop_duplicates(subset=["property_id"]).sort_values(["account_name", "property_name"]).reset_index(drop=True)


@st.cache_data(ttl=600, show_spinner=False)
def fetch_ga4_reference_data(
    property_id: str,
    token_file: str,
    lookback_days: int = 30,
    top_events_limit: int = 200,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    pid = str(property_id).strip()
    if not pid:
        raise RuntimeError("GA4 Property ID가 비어 있습니다.")

    token_path = _resolve_project_path(token_file)
    access_token = get_google_access_token(token_path)
    base = f"https://analyticsdata.googleapis.com/v1beta/properties/{pid}"

    report = _ga4_api_request(
        method="POST",
        url=f"{base}:runReport",
        access_token=access_token,
        body={
            "dateRanges": [{"startDate": f"{int(lookback_days)}daysAgo", "endDate": "today"}],
            "dimensions": [{"name": "eventName"}],
            "metrics": [{"name": "eventCount"}],
            "orderBys": [{"metric": {"metricName": "eventCount"}, "desc": True}],
            "limit": int(top_events_limit),
        },
    )

    event_rows: List[Dict[str, object]] = []
    for row in report.get("rows", []) or []:
        dims = row.get("dimensionValues", []) or []
        mets = row.get("metricValues", []) or []
        event_name = str(dims[0].get("value", "")).strip() if dims else ""
        event_count_raw = str(mets[0].get("value", "0")).strip() if mets else "0"
        if not event_name:
            continue
        try:
            event_count = int(float(event_count_raw))
        except Exception:
            event_count = 0
        event_rows.append({"event_name": event_name, "event_count_30d": event_count})
    events_df = pd.DataFrame(event_rows)
    if not events_df.empty:
        events_df = events_df.sort_values("event_count_30d", ascending=False).reset_index(drop=True)

    metadata = _ga4_api_request(
        method="GET",
        url=f"{base}/metadata",
        access_token=access_token,
    )
    param_rows: List[Dict[str, str]] = []
    for dim in metadata.get("dimensions", []) or []:
        api_name = str(dim.get("apiName", "")).strip()
        if not api_name.startswith("customEvent:"):
            continue
        param_name = api_name.split("customEvent:", 1)[1].strip()
        if not param_name:
            continue
        param_rows.append(
            {
                "parameter_name": param_name,
                "api_name": api_name,
                "description": str(dim.get("description", "")).strip() or "-",
            }
        )
    params_df = pd.DataFrame(param_rows)
    if not params_df.empty:
        params_df = params_df.drop_duplicates(subset=["parameter_name"]).sort_values("parameter_name").reset_index(drop=True)

    return events_df, params_df


def build_api_aggregate_results(
    events_df: pd.DataFrame,
    required_events: List[str],
    required_params_by_event: Dict[str, List[str]],
    selected_params: List[str],
) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    work = events_df.copy()
    if "event_name" not in work.columns:
        work["event_name"] = ""
    if "event_count_30d" not in work.columns:
        work["event_count_30d"] = 0

    event_counts = (
        work.groupby("event_name")["event_count_30d"].sum().to_dict()
        if not work.empty
        else {}
    )

    for event_name in [e.strip() for e in required_events if e.strip()]:
        count = int(event_counts.get(event_name, 0))
        rows.append(
            {
                "rule_id": f"required_event:{event_name}",
                "rule_name": f"필수 이벤트 존재(API): {event_name}",
                "status": "PASS" if count > 0 else "FAIL",
                "detail": f"30일 이벤트 수 {count}건",
                "fail_count": 0 if count > 0 else 1,
                "total_count": 1,
            }
        )

    for event_name, params in required_params_by_event.items():
        for param_name in [p.strip() for p in params if p.strip()]:
            rows.append(
                {
                    "rule_id": f"required_param:{event_name}:{param_name}",
                    "rule_name": f"필수 param 체크(API): {event_name}.{param_name}",
                    "status": "WARN",
                    "detail": "API 집계 모드에서는 이벤트별 param null/누락을 직접 계산하지 않습니다.",
                    "fail_count": 0,
                    "total_count": 1,
                }
            )

    null_targets = list(dict.fromkeys([p for p in selected_params if p.strip()] or ["transaction_id", "value", "currency"]))
    for param_name in null_targets:
        rows.append(
            {
                "rule_id": f"null_ratio:{param_name}",
                "rule_name": f"null 비율 체크(API): {param_name}",
                "status": "WARN",
                "detail": "API 집계 모드에서는 null 비율을 직접 계산하지 않습니다.",
                "fail_count": 0,
                "total_count": 1,
            }
        )

    rows.append(
        {
            "rule_id": "duplicate_transaction_id",
            "rule_name": "중복 transaction_id 체크(API)",
            "status": "WARN",
            "detail": "API 집계 모드에서는 transaction_id 중복을 직접 계산하지 않습니다.",
            "fail_count": 0,
            "total_count": 1,
        }
    )

    return pd.DataFrame(rows)


def build_api_event_preview_df(events_df: pd.DataFrame) -> pd.DataFrame:
    if events_df.empty:
        return pd.DataFrame(columns=["event_name", "event_count_30d", "event_date", "event_timestamp", "user_pseudo_id"])
    out = events_df.copy()
    out["event_date"] = pd.Timestamp.now(tz=LOCAL_TZ).date()
    out["event_timestamp"] = pd.NaT
    out["user_pseudo_id"] = pd.NA
    return out


def format_local_time(value: object) -> str:
    ts = pd.to_datetime(value, errors="coerce", utc=True)
    if pd.isna(ts):
        return "-"
    return ts.tz_convert(LOCAL_TZ).strftime("%H:%M:%S")


def filter_events_for_active_session(
    debug_df: pd.DataFrame,
    session_id: str,
    session_started_at: str = "",
) -> pd.DataFrame:
    if debug_df.empty:
        return debug_df
    sid = str(session_id or "").strip()
    out = debug_df.copy()
    if sid:
        row_sid = out.get("session_id", pd.Series(index=out.index, dtype="object")).astype(str).str.strip()
        param_sid = out.get("params", pd.Series(index=out.index, dtype="object")).apply(
            lambda p: str(p.get("qa_debug_session_id", "")).strip() if isinstance(p, dict) else ""
        )
        # session_id 컬럼이 있으면 우선 신뢰하고, 없을 때만 param sid를 보조로 사용한다.
        sid_mask = row_sid.eq(sid) | (row_sid.eq("") & param_sid.eq(sid))
        out = out[sid_mask].copy()

    started_text = str(session_started_at or "").strip()
    if started_text and not out.empty:
        started_ts = pd.to_datetime(started_text, errors="coerce", utc=True)
        if pd.notna(started_ts):
            cap_ts = pd.to_datetime(out.get("captured_at"), errors="coerce", utc=True)
            out = out[cap_ts >= started_ts].copy()
    return out


def normalize_uploaded_df(df: pd.DataFrame, requested_params: List[str]) -> pd.DataFrame:
    if any(col in df.columns for col in ["eventName", "dateHourMinute", "userPseudoId"]):
        return flatten_events(df, requested_params)

    rename_map = {col: normalize_param_name(str(col)) for col in df.columns}
    out = df.rename(columns=rename_map).copy()

    if "event_timestamp" in out.columns:
        out["event_timestamp"] = pd.to_datetime(out["event_timestamp"], errors="coerce")
    if "event_date" in out.columns:
        out["event_date"] = pd.to_datetime(out["event_date"], errors="coerce").dt.date

    expected = [
        "event_name",
        "event_date",
        "event_timestamp",
        "user_pseudo_id",
        "transaction_id",
        "value",
        "currency",
    ]
    expected.extend(normalize_param_name(p) for p in requested_params)
    for col in expected:
        if col not in out.columns:
            out[col] = pd.NA

    return out


def debug_events_to_qa_df(debug_df: pd.DataFrame, requested_params: List[str]) -> pd.DataFrame:
    if debug_df.empty:
        return normalize_uploaded_df(pd.DataFrame(), requested_params)

    rows: List[Dict[str, object]] = []
    req_norm = [normalize_param_name(p) for p in requested_params if str(p).strip()]
    for _, row in debug_df.iterrows():
        event_name = str(row.get("event_name", "")).strip()
        if not event_name:
            continue
        params = row.get("params")
        if not isinstance(params, dict):
            params = {}
        ts = pd.to_datetime(row.get("captured_at"), errors="coerce")
        out_row: Dict[str, object] = {
            "event_name": event_name,
            "event_timestamp": ts,
            "event_date": ts.date() if pd.notna(ts) else pd.NaT,
            "user_pseudo_id": str(params.get("user_pseudo_id", "")).strip(),
            "event_count": 1,
            "page_location": row.get("page_url", ""),
            "qa_debug_session_id": str(
                row.get("session_id", params.get("qa_debug_session_id", ""))
            ).strip(),
        }
        for key, value in params.items():
            out_row[normalize_param_name(str(key))] = value
        for key in req_norm:
            if key not in out_row:
                out_row[key] = pd.NA
        rows.append(out_row)

    out_df = pd.DataFrame(rows)
    if out_df.empty:
        return normalize_uploaded_df(pd.DataFrame(), requested_params)
    return normalize_uploaded_df(out_df, requested_params)


def describe_param(param_name: str) -> str:
    desc_map = {
        "page_location": "현재 페이지 URL",
        "page_title": "페이지 제목",
        "transaction_id": "구매/주문 식별자",
        "value": "매출/가치 값",
        "currency": "통화 코드",
        "qa_debug_session_id": "테스트 세션 식별자",
        "engagement_time_msec": "페이지 체류 시간(ms)",
        "session_id": "세션 ID",
        "session_engaged": "참여 세션 여부",
        "menu": "클릭 메뉴명",
        "click_text": "클릭 텍스트",
    }
    return desc_map.get(param_name, "-")


def extract_gtm_tag_name(params: Dict[str, object]) -> str:
    for key in ("gtm.tagName", "gtm_tag_name", "tag_name", "tag"):
        val = params.get(key)
        if str(val).strip():
            return str(val).strip()
    triggers = str(params.get("gtm.triggers", "")).strip()
    if triggers:
        return f"trigger:{triggers}"
    return "-"


def evaluate_event_status(
    event_name: str,
    params: Dict[str, object],
    allowed_events: List[str] | None = None,
    unknown_event_policy: str = "정보",
) -> str:
    status = "OK"
    allowed_set = {e.strip() for e in (allowed_events or []) if str(e).strip()}
    internal_prefixes = ("gtm.",)
    internal_events = {"set_user_property"}

    if not event_name:
        return "ERROR"

    if (
        allowed_set
        and event_name not in allowed_set
        and not event_name.startswith(internal_prefixes)
        and event_name not in internal_events
    ):
        if unknown_event_policy == "경고":
            status = "WARN"
        elif unknown_event_policy == "정보":
            status = "INFO"

    if event_name == "purchase":
        if not str(params.get("value", "")).strip() and status == "OK":
            status = "WARN"
        if not str(params.get("transaction_id", "")).strip():
            status = "ERROR"

    return status


def build_event_param_kv_view(
    debug_df: pd.DataFrame,
    max_rows: int = 2000,
    allowed_events: List[str] | None = None,
    unknown_event_policy: str = "정보",
) -> pd.DataFrame:
    cols = [
        "시간",
        "이벤트이름",
        "이벤트에 딸린 파라미터",
        "디스크립션",
        "값",
        "상태",
        "원래 수집될 수 있는 모든 값",
        "gtm태그명",
    ]
    if debug_df.empty:
        return pd.DataFrame(columns=cols)

    work = debug_df.copy()
    if "captured_at" in work.columns:
        work["captured_at"] = pd.to_datetime(work["captured_at"], errors="coerce")
        work = work.sort_values("captured_at", ascending=False)

    rows: List[Dict[str, object]] = []
    for _, row in work.iterrows():
        event_name = str(row.get("event_name", "")).strip()
        if not event_name:
            continue
        params = row.get("params")
        if not isinstance(params, dict):
            params = {}

        all_values = dict(params)
        if str(row.get("measurement_id", "")).strip() and "measurement_id" not in all_values:
            all_values["measurement_id"] = str(row.get("measurement_id", "")).strip()
        if str(row.get("client_id", "")).strip() and "client_id" not in all_values:
            all_values["client_id"] = str(row.get("client_id", "")).strip()

        all_values_text = json.dumps(all_values, ensure_ascii=False, sort_keys=True)
        gtm_tag_name = extract_gtm_tag_name(all_values)
        ts_text = format_local_time(row.get("captured_at"))
        status = evaluate_event_status(
            event_name=event_name,
            params=all_values,
            allowed_events=allowed_events,
            unknown_event_policy=unknown_event_policy,
        )

        if not all_values:
            rows.append(
                {
                    "시간": ts_text,
                    "이벤트이름": event_name,
                    "이벤트에 딸린 파라미터": "-",
                    "디스크립션": "-",
                    "값": "-",
                    "상태": status,
                    "원래 수집될 수 있는 모든 값": "{}",
                    "gtm태그명": gtm_tag_name,
                }
            )
        else:
            for key, value in all_values.items():
                val_text = str(value).strip()
                if not val_text:
                    continue
                rows.append(
                    {
                        "시간": ts_text,
                        "이벤트이름": event_name,
                        "이벤트에 딸린 파라미터": str(key),
                        "디스크립션": describe_param(str(key)),
                        "값": val_text,
                        "상태": status,
                        "원래 수집될 수 있는 모든 값": all_values_text,
                        "gtm태그명": gtm_tag_name,
                    }
                )
                if len(rows) >= max_rows:
                    break
        if len(rows) >= max_rows:
            break

    out = pd.DataFrame(rows)
    if out.empty:
        return pd.DataFrame(columns=cols)
    return out


SYSTEM_PARAM_KEYS = {
    "measurement_id",
    "client_id",
    "protocol_version",
    "user_language",
    "architecture",
    "bitness",
    "browser",
    "platform_version",
    "screen_resolution",
    "random_page_hash",
    "google_consent_default",
    "dma",
    "pscdl",
    "are",
    "frm",
    "npa",
    "tag_exp",
    "tfd",
    "event_usage",
    "hit_counter",
    "session_count",
    "session_engaged",
    "session_id",
}

GA4_RESERVED_EVENT_NAMES = {
    "ad_impression",
    "app_clear_data",
    "app_exception",
    "app_install",
    "app_remove",
    "app_store_refund",
    "app_store_subscription_cancel",
    "app_store_subscription_convert",
    "app_store_subscription_renew",
    "click",
    "dynamic_link_app_open",
    "dynamic_link_app_update",
    "dynamic_link_first_open",
    "error",
    "exception",
    "file_download",
    "first_open",
    "first_visit",
    "form_start",
    "form_submit",
    "in_app_purchase",
    "notification_dismiss",
    "notification_foreground",
    "notification_open",
    "notification_receive",
    "os_update",
    "page_view",
    "screen_view",
    "scroll",
    "session_start",
    "user_engagement",
    "video_complete",
    "video_progress",
    "video_start",
    "view_search_results",
}
GA4_DEFAULT_PARAM_KEYS = {
    "app_id",
    "batch_ordering_id",
    "batch_page_id",
    "batch_event_index",
    "campaign",
    "campaign_id",
    "campaign_content",
    "campaign_medium",
    "campaign_name",
    "campaign_source",
    "campaign_term",
    "client_id",
    "content_group",
    "debug_mode",
    "engagement_time_msec",
    "firebase_conversion",
    "firebase_event_origin",
    "firebase_screen",
    "firebase_screen_class",
    "firebase_screen_id",
    "ga_session_id",
    "ga_session_number",
    "ignore_referrer",
    "language",
    "medium",
    "page_hostname",
    "page_location",
    "page_referrer",
    "page_title",
    "screen_resolution",
    "session_engaged",
    "session_id",
    "session_number",
    "source",
    "term",
    "user_agent",
    "user_id",
    "user_pseudo_id",
}
GA4_TECHNICAL_PARAM_PREFIXES = (
    "ga_",
    "google_",
    "gtm.",
    "qa_",
    "uaa",
    "uab",
    "uap",
    "uapv",
    "uaw",
    "up.",
)

MISSING_VALUE_TOKENS = {
    "",
    "(not set)",
    "not set",
    "--omitted--",
    "null",
    "none",
    "undefined",
    "nan",
    "(none)",
    "n/a",
    "na",
    "false",
}

PRIMARY_PARAM_PRIORITY = [
    "transaction_id",
    "value",
    "currency",
    "menu",
    "click_text",
    "page_location",
    "page_title",
    "qa_debug_session_id",
]


def _is_system_param_key(key: str) -> bool:
    k = str(key or "").strip().lower()
    if not k:
        return True
    if k in SYSTEM_PARAM_KEYS:
        return True
    if k.startswith("gtm."):
        return True
    return False


def _is_hidden_default_param_key(key: str) -> bool:
    raw_key = str(key or "").strip().lower()
    normalized = normalize_param_name(str(key or "")).strip().lower()
    if not raw_key and not normalized:
        return True
    if _is_system_param_key(raw_key) or _is_system_param_key(normalized):
        return True
    if raw_key in EVENT_DEFINITION_HIDDEN_KEYS or normalized in EVENT_DEFINITION_HIDDEN_KEYS:
        return True
    if raw_key in GA4_DEFAULT_PARAM_KEYS or normalized in GA4_DEFAULT_PARAM_KEYS:
        return True
    if raw_key.startswith("session_") or normalized.startswith("session_"):
        return True
    if raw_key.startswith("client_") or normalized.startswith("client_"):
        return True
    if raw_key.startswith(GA4_TECHNICAL_PARAM_PREFIXES) or normalized.startswith(GA4_TECHNICAL_PARAM_PREFIXES):
        return True
    return False


def _has_schema_definition(event_name: str, schemas: Dict[str, Dict[str, object]]) -> bool:
    name = str(event_name or "").strip()
    if not name:
        return False
    return name in schemas


def _schema_parameter_columns(event_name: str, schemas: Dict[str, Dict[str, object]]) -> List[str]:
    schema_obj = schemas.get(str(event_name or "").strip(), {})
    if not isinstance(schema_obj, dict):
        return []
    ordered: List[str] = []
    for key in (
        list(schema_obj.get("required", []))
        + list(schema_obj.get("optional", []))
        + list(schema_obj.get("parameter_columns", []))
    ):
        key_text = str(key).strip()
        if not key_text or key_text in ordered:
            continue
        if _is_hidden_default_param_key(key_text):
            continue
        ordered.append(key_text)
    return ordered


def _is_custom_event_name(event_name: str, schemas: Dict[str, Dict[str, object]]) -> bool:
    name = str(event_name or "").strip()
    if not name:
        return False
    if _has_schema_definition(name, schemas):
        return True
    return name.lower() not in GA4_RESERVED_EVENT_NAMES


def _schema_custom_param_keys(event_name: str, schemas: Dict[str, Dict[str, object]]) -> List[str]:
    return _schema_parameter_columns(event_name, schemas)


def _sort_custom_param_keys(keys: List[str]) -> List[str]:
    priority_order = {name: idx for idx, name in enumerate(EVENT_DEFINITION_PARAM_PRIORITY)}
    return sorted(
        list(dict.fromkeys([str(key).strip() for key in keys if str(key).strip()])),
        key=lambda key: (priority_order.get(key, 10_000), key),
    )


def _is_meaningful_custom_param_key(key: str) -> bool:
    key_text = str(key or "").strip().lower()
    if not key_text:
        return False
    if key_text in {item.lower() for item in EVENT_DEFINITION_PARAM_PRIORITY}:
        return True
    meaningful_tokens = (
        "button",
        "section",
        "content",
        "banner",
        "category",
        "filter",
        "tab",
        "page_id",
        "page_link",
        "extra",
        "title",
        "index",
    )
    if any(token in key_text for token in meaningful_tokens):
        return True
    if key_text.endswith("_id") or key_text.endswith("_name") or key_text.endswith("_type") or key_text.endswith("_value"):
        return True
    return False


def split_custom_param_items(custom_items: List[Tuple[str, str]]) -> Tuple[List[Tuple[str, str]], List[Tuple[str, str]]]:
    if not custom_items:
        return [], []
    ordered = [(str(key).strip(), str(value).strip()) for key, value in custom_items if str(key).strip() and str(value).strip()]
    priority_names = list(dict.fromkeys(EVENT_DEFINITION_PARAM_PRIORITY))
    top_items: List[Tuple[str, str]] = []
    for key in priority_names:
        for item_key, item_value in ordered:
            if item_key == key and (item_key, item_value) not in top_items:
                top_items.append((item_key, item_value))
    if not top_items:
        top_items = ordered[:6]
    else:
        top_items = top_items[:8]
    extra_items = [item for item in ordered if item not in top_items]
    return top_items, extra_items


def extract_custom_event_payload(
    event_name: str,
    params: Dict[str, object],
    schemas: Dict[str, Dict[str, object]],
) -> Tuple[str, List[Tuple[str, str]]]:
    normalized_event = str(event_name or "").strip()
    if not normalized_event:
        return "", []
    if not _is_custom_event_name(normalized_event, schemas):
        return "", []

    custom_keys: List[str] = []
    if _has_schema_definition(normalized_event, schemas):
        custom_keys = _schema_custom_param_keys(normalized_event, schemas)
    else:
        for key in params.keys():
            key_text = str(key).strip()
            if not key_text or _is_hidden_default_param_key(key_text):
                continue
            if not _is_meaningful_custom_param_key(key_text):
                continue
            custom_keys.append(key_text)
    ordered_keys = _sort_custom_param_keys(custom_keys)
    custom_items: List[Tuple[str, str]] = []
    for key in ordered_keys:
        if key not in params:
            continue
        value_text = _to_text_value(params.get(key, ""))
        if not value_text:
            continue
        custom_items.append((key, value_text))
    if not custom_items and not _has_schema_definition(normalized_event, schemas):
        return "", []
    return normalized_event, custom_items


def filter_debug_df_for_event_scope(
    debug_df: pd.DataFrame,
    scoped_event_names: set[str] | None = None,
) -> pd.DataFrame:
    if debug_df.empty or scoped_event_names is None:
        return debug_df
    scope = {str(v).strip() for v in scoped_event_names if str(v).strip()}
    if not scope:
        return debug_df.iloc[0:0].copy()
    work = debug_df.copy()
    source_series = work.get("source", pd.Series(index=work.index, dtype=str)).astype(str).str.strip()
    event_series = work.get("event_name", pd.Series(index=work.index, dtype=str)).astype(str).str.strip()
    keep_mask = (~source_series.eq("ga_hit")) | event_series.isin(scope)
    return work[keep_mask].copy()


def _to_text_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    text = str(value).strip()
    return text


def _is_missing_like_value(value_text: str) -> bool:
    norm = str(value_text).strip().lower()
    return norm in MISSING_VALUE_TOKENS


def _is_false_like_value(value_text: str) -> bool:
    return str(value_text).strip().lower() in {"false"}


def _is_zero_like_value(value_text: str) -> bool:
    try:
        return float(str(value_text).strip()) == 0.0
    except Exception:
        return False


def validate_param_format(param_name: str, value_text: str) -> Tuple[bool, str]:
    key = normalize_param_name(str(param_name))
    value = str(value_text).strip()
    if not value:
        return False, "빈 값"

    if key == "currency":
        return (bool(re.fullmatch(r"[A-Z]{3}", value)), "3자리 대문자 통화코드 필요")
    if key == "value":
        try:
            float(value)
            return True, "숫자형"
        except Exception:
            return False, "숫자형 아님"
    if key in {"page_location", "page_referrer"}:
        return (value.startswith("http://") or value.startswith("https://"), "URL 형식 필요")
    if key == "qa_debug_session_id":
        ok = value.startswith("dbg_") or value.startswith("qa_debug_session_")
        return (ok, "qa 디버그 세션 형식 아님")
    if key.endswith("_id"):
        return (len(value) >= 2, "id 값 길이 부족")
    return True, "기본 검증 통과"


def build_suspicious_param_profile(debug_df: pd.DataFrame) -> Dict[str, str]:
    if debug_df.empty:
        return {}

    key_values: Dict[str, List[str]] = {}
    for _, row in debug_df.iterrows():
        params = row.get("params")
        if not isinstance(params, dict):
            continue
        for key, value in params.items():
            key_text = str(key).strip()
            if not key_text or _is_system_param_key(key_text):
                continue
            value_text = _to_text_value(value)
            if not value_text:
                continue
            key_values.setdefault(key_text, []).append(value_text)

    profile: Dict[str, str] = {}
    for key, values in key_values.items():
        if not values:
            continue
        reasons: List[str] = []
        norms = [v.strip().lower() for v in values if v.strip()]
        uniq_norm = sorted(set(norms))
        if values and all(_is_false_like_value(v) for v in values):
            reasons.append("모든 이벤트에서 false")
        if values and all(_is_zero_like_value(v) for v in values):
            reasons.append("모든 이벤트에서 0")
        if len(values) >= 2 and len(uniq_norm) == 1:
            reasons.append(f"모든 이벤트에서 동일값({values[0]})")
        if reasons:
            profile[key] = " / ".join(dict.fromkeys(reasons))
    return profile


def _collect_auto_context_rows(debug_df: pd.DataFrame) -> List[Dict[str, object]]:
    if debug_df.empty:
        return []
    work = debug_df.copy()
    if "captured_at" in work.columns:
        work["captured_at"] = pd.to_datetime(work["captured_at"], errors="coerce")
        work = work.sort_values("captured_at")
    rows: List[Dict[str, object]] = []
    src_series = work.get("source", pd.Series(index=work.index, dtype=str)).astype(str).str.strip()
    for _, row in work[src_series.eq("auto_crawl")].iterrows():
        params = row.get("params", {})
        params = params if isinstance(params, dict) else {}
        rows.append(
            {
                "captured_at": row.get("captured_at"),
                "page_id": _extract_page_id(_to_text_value(row.get("page_url", "")), params),
                "section_name": _to_text_value(params.get("section_name", "")),
                "target_id": _to_text_value(params.get("target_id", "")),
                "selector": _to_text_value(params.get("selector", "")),
                "class_attribute": _to_text_value(params.get("class_attribute", "")),
                "screen_state": _to_text_value(params.get("screen_state", params.get("screen_name", ""))),
                "screen_name": _to_text_value(params.get("screen_name", params.get("page_id", ""))),
                "selector_screenshot": _to_text_value(params.get("selector_screenshot", params.get("screenshot_path", ""))),
                "screenshot_path": _to_text_value(params.get("screenshot_path", params.get("selector_screenshot", ""))),
                "raw_screenshot_file": _to_text_value(params.get("raw_screenshot_file", "")),
                "annotation_no": _to_text_value(params.get("annotation_no", params.get("auto_click_index", ""))),
                "bbox_x": _to_text_value(params.get("bbox_x", 0)),
                "bbox_y": _to_text_value(params.get("bbox_y", 0)),
                "bbox_width": _to_text_value(params.get("bbox_width", 0)),
                "bbox_height": _to_text_value(params.get("bbox_height", 0)),
            }
        )
    return rows


def debug_events_to_full_csv_bytes(debug_df: pd.DataFrame) -> bytes:
    if debug_df.empty:
        return b""
    work = debug_df.copy()
    if "params" in work.columns:
        work["params_json"] = work["params"].apply(
            lambda x: json.dumps(x, ensure_ascii=False, sort_keys=True) if isinstance(x, dict) else "{}"
        )
        work = work.drop(columns=["params"])
    return work.to_csv(index=False).encode("utf-8-sig")


def debug_file_to_full_csv_bytes(output_file: Path) -> bytes:
    path = Path(output_file)
    if not path.exists():
        return b""
    try:
        full_df = pd.read_json(path, lines=True)
    except Exception:
        return b""
    return debug_events_to_full_csv_bytes(full_df)


def _serialize_realtime_events_for_export(rt_events: pd.DataFrame) -> pd.DataFrame:
    if rt_events.empty:
        return pd.DataFrame()
    work = rt_events.copy()
    for col in ["정상 그룹", "값없음 그룹", "의심 그룹", "시스템 그룹"]:
        if col in work.columns:
            work[col] = work[col].apply(
                lambda x: json.dumps(x, ensure_ascii=False, sort_keys=True) if isinstance(x, dict) else "{}"
            )
    if "captured_at" in work.columns:
        work["captured_at"] = pd.to_datetime(work["captured_at"], errors="coerce").astype(str)
    ordered_cols = [
        "group_session_id",
        "시간",
        "이벤트",
        "대표 파라미터",
        "대표 값",
        "상태",
        "정상 그룹",
        "값없음 그룹",
        "의심 그룹",
        "시스템 그룹",
        "captured_at",
    ]
    keep_cols = [c for c in ordered_cols if c in work.columns]
    return work[keep_cols]


def realtime_events_to_csv_bytes(rt_events: pd.DataFrame) -> bytes:
    if rt_events.empty:
        return b""
    export_df = _serialize_realtime_events_for_export(rt_events)
    if export_df.empty:
        return b""
    return export_df.to_csv(index=False).encode("utf-8-sig")


def realtime_events_to_excel_bytes(rt_events: pd.DataFrame) -> bytes:
    if rt_events.empty:
        return b""
    export_df = _serialize_realtime_events_for_export(rt_events)
    if export_df.empty:
        return b""
    out = io.BytesIO()
    try:
        with pd.ExcelWriter(out) as writer:
            export_df.to_excel(writer, index=False, sheet_name="realtime_events")
        return out.getvalue()
    except Exception:
        return b""


def summarize_event_conclusion(row_ev: Dict[str, object]) -> Dict[str, object]:
    normal_group = row_ev.get("정상 그룹", {})
    missing_group = row_ev.get("값없음 그룹", {})
    suspicious_group = row_ev.get("의심 그룹", {})
    system_group = row_ev.get("시스템 그룹", {})

    normal_count = len(normal_group) if isinstance(normal_group, dict) else 0
    missing_count = len(missing_group) if isinstance(missing_group, dict) else 0
    suspicious_count = len(suspicious_group) if isinstance(suspicious_group, dict) else 0
    system_count = len(system_group) if isinstance(system_group, dict) else 0

    primary_key = str(row_ev.get("대표 파라미터", "-")).strip()
    primary_value = str(row_ev.get("대표 값", "-")).strip()
    if (
        primary_key
        and primary_key != "-"
        and primary_value
        and primary_value != "-"
        and not _is_system_param_key(primary_key)
        and not _is_missing_like_value(primary_value)
    ):
        normal_count += 1

    critical_count = 1 if str(row_ev.get("상태", "")).upper() == "ERROR" else 0
    caution_count = missing_count + suspicious_count
    info_count = system_count

    if critical_count > 0:
        message = "이 이벤트는 오류 가능성이 있어 우선 수정이 필요합니다."
    elif caution_count > 0 and suspicious_count > 0:
        message = "이 이벤트는 기능적으로는 정상, 다만 불필요 파라미터 다수 포함 가능성이 있습니다."
    elif caution_count > 0:
        message = "이 이벤트는 기능적으로는 정상, 다만 비어 있는 값 확인이 필요합니다."
    else:
        message = "이 이벤트는 기능적으로 정상입니다."

    return {
        "critical": int(critical_count),
        "caution": int(caution_count),
        "info": int(info_count),
        "ok": int(normal_count),
        "message": message,
    }


def summarize_realtime_quality(rt_events: pd.DataFrame) -> Dict[str, object]:
    if rt_events.empty:
        return {
            "score": 0,
            "critical": 0,
            "caution": 0,
            "info": 0,
            "ok": 0,
            "event_count": 0,
            "message": "실시간 데이터가 없어 점수를 계산할 수 없습니다.",
        }

    total_critical = 0
    total_caution = 0
    total_info = 0
    total_ok = 0
    for row_ev in rt_events.to_dict("records"):
        event_summary = summarize_event_conclusion(row_ev)
        total_critical += int(event_summary["critical"])
        total_caution += int(event_summary["caution"])
        total_info += int(event_summary["info"])
        total_ok += int(event_summary["ok"])

    signals = total_critical + total_caution + total_info + total_ok
    if signals <= 0:
        score = 100
    else:
        risk_ratio = (
            (1.0 * total_critical) + (0.45 * total_caution) + (0.1 * total_info)
        ) / float(signals)
        score = int(round(max(0.0, min(1.0, 1.0 - risk_ratio)) * 100))

    if score >= 85:
        message = "기능적으로 양호한 상태입니다."
    elif score >= 70:
        message = "전반적으로 정상이나 일부 점검이 필요합니다."
    elif score >= 50:
        message = "주의 항목이 많아 우선 점검이 필요합니다."
    else:
        message = "치명/주의 항목 비중이 높아 즉시 점검이 필요합니다."

    return {
        "score": score,
        "critical": int(total_critical),
        "caution": int(total_caution),
        "info": int(total_info),
        "ok": int(total_ok),
        "event_count": int(len(rt_events)),
        "message": message,
    }


def build_realtime_event_rows(
    debug_df: pd.DataFrame,
    allowed_events: List[str] | None = None,
    unknown_event_policy: str = "정보",
    schemas: Dict[str, Dict[str, object]] | None = None,
    scoped_event_names: set[str] | None = None,
) -> pd.DataFrame:
    cols = [
        "group_session_id",
        "시간",
        "이벤트",
        "대표 파라미터",
        "대표 값",
        "상태",
        "status",
        "정상 그룹",
        "값없음 그룹",
        "의심 그룹",
        "시스템 그룹",
        "전체 파라미터",
        "custom_event",
        "custom_parameters",
        "target_id",
        "selector",
        "screen_name",
        "page_id",
        "section_name",
        "screenshot_path",
        "selector_screenshot",
        "raw_screenshot_file",
        "annotation_no",
        "bbox_x",
        "bbox_y",
        "bbox_width",
        "bbox_height",
        "bounding_box",
        "captured_at",
    ]
    if debug_df.empty:
        return pd.DataFrame(columns=cols)

    work = debug_df.copy()
    if "captured_at" in work.columns:
        work["captured_at"] = pd.to_datetime(work["captured_at"], errors="coerce")
        work = work.sort_values("captured_at", ascending=False)

    suspicious_profile = build_suspicious_param_profile(work)
    schema_store = schemas if isinstance(schemas, dict) else {}
    event_scope = {str(v).strip() for v in scoped_event_names if str(v).strip()} if scoped_event_names else None
    auto_ctx_rows = _collect_auto_context_rows(work)

    rows: List[Dict[str, object]] = []
    for _, row in work.iterrows():
        event_name = str(row.get("event_name", "")).strip()
        if not event_name:
            continue
        raw_params = row.get("params")
        params = raw_params if isinstance(raw_params, dict) else {}
        all_params = dict(params)

        measurement_id = _to_text_value(row.get("measurement_id", ""))
        client_id = _to_text_value(row.get("client_id", ""))
        if measurement_id and "measurement_id" not in all_params:
            all_params["measurement_id"] = measurement_id
        if client_id and "client_id" not in all_params:
            all_params["client_id"] = client_id
        custom_event_name, custom_items = extract_custom_event_payload(event_name, all_params, schema_store)
        if not custom_event_name:
            continue
        if event_scope is not None and custom_event_name not in event_scope:
            continue
        custom_params = {key: value for key, value in custom_items}

        normal_group: Dict[str, Dict[str, str]] = {}
        missing_group: Dict[str, Dict[str, str]] = {}
        suspicious_group: Dict[str, Dict[str, str]] = {}
        system_group: Dict[str, Dict[str, str]] = {}
        for key, value in custom_params.items():
            key_text = str(key).strip()
            value_text = _to_text_value(value)
            if not key_text:
                continue
            if _is_missing_like_value(value_text):
                missing_group[key_text] = {"값": value_text or "-", "사유": "미입력/기본값"}
                continue
            if key_text in suspicious_profile:
                suspicious_group[key_text] = {"값": value_text, "사유": suspicious_profile[key_text]}
                continue
            valid_ok, valid_msg = validate_param_format(key_text, value_text)
            if not valid_ok:
                missing_group[key_text] = {"값": value_text, "사유": f"형식 오류: {valid_msg}"}
                continue
            normal_group[key_text] = {"값": value_text, "사유": "형식 검증 통과"}

        primary_key = "-"
        primary_value = "-"
        for cand in PRIMARY_PARAM_PRIORITY:
            if cand in normal_group:
                primary_key = cand
                primary_value = normal_group[cand]["값"]
                break
        if primary_key == "-" and normal_group:
            first_key = next(iter(normal_group.keys()))
            primary_key = first_key
            primary_value = normal_group[first_key]["값"]
        if primary_key == "-" and suspicious_group:
            first_key = next(iter(suspicious_group.keys()))
            primary_key = first_key
            primary_value = suspicious_group[first_key]["값"]
        if primary_key == "-" and missing_group:
            first_key = next(iter(missing_group.keys()))
            primary_key = first_key
            primary_value = missing_group[first_key]["값"]

        if primary_key in normal_group:
            normal_group = {k: v for k, v in normal_group.items() if k != primary_key}
        if primary_key in suspicious_group:
            suspicious_group = {k: v for k, v in suspicious_group.items() if k != primary_key}
        if primary_key in missing_group:
            missing_group = {k: v for k, v in missing_group.items() if k != primary_key}

        group_session_id = _to_text_value(row.get("session_id", "")) or _to_text_value(params.get("qa_debug_session_id", "")) or "-"
        status = evaluate_event_status(
            event_name=custom_event_name,
            params=custom_params,
            allowed_events=allowed_events,
            unknown_event_policy=unknown_event_policy,
        )
        page_id = _extract_page_id(_to_text_value(row.get("page_url", "")), all_params)
        auto_ctx = _find_nearest_auto_context(auto_ctx_rows, row.get("captured_at"), page_id)
        meta = _build_event_meta(row, all_params, auto_ctx=auto_ctx)
        bounding_box = {
            "x": int(pd.to_numeric(meta.get("bbox_x", 0), errors="coerce") or 0),
            "y": int(pd.to_numeric(meta.get("bbox_y", 0), errors="coerce") or 0),
            "width": int(pd.to_numeric(meta.get("bbox_width", 0), errors="coerce") or 0),
            "height": int(pd.to_numeric(meta.get("bbox_height", 0), errors="coerce") or 0),
        }

        rows.append(
            {
                "group_session_id": group_session_id,
                "시간": format_local_time(row.get("captured_at")),
                "이벤트": custom_event_name,
                "대표 파라미터": primary_key,
                "대표 값": primary_value,
                "상태": status,
                "status": status,
                "정상 그룹": normal_group,
                "값없음 그룹": missing_group,
                "의심 그룹": suspicious_group,
                "시스템 그룹": system_group,
                "전체 파라미터": custom_params,
                "원본 payload": all_params,
                "custom_event": custom_event_name,
                "custom_parameters": " | ".join([f"{key}={value}" for key, value in custom_items]) if custom_items else "-",
                "target_id": str(meta.get("target_id", "-")).strip() or "-",
                "selector": str(meta.get("selector", "-")).strip() or "-",
                "screen_name": str(meta.get("screen_name", meta.get("page_id", "-"))).strip() or "-",
                "page_id": str(meta.get("page_id", "-")).strip() or "-",
                "section_name": str(meta.get("section_name", "-")).strip() or "-",
                "screenshot_path": str(meta.get("screenshot_path", meta.get("screenshot_file", ""))).strip() or "-",
                "selector_screenshot": str(meta.get("screenshot_file", "")).strip() or "-",
                "raw_screenshot_file": str(meta.get("raw_screenshot_file", "")).strip() or "-",
                "annotation_no": str(meta.get("annotation_no", "-")).strip() or "-",
                "bbox_x": bounding_box["x"],
                "bbox_y": bounding_box["y"],
                "bbox_width": bounding_box["width"],
                "bbox_height": bounding_box["height"],
                "bounding_box": json.dumps(bounding_box, ensure_ascii=False),
                "captured_at": row.get("captured_at"),
            }
        )

    out = pd.DataFrame(rows)
    if out.empty:
        return pd.DataFrame(columns=cols)
    return out


def build_realtime_funnel_progress(debug_df: pd.DataFrame, steps: List[str]) -> pd.DataFrame:
    if debug_df.empty or not steps:
        return pd.DataFrame(columns=["step", "상태", "최초 감지 시각"])

    work = debug_df.copy()
    if "captured_at" in work.columns:
        work["captured_at"] = pd.to_datetime(work["captured_at"], errors="coerce")
        work = work.sort_values("captured_at")

    first_seen: Dict[str, str] = {}
    ordered_events: List[str] = []
    for _, row in work.iterrows():
        event_name = str(row.get("event_name", "")).strip()
        if not event_name:
            continue
        ordered_events.append(event_name)
        if event_name not in first_seen:
            first_seen[event_name] = format_local_time(row.get("captured_at"))

    pointer = 0
    for ev in ordered_events:
        if pointer < len(steps) and ev == steps[pointer]:
            pointer += 1

    rows: List[Dict[str, str]] = []
    observed_set = set(ordered_events)
    for idx, step in enumerate(steps):
        if idx < pointer:
            status = "완료"
        elif step in observed_set:
            status = "감지됨(순서대기)"
        else:
            status = "대기"
        rows.append(
            {
                "step": step,
                "상태": status,
                "최초 감지 시각": first_seen.get(step, "-"),
            }
        )
    return pd.DataFrame(rows)


def build_qa_results_view(results_df: pd.DataFrame) -> pd.DataFrame:
    if results_df.empty:
        return results_df

    view = results_df.copy()
    view["구분"] = "기타"
    view["이벤트"] = ""
    view["파라미터"] = ""

    for idx, rule_id in view["rule_id"].astype(str).items():
        if rule_id.startswith("required_event:"):
            event_name = rule_id.split(":", 1)[1]
            view.at[idx, "구분"] = "필수 이벤트"
            view.at[idx, "이벤트"] = event_name
        elif rule_id.startswith("required_param:"):
            parts = rule_id.split(":", 2)
            if len(parts) == 3:
                view.at[idx, "구분"] = "필수 파라미터"
                view.at[idx, "이벤트"] = parts[1]
                view.at[idx, "파라미터"] = parts[2]
        elif rule_id.startswith("null_ratio:"):
            param_name = rule_id.split(":", 1)[1]
            view.at[idx, "구분"] = "Null 비율"
            view.at[idx, "파라미터"] = param_name
        elif rule_id == "value_numeric":
            view.at[idx, "구분"] = "타입 체크"
            view.at[idx, "이벤트"] = "purchase"
            view.at[idx, "파라미터"] = "value"
        elif rule_id == "duplicate_transaction_id":
            view.at[idx, "구분"] = "중복 체크"
            view.at[idx, "파라미터"] = "transaction_id"
        elif rule_id in {"purchase_without_value", "purchase_without_add_to_cart", "purchase_without_sign_up"}:
            view.at[idx, "구분"] = "비즈니스 룰"
        elif rule_id == "funnel_order":
            view.at[idx, "구분"] = "퍼널 순서"
        elif rule_id == "missing_event_timestamp":
            view.at[idx, "구분"] = "로그 품질"
            view.at[idx, "파라미터"] = "event_timestamp"

    ordered_cols = [
        "구분",
        "이벤트",
        "파라미터",
        "status",
        "detail",
        "rule_name",
        "rule_id",
        "fail_count",
        "total_count",
    ]
    keep_cols = [col for col in ordered_cols if col in view.columns]
    return view[keep_cols]


def build_event_preview(df: pd.DataFrame, requested_params: List[str]) -> pd.DataFrame:
    if df.empty:
        return df

    requested_norm = [normalize_param_name(p) for p in requested_params if p.strip()]
    base_cols = ["event_name", "event_date", "event_timestamp", "user_pseudo_id"]
    priority_params = ["transaction_id", "qa_debug_session_id", "value", "currency", "page_location"]
    param_cols = list(dict.fromkeys(priority_params + requested_norm))

    ordered_cols: List[str] = []
    for col in base_cols + param_cols:
        if col in df.columns:
            ordered_cols.append(col)
    for col in df.columns:
        if col not in ordered_cols:
            ordered_cols.append(col)

    return df[ordered_cols]


def resolve_debug_output_file(
    session_id: str,
    snapshot: Dict[str, object],
    state_output_file: str,
) -> Path | None:
    candidates: List[Path] = []
    snap_path_text = str(snapshot.get("output_file", "")).strip() if snapshot else ""
    if snap_path_text:
        candidates.append(Path(snap_path_text))

    state_path_text = str(state_output_file or "").strip()
    if state_path_text:
        candidates.append(Path(state_path_text))

    sid = str(session_id or "").strip()
    if sid:
        candidates.append(get_active_project_paths()["qa_sessions_dir"] / sid / "debug_stream.jsonl")
        candidates.append(Path(f"data/debug_stream/{sid}.jsonl"))

    for cand in candidates:
        if cand.exists():
            return cand
    return candidates[0] if candidates else None


def restore_latest_project_session_state() -> None:
    current_sid = str(st.session_state.get("qa_debug_session_id", "")).strip()
    current_output = str(st.session_state.get("qa_debug_output_file", "")).strip()
    if current_sid and current_output and Path(current_output).exists():
        return
    try:
        sess_df = list_recent_sessions(get_active_test_log_db_path(), limit=1)
    except Exception:
        sess_df = pd.DataFrame()
    if not sess_df.empty:
        sid = str(sess_df.iloc[0].get("session_id", "")).strip()
        if sid:
            snap = get_session(get_active_test_log_db_path(), sid)
            output_file = get_active_project_paths()["qa_sessions_dir"] / sid / "debug_stream.jsonl"
            if not output_file.exists():
                fallback = Path(f"data/debug_stream/{sid}.jsonl")
                if fallback.exists():
                    output_file = fallback
            if output_file.exists():
                st.session_state["qa_debug_session_id"] = sid
                st.session_state["qa_debug_output_file"] = str(output_file)
                st.session_state["qa_debug_started_at"] = str(snap.get("started_at", "")).strip()
                return

    candidates: List[Path] = []
    try:
        candidates.extend(
            sorted(
                get_active_project_paths()["qa_sessions_dir"].glob("*/debug_stream.jsonl"),
                key=lambda p: p.stat().st_mtime,
                reverse=True,
            )
        )
    except Exception:
        pass
    try:
        candidates.extend(
            sorted(
                Path("data/debug_stream").glob("*.jsonl"),
                key=lambda p: p.stat().st_mtime,
                reverse=True,
            )
        )
    except Exception:
        pass
    for candidate in candidates:
        if not candidate.exists():
            continue
        sid = candidate.parent.name if candidate.parent.name != "debug_stream" else candidate.stem
        if not sid:
            continue
        st.session_state["qa_debug_session_id"] = sid
        st.session_state["qa_debug_output_file"] = str(candidate)
        if not str(st.session_state.get("qa_debug_started_at", "")).strip():
            st.session_state["qa_debug_started_at"] = ""
        return


def inject_case_from_page_location(df: pd.DataFrame, case_param: str) -> pd.DataFrame:
    param = case_param.strip()
    if df.empty or not param or "page_location" not in df.columns:
        return df

    out = df.copy()
    if param not in out.columns:
        out[param] = pd.NA

    missing_mask = out[param].isna() | (out[param].astype(str).str.strip() == "")
    if not missing_mask.any():
        return out

    def extract_value(page_url: object) -> str | None:
        if page_url is None or pd.isna(page_url):
            return None
        try:
            parsed = urlparse(str(page_url))
            values = parse_qs(parsed.query).get(param, [])
            if values:
                return values[0]
            return None
        except Exception:
            return None

    out.loc[missing_mask, param] = out.loc[missing_mask, "page_location"].apply(extract_value)
    return out


def _safe_params_dict(value: object) -> Dict[str, object]:
    if isinstance(value, dict):
        return value
    return {}


def source_to_view_label(source: str) -> str:
    src = (source or "").strip()
    if src == "ga_hit":
        return "히트(collect)"
    if src == "gtag":
        return "발화(gtag)"
    if src == "dataLayer.push":
        return "발화(dataLayer)"
    return src or "-"


def filter_debug_events_by_view(debug_df: pd.DataFrame, capture_view: str) -> pd.DataFrame:
    if debug_df.empty or "source" not in debug_df.columns:
        return debug_df
    src_series = debug_df["source"].astype(str).str.strip()
    if capture_view == "히트(collect)":
        hit_df = debug_df[src_series == "ga_hit"].copy()
        if not hit_df.empty:
            return hit_df
        return debug_df[src_series == "auto_crawl"].copy()
    if capture_view == "발화(dataLayer/gtag)":
        return debug_df[src_series != "ga_hit"].copy()
    return debug_df.copy()


def summarize_major_params(params: Dict[str, object]) -> str:
    if not params:
        return ""
    priority = ["measurement_id", "transaction_id", "value", "currency", "qa_debug_session_id"]
    parts: List[str] = []
    for key in priority:
        if key in params and str(params.get(key, "")).strip():
            parts.append(f"{key}={params[key]}")
    if not parts:
        for key, value in params.items():
            if str(value).strip():
                parts.append(f"{key}={value}")
            if len(parts) >= 3:
                break
    return ", ".join(parts)


def build_realtime_timeline_view(
    debug_df: pd.DataFrame,
    allowed_events: List[str] | None = None,
    unknown_event_policy: str = "정보",
) -> pd.DataFrame:
    if debug_df.empty:
        return pd.DataFrame(columns=["시간", "수집원", "event_name", "주요 파라미터", "상태", "상세"])

    work = debug_df.copy()
    if "captured_at" in work.columns:
        work["captured_at"] = pd.to_datetime(work["captured_at"], errors="coerce")
        work = work.sort_values("captured_at")

    seen_steps: set[str] = set()
    allowed_set = {e.strip() for e in (allowed_events or []) if str(e).strip()}
    internal_prefixes = ("gtm.",)
    internal_events = {"set_user_property"}
    out_rows: List[Dict[str, str]] = []

    for _, row in work.iterrows():
        event_name = str(row.get("event_name", "")).strip()
        params = _safe_params_dict(row.get("params"))
        measurement_id = str(row.get("measurement_id", "")).strip()
        if measurement_id and "measurement_id" not in params:
            params = dict(params)
            params["measurement_id"] = measurement_id
        issues: List[str] = []
        status = "OK"

        if not event_name:
            status = "ERROR"
            issues.append("event_name 없음")
        elif (
            allowed_set
            and event_name not in allowed_set
            and not event_name.startswith(internal_prefixes)
            and event_name not in internal_events
        ):
            if unknown_event_policy == "경고":
                status = "WARN" if status == "OK" else status
                issues.append("정의되지 않은 이벤트명")
            elif unknown_event_policy == "정보":
                if status == "OK":
                    status = "INFO"
                issues.append("정의되지 않은 이벤트명")

        if event_name == "purchase":
            if not str(params.get("value", "")).strip():
                status = "WARN" if status == "OK" else status
                issues.append("purchase에 value 없음")
            if not str(params.get("transaction_id", "")).strip():
                status = "ERROR"
                issues.append("purchase에 transaction_id 없음")
            if "add_to_cart" not in seen_steps:
                status = "WARN" if status == "OK" else status
                issues.append("add_to_cart 없이 purchase")

        if event_name:
            seen_steps.add(event_name)

        ts_text = format_local_time(row.get("captured_at"))

        out_rows.append(
            {
                "시간": ts_text,
                "수집원": source_to_view_label(str(row.get("source", ""))),
                "event_name": event_name or "(empty)",
                "주요 파라미터": summarize_major_params(params),
                "상태": status,
                "상세": " | ".join(issues) if issues else "-",
            }
        )

    out = pd.DataFrame(out_rows)
    return out.iloc[::-1].reset_index(drop=True)


def style_realtime_timeline(df: pd.DataFrame):
    def _row_style(row):
        status = str(row.get("상태", ""))
        if status == "ERROR":
            return ["background-color: #ffebee"] * len(row)
        if status == "WARN":
            return ["background-color: #fff8e1"] * len(row)
        if status == "INFO":
            return ["background-color: #e3f2fd"] * len(row)
        return [""] * len(row)

    return df.style.apply(_row_style, axis=1)


def to_user_error_message(exc: Exception) -> str:
    msg = str(exc).strip()
    low = msg.lower()

    if "playwright" in low and ("not found" in low or "install" in low):
        return "브라우저 디버깅 의존성 오류: `pip install playwright && playwright install chromium` 실행 후 다시 시도하세요."
    if "디버깅 브라우저 실행에 실패했습니다" in msg:
        return (
            "디버깅 브라우저 실행 실패: EC2에서 아래 순서로 복구하세요.\n"
            "1) `cd /opt/ga4-qa-mvp && source .venv/bin/activate`\n"
            "2) `playwright install --with-deps chromium`\n"
            "3) `sudo systemctl restart ga4-qa-mvp`"
        )
    return msg if msg else "알 수 없는 오류가 발생했습니다."


process_google_oauth_callback_if_present()


source = "실시간 디버깅 스트림"
default_required_events = ""
default_param_text = ""
default_funnel_steps = ""
default_scenario_steps = "view_item,add_to_cart,begin_checkout,purchase"
default_params = "transaction_id,value,currency,qa_debug_session_id"
scenario_key_modes = {
    "transaction_id 기준": "transaction_id",
    "user_pseudo_id 기준": "user_pseudo_id",
    "qa_debug_session_id 기준": "qa_debug_session_id",
    "키 없이 집계형": "none",
}
risk_level_ko = {"Low": "낮음", "Medium": "중간", "High": "높음"}
scenario_mode_ko = {"key": "키 기반", "user": "사용자 기반", "aggregate": "집계형", "empty": "데이터 없음"}
debug_case_param_name = "qa_debug_session_id"

if "qa_debug_session_id" not in st.session_state:
    st.session_state["qa_debug_session_id"] = ""
if "qa_debug_output_file" not in st.session_state:
    st.session_state["qa_debug_output_file"] = ""
if "qa_debug_started_at" not in st.session_state:
    st.session_state["qa_debug_started_at"] = ""
if "qa_debug_target_url" not in st.session_state:
    st.session_state["qa_debug_target_url"] = ""
if "qa_tester_name" not in st.session_state:
    st.session_state["qa_tester_name"] = ""
if "qa_tester_note" not in st.session_state:
    st.session_state["qa_tester_note"] = ""
if "unknown_event_policy" not in st.session_state:
    st.session_state["unknown_event_policy"] = "정보"
if "realtime_panel_auto_refresh" not in st.session_state:
    st.session_state["realtime_panel_auto_refresh"] = False
if "realtime_panel_refresh_interval" not in st.session_state:
    st.session_state["realtime_panel_refresh_interval"] = 2
if "realtime_last_refresh_at" not in st.session_state:
    st.session_state["realtime_last_refresh_at"] = ""
if "qa_report_client_secret_file" not in st.session_state:
    st.session_state["qa_report_client_secret_file"] = get_config_value("GA4_CLIENT_SECRETS_FILE", "client_secret.json")
if "qa_report_token_file" not in st.session_state:
    st.session_state["qa_report_token_file"] = get_config_value("GA4_TOKEN_FILE", "token.json")
if "qa_report_property_options" not in st.session_state:
    st.session_state["qa_report_property_options"] = []
if "qa_oauth_auth_url" not in st.session_state:
    st.session_state["qa_oauth_auth_url"] = ""
if "qa_oauth_state" not in st.session_state:
    st.session_state["qa_oauth_state"] = ""
if "qa_oauth_notice" not in st.session_state:
    st.session_state["qa_oauth_notice"] = ""
if "qa_oauth_error" not in st.session_state:
    st.session_state["qa_oauth_error"] = ""
if "qa_ingest_server_error" not in st.session_state:
    st.session_state["qa_ingest_server_error"] = ""
if "qa_ingest_local_ok" not in st.session_state:
    st.session_state["qa_ingest_local_ok"] = False
if "qa_access_authenticated" not in st.session_state:
    st.session_state["qa_access_authenticated"] = False
if "qa_access_role" not in st.session_state:
    st.session_state["qa_access_role"] = "guest"
if "qa_access_user" not in st.session_state:
    st.session_state["qa_access_user"] = ""
if "qa_oauth_go_now_url" not in st.session_state:
    st.session_state["qa_oauth_go_now_url"] = ""
if "qa_oauth_resume_notice" not in st.session_state:
    st.session_state["qa_oauth_resume_notice"] = ""
if "qa_ui_focus_after_oauth" not in st.session_state:
    st.session_state["qa_ui_focus_after_oauth"] = ""
if "qa_project_slug" not in st.session_state:
    st.session_state["qa_project_slug"] = DEFAULT_PROJECT_SLUG
if "qa_project_domain" not in st.session_state:
    st.session_state["qa_project_domain"] = ""
if "qa_new_project_name" not in st.session_state:
    st.session_state["qa_new_project_name"] = ""
if "qa_new_project_domain" not in st.session_state:
    st.session_state["qa_new_project_domain"] = ""
if "qa_project_notice" not in st.session_state:
    st.session_state["qa_project_notice"] = ""
if "qa_environment" not in st.session_state:
    st.session_state["qa_environment"] = "prod"
if "qa_browser_name" not in st.session_state:
    st.session_state["qa_browser_name"] = "Chrome"
if "qa_viewport_preset" not in st.session_state:
    st.session_state["qa_viewport_preset"] = "Desktop 1440 x 900"
if "qa_mode" not in st.session_state:
    st.session_state["qa_mode"] = "전체 이벤트 테스트"
if "qa_definition_validation_file_name" not in st.session_state:
    st.session_state["qa_definition_validation_file_name"] = ""
if "qa_definition_validation_error" not in st.session_state:
    st.session_state["qa_definition_validation_error"] = ""
if "qa_uploaded_definition_schemas" not in st.session_state:
    st.session_state["qa_uploaded_definition_schemas"] = {}
if "qa_uploaded_definition_summary" not in st.session_state:
    st.session_state["qa_uploaded_definition_summary"] = []
if "qa_definition_validation_debug" not in st.session_state:
    st.session_state["qa_definition_validation_debug"] = {}
if "qa_definition_validation_loaded_key" not in st.session_state:
    st.session_state["qa_definition_validation_loaded_key"] = ""
if "qa_definition_validation_target_mode" not in st.session_state:
    st.session_state["qa_definition_validation_target_mode"] = "선택하세요"
if "qa_definition_validation_target_mode_prev" not in st.session_state:
    st.session_state["qa_definition_validation_target_mode_prev"] = "선택하세요"
if "qa_definition_validation_execution_mode" not in st.session_state:
    st.session_state["qa_definition_validation_execution_mode"] = ""
if "qa_definition_validation_execution_session_id" not in st.session_state:
    st.session_state["qa_definition_validation_execution_session_id"] = ""
if "qa_scenario_template" not in st.session_state:
    st.session_state["qa_scenario_template"] = "상품 구매 흐름"
if "qa_auto_crawl_max_clicks" not in st.session_state:
    st.session_state["qa_auto_crawl_max_clicks"] = int(get_config_value("QA_AUTO_CRAWL_MAX_CLICKS", "80") or "80")
if "qa_auto_crawl_click_interval_ms" not in st.session_state:
    st.session_state["qa_auto_crawl_click_interval_ms"] = int(
        get_config_value("QA_AUTO_CRAWL_CLICK_INTERVAL_MS", "1200") or "1200"
    )
if "qa_auto_crawl_wait_after_click_ms" not in st.session_state:
    st.session_state["qa_auto_crawl_wait_after_click_ms"] = int(
        get_config_value("QA_AUTO_CRAWL_WAIT_AFTER_CLICK_MS", "1200") or "1200"
    )
if "qa_auto_block_link_nav" not in st.session_state:
    st.session_state["qa_auto_block_link_nav"] = is_truthy(get_config_value("QA_AUTO_BLOCK_LINK_NAV", "1"))
if "qa_single_page_only" not in st.session_state:
    st.session_state["qa_single_page_only"] = is_truthy(get_config_value("QA_SINGLE_PAGE_ONLY", "1"))
if "qa_template_seed" not in st.session_state:
    st.session_state["qa_template_seed"] = ""
if "qa_test_setup_snapshot" not in st.session_state:
    st.session_state["qa_test_setup_snapshot"] = {}
if "qa_preserve_test_setup_during_run" not in st.session_state:
    st.session_state["qa_preserve_test_setup_during_run"] = False

ensure_project_structure(st.session_state.get("qa_project_slug", DEFAULT_PROJECT_SLUG))
restore_definition_validation_cache_if_needed()
restore_test_setup_snapshot_if_needed(force=bool(st.session_state.get("qa_preserve_test_setup_during_run", False)))

enforce_access_gate()

ensure_err = ""
# 기본은 독립 systemd 수집기(ga4-qa-ingest.service)를 사용한다.
ingest_embedded_mode = is_truthy(get_config_value("QA_INGEST_EMBEDDED_MODE", "0"))
if ingest_embedded_mode:
    try:
        ensure_ingest_server(host="127.0.0.1", port=8600)
    except Exception as exc:
        ensure_err = str(exc)

local_health_ok = probe_ingest_health("http://127.0.0.1:8600/qa/health")

st.session_state["qa_ingest_local_ok"] = bool(local_health_ok)
st.session_state["qa_ingest_server_ok"] = bool(local_health_ok)
if st.session_state["qa_ingest_server_ok"]:
    st.session_state["qa_ingest_server_error"] = ""
else:
    detail_bits: List[str] = []
    if ensure_err:
        detail_bits.append(f"ensure: {ensure_err}")
    detail_bits.append(f"local_health({'ok' if local_health_ok else 'fail'}): http://127.0.0.1:8600/qa/health")
    st.session_state["qa_ingest_server_error"] = " | ".join(detail_bits)

if not str(st.session_state.get("qa_tester_name", "")).strip():
    auto_tester = str(st.session_state.get("qa_access_user", "")).strip()
    if auto_tester:
        st.session_state["qa_tester_name"] = auto_tester

restore_latest_project_session_state()

scenario_enabled = True
scenario_steps_text = default_scenario_steps
scenario_key_mode_label = "transaction_id 기준"
scenario_key_field = "transaction_id"
scenario_key_value = ""
realtime_debug_start_clicked = False
realtime_debug_stop_clicked = False

with st.sidebar:
    st.header("QA Workflow")
    st.caption(
        f"접속: {st.session_state.get('qa_access_user', '-') or '-'} "
        f"({st.session_state.get('qa_access_role', 'guest')})"
    )
    if st.button("로그아웃", key="qa_access_logout_btn"):
        log_ui_action("access_logout")
        st.session_state["qa_access_authenticated"] = False
        st.session_state["qa_access_role"] = "guest"
        st.session_state["qa_access_user"] = ""
        st.rerun()
    with st.expander("1) Test Setup", expanded=True):
        project_notice = str(st.session_state.get("qa_project_notice", "")).strip()
        if project_notice:
            st.info(project_notice)
            st.session_state["qa_project_notice"] = ""
        project_options = list_workspace_projects()
        current_project = _slugify_project_name(st.session_state.get("qa_project_slug", DEFAULT_PROJECT_SLUG))
        if current_project not in project_options:
            project_options = sorted(list(dict.fromkeys(project_options + [current_project])))
        selected_project = st.selectbox(
            "Project",
            options=project_options,
            index=project_options.index(current_project) if current_project in project_options else 0,
            key="qa_project_slug_selectbox",
        )
        st.session_state["qa_project_slug"] = _slugify_project_name(selected_project)
        st.selectbox("Environment", options=QA_ENVIRONMENT_OPTIONS, key="qa_environment")
        st.text_input(
            "Target URL",
            value=st.session_state.get("qa_debug_target_url", ""),
            key="qa_debug_target_url",
            placeholder="예: https://www.musinsa.com/main/beauty/recommend?gf=A",
        )
        st.selectbox("Browser", options=QA_BROWSER_OPTIONS, key="qa_browser_name")
        st.selectbox("Viewport", options=list(QA_VIEWPORT_PRESETS.keys()), key="qa_viewport_preset")
        st.text_input(
            "Tester",
            value=st.session_state.get("qa_tester_name", ""),
            key="qa_tester_name",
            placeholder="자동 채움 가능",
        )
        st.text_area(
            "Memo",
            value=st.session_state.get("qa_tester_note", ""),
            key="qa_tester_note",
            placeholder="선택 입력",
            height=80,
        )
        st.caption("Project / Environment / URL만 입력하면 자동수집 중심으로 실행됩니다.")

    with st.expander("2) QA Mode / Scenario", expanded=True):
        current_mode = str(st.session_state.get("qa_mode", "전체 이벤트 테스트")).strip()
        if current_mode not in QA_MODE_OPTIONS:
            current_mode = "전체 이벤트 테스트"
        selected_mode = st.radio(
            "QA Mode",
            options=QA_MODE_OPTIONS,
            index=QA_MODE_OPTIONS.index(current_mode),
            key="qa_mode",
        )
        st.caption("이벤트 수집 방식은 기본적으로 Auto Crawl입니다.")

        if selected_mode == "시나리오 테스트":
            template_name = st.selectbox(
                "Scenario template",
                options=list(QA_SCENARIO_TEMPLATES.keys()),
                key="qa_scenario_template",
            )
            template = QA_SCENARIO_TEMPLATES[template_name]
            scenario_enabled = True
            scenario_steps_text = str(template["steps"])
            required_event_text = str(template["required_events"])
            funnel_text = str(template["funnel_steps"])
            scenario_key_mode_label = str(template["key_mode_label"])
            scenario_key_field = scenario_key_modes[scenario_key_mode_label]
            scenario_key_value = ""
            template_seed = f"{selected_mode}:{template_name}"
            if st.session_state.get("qa_template_seed", "") != template_seed:
                st.session_state["required_event_text_input"] = required_event_text
                st.session_state["qa_template_seed"] = template_seed
            st.caption(f"자동 채움 이벤트: {required_event_text}")
        elif selected_mode == "Tracking plan 검증":
            scenario_enabled = False
            scenario_steps_text = ""
            required_event_text = ""
            funnel_text = default_funnel_steps
            scenario_key_mode_label = "qa_debug_session_id 기준"
            scenario_key_field = scenario_key_modes[scenario_key_mode_label]
            scenario_key_value = ""
            st.session_state["qa_template_seed"] = selected_mode
            st.caption("자동수집 후 Tracking Plan / Event Definition Export를 바로 비교하는 모드입니다.")
        elif selected_mode == "정의서 검증":
            scenario_enabled = False
            scenario_steps_text = ""
            required_event_text = ""
            funnel_text = default_funnel_steps
            scenario_key_mode_label = "qa_debug_session_id 기준"
            scenario_key_field = scenario_key_modes[scenario_key_mode_label]
            scenario_key_value = ""
            st.session_state["qa_template_seed"] = selected_mode
            tpl1, tpl2 = st.columns(2)
            with tpl1:
                st.download_button(
                    "양식 다운로드 (XLSX)",
                    data=build_definition_template_xlsx_bytes(),
                    file_name="event_definition_template.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    help="현재 정의서 검증 파서가 바로 읽을 수 있는 업로드 양식입니다.",
                )
            with tpl2:
                st.download_button(
                    "양식 다운로드 (CSV)",
                    data=build_definition_template_csv_bytes(),
                    file_name="event_definition_template.csv",
                    mime="text/csv",
                    help="간단 편집용 CSV 양식입니다.",
                )
            st.caption("양식은 상단 이벤트 목록 + 하단 매개변수 매트릭스 구조입니다.")
            st.markdown("**1단계: 정의서 업로드**")
            definition_file = st.file_uploader(
                "이벤트 정의서 업로드",
                type=["csv", "xlsx", "xls"],
                key="qa_definition_validation_file",
                accept_multiple_files=False,
                help="event_definition.csv 또는 event_name/required/optional 형식의 파일을 올리면 그 기준으로만 PASS/FAIL 검증합니다.",
            )
            has_cached_definition = bool(st.session_state.get("qa_uploaded_definition_schemas", {}))
            if definition_file is None and not has_cached_definition:
                st.session_state["qa_definition_validation_file_name"] = ""
                st.session_state["qa_definition_validation_error"] = ""
                st.session_state["qa_uploaded_definition_schemas"] = {}
                st.session_state["qa_uploaded_definition_summary"] = []
                st.session_state["qa_definition_validation_debug"] = {}
                st.session_state["qa_definition_validation_loaded_key"] = ""
                st.session_state["qa_definition_validation_target_mode"] = "선택하세요"
                st.session_state["qa_definition_validation_target_mode_prev"] = "선택하세요"
                st.session_state["qa_definition_validation_execution_mode"] = ""
                st.session_state["qa_definition_validation_execution_session_id"] = ""
                st.info("정의서 파일을 업로드하면 그 파일의 이벤트/매개변수만 검증합니다.")
            else:
                uploaded_schemas: Dict[str, Dict[str, object]] = st.session_state.get("qa_uploaded_definition_schemas", {})
                uploaded_summary_df = pd.DataFrame(st.session_state.get("qa_uploaded_definition_summary", []))
                if definition_file is None and uploaded_schemas:
                    loaded_file_name = str(st.session_state.get("qa_definition_validation_file_name", "")).strip()
                    if loaded_file_name:
                        st.caption(f"현재 로드된 정의서: {loaded_file_name}")
                if definition_file is not None:
                    try:
                        uploaded_schemas, uploaded_summary_df, upload_debug = load_uploaded_definition_schemas(definition_file)
                        loaded_key = (
                            f"{str(definition_file.name).strip()}|"
                            f"{len(uploaded_schemas)}|"
                            f"{','.join(sorted(list(uploaded_schemas.keys()))[:5])}"
                        )
                        if st.session_state.get("qa_definition_validation_loaded_key", "") != loaded_key:
                            st.session_state["qa_definition_validation_target_mode"] = "선택하세요"
                            st.session_state["qa_definition_validation_target_mode_prev"] = "선택하세요"
                            st.session_state["qa_definition_validation_execution_mode"] = ""
                            st.session_state["qa_definition_validation_execution_session_id"] = ""
                        st.session_state["qa_definition_validation_file_name"] = str(definition_file.name).strip()
                        st.session_state["qa_definition_validation_error"] = ""
                        st.session_state["qa_uploaded_definition_schemas"] = uploaded_schemas
                        st.session_state["qa_uploaded_definition_summary"] = uploaded_summary_df.to_dict("records")
                        st.session_state["qa_definition_validation_debug"] = upload_debug
                        st.session_state["qa_definition_validation_loaded_key"] = loaded_key
                        save_definition_validation_cache(
                            {
                                "qa_mode": "정의서 검증",
                                "file_name": str(definition_file.name).strip(),
                                "schemas": uploaded_schemas,
                                "summary": uploaded_summary_df.to_dict("records"),
                                "debug": upload_debug,
                                "loaded_key": loaded_key,
                                "target_mode": st.session_state.get("qa_definition_validation_target_mode", "선택하세요"),
                                "execution_mode": "",
                                "execution_session_id": "",
                            }
                        )
                    except Exception as exc:
                        st.session_state["qa_definition_validation_file_name"] = str(definition_file.name).strip()
                        st.session_state["qa_definition_validation_error"] = to_user_error_message(exc)
                        st.session_state["qa_uploaded_definition_schemas"] = {}
                        st.session_state["qa_uploaded_definition_summary"] = []
                        st.session_state["qa_definition_validation_debug"] = getattr(exc, "debug_info", {}) or {}
                        st.session_state["qa_definition_validation_execution_mode"] = ""
                        st.session_state["qa_definition_validation_execution_session_id"] = ""
                        st.error(f"정의서 업로드 실패: {to_user_error_message(exc)}")
                uploaded_schemas = st.session_state.get("qa_uploaded_definition_schemas", {})
                uploaded_summary_df = pd.DataFrame(st.session_state.get("qa_uploaded_definition_summary", []))
                if uploaded_schemas:
                    schema_param_keys: set[str] = set()
                    for schema_obj in uploaded_schemas.values():
                        if isinstance(schema_obj, dict):
                            for key in list(schema_obj.get("parameter_columns", [])) + list(schema_obj.get("required", [])) + list(schema_obj.get("optional", [])):
                                key_text = str(key).strip()
                                if key_text:
                                    schema_param_keys.add(key_text)
                    sample_events = list(uploaded_schemas.keys())[:5]
                    st.success("검증 준비 완료")
                    u1, u2, u3, u4 = st.columns(4)
                    u1.metric("파일명", str(st.session_state.get("qa_definition_validation_file_name", "")).strip() or "-")
                    u2.metric("파싱 상태", "성공")
                    u3.metric("인식 이벤트 수", int(len(uploaded_summary_df)))
                    u4.metric("인식 파라미터 수", int(len(schema_param_keys)))
                    if sample_events:
                        st.caption(f"샘플 이벤트명: {', '.join(sample_events)}")
                debug_view = st.session_state.get("qa_definition_validation_debug", {})
                if debug_view:
                    with st.expander("업로드 진단 정보", expanded=bool(st.session_state.get("qa_definition_validation_error", ""))):
                        d1, d2, d3 = st.columns(3)
                        d1.metric("raw rows", int(debug_view.get("raw_rows", 0)))
                        d2.metric("raw cols", int(debug_view.get("raw_cols", 0)))
                        d3.metric("parsed events", int(debug_view.get("parsed_event_count", 0)))
                        st.caption(
                            f"detected event header row: {debug_view.get('event_header_row_index', '-')}, "
                            f"param header row: {debug_view.get('param_header_row_index', '-')}, "
                            f"fallback used: {debug_view.get('fallback_used', False)}"
                        )
                        sample_names = debug_view.get("sample_event_names", [])
                        if sample_names:
                            st.caption(f"sample event_name: {', '.join([str(v) for v in sample_names[:5]])}")
                        failure_stage = str(debug_view.get("failure_stage", "")).strip()
                        final_reason = str(debug_view.get("final_validation_reason", "")).strip()
                        if failure_stage or final_reason:
                            st.caption(f"final validation: {failure_stage or 'ok'} | {final_reason or 'ok'}")
                        read_attempts = debug_view.get("read_attempts", [])
                        if read_attempts:
                            st.caption(f"read attempts: {', '.join([str(v) for v in read_attempts])}")
                        sheet_debugs = debug_view.get("sheet_debugs", [])
                        if sheet_debugs:
                            st.dataframe(pd.DataFrame(sheet_debugs), use_container_width=True, height=200)
                if uploaded_schemas:
                    st.markdown("**2단계: 검증 대상 선택**")
                    target_mode = st.selectbox(
                        "검증 대상",
                        options=["선택하세요", "기존 세션으로 검증", "새 Run QA Test 실행 후 검증"],
                        key="qa_definition_validation_target_mode",
                    )
                    if st.session_state.get("qa_definition_validation_target_mode_prev", "선택하세요") != target_mode:
                        st.session_state["qa_definition_validation_execution_mode"] = ""
                        st.session_state["qa_definition_validation_execution_session_id"] = ""
                        st.session_state["qa_definition_validation_target_mode_prev"] = target_mode
                        save_definition_validation_cache(
                            {
                                "qa_mode": "정의서 검증",
                                "file_name": str(st.session_state.get("qa_definition_validation_file_name", "")).strip(),
                                "schemas": st.session_state.get("qa_uploaded_definition_schemas", {}),
                                "summary": st.session_state.get("qa_uploaded_definition_summary", []),
                                "debug": st.session_state.get("qa_definition_validation_debug", {}),
                                "loaded_key": st.session_state.get("qa_definition_validation_loaded_key", ""),
                                "target_mode": target_mode,
                                "execution_mode": "",
                                "execution_session_id": "",
                            }
                        )
                    if target_mode == "기존 세션으로 검증":
                        st.caption("현재 불러온 세션이 있으면 그 세션을 정의서 기준으로 재검증합니다.")
                    elif target_mode == "새 Run QA Test 실행 후 검증":
                        st.caption("3단계에서 새 Run QA Test를 실행하면 정의서 기준으로 새 세션 수집 후 검증합니다.")
                    else:
                        st.caption("다음 단계로 가려면 검증 대상을 먼저 선택하세요.")
            st.caption("업로드한 정의서에 있는 custom event / custom parameter만 검증합니다.")
        else:
            scenario_enabled = False
            scenario_steps_text = ""
            required_event_text = ""
            funnel_text = default_funnel_steps
            scenario_key_mode_label = "qa_debug_session_id 기준"
            scenario_key_field = scenario_key_modes[scenario_key_mode_label]
            scenario_key_value = ""
            st.session_state["qa_template_seed"] = selected_mode
            st.caption("대표 UI 패턴 기준으로 이벤트를 자동 수집합니다.")

    with st.expander("3) Run Test", expanded=True):
        st.caption("1. URL 입력  2. QA Mode 선택  3. Run QA Test  4. 결과 확인")
        definition_target_mode = str(st.session_state.get("qa_definition_validation_target_mode", "선택하세요")).strip()
        definition_file_ready = bool(st.session_state.get("qa_uploaded_definition_schemas", {}))
        run_button_label = "Run QA Test"
        run_button_help = None
        run_button_disabled = False
        if selected_mode == "정의서 검증":
            if not definition_file_ready:
                run_button_disabled = True
                run_button_help = "1단계에서 정의서 업로드를 완료해야 합니다."
            elif definition_target_mode != "새 Run QA Test 실행 후 검증":
                run_button_disabled = True
                run_button_help = "2단계에서 '새 Run QA Test 실행 후 검증'을 선택해야 합니다."
            else:
                run_button_label = "정의서 기준으로 새 세션 수집 후 검증"
                run_button_help = "업로드한 정의서 기준으로 새 세션을 수집한 뒤 PASS / FAIL 검증을 수행합니다."
        dc1, dc2 = st.columns(2)
        with dc1:
            realtime_debug_start_clicked = st.button(
                run_button_label,
                type="primary",
                disabled=run_button_disabled,
                help=run_button_help,
            )
        with dc2:
            realtime_debug_stop_clicked = st.button("Stop Test")
        if selected_mode == "정의서 검증":
            if not definition_file_ready:
                st.caption("정의서 업로드 후 검증 대상을 선택하면 실행할 수 있습니다.")
            elif definition_target_mode == "기존 세션으로 검증":
                st.caption("기존 세션 재검증이 선택되어 있습니다. 새 테스트 실행 없이 현재 세션을 정의서 기준으로 검증합니다.")
            elif definition_target_mode == "새 Run QA Test 실행 후 검증":
                st.caption("버튼을 누르면 정의서 기준으로 새 세션 수집 후 검증합니다.")
            else:
                st.caption("검증 대상을 먼저 선택하세요.")
        active_debug_session_id = st.session_state.get("qa_debug_session_id", "").strip()
        debug_snapshot = get_debug_session_snapshot(active_debug_session_id) if active_debug_session_id else {}
        resolved_output_path = resolve_debug_output_file(
            session_id=active_debug_session_id,
            snapshot=debug_snapshot,
            state_output_file=st.session_state.get("qa_debug_output_file", ""),
        )
        recovered_from_file = bool((not debug_snapshot) and resolved_output_path and resolved_output_path.exists())
        if debug_snapshot or recovered_from_file:
            sid_for_view = (
                str(debug_snapshot.get("session_id", "")).strip() if debug_snapshot else active_debug_session_id
            )
            status_for_view = str(debug_snapshot.get("status", "")).strip() if debug_snapshot else "recovered(file)"
            captured_count_for_view = int(debug_snapshot.get("captured_events", 0)) if debug_snapshot else "-"
            st.caption(
                f"세션: {sid_for_view} | "
                f"상태: {status_for_view or '-'} | "
                f"캡처: {captured_count_for_view}건"
            )
            st.caption(f"표시 시간대: {DISPLAY_TZ_NAME}")
            tester_name_view = str(debug_snapshot.get("tester_name", "")).strip() if debug_snapshot else ""
            if tester_name_view:
                st.caption(f"테스터: {tester_name_view}")
            if debug_snapshot and debug_snapshot.get("last_error", "").strip():
                last_runtime_error = str(debug_snapshot.get("last_error", "")).strip()
                if "접근 제한/봇 확인 페이지" in last_runtime_error:
                    st.warning(f"디버깅 런타임 알림: {last_runtime_error}")
                else:
                    st.error(f"디버깅 런타임 오류: {last_runtime_error}")
            run_settings = debug_snapshot.get("run_settings", {}) if isinstance(debug_snapshot, dict) else {}
            mode_label = str(run_settings.get("qa_mode", st.session_state.get("qa_mode", ""))).strip()
            viewport_label = str(st.session_state.get("qa_viewport_preset", "")).strip()
            if mode_label:
                st.caption(f"모드: {mode_label} | Viewport: {viewport_label or '-'}")
            setup_snapshot = st.session_state.get("qa_test_setup_snapshot", {})
            initial_target_url = (
                str(debug_snapshot.get("target_url", "")).strip()
                if debug_snapshot
                else str(setup_snapshot.get("qa_debug_target_url", "")).strip()
            )
            if initial_target_url:
                st.caption(f"초기 테스트 URL: {initial_target_url}")
            st.caption("결과는 오른쪽 패널의 QA Overview / Export에서 확인합니다.")
        else:
            st.caption("실행 후 자동수집이 끝나면 결과 패널에서 바로 확인할 수 있습니다.")

    with st.expander("4) Advanced", expanded=False):
        if st.session_state.get("qa_ingest_server_ok", False):
            st.caption("수집기 상태: OK (local:127.0.0.1:8600)")
        else:
            st.error("수집기 상태: 실패")
            last_ingest_err = str(st.session_state.get("qa_ingest_server_error", "")).strip()
            if last_ingest_err:
                st.caption(f"원인: {last_ingest_err}")
        analytics_proxy_on = str(os.getenv("QA_ANALYTICS_PROXY_ENABLED", "0")).strip()
        st.caption(f"Analytics Proxy: {'ON' if analytics_proxy_on in {'1', 'true', 'True'} else 'OFF'}")
        st.number_input("Auto Crawl Max Clicks", min_value=1, max_value=500, key="qa_auto_crawl_max_clicks")
        st.number_input(
            "Click Interval (ms)",
            min_value=100,
            max_value=10000,
            step=100,
            key="qa_auto_crawl_click_interval_ms",
        )
        st.number_input(
            "Wait After Click (ms)",
            min_value=100,
            max_value=10000,
            step=100,
            key="qa_auto_crawl_wait_after_click_ms",
        )
        st.checkbox("링크 이동 차단", key="qa_auto_block_link_nav")
        st.checkbox("시작 페이지 범위만 수집", key="qa_single_page_only")
        today = date.today()
        start_date = st.date_input("시작일", value=today - timedelta(days=1))
        end_date = st.date_input("종료일", value=today)
        param_text = st.text_area("조회할 param (쉼표 구분)", value=default_params)
        st.caption("예: transaction_id,value,currency,qa_debug_session_id")
        required_event_default = required_event_text if selected_mode == "시나리오 테스트" else default_required_events
        if selected_mode == "정의서 검증":
            definition_summary_rows = st.session_state.get("qa_uploaded_definition_summary", [])
            definition_events = [str(row.get("event_name", "")).strip() for row in definition_summary_rows if str(row.get("event_name", "")).strip()]
            st.text_area(
                "검증 대상 이벤트",
                value=", ".join(definition_events),
                height=90,
                disabled=True,
            )
            st.caption("정의서 검증 모드에서는 업로드한 파일의 이벤트 목록만 검증합니다.")
        else:
            required_event_text = st.text_area("필수 이벤트", value=required_event_default, key="required_event_text_input")
            st.caption("비워두면 필수 이벤트 체크를 생략합니다. 필요할 때만 입력하세요.")
        unknown_event_policy = st.selectbox(
            "실시간 미정의 이벤트 처리",
            options=["정보", "경고", "무시"],
            index=["정보", "경고", "무시"].index(st.session_state.get("unknown_event_policy", "정보")),
            help="필수 이벤트 목록에 없는 이벤트를 실시간 타임라인에서 어떻게 표시할지 선택합니다.",
            key="unknown_event_policy",
        )
        st.caption("사이트 상황에 맞는 이벤트명을 쉼표로 자유 입력하세요.")
        required_param_text = st.text_area(
            "필수 param 규칙 (event: p1,p2)",
            value=default_param_text,
            height=120,
        )
        st.caption("비워두면 필수 파라미터 체크를 생략합니다.")
        funnel_default = funnel_text if selected_mode == "시나리오 테스트" else default_funnel_steps
        funnel_text = st.text_input("퍼널 시퀀스", value=funnel_default)
        null_threshold = st.slider("null 경고 임계치", min_value=0.05, max_value=0.9, value=0.2, step=0.05)
        max_rows = st.number_input("최대 조회 행", min_value=1000, max_value=500000, value=50000, step=1000)

    with st.expander("5) Admin", expanded=False):
        p1, p2 = st.columns(2)
        with p1:
            st.text_input("새 Project", key="qa_new_project_name", placeholder="예: musinsa")
        with p2:
            st.text_input("도메인", key="qa_new_project_domain", placeholder="예: musinsa.com")
        if st.button("Project 생성", key="qa_create_project_btn"):
            new_name = str(st.session_state.get("qa_new_project_name", "")).strip()
            if not new_name:
                st.session_state["qa_project_notice"] = "새 Project 이름을 먼저 입력하세요."
                st.rerun()
            new_slug = _slugify_project_name(new_name)
            ensure_project_structure(new_slug)
            existed = new_slug in project_options
            st.session_state["qa_project_slug"] = new_slug
            if existed:
                st.session_state["qa_project_notice"] = f"이미 존재하는 Project를 선택했습니다: {new_slug}"
            else:
                st.session_state["qa_project_notice"] = f"Project 생성 완료: {new_slug}"
            st.rerun()
        if st.button("이전 테스트 데이터 가져오기", key="qa_import_legacy_btn"):
            import_counts = import_legacy_data_to_project(st.session_state.get("qa_project_slug", DEFAULT_PROJECT_SLUG))
            st.success(
                "가져오기 완료 | "
                f"session {import_counts['sessions']}개, "
                f"event {import_counts['events']}개, "
                f"ui_action {import_counts['actions']}개, "
                f"stream {import_counts['streams']}개, "
                f"schema {import_counts['schemas']}개"
            )
            st.rerun()
        project_paths = get_active_project_paths()
        st.caption(f"Workspace 경로: {project_paths['project_root']}")
        sess_df = pd.DataFrame()
        try:
            sess_df = list_recent_sessions(get_active_test_log_db_path(), limit=10)
            if sess_df.empty:
                st.caption("QA Sessions: 없음")
            else:
                st.caption(f"QA Sessions: 최근 {len(sess_df)}개")
                st.dataframe(sess_df[["session_id", "status", "captured_events"]], use_container_width=True, height=180)
        except Exception:
            st.caption("QA Sessions 목록을 불러오지 못했습니다.")
        if not sess_df.empty:
            load_options = sess_df["session_id"].astype(str).tolist()
            load_sid = st.selectbox("불러올 QA Session", options=load_options, key="qa_load_session_id")
            if st.button("선택 세션 불러오기", key="qa_load_session_btn"):
                sid = str(load_sid).strip()
                snap = get_session(get_active_test_log_db_path(), sid)
                output_file = project_paths["qa_sessions_dir"] / sid / "debug_stream.jsonl"
                if not output_file.exists():
                    fallback = Path(f"data/debug_stream/{sid}.jsonl")
                    if fallback.exists():
                        output_file = fallback
                st.session_state["qa_debug_session_id"] = sid
                st.session_state["qa_debug_output_file"] = str(output_file)
                st.session_state["qa_debug_started_at"] = str(snap.get("started_at", "")).strip()
                st.success(f"세션 불러오기 완료: {sid}")
                st.rerun()

st.caption("왼쪽에서는 Test Setup / QA Mode / Run Test만 조작하면 되고, 결과 확인은 메인 패널에서 진행합니다.")

resume_notice = str(st.session_state.get("qa_oauth_resume_notice", "")).strip()
if resume_notice:
    st.info(resume_notice)
    st.session_state["qa_oauth_resume_notice"] = ""

if realtime_debug_start_clicked:
    try:
        save_test_setup_snapshot()
        normalized_target_url = normalize_debug_target_url(st.session_state.get("qa_debug_target_url", ""))
        if not normalized_target_url:
            st.error("Target URL 형식이 올바르지 않습니다. 예: https://datanugget.io/")
            log_ui_action("debug_start_error", {"error": "invalid_target_url"})
            st.stop()
        log_ui_action(
            "debug_start_click",
            {
                "target_url": normalized_target_url,
                "qa_mode": st.session_state.get("qa_mode", "전체 이벤트 테스트"),
                "environment": st.session_state.get("qa_environment", "prod"),
            },
        )
        previous_sid = st.session_state.get("qa_debug_session_id", "").strip()
        if previous_sid:
            stop_debug_session(previous_sid)
        debug_session_id = f"dbg_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S_%f')}_{uuid4().hex[:8]}"
        active_paths = get_active_project_paths()
        debug_file = active_paths["qa_sessions_dir"] / debug_session_id / "debug_stream.jsonl"
        debug_file.parent.mkdir(parents=True, exist_ok=True)
        run_settings = build_debug_run_settings()
        snapshot = start_debug_session(
            session_id=debug_session_id,
            target_url=normalized_target_url,
            output_file=debug_file,
            tester_name=st.session_state.get("qa_tester_name", "").strip(),
            tester_note=st.session_state.get("qa_tester_note", "").strip(),
            db_path=get_active_test_log_db_path(),
            launch_browser=True,
            run_settings=run_settings,
        )
        st.session_state["qa_debug_session_id"] = debug_session_id
        st.session_state["qa_debug_output_file"] = str(debug_file)
        st.session_state["qa_debug_started_at"] = str(snapshot.get("started_at", "")).strip()
        st.session_state["qa_preserve_test_setup_during_run"] = True
        st.session_state["realtime_panel_auto_refresh"] = True
        definition_target_names = run_settings.get("definition_event_names", [])
        if (
            st.session_state.get("qa_mode", "전체 이벤트 테스트") == "정의서 검증"
            and str(st.session_state.get("qa_definition_validation_target_mode", "")).strip() == "새 Run QA Test 실행 후 검증"
        ):
            st.session_state["qa_definition_validation_execution_mode"] = "새 테스트 실행 후 검증"
            st.session_state["qa_definition_validation_execution_session_id"] = debug_session_id
            save_definition_validation_cache(
                {
                    "qa_mode": "정의서 검증",
                    "file_name": str(st.session_state.get("qa_definition_validation_file_name", "")).strip(),
                    "schemas": st.session_state.get("qa_uploaded_definition_schemas", {}),
                    "summary": st.session_state.get("qa_uploaded_definition_summary", []),
                    "debug": st.session_state.get("qa_definition_validation_debug", {}),
                    "loaded_key": st.session_state.get("qa_definition_validation_loaded_key", ""),
                    "target_mode": st.session_state.get("qa_definition_validation_target_mode", "선택하세요"),
                    "execution_mode": "새 테스트 실행 후 검증",
                    "execution_session_id": debug_session_id,
                }
            )
        log_ui_action("debug_start_success", {"session_id": debug_session_id})
        success_message = (
            f"QA Test 시작: qa_debug_session_id={debug_session_id} "
            f"(모드: {run_settings.get('qa_mode', '-')}, 수집: Auto Crawl)"
        )
        if isinstance(definition_target_names, list) and definition_target_names:
            success_message += f" | 정의서 이벤트 {len(definition_target_names)}개 감지 시 조기 종료"
        st.success(success_message)
        novnc_public_url = get_config_value("QA_NOVNC_PUBLIC_URL", "").strip()
        if (
            st.session_state.get("qa_access_role", "guest") == "admin"
            and novnc_public_url
            and is_truthy(get_config_value("QA_OPEN_NOVNC_ON_START", "0"))
        ):
            auto_open_popup_window(novnc_public_url, popup_name=f"qa_debug_popup_{debug_session_id}")
            st.caption("원격 디버그 팝업(noVNC) 자동 열기를 시도했습니다. 차단되면 링크를 직접 열어주세요.")
    except Exception as exc:
        log_ui_action("debug_start_error", {"error": str(exc)})
        st.error(f"QA Test 시작 실패: {to_user_error_message(exc)}")

if realtime_debug_stop_clicked:
    log_ui_action("debug_stop_click")
    sid = st.session_state.get("qa_debug_session_id", "").strip()
    if sid:
        stop_debug_session(sid)
        st.session_state["qa_preserve_test_setup_during_run"] = False
        log_ui_action("debug_stop_success", {"session_id": sid})
        st.success("테스트 종료 요청을 보냈습니다.")
    else:
        st.info("종료할 테스트 세션이 없습니다.")

realtime_tab, report_tab, schema_tab = st.tabs(
    ["1단: 실시간 테스트 화면", "2단: QA 리포트 화면", "3단: 이벤트 기준표 탭"]
)
with realtime_tab:
    restore_latest_project_session_state()
    st.markdown("### QA Workflow")
    wf1, wf2, wf3, wf4 = st.columns(4)
    wf1.info("1. Test Events\n사이트 클릭으로 이벤트 발생")
    wf2.info("2. Define Rules\n이벤트별 필수/선택 규칙 정의")
    wf3.info("3. Validate Events\npayload와 기준표 비교")
    wf4.info("4. Confirm Rules\n수정 후 다시 클릭해 재검증")

    active_debug_session_id = st.session_state.get("qa_debug_session_id", "").strip()
    debug_snapshot = get_debug_session_snapshot(active_debug_session_id) if active_debug_session_id else {}
    resolved_output_path = resolve_debug_output_file(
        session_id=active_debug_session_id,
        snapshot=debug_snapshot,
        state_output_file=st.session_state.get("qa_debug_output_file", ""),
    )
    recovered_from_file = bool((not debug_snapshot) and resolved_output_path and resolved_output_path.exists())
    st.session_state["realtime_last_refresh_at"] = pd.Timestamp.now(tz=LOCAL_TZ).strftime("%H:%M:%S")

    ctrl1, ctrl2, ctrl3, ctrl4 = st.columns([1.3, 1, 1.2, 2.2])
    with ctrl1:
        realtime_manual_refresh_clicked = st.button("실시간 데이터 새로고침", key="realtime_manual_refresh")
    with ctrl2:
        st.checkbox("자동 갱신", key="realtime_panel_auto_refresh")
    with ctrl3:
        st.select_slider(
            "간격(초)",
            options=[1, 2, 3, 5],
            key="realtime_panel_refresh_interval",
        )
    with ctrl4:
        st.caption(f"마지막 갱신: {st.session_state.get('realtime_last_refresh_at', '-')}")

    if realtime_manual_refresh_clicked:
        log_ui_action("realtime_manual_refresh")
        st.rerun()

    sid_for_view = "-"
    status_for_view = "-"
    captured_for_view = 0
    raw_captured_for_view = 0
    ga_hit_count_for_view = 0
    auto_crawl_count_for_view = 0
    timeline_df_active = pd.DataFrame()
    timeline_df = pd.DataFrame()
    rt_events = pd.DataFrame()
    st.session_state["schema_rt_events_cache"] = []

    if debug_snapshot or recovered_from_file:
        if recovered_from_file:
            st.info("세션 객체는 없지만 파일 기준으로 이어서 표시 중입니다. (상태: recovered(file))")
        debug_output_file = str(resolved_output_path) if resolved_output_path else ""
        timeline_df_raw = load_debug_events(Path(debug_output_file), limit=3000) if debug_output_file else pd.DataFrame()
        session_started_at = (
            str(debug_snapshot.get("started_at", "")).strip()
            if debug_snapshot
            else str(st.session_state.get("qa_debug_started_at", "")).strip()
        )
        timeline_df_active = filter_events_for_active_session(
            timeline_df_raw,
            session_id=active_debug_session_id,
            session_started_at=session_started_at,
        )
        timeline_df = filter_debug_events_by_view(timeline_df_active, "히트(collect)")
        sid_for_view = (
            str(debug_snapshot.get("session_id", "")).strip() if debug_snapshot else active_debug_session_id
        ) or "-"
        status_for_view = (
            str(debug_snapshot.get("status", "")).strip() if debug_snapshot else "recovered(file)"
        ) or "-"
        captured_from_snapshot = int(debug_snapshot.get("captured_events", 0)) if debug_snapshot else 0
        captured_from_file = int(len(timeline_df))
        captured_for_view = max(captured_from_snapshot, captured_from_file)
        raw_captured_for_view = int(len(timeline_df_active))
        if not timeline_df_active.empty and "source" in timeline_df_active.columns:
            source_series = timeline_df_active["source"].astype(str).str.strip()
            ga_hit_count_for_view = int(source_series.eq("ga_hit").sum())
            auto_crawl_count_for_view = int(source_series.eq("auto_crawl").sum())
    else:
        st.info("활성 디버깅 세션이 없습니다. 사이드바에서 `디버깅 모드 시작`을 실행하세요.")

    selected_mode = str(st.session_state.get("qa_mode", "전체 이벤트 테스트")).strip()
    persisted_schema_store = load_event_schemas()
    uploaded_definition_schemas = st.session_state.get("qa_uploaded_definition_schemas", {})
    uploaded_definition_summary_df = pd.DataFrame(st.session_state.get("qa_uploaded_definition_summary", []))
    definition_upload_error = str(st.session_state.get("qa_definition_validation_error", "")).strip()
    definition_mode_active = selected_mode == "정의서 검증"
    definition_mode_ready = (
        definition_mode_active
        and isinstance(uploaded_definition_schemas, dict)
        and bool(uploaded_definition_schemas)
    )
    definition_target_mode = str(st.session_state.get("qa_definition_validation_target_mode", "선택하세요")).strip()
    definition_session_available = bool(debug_snapshot or recovered_from_file)
    definition_existing_session_ready = definition_mode_ready and definition_target_mode == "기존 세션으로 검증" and definition_session_available
    definition_new_run_ready = (
        definition_mode_ready
        and definition_target_mode == "새 Run QA Test 실행 후 검증"
        and definition_session_available
        and str(st.session_state.get("qa_definition_validation_execution_mode", "")).strip() == "새 테스트 실행 후 검증"
        and str(st.session_state.get("qa_definition_validation_execution_session_id", "")).strip() == str(active_debug_session_id or sid_for_view).strip()
    )
    definition_validation_active = (
        (not definition_mode_active)
        or definition_existing_session_ready
        or definition_new_run_ready
    )
    definition_validation_method = (
        "기존 세션 재검증"
        if definition_existing_session_ready
        else "새 테스트 실행 후 검증"
        if definition_new_run_ready
        else ""
    )
    schema_store = uploaded_definition_schemas if definition_mode_ready else ({} if definition_mode_active else persisted_schema_store)
    scoped_event_names = set(schema_store.keys()) if definition_mode_ready else None
    timeline_df_active_for_validation = (
        filter_debug_df_for_event_scope(timeline_df_active, scoped_event_names)
        if definition_mode_ready and definition_validation_active
        else timeline_df_active
    )
    timeline_df_for_validation = (
        filter_debug_events_by_view(timeline_df_active_for_validation, "히트(collect)")
        if definition_mode_ready and definition_validation_active
        else timeline_df
    )
    allowed_events_rt = (
        list(schema_store.keys())
        if definition_mode_ready
        else parse_csv_list(st.session_state.get("required_event_text_input", default_required_events))
    )
    if definition_mode_ready and uploaded_definition_summary_df.empty:
        uploaded_definition_summary_df = build_schema_catalog_df(schema_store)
    if definition_mode_active and not definition_validation_active:
        timeline_df_active_for_validation = pd.DataFrame()
        timeline_df_for_validation = pd.DataFrame()
    if not definition_mode_active and not timeline_df_for_validation.empty:
        rt_events = build_realtime_event_rows(
            timeline_df_for_validation,
            allowed_events=allowed_events_rt,
            unknown_event_policy=st.session_state.get("unknown_event_policy", "정보"),
            schemas=schema_store,
            scoped_event_names=scoped_event_names,
        )
    elif definition_mode_ready and definition_validation_active and not timeline_df_for_validation.empty:
        rt_events = build_realtime_event_rows(
            timeline_df_for_validation,
            allowed_events=allowed_events_rt,
            unknown_event_policy="무시",
            schemas=schema_store,
            scoped_event_names=scoped_event_names,
        )
    if not rt_events.empty:
        st.session_state["schema_rt_events_cache"] = rt_events.to_dict("records")
    issue_df = build_issue_summary(rt_events, schema_store)
    tracking_plan_df = build_tracking_plan_df(rt_events, schema_store)
    schema_summary = summarize_schema_validation(rt_events, schema_store)
    rt_summary = summarize_realtime_quality(rt_events)
    parameter_validation_df_all = build_parameter_validation_export_df(rt_events, schema_store)
    order_validation_df_all = build_event_order_validation_df(rt_events, parse_csv_list(funnel_text))
    if definition_mode_ready:
        issue_df = build_definition_validation_issue_df(rt_events, schema_store, issue_df)
    event_schema_qa_all = (
        build_definition_validation_report_df(
            rt_events=rt_events,
            schemas=schema_store,
            issue_df=issue_df,
            parameter_validation_df=parameter_validation_df_all,
            order_validation_df=order_validation_df_all,
        )
        if definition_mode_ready
        else build_event_qa_report_export_df(
            rt_events=rt_events,
            issue_df=issue_df,
            schemas=schema_store,
            parameter_validation_df=parameter_validation_df_all,
            order_validation_df=order_validation_df_all,
        )
    )
    schema_summary_view = summarize_event_qa_report(event_schema_qa_all) if definition_mode_ready else schema_summary

    event_names_seen = (
        sorted(rt_events["이벤트"].astype(str).str.strip().unique().tolist()) if not rt_events.empty else []
    )
    missing_schema_count = 0
    if not definition_mode_ready:
        for ev in event_names_seen:
            schema_obj = schema_store.get(ev, {})
            has_schema = isinstance(schema_obj, dict) and bool(schema_obj.get("required", []) or schema_obj.get("optional", []))
            if not has_schema:
                missing_schema_count += 1
    missing_parameter_count = 0
    if not parameter_validation_df_all.empty:
        pwork = parameter_validation_df_all.copy()
        pwork["issue_type"] = pwork["issue_type"].astype(str)
        pwork["severity"] = pwork["severity"].astype(str)
        missing_parameter_count = int(
            (
                (pwork["severity"] == "FAIL")
                & (pwork["issue_type"].isin(["required_missing", "type_validation_fail"]))
            ).sum()
        )
    flow_issue_count = 0
    if not order_validation_df_all.empty:
        owork = order_validation_df_all.copy()
        owork["status"] = owork["status"].astype(str)
        flow_issue_count = int((owork["status"].isin(["FAIL", "WARN"])).sum())
    definition_missing_event_count = 0
    if definition_mode_ready and not event_schema_qa_all.empty:
        qwork = event_schema_qa_all.copy()
        qwork["matched_capture_count"] = pd.to_numeric(qwork.get("matched_capture_count", 0), errors="coerce").fillna(0).astype(int)
        qwork["result"] = qwork.get("result", pd.Series(dtype=str)).astype(str).str.upper()
        definition_missing_event_count = int(((qwork["matched_capture_count"] == 0) & qwork["result"].eq("FAIL")).sum())
    definition_param_count = 0
    if definition_mode_ready:
        schema_param_keys: set[str] = set()
        for schema_obj in schema_store.values():
            if not isinstance(schema_obj, dict):
                continue
            for key in list(schema_obj.get("required", [])) + list(schema_obj.get("optional", [])) + list(schema_obj.get("parameter_columns", [])):
                key_text = str(key).strip()
                if key_text:
                    schema_param_keys.add(key_text)
        definition_param_count = len(schema_param_keys)

    section_overview, section_issues, section_detail, section_export = st.tabs(
        ["1. QA Overview", "2. Issues (문제 이벤트)", "3. Event Detail", "4. Export / Event Definition"]
    )

    with section_overview:
        st.caption(
            f"세션: {sid_for_view} | 상태: {status_for_view} | 캡처: {captured_for_view}건 | 시간대: {DISPLAY_TZ_NAME}"
        )
        if definition_mode_active:
            definition_file_name = str(st.session_state.get("qa_definition_validation_file_name", "")).strip()
            if definition_file_name:
                st.caption(f"정의서 검증 파일: {definition_file_name}")
            if definition_upload_error:
                st.error(f"정의서 로드 오류: {definition_upload_error}")
            elif not definition_mode_ready:
                st.warning("정의서 검증 모드입니다. QA Mode에서 정의서 파일을 업로드해야 PASS / FAIL 검증을 시작할 수 있습니다.")
            elif definition_target_mode == "선택하세요":
                st.info("1단계 정의서 업로드가 완료되었습니다. 2단계에서 '기존 세션으로 검증' 또는 '새 Run QA Test 실행 후 검증'을 선택하세요.")
            elif definition_target_mode == "기존 세션으로 검증" and not definition_session_available:
                st.info("기존 세션으로 검증이 선택되었습니다. 먼저 불러올 세션을 선택하거나 새 세션을 준비하세요.")
            elif definition_target_mode == "새 Run QA Test 실행 후 검증" and not definition_new_run_ready:
                st.info("정의서 기준 새 테스트 실행 전입니다. 3단계에서 '정의서 기준으로 새 세션 수집 후 검증' 버튼을 눌러주세요.")
            elif not uploaded_definition_summary_df.empty:
                st.caption(f"정의서 기준 이벤트 {len(uploaded_definition_summary_df)}개 | 파라미터 {definition_param_count}개")
        if raw_captured_for_view:
            st.caption(
                f"원본 로그 {raw_captured_for_view}건 | GA hit {ga_hit_count_for_view}건 | auto crawl {auto_crawl_count_for_view}건"
            )
        if rt_events.empty:
            if definition_mode_active and not definition_mode_ready:
                st.info("정의서 업로드 전에는 업로드 파일 기준 검증 결과를 계산하지 않습니다.")
            elif definition_mode_active and not definition_validation_active:
                st.info("정의서 검증 준비 완료 상태입니다. 검증 대상 선택 및 실행 후 결과가 표시됩니다.")
            elif raw_captured_for_view > 0 and ga_hit_count_for_view == 0:
                st.warning("자동수집 로그는 있지만 아직 GA hit(collect)가 없습니다. 브라우저/수집기 연결 상태를 확인하세요.")
            elif raw_captured_for_view > 0:
                if definition_mode_ready:
                    st.info("GA4 hit는 있었지만 업로드한 정의서 기준 custom event는 아직 수집되지 않았습니다.")
                else:
                    st.info("GA4 hit는 있지만 기본/기술 이벤트를 제외하고 나면 표시할 custom event가 없습니다.")
            else:
                st.info("아직 캡처된 이벤트가 없습니다.")
        else:
            st.subheader("QA Summary")
            if definition_mode_active and definition_validation_active:
                st.info(
                    " | ".join(
                        [
                            f"기준 정의서: {str(st.session_state.get('qa_definition_validation_file_name', '-') or '-')}",
                            f"검증 대상 세션: {sid_for_view}",
                            f"검증 방식: {definition_validation_method or '-'}",
                            f"수집 건수: {len(rt_events)}",
                            f"정의서 이벤트 수: {len(schema_store)}",
                            f"정의서 파라미터 수: {definition_param_count}",
                        ]
                    )
                )
            m0, m1, m2, m3, m4, m5 = st.columns(6)
            if definition_mode_ready:
                m0.metric("Definition events", int(len(event_schema_qa_all)))
            else:
                m0.metric("Events captured", int(captured_for_view))
            m1.metric("실시간 품질 점수", f"{rt_summary['score']} / 100")
            m2.metric("🔴 치명 오류", int(rt_summary["critical"]))
            m3.metric("🟡 주의 필요", int(rt_summary["caution"]))
            m4.metric("⚪ 참고 정보", int(rt_summary["info"]))
            m5.metric("🟢 정상", int(rt_summary["ok"]))
            st.caption(
                f"기준표 QA 상태 | PASS {schema_summary_view['PASS']} · WARN {schema_summary_view['WARN']} · FAIL {schema_summary_view['FAIL']}"
            )
            q1, q2, q3 = st.columns(3)
            if definition_mode_ready:
                q1.metric("Missing rows", int(definition_missing_event_count))
            else:
                q1.metric("Missing schema", int(missing_schema_count))
            q2.metric("Missing parameter", int(missing_parameter_count))
            q3.metric("Flow issues", int(flow_issue_count))
            if definition_mode_ready:
                st.caption(f"업로드 정의서 이벤트 {len(event_schema_qa_all)}건 기준 PASS / FAIL 요약입니다.")
                show_cols = ["no", "event_name", "description", "expected_identification", "matched_identification", "result", "test_count", "parameter_issue"]
                keep_cols = [c for c in show_cols if c in event_schema_qa_all.columns]
                st.dataframe(event_schema_qa_all[keep_cols], use_container_width=True, height=320)
            else:
                st.caption(
                    f"대상 이벤트 {rt_summary['event_count']}건 기준 | {rt_summary['message']}"
                )
                quick_cols = ["시간", "이벤트", "대표 파라미터", "대표 값", "상태"]
                st.dataframe(rt_events[quick_cols].head(30), use_container_width=True, height=320)

    with section_issues:
        if definition_mode_ready:
            st.caption("업로드한 정의서를 기준으로 PASS / FAIL이 아닌 이벤트만 모아 보여줍니다.")
        else:
            st.caption("문제가 있는 이벤트를 확인하고 여기서 바로 이벤트 기준표를 만들 수 있습니다.")
        if issue_df.empty:
            st.success("현재 문제 이벤트가 없습니다.")
        else:
            issue_event_names = sorted(issue_df["event_name"].astype(str).unique().tolist())
            if not definition_mode_ready:
                st.markdown("**이벤트 기준표 만들기 (문제 이벤트 기준)**")
                include_system_params_for_schema = st.checkbox(
                    "기술 파라미터까지 포함해서 기준표 만들기",
                    value=False,
                    key="issue_include_system_params",
                    help="기본은 measurement_id/client_id 같은 시스템 파라미터를 제외합니다.",
                )
                selected_issue_events = st.multiselect(
                    "기준표 생성 대상 이벤트",
                    options=issue_event_names,
                    default=issue_event_names,
                    key="issue_schema_targets",
                )
                ia1, ia2 = st.columns([1.2, 1.2])
                with ia1:
                    if st.button("선택 이벤트 기준표 생성", key="issue_generate_selected"):
                        generated = 0
                        total_excluded = 0
                        for event_name in selected_issue_events:
                            inferred = infer_event_schema_from_rows(
                                rt_events,
                                event_name,
                                include_system_params=include_system_params_for_schema,
                            )
                            if int(inferred.get("sample_size", 0)) <= 0:
                                continue
                            schema_store[event_name] = {
                                "required": list(inferred.get("required", [])),
                                "optional": list(inferred.get("optional", [])),
                                "updated_at": datetime.now(tz=ZoneInfo("UTC")).isoformat(),
                            }
                            generated += 1
                            total_excluded += int(inferred.get("excluded_system_params", 0))
                        save_event_schemas(schema_store)
                        st.success(f"선택 이벤트 기준표 {generated}개 생성 완료")
                        if not include_system_params_for_schema:
                            st.info(f"시스템/기술 파라미터 제외: {total_excluded}건")
                        st.rerun()
                with ia2:
                    if st.button("문제 이벤트 전체 기준표 생성", key="issue_generate_all"):
                        generated = 0
                        total_excluded = 0
                        for event_name in issue_event_names:
                            inferred = infer_event_schema_from_rows(
                                rt_events,
                                event_name,
                                include_system_params=include_system_params_for_schema,
                            )
                            if int(inferred.get("sample_size", 0)) <= 0:
                                continue
                            schema_store[event_name] = {
                                "required": list(inferred.get("required", [])),
                                "optional": list(inferred.get("optional", [])),
                                "updated_at": datetime.now(tz=ZoneInfo("UTC")).isoformat(),
                            }
                            generated += 1
                            total_excluded += int(inferred.get("excluded_system_params", 0))
                        save_event_schemas(schema_store)
                        st.success(f"문제 이벤트 전체 기준표 {generated}개 생성 완료")
                        if not include_system_params_for_schema:
                            st.info(f"시스템/기술 파라미터 제외: {total_excluded}건")
                        st.rerun()

            st.dataframe(issue_df, use_container_width=True, height=320)
            st.markdown("**기준표 비교 QA 결과 (이벤트별)**")
            if event_schema_qa_all.empty:
                st.info("비교할 데이터가 없습니다.")
            else:
                show_cols = [
                    "no",
                    "event_name",
                    "description",
                    "expected_identification",
                    "matched_identification",
                    "schema_status",
                    "parameter_issue",
                    "test_count",
                    "parameter_fail_count",
                    "parameter_warn_count",
                    "order_fail_count",
                    "order_warn_count",
                ]
                keep_cols = [c for c in show_cols if c in event_schema_qa_all.columns]
                st.dataframe(event_schema_qa_all[keep_cols], use_container_width=True, height=240)

    with section_detail:
        st.caption("이벤트를 선택해 payload 확인 → schema 검증 순서로 점검하세요.")
        if rt_events.empty:
            st.info("확인할 이벤트가 없습니다.")
        else:
            event_names = sorted(rt_events["이벤트"].astype(str).unique().tolist())
            selected_event_name = st.selectbox("이벤트 선택", options=event_names, key="detail_event_name")
            event_rows = rt_events[rt_events["이벤트"].astype(str) == selected_event_name].copy()
            event_rows["captured_at"] = pd.to_datetime(event_rows["captured_at"], errors="coerce")
            event_rows = event_rows.sort_values("captured_at", ascending=False).reset_index(drop=True)
            option_labels = [
                f"{to_circled_number(idx + 1)} {selected_event_name} | {row.get('시간', '-')} | {row.get('상태', '-')}"
                for idx, row in event_rows.iterrows()
            ]
            selected_option = st.selectbox("발생 건 선택", options=option_labels, key="detail_event_occurrence")
            selected_idx = option_labels.index(selected_option)
            row_ev = event_rows.iloc[selected_idx].to_dict()

            c1, c2, c3 = st.columns([2, 2, 2])
            c1.metric("이벤트", str(row_ev.get("이벤트", "-")))
            c2.metric("수집 시각", str(row_ev.get("시간", "-")))
            c3.metric("상태", str(row_ev.get("상태", "-")))

            all_params = row_ev.get("전체 파라미터", {})
            raw_payload = row_ev.get("원본 payload", {})
            st.markdown("**Custom Parameters**")
            if isinstance(all_params, dict) and all_params:
                payload_rows = [{"custom_parameter": str(k), "custom_value": _to_text_value(v)} for k, v in all_params.items()]
                st.table(pd.DataFrame(payload_rows))
            else:
                st.caption("기본 노출할 custom parameter가 없습니다.")

            extra_custom_summary = str(row_ev.get("extra_custom_parameters", "-")).strip()
            if extra_custom_summary and extra_custom_summary != "-":
                with st.expander("추가 custom parameter 보기", expanded=False):
                    extra_rows = []
                    for item in extra_custom_summary.split(" | "):
                        if "=" not in item:
                            continue
                        key, value = item.split("=", 1)
                        extra_rows.append({"custom_parameter": key, "custom_value": value})
                    if extra_rows:
                        st.table(pd.DataFrame(extra_rows))

            if isinstance(raw_payload, dict) and raw_payload and raw_payload != all_params:
                with st.expander("원본 payload 보기", expanded=False):
                    raw_rows = [{"param": str(k), "value": _to_text_value(v)} for k, v in raw_payload.items()]
                    st.table(pd.DataFrame(raw_rows))

            schema_check = validate_schema_for_event(
                event_name=str(row_ev.get("이벤트", "")),
                params=all_params if isinstance(all_params, dict) else {},
                schemas=schema_store,
            )
            schema_status = str(schema_check.get("status", "WARN")).upper()
            st.markdown(f"**기준표 검증 (Schema Validation)**: `{schema_status}` · {schema_check.get('message', '-')}")
            schema_checks = schema_check.get("checks", [])
            if isinstance(schema_checks, list) and schema_checks:
                check_rows = []
                for item in schema_checks:
                    if not isinstance(item, dict):
                        continue
                    check_rows.append(
                        {
                            "group": str(item.get("group", "-")),
                            "param": str(item.get("param", "-")),
                            "result": "PASS" if bool(item.get("ok", False)) else "FAIL",
                            "value": str(item.get("value", "-")),
                        }
                    )
                if check_rows:
                    st.table(pd.DataFrame(check_rows))

    with section_export:
        st.caption("Export는 QA 문서, 이벤트 로그, event definition bundle 형태로 제공합니다.")
        if definition_mode_active and not definition_mode_ready:
            st.info("정의서 파일을 업로드하면 그 파일 기준의 QA Report / Bundle을 생성할 수 있습니다.")
        elif timeline_df_active_for_validation.empty:
            st.info("내보낼 실시간 데이터가 없습니다.")
        else:
            browser_timeline = build_realtime_timeline_view(
                timeline_df_for_validation,
                allowed_events=allowed_events_rt,
                unknown_event_policy="무시" if definition_mode_ready else st.session_state.get("unknown_event_policy", "정보"),
            )
            browser_errors_df = (
                browser_timeline[browser_timeline["상태"].isin(["WARN", "ERROR"])].copy()
                if not browser_timeline.empty
                else pd.DataFrame(columns=["시간", "수집원", "event_name", "주요 파라미터", "상태", "상세"])
            )
            event_logs_df = build_event_logs_export_df(rt_events)
            parameter_validation_df = build_parameter_validation_export_df(rt_events, schema_store)
            order_validation_df = build_event_order_validation_df(rt_events, parse_csv_list(funnel_text))
            event_qa_result_df = event_schema_qa_all.copy()
            event_logs_csv = dataframe_to_csv_bytes(event_logs_df)
            event_definition_preview_df, event_support_preview_df, _ = build_event_definition_exports(
                timeline_df_active_for_validation,
                schema_store,
            )
            event_detail_preview_df = build_event_review_detail_df(timeline_df_active_for_validation, schema_store).drop(
                columns=["raw_payload_json"],
                errors="ignore",
            )
            definition_validation_summary_preview_df, definition_validation_matches_preview_df, _ = (
                build_definition_validation_frames(rt_events, schema_store)
                if definition_mode_ready
                else (pd.DataFrame(), pd.DataFrame(), pd.DataFrame())
            )
            qa_result_export_df = build_qa_result_export_df(
                base_df=event_qa_result_df,
                event_detail_df=build_event_review_detail_df(timeline_df_active_for_validation, schema_store),
                schemas=schema_store,
                definition_matches_df=definition_validation_matches_preview_df if definition_mode_ready else None,
            )
            qa_export_running = str(status_for_view).strip() in {"running", "stopping", "recovered(file)"}
            qa_export_label = "QA 결과 파일 (진행 중)" if qa_export_running else "QA 결과 파일 (완료)"
            qa_export_caption = (
                "테스트 진행 중: 현재까지의 QA 결과를 바로 다운로드할 수 있습니다."
                if qa_export_running
                else "테스트 완료: 최종 QA 결과 파일을 다운로드할 수 있습니다."
            )
            qa_report_xlsx = dataframes_to_excel_bytes(
                {
                    "QA Result": qa_result_export_df,
                    "Parameter Validation": parameter_validation_df,
                    "Browser Errors": browser_errors_df,
                    "Event Order Validation": order_validation_df,
                }
            )
            qa_review_xlsx = build_qa_review_excel_bytes(
                event_definition_df=event_definition_preview_df,
                event_detail_df=build_event_review_detail_df(timeline_df_active_for_validation, schema_store),
                definition_validation_summary_df=definition_validation_summary_preview_df,
                bundle_scope="definition_validation" if definition_mode_ready else "custom_event_review",
            )
            log_definition_export_zip = build_log_definition_export_zip_bytes(
                event_definition_df=event_definition_preview_df,
                event_detail_df=build_event_review_detail_df(timeline_df_active_for_validation, schema_store),
                schemas=schema_store,
                definition_matches_df=definition_validation_matches_preview_df if definition_mode_ready else None,
            )
            target_url_for_bundle = (
                str(st.session_state.get("qa_debug_target_url", "")).strip()
                or str(debug_snapshot.get("target_url", "")).strip()
                or (
                    str(timeline_df_active_for_validation.iloc[-1].get("page_url", "")).strip()
                    if not timeline_df_active_for_validation.empty
                    else ""
                )
            )
            bundle_cache_key = (
                f"{sid_for_view}|{len(timeline_df_active_for_validation)}|"
                f"{str(timeline_df_active_for_validation['captured_at'].max()) if 'captured_at' in timeline_df_active_for_validation.columns and not timeline_df_active_for_validation.empty else ''}|"
                f"{selected_mode}|{definition_target_mode}|"
                f"{str(st.session_state.get('qa_definition_validation_file_name', '')).strip()}|"
                f"{','.join(sorted(list(scoped_event_names or set())))}"
            )
            if st.session_state.get("qa_event_definition_bundle_cache_key", "") != bundle_cache_key:
                st.session_state["qa_event_definition_bundle_cache_key"] = bundle_cache_key
                basic_bundle_bytes, debug_bundle_bytes, _, _, preview_bundle_meta = save_event_definition_bundle(
                    raw_debug_df=timeline_df_active_for_validation,
                    schemas=schema_store,
                    qa_report_xlsx=qa_report_xlsx,
                    target_url=target_url_for_bundle,
                    session_id=active_debug_session_id or sid_for_view,
                    scoped_event_names=scoped_event_names if definition_mode_ready else None,
                    definition_file_name=str(st.session_state.get("qa_definition_validation_file_name", "")).strip(),
                    bundle_scope="definition_validation" if definition_mode_ready else "custom_event_review",
                    persist_run=False,
                )
                st.session_state["qa_event_definition_bundle_bytes"] = basic_bundle_bytes
                st.session_state["qa_event_definition_debug_bundle_bytes"] = debug_bundle_bytes
                st.session_state["qa_event_definition_bundle_meta"] = preview_bundle_meta
            if "qa_event_definition_debug_bundle_bytes" not in st.session_state:
                st.session_state["qa_event_definition_debug_bundle_bytes"] = b""

            st.markdown("### Export")
            st.caption("기본 사용자는 QA 결과와 로그정의서 생성 파일만 확인하면 됩니다.")
            b1, b2 = st.columns(2)
            with b1:
                st.download_button(
                    qa_export_label,
                    data=qa_report_xlsx,
                    file_name="qa_result.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    disabled=not bool(qa_report_xlsx),
                    help="사람이 QA 상태만 확인하는 결과 엑셀",
                )
                st.caption(qa_export_caption)
            with b2:
                st.download_button(
                    "로그정의서 생성",
                    data=log_definition_export_zip,
                    file_name="event_definition_export.zip",
                    mime="application/zip",
                    disabled=not bool(log_definition_export_zip),
                    help="annotations/ + event_definition_review.html + event_definition.csv + event_support.csv 구조의 검수용 로그정의서 폴더",
                )
                st.caption("검수 가능한 로그정의서 폴더 구조 ZIP")

            with st.expander("고급 Export", expanded=False):
                st.caption("디버그 로그와 번들은 필요할 때만 사용하세요.")
                g1, g2, g3, g4 = st.columns(4)
                with g1:
                    st.download_button(
                        "Event Logs",
                        data=event_logs_csv,
                        file_name="event_logs.csv",
                        mime="text/csv",
                        disabled=not bool(event_logs_csv),
                        help="테스트 중 수집된 이벤트 로그",
                    )
                with g2:
                    bundle_bytes = st.session_state.get("qa_event_definition_bundle_bytes", b"")
                    st.download_button(
                        "Result Bundle",
                        data=bundle_bytes,
                        file_name="qa_result_bundle.zip",
                        mime="application/zip",
                        disabled=not bool(bundle_bytes),
                        help="최종 리포트 + 요약 CSV + 상세 디버그 CSV만 포함한 기본 번들",
                    )
                with g3:
                    debug_bundle_bytes = st.session_state.get("qa_event_definition_debug_bundle_bytes", b"")
                    st.download_button(
                        "Debug Bundle",
                        data=debug_bundle_bytes,
                        file_name="qa_debug_bundle.zip",
                        mime="application/zip",
                        disabled=not bool(debug_bundle_bytes),
                        help="중간 산출물과 메타데이터까지 포함한 디버그 번들",
                    )
                with g4:
                    save_snapshot_clicked = st.button(
                        "저장본 저장",
                        key="generate_event_definition_bundle",
                        disabled=event_definition_preview_df.empty,
                        help="현재 결과를 Run 저장본으로 보관합니다.",
                    )
                    if save_snapshot_clicked:
                        bundle_bytes, debug_bundle_bytes, _, _, bundle_meta = save_event_definition_bundle(
                            raw_debug_df=timeline_df_active_for_validation,
                            schemas=schema_store,
                            qa_report_xlsx=qa_report_xlsx,
                            target_url=target_url_for_bundle,
                            session_id=active_debug_session_id or sid_for_view,
                            scoped_event_names=scoped_event_names if definition_mode_ready else None,
                            definition_file_name=str(st.session_state.get("qa_definition_validation_file_name", "")).strip(),
                            bundle_scope="definition_validation" if definition_mode_ready else "custom_event_review",
                            persist_run=True,
                        )
                        st.session_state["qa_event_definition_bundle_bytes"] = bundle_bytes
                        st.session_state["qa_event_definition_debug_bundle_bytes"] = debug_bundle_bytes
                        st.session_state["qa_event_definition_bundle_meta"] = bundle_meta
                        st.success(f"저장본 저장 완료: {bundle_meta.get('run_id', '-')}")
                bundle_meta = st.session_state.get("qa_event_definition_bundle_meta", {})
                if bundle_meta:
                    st.caption(
                        f"run_id={bundle_meta.get('run_id', '-')} | "
                        f"rows={bundle_meta.get('definition_row_count', bundle_meta.get('event_count', 0))} | "
                        f"scope={bundle_meta.get('bundle_scope', '-')} | "
                        f"auto_baseline={bundle_meta.get('auto_baseline', False)}"
                    )

            st.markdown("**Event Definition 미리보기**")
            if event_definition_preview_df.empty:
                if ga_hit_count_for_view == 0 and auto_crawl_count_for_view > 0:
                    st.warning("auto crawl 로그만 있고 GA hit가 없어 event definition CSV를 만들 수 없습니다.")
                elif ga_hit_count_for_view > 0:
                    st.info("GA4 hit 중 기본/기술 이벤트를 제외하면 산출할 custom event definition이 없습니다.")
                else:
                    st.info("Event Definition을 만들 데이터가 없습니다.")
            else:
                st.dataframe(event_definition_preview_df, use_container_width=True, height=320)
            st.markdown("**Annotation / Support 미리보기**")
            if event_support_preview_df.empty:
                st.caption("보조 annotation 정보가 없습니다.")
            else:
                st.dataframe(event_support_preview_df, use_container_width=True, height=220)
            st.markdown("**이벤트별 상세 검수 미리보기**")
            if event_detail_preview_df.empty:
                st.caption("이벤트별 상세 검수 산출물이 없습니다.")
            else:
                st.dataframe(event_detail_preview_df, use_container_width=True, height=260)

            st.markdown("**Version History**")
            run_history_df = list_event_definition_runs(target_url_for_bundle)
            if run_history_df.empty:
                st.caption("저장된 Run이 없습니다. 저장본 저장 버튼으로 첫 Run을 만드세요.")
            else:
                history_show_cols = ["run_id", "saved_at", "state", "event_count", "session_id"]
                st.dataframe(run_history_df[history_show_cols], use_container_width=True, height=220)

                baseline_candidates = run_history_df["run_id"].astype(str).tolist()
                default_baseline_idx = 0
                for idx, state in enumerate(run_history_df["state"].astype(str).tolist()):
                    if "Baseline" in state:
                        default_baseline_idx = idx
                        break
                selected_baseline_run = st.selectbox(
                    "Baseline Run",
                    options=baseline_candidates,
                    index=default_baseline_idx,
                    key="qa_compare_baseline_run",
                )
                selected_candidate_run = st.selectbox(
                    "Compare Run",
                    options=baseline_candidates,
                    index=0,
                    key="qa_compare_candidate_run",
                )
                h1, h2 = st.columns([1.1, 2.4])
                with h1:
                    if st.button("선택 Run을 Baseline으로 승격", key="qa_promote_baseline_btn"):
                        set_event_definition_baseline(target_url_for_bundle, selected_baseline_run)
                        st.success(f"Baseline 변경 완료: {selected_baseline_run}")
                        st.rerun()
                with h2:
                    st.caption("Run 저장 후 Baseline/Latest/Candidate 상태로 관리하고, 두 버전 간 차이를 바로 비교합니다.")

                baseline_detail_df = load_event_definition_run_detail_df(target_url_for_bundle, selected_baseline_run)
                candidate_detail_df = load_event_definition_run_detail_df(target_url_for_bundle, selected_candidate_run)
                compare_df = compare_event_definition_runs(baseline_detail_df, candidate_detail_df)
                c1, c2, c3 = st.columns(3)
                c1.metric("Added", int((compare_df["change_type"] == "added").sum()) if not compare_df.empty else 0)
                c2.metric("Removed", int((compare_df["change_type"] == "removed").sum()) if not compare_df.empty else 0)
                c3.metric("Changed", int((compare_df["change_type"] == "changed").sum()) if not compare_df.empty else 0)
                if compare_df.empty:
                    st.caption("선택한 두 버전의 custom event 차이가 없습니다.")
                else:
                    st.dataframe(compare_df, use_container_width=True, height=260)

    if (
        st.session_state.get("realtime_panel_auto_refresh", False)
        and status_for_view in {"running", "stopping", "recovered(file)"}
    ):
        interval_sec = int(st.session_state.get("realtime_panel_refresh_interval", 2))
        components.html(
            f"<script>setTimeout(() => window.parent.location.reload(), {interval_sec * 1000});</script>",
            height=0,
            width=0,
        )
        st.caption(f"실시간 자동 갱신 동작 중 ({interval_sec}초 간격)")

with report_tab:
    if str(st.session_state.get("qa_ui_focus_after_oauth", "")).strip() == "report":
        st.success("로그인 완료. QA 리포트 화면에서 이어서 작업하세요.")
        st.session_state["qa_ui_focus_after_oauth"] = ""
    st.subheader("Google 로그인 및 속성 선택")
    c_auth1, c_auth2 = st.columns([2, 2])
    with c_auth1:
        st.text_input(
            "OAuth Client Secret 파일",
            value=st.session_state.get("qa_report_client_secret_file", get_config_value("GA4_CLIENT_SECRETS_FILE", "client_secret.json")),
            key="qa_report_client_secret_file",
            placeholder="예: client_secret.json",
        )
    with c_auth2:
        st.text_input(
            "OAuth 토큰 파일",
            value=st.session_state.get("qa_report_token_file", get_config_value("GA4_TOKEN_FILE", "token.json")),
            key="qa_report_token_file",
            placeholder="예: token.json",
        )
    redirect_preview = ""
    try:
        redirect_preview = get_google_oauth_redirect_uri(
            st.session_state.get("qa_report_client_secret_file", "").strip() or "client_secret.json"
        )
    except Exception:
        redirect_preview = ""
    if redirect_preview:
        st.caption(f"현재 OAuth Redirect URI: {redirect_preview}")
        st.caption("Google Cloud Console의 Authorized redirect URI에 위 주소가 정확히 등록되어 있어야 합니다.")

    b_auth1, b_auth2 = st.columns([1.3, 1.3])
    with b_auth1:
        oauth_start_clicked = st.button("Google 로그인 시작", key="qa_oauth_start")
    with b_auth2:
        property_refresh_clicked = st.button("속성 리스트 불러오기", key="qa_property_refresh")

    oauth_notice = str(st.session_state.get("qa_oauth_notice", "")).strip()
    oauth_error = str(st.session_state.get("qa_oauth_error", "")).strip()
    if oauth_notice:
        st.success(oauth_notice)
    if oauth_error:
        st.error(oauth_error)

    if oauth_start_clicked:
        log_ui_action("oauth_start_click")
        try:
            oauth_state = uuid4().hex
            auth_url = build_google_oauth_url(
                client_secrets_file=st.session_state.get("qa_report_client_secret_file", "").strip() or "client_secret.json",
                state=oauth_state,
            )
            save_oauth_context(
                oauth_state,
                {
                    "qa_project_slug": st.session_state.get("qa_project_slug", DEFAULT_PROJECT_SLUG),
                    "qa_project_domain": st.session_state.get("qa_project_domain", ""),
                    "qa_debug_session_id": st.session_state.get("qa_debug_session_id", ""),
                    "qa_debug_output_file": st.session_state.get("qa_debug_output_file", ""),
                    "qa_debug_started_at": st.session_state.get("qa_debug_started_at", ""),
                    "qa_debug_target_url": st.session_state.get("qa_debug_target_url", ""),
                    "qa_tester_name": st.session_state.get("qa_tester_name", ""),
                    "qa_tester_note": st.session_state.get("qa_tester_note", ""),
                    "required_event_text_input": st.session_state.get("required_event_text_input", ""),
                    "unknown_event_policy": st.session_state.get("unknown_event_policy", "정보"),
                    "qa_report_property_id": st.session_state.get("qa_report_property_id", ""),
                },
            )
            st.session_state["qa_oauth_state"] = oauth_state
            st.session_state["qa_oauth_auth_url"] = auth_url
            st.session_state["qa_oauth_go_now_url"] = auth_url
            st.session_state["qa_oauth_notice"] = ""
            st.session_state["qa_oauth_error"] = ""
            st.rerun()
        except Exception as exc:
            log_ui_action("oauth_start_error", {"error": str(exc)})
            st.session_state["qa_oauth_error"] = to_user_error_message(exc)
            st.session_state["qa_oauth_auth_url"] = ""
            st.session_state["qa_oauth_go_now_url"] = ""

    go_now_url = str(st.session_state.pop("qa_oauth_go_now_url", "")).strip()
    if go_now_url:
        safe_url = json.dumps(go_now_url)
        st.info("현재 창에서 Google 로그인 페이지로 이동합니다...")
        components.html(
            f"""
            <script>
              (() => {{
                const url = {safe_url};
                try {{
                  if (window.top && window.top.location) {{
                    window.top.location.replace(url);
                    return;
                  }}
                }} catch (e) {{}}
                try {{
                  if (window.parent && window.parent.location) {{
                    window.parent.location.replace(url);
                    return;
                  }}
                }} catch (e) {{}}
                window.location.replace(url);
              }})();
            </script>
            """,
            height=0,
            width=0,
        )
        st.markdown(
            f'<a href="{go_now_url}" target="_self">자동 이동이 안 되면 같은 창에서 Google 로그인 열기</a>',
            unsafe_allow_html=True,
        )
        st.stop()

    pending_auth_url = str(st.session_state.get("qa_oauth_auth_url", "")).strip()
    if pending_auth_url:
        st.markdown(
            f'<a href="{pending_auth_url}" target="_self">같은 창에서 Google 로그인 페이지 열기</a>',
            unsafe_allow_html=True,
        )
        st.caption("로그인 완료 후 앱으로 돌아오면 토큰이 저장되고 속성 리스트를 불러올 수 있습니다.")

    if property_refresh_clicked:
        log_ui_action("ga_property_refresh_click")
        try:
            with st.spinner("GA4 속성 리스트를 조회 중입니다..."):
                prop_df = fetch_ga4_property_list(
                    token_file=st.session_state.get("qa_report_token_file", "").strip() or "token.json"
                )
            st.session_state["qa_report_property_options"] = prop_df.to_dict("records")
            log_ui_action("ga_property_refresh_success", {"count": int(len(prop_df))})
            st.success(f"속성 {len(prop_df)}개를 불러왔습니다.")
        except Exception as exc:
            log_ui_action("ga_property_refresh_error", {"error": str(exc)})
            st.error(to_user_error_message(exc))

    property_options_df = pd.DataFrame(st.session_state.get("qa_report_property_options", []))
    if property_options_df.empty:
        st.text_input(
            "GA4 Property ID (수동 입력)",
            value=st.session_state.get("qa_report_property_id", get_config_value("GA4_PROPERTY_ID", "")),
            key="qa_report_property_id",
            placeholder="예: 123456789",
        )
    else:
        option_labels = property_options_df["label"].astype(str).tolist()
        option_ids = property_options_df["property_id"].astype(str).tolist()
        current_pid = str(st.session_state.get("qa_report_property_id", "")).strip()
        default_idx = option_ids.index(current_pid) if current_pid in option_ids else 0
        selected_label = st.selectbox("GA4 속성 선택", options=option_labels, index=default_idx)
        selected_row = property_options_df[property_options_df["label"] == selected_label].iloc[0]
        st.session_state["qa_report_property_id"] = str(selected_row["property_id"])
        st.caption(
            f"선택된 속성: {selected_row['property_name']} ({selected_row['property_id']}) / 계정: {selected_row['account_name']}"
        )

    st.divider()
    st.subheader("최근 30일 API 이벤트/매개변수")

    api_fetch_clicked = st.button("API 목록 조회(최근 30일)", key="qa_report_fetch_api")
    if api_fetch_clicked:
        log_ui_action("ga_api_ref_fetch_click")
        try:
            with st.spinner("GA4 API에서 이벤트/매개변수 목록을 조회 중입니다..."):
                api_events_df, api_params_df = fetch_ga4_reference_data(
                    property_id=st.session_state.get("qa_report_property_id", "").strip(),
                    token_file=st.session_state.get("qa_report_token_file", "").strip() or "token.json",
                    lookback_days=30,
                    top_events_limit=300,
                )
            st.session_state["qa_report_api_events"] = api_events_df.to_dict("records")
            st.session_state["qa_report_api_params"] = api_params_df.to_dict("records")
            st.session_state["qa_report_api_fetched_at"] = pd.Timestamp.now(tz=LOCAL_TZ).strftime("%Y-%m-%d %H:%M:%S")
            log_ui_action(
                "ga_api_ref_fetch_success",
                {"event_count": int(len(api_events_df)), "param_count": int(len(api_params_df))},
            )
            st.success(
                f"API 조회 완료: 이벤트 {len(api_events_df)}개, 매개변수 {len(api_params_df)}개"
            )
        except Exception as exc:
            log_ui_action("ga_api_ref_fetch_error", {"error": str(exc)})
            st.error(to_user_error_message(exc))

    api_events_view = pd.DataFrame(st.session_state.get("qa_report_api_events", []))
    api_params_view = pd.DataFrame(st.session_state.get("qa_report_api_params", []))
    fetched_at = st.session_state.get("qa_report_api_fetched_at", "").strip()
    if fetched_at:
        st.caption(f"마지막 API 조회 시각: {fetched_at}")

    events_col, params_col = st.columns(2)
    with events_col:
        st.markdown("**이벤트 이름(최근 30일 eventCount 기준)**")
        if api_events_view.empty:
            st.info("API 이벤트 목록이 없습니다. `API 목록 조회(최근 30일)`를 실행하세요.")
        else:
            st.dataframe(api_events_view, use_container_width=True, height=220)
    with params_col:
        st.markdown("**매개변수 이름(등록된 customEvent 메타데이터)**")
        if api_params_view.empty:
            st.info("API 매개변수 목록이 없습니다. `API 목록 조회(최근 30일)`를 실행하세요.")
        else:
            st.dataframe(api_params_view, use_container_width=True, height=220)
            st.caption("참고: 매개변수 목록은 Property 메타데이터 기준입니다.")

    selected_event_options = (
        api_events_view["event_name"].dropna().astype(str).str.strip().tolist()
        if not api_events_view.empty and "event_name" in api_events_view.columns
        else []
    )
    selected_param_options = (
        api_params_view["parameter_name"].dropna().astype(str).str.strip().tolist()
        if not api_params_view.empty and "parameter_name" in api_params_view.columns
        else []
    )
    st.multiselect(
        "핵심 이벤트 선택 (선택 시 이 이벤트 기준으로만 리포트 집계)",
        options=selected_event_options,
        key="qa_report_selected_events",
    )
    st.multiselect(
        "조회 파라미터 선택 (리포트 미리보기 컬럼에 포함)",
        options=selected_param_options,
        key="qa_report_selected_params",
    )
    report_source_mode = st.radio(
        "리포트 판정 데이터 소스",
        options=["GA4 API 집계(최근 30일)", "브라우저 히트(디버깅 세션)"],
        horizontal=True,
        key="qa_report_source_mode",
    )

    st.divider()
    run_clicked = st.button("QA 리포트 생성", type="primary")
    if run_clicked:
        log_ui_action(
            "qa_report_run_click",
            {"source_mode": st.session_state.get("qa_report_source_mode", "")},
        )
        requested_params = parse_csv_list(param_text)
        selected_events = [
            str(v).strip() for v in st.session_state.get("qa_report_selected_events", []) if str(v).strip()
        ]
        selected_params = [
            str(v).strip() for v in st.session_state.get("qa_report_selected_params", []) if str(v).strip()
        ]
        for param_name in selected_params:
            if param_name not in requested_params:
                requested_params.append(param_name)

        required_events = selected_events if selected_events else parse_csv_list(required_event_text)
        required_param_map = parse_required_params(required_param_text)
        funnel_steps = parse_csv_list(funnel_text)
        live_case_field = debug_case_param_name

        if live_case_field and live_case_field not in requested_params:
            requested_params.append(live_case_field)
        if scenario_enabled and scenario_key_field not in {"none", "user_pseudo_id"}:
            if scenario_key_field not in requested_params:
                requested_params.append(scenario_key_field)

        config_notes: List[str] = []
        if selected_events:
            config_notes.append(
                f"핵심 이벤트 {len(selected_events)}개 선택 기준으로 리포트를 집계합니다."
            )
        if not required_events:
            config_notes.append("필수 이벤트 미설정: 해당 체크를 생략합니다.")
        if not required_param_map:
            config_notes.append("필수 파라미터 규칙 미설정: 해당 체크를 생략합니다.")
        if not funnel_steps:
            config_notes.append("퍼널 시퀀스 미설정: 퍼널 기반 체크는 제한됩니다.")

        try:
            if report_source_mode == "브라우저 히트(디버깅 세션)":
                debug_file = Path(st.session_state.get("qa_debug_output_file", "").strip())
                if not debug_file.exists():
                    st.error("디버깅 스트림 데이터가 없습니다. 사이드바에서 '디버깅 모드 시작' 후 테스트를 진행하세요.")
                    st.stop()

                raw_df_all = load_debug_events(debug_file, limit=int(max_rows))
                raw_df_active = filter_events_for_active_session(
                    raw_df_all,
                    session_id=st.session_state.get("qa_debug_session_id", "").strip(),
                    session_started_at=st.session_state.get("qa_debug_started_at", "").strip(),
                )
                raw_df = filter_debug_events_by_view(raw_df_active, "히트(collect)")
                if raw_df.empty:
                    st.error(
                        "실제 수집 히트(collect)가 없습니다. 디버깅 모드 시작 후 이벤트를 발생시키고, "
                        "기존 세션이라면 새 세션으로 다시 시작하세요."
                    )
                    st.stop()

                qa_df = debug_events_to_qa_df(raw_df, requested_params)
                if "event_timestamp" in qa_df.columns and not qa_df.empty:
                    ts = pd.to_datetime(qa_df["event_timestamp"], errors="coerce", utc=True)
                    start_local = pd.Timestamp(start_date).tz_localize(LOCAL_TZ)
                    end_local = (pd.Timestamp(end_date) + pd.Timedelta(days=1)).tz_localize(LOCAL_TZ)
                    start_ts = start_local.tz_convert("UTC")
                    end_ts = end_local.tz_convert("UTC")
                    qa_df = qa_df[(ts >= start_ts) & (ts < end_ts)].copy()
                    if qa_df.empty:
                        st.error("선택한 기간에 해당하는 수집 히트가 없습니다. 기간을 조정하세요.")
                        st.stop()
                if selected_events:
                    qa_df = qa_df[qa_df["event_name"].astype(str).isin(selected_events)].copy()
                    if qa_df.empty:
                        st.error("선택한 핵심 이벤트 기준으로는 수집 히트가 없습니다. 이벤트 선택을 조정하세요.")
                        st.stop()

                results_df = run_qa_rules(
                    qa_df,
                    required_events=required_events,
                    required_params_by_event=required_param_map,
                    null_threshold=float(null_threshold),
                    funnel_steps=funnel_steps,
                )
            else:
                property_id = st.session_state.get("qa_report_property_id", "").strip()
                token_file = st.session_state.get("qa_report_token_file", "").strip() or "token.json"
                if not property_id:
                    st.error("GA4 속성을 먼저 선택하세요.")
                    st.stop()
                api_events_df, _ = fetch_ga4_reference_data(
                    property_id=property_id,
                    token_file=token_file,
                    lookback_days=30,
                    top_events_limit=300,
                )
                if selected_events:
                    api_events_df = api_events_df[api_events_df["event_name"].astype(str).isin(selected_events)].copy()
                qa_df = build_api_event_preview_df(api_events_df)
                if qa_df.empty:
                    st.error("선택한 핵심 이벤트 기준으로 30일 집계 데이터가 없습니다.")
                    st.stop()
                results_df = build_api_aggregate_results(
                    api_events_df,
                    required_events=required_events,
                    required_params_by_event=required_param_map,
                    selected_params=selected_params,
                )
                config_notes.append("현재 리포트는 GA4 API 30일 집계 데이터 기준으로 생성되었습니다.")

            results_view = build_qa_results_view(results_df)
            issue_property = st.session_state.get("qa_report_property_id", "").strip() or "browser_intercept_stream"
            new_issues = upsert_auto_issues(results_df, property_id=issue_property, issue_path=DEFAULT_ISSUE_PATH)

            if config_notes:
                for note in config_notes:
                    st.info(note)

            st.subheader("저장 후 오류")
            st.caption(f"실제 수집 히트(collect) 기준 판정 결과 | 신규 이슈 저장: {new_issues}건")
            st.dataframe(results_view, use_container_width=True, height=260)

            st.subheader("null 비율")
            null_view = results_view[results_view["구분"] == "Null 비율"] if not results_view.empty else results_view
            if null_view.empty:
                st.success("null 비율 검사 항목이 없습니다.")
            else:
                st.dataframe(null_view, use_container_width=True, height=220)

            st.subheader("중복 검사")
            dup_view = results_view[results_view["rule_id"].astype(str).str.startswith("duplicate_")] if not results_view.empty else results_view
            if dup_view.empty:
                st.success("중복 검사 항목이 없습니다.")
            else:
                st.dataframe(dup_view, use_container_width=True, height=180)

            st.subheader("정규화 이벤트 미리보기")
            st.dataframe(build_event_preview(qa_df, requested_params), use_container_width=True, height=320)
            log_ui_action(
                "qa_report_run_success",
                {
                    "source_mode": report_source_mode,
                    "row_count": int(len(qa_df)),
                    "rule_count": int(len(results_df)),
                },
            )

        except Exception as exc:
            log_ui_action("qa_report_run_error", {"error": str(exc)})
            st.error(to_user_error_message(exc))
            with st.expander("기술 상세"):
                st.code(str(exc))

    st.subheader("이슈 히스토리")
    issues_df = issues_to_dataframe(DEFAULT_ISSUE_PATH)
    st.dataframe(issues_df, use_container_width=True, height=260)

    open_issues = issues_df[issues_df["status"] == "open"] if not issues_df.empty else issues_df
    if not open_issues.empty:
        st.markdown("**이슈 해결 기록 추가**")
        issue_options = open_issues["issue_id"].tolist()
        issue_id = st.selectbox("해결할 issue_id", options=issue_options)
        resolution = st.text_input("해결 내용", placeholder="예: checkout dataLayer 누락 수정")
        if st.button("이슈 해결 처리"):
            if not resolution.strip():
                st.error("해결 내용을 입력하세요.")
            else:
                ok = resolve_issue(issue_id=issue_id, resolution=resolution, issue_path=DEFAULT_ISSUE_PATH)
                if ok:
                    log_ui_action("issue_resolved", {"issue_id": str(issue_id)})
                    st.success("이슈를 resolved 상태로 저장했습니다.")
                    st.rerun()
                else:
                    st.error("issue_id를 찾지 못했습니다.")
    else:
        st.caption("현재 open 상태 이슈가 없습니다.")

with schema_tab:
    st.subheader("Event Rule Manager (이벤트 기준표)")
    st.caption(
        f"Project: {get_active_project_slug()} | 이벤트별 필수/선택 파라미터 기준표를 관리합니다."
    )

    cached_rt_events = pd.DataFrame(st.session_state.get("schema_rt_events_cache", []))
    schema_store = load_event_schemas()

    captured_events = []
    if not cached_rt_events.empty and "이벤트" in cached_rt_events.columns:
        captured_events = sorted(
            list({str(v).strip() for v in cached_rt_events["이벤트"].dropna().tolist() if str(v).strip()})
        )
    saved_events = sorted(list(schema_store.keys()))
    selectable_events = sorted(list(dict.fromkeys(captured_events + saved_events)))

    if not selectable_events:
        st.info("기준표를 만들 이벤트가 없습니다. 먼저 실시간 테스트에서 이벤트를 수집하세요.")
    else:
        selected_event = st.selectbox("이벤트 선택", options=selectable_events, key="schema_manager_selected_event")
        include_system_params_schema_tab = st.checkbox(
            "기술 파라미터까지 포함",
            value=False,
            key="schema_tab_include_system_params",
            help="기본은 시스템/기술 파라미터를 제외합니다.",
        )
        inferred = infer_event_schema_from_rows(
            cached_rt_events,
            selected_event,
            include_system_params=include_system_params_schema_tab,
        )
        existing_schema = schema_store.get(selected_event, {})
        all_params = sorted(
            list(
                dict.fromkeys(
                    list(inferred.get("required", []))
                    + list(inferred.get("optional", []))
                    + [str(v).strip() for v in existing_schema.get("required", []) if str(v).strip()]
                    + [str(v).strip() for v in existing_schema.get("optional", []) if str(v).strip()]
                )
            )
        )

        c1, c2 = st.columns([1.2, 1.2])
        with c1:
            if st.button("선택 이벤트 기준표 자동 생성", key="schema_generate_for_selected"):
                required_auto = list(inferred.get("required", []))
                optional_auto = list(inferred.get("optional", []))
                schema_store[selected_event] = {
                    "required": required_auto,
                    "optional": optional_auto,
                    "updated_at": datetime.now(tz=ZoneInfo("UTC")).isoformat(),
                }
                save_event_schemas(schema_store)
                st.success(
                    f"{selected_event}: 필수 {len(required_auto)}개 / 선택 {len(optional_auto)}개 저장"
                )
                if not include_system_params_schema_tab:
                    st.info(f"시스템/기술 파라미터 제외: {int(inferred.get('excluded_system_params', 0))}건")
                st.rerun()
        with c2:
            st.caption(
                f"자동 생성 기준 샘플 수: {int(inferred.get('sample_size', 0))}건 | "
                f"제외된 기술 파라미터: {int(inferred.get('excluded_system_params', 0))}건"
            )

        current_required = [str(v).strip() for v in existing_schema.get("required", []) if str(v).strip()]
        current_optional = [str(v).strip() for v in existing_schema.get("optional", []) if str(v).strip()]

        required_keys = st.multiselect(
            "required",
            options=all_params,
            default=[v for v in current_required if v in all_params],
            key=f"schema_required_{selected_event}",
        )
        optional_keys = st.multiselect(
            "optional",
            options=[v for v in all_params if v not in required_keys],
            default=[v for v in current_optional if v in all_params and v not in required_keys],
            key=f"schema_optional_{selected_event}",
        )

        b1, b2 = st.columns([1.2, 1.2])
        with b1:
            if st.button("기준표 저장", key=f"schema_save_{selected_event}"):
                schema_store[selected_event] = {
                    "required": required_keys,
                    "optional": optional_keys,
                    "updated_at": datetime.now(tz=ZoneInfo("UTC")).isoformat(),
                }
                save_event_schemas(schema_store)
                st.success(f"{selected_event} 기준표를 저장했습니다.")
                st.rerun()
        with b2:
            if st.button("기준표 삭제", key=f"schema_delete_{selected_event}"):
                if selected_event in schema_store:
                    del schema_store[selected_event]
                    save_event_schemas(schema_store)
                    st.success(f"{selected_event} 기준표를 삭제했습니다.")
                    st.rerun()

        preview = validate_schema_for_event(
            event_name=selected_event,
            params={k: "sample" for k in required_keys},
            schemas={selected_event: {"required": required_keys, "optional": optional_keys}},
        )
        st.caption(f"기준표 미리보기: {preview.get('status', 'WARN')} · {preview.get('message', '-')}")

    st.divider()
    st.markdown("**저장된 기준표 목록**")
    schema_rows = []
    for event_name, schema in sorted(schema_store.items()):
        required = [str(v).strip() for v in schema.get("required", []) if str(v).strip()]
        optional = [str(v).strip() for v in schema.get("optional", []) if str(v).strip()]
        schema_rows.append(
            {
                "event_name": event_name,
                "required_count": len(required),
                "optional_count": len(optional),
                "required": ", ".join(required),
                "optional": ", ".join(optional),
                "updated_at": str(schema.get("updated_at", "")).strip(),
            }
        )
    if schema_rows:
        st.dataframe(pd.DataFrame(schema_rows), use_container_width=True, height=280)
    else:
        st.info("저장된 기준표가 없습니다.")

    st.divider()
    st.markdown("**Action/Object 매핑표**")
    st.caption("tracking_plan의 action/object/page/section을 이벤트별로 고정합니다.")
    map_col1, map_col2 = st.columns([1.2, 2.0])
    with map_col1:
        st.download_button(
            "매핑표 템플릿 다운로드",
            data=action_map_template_csv_bytes(),
            file_name="action_object_map_template.csv",
            mime="text/csv",
        )
    with map_col2:
        st.caption("필수 컬럼: event_name, action, object (나머지는 선택)")
    upload_file = st.file_uploader(
        "매핑표 업로드 (CSV)",
        type=["csv"],
        key="action_object_map_upload",
        accept_multiple_files=False,
    )
    if st.button("업로드 매핑표 저장", key="save_action_object_map_btn"):
        if upload_file is None:
            st.warning("업로드할 CSV 파일을 먼저 선택하세요.")
        else:
            try:
                uploaded_df = pd.read_csv(upload_file)
                normalized_df = _normalize_action_map_df(uploaded_df)
                if normalized_df.empty:
                    st.error("유효한 event_name 행이 없습니다. 템플릿 형식으로 다시 업로드하세요.")
                else:
                    save_action_object_map(normalized_df)
                    st.success(f"매핑표 저장 완료: {len(normalized_df)}개 이벤트")
                    st.rerun()
            except Exception as exc:
                st.error(f"매핑표 업로드 실패: {to_user_error_message(exc)}")

    current_map_df = load_action_object_map()
    if current_map_df.empty:
        st.info("저장된 매핑표가 없습니다. 기본 규칙(이벤트명 추정)으로 tracking_plan을 생성합니다.")
    else:
        st.dataframe(current_map_df, use_container_width=True, height=240)
